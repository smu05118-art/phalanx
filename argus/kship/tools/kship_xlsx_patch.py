#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""kship_xlsx_patch — 레퍼런스 subQ 모델(xlsx)의 '박힌' FnGuide 시트를 제자리 패치.

세진중공업·삼성중공업·HD현대미포의 애널리스트 모델 원본(`*_subQ_orig.xlsx`, 사용자 사유 파일,
비공개 `KSHIP_REFERENCE_DIR`)은 도형·스파크라인·차트 때문에 openpyxl 로 열어
저장하면 깨진다. 그래서 **쓰기는 zipfile + XML 문자열 수준**에서만 한다 — 원본 zip 멤버를 바이트
그대로 옮기고, BS 시트(`BS연결`/`BS별도`/세진의 `BS일승(별도)`·`BS동방선기(별도)`)의 새 기간 열에
fin json(kship_fin.py 산출, 백만원) 값을 셀로 끼워 넣는다.

규약(MODEL_SPEC 2-6):
- 원본 `*_orig.xlsx` 는 절대 덮어쓰지 않는다. 출력은 `<이름>_<stock>_subQ_<분기>.xlsx`.
- 기존 셀은 한 바이트도 바꾸지 않는다. 값이 이미 있는 셀(애널리스트 가이던스 상수·수식)은
  건너뛰고 `conflicts` 로 보고한다. 예외는 명시된 세 곳: 행3 UPDATE 스탬프, `종가` 이름정의 셀과
  그 라벨(prices.json 있을 때), 행4 의 비어 있던 기간 라벨.
- 새 기간 열의 위치는 행1 라벨('4Q23', 2023, '1Q24'…)로 찾는다 — 열 문자를 하드코딩하지 않는다.
- 문자열 라벨('2023.12A', 'UPDATE: 26-09-30', '종가(09월 30일)')은 sharedStrings 끝에 추가한다.
- 셀 스타일(s)은 같은 행의 직전 셀 것을 복사한다. 비어 있던 자리 셀(`<c r="CS5" s="77"/>`)은
  그 셀의 s 를 그대로 쓴다.
- IS 분기 값은 **`is_ytd_diff` 우선**(FnGuide 관행 = 누적 차분: Q1=누적, Qn=누적−전누적, Q4=연간−3Q누적;
  MODEL_SPEC 5-4·결정 ⓐ). 그 계정·분기에 `is_ytd_diff` 가 없으면 `is`(보고서 3개월 열 — 주석으로 채운 이자수익 등은
  3개월 값만 있어 자동으로 이 길을 탄다). `--is-convention 3m` 이면 예전처럼 항상 `is`. 둘이 1백만원 초과로 다른
  셀(후속 보고서의 전기 재작성)은 `sheets[].restated_cells[]`(is·ytd_diff·diff) 에, 적용 관행은 `is_convention` 에 남긴다.
- 연간 열(YYYY.12A): BS 는 Q4 시점값, IS 는 `is_ytd[Q4]`(없으면 4분기 합), CF 는 `cf_ytd[Q4]`
  (없으면 4분기 합), 주식수·배당은 Q4 값. 관행과 무관.
- workbook.xml `<calcPr … fullCalcOnLoad="1"/>` — 열 때 전체 재계산(subQ 의 VLOOKUP 이 새 열을 읽는다).
- 옵트인 두 가지(기본 꺼짐):
  · `--overwrite-placeholders` — 새 기간 열에 애널리스트가 미발표 분기용으로 넣어 둔 가이던스·잠정치 셀(삼성重 4Q23
    11셀·미포 1Q25 9셀)을 fin 실적으로 바꾸고 원문을 보고 `sheets[].replaced[].was` 에 남긴다. 공유수식 그룹은 그룹
    전 셀이 함께 바뀔 때만(앵커만 지우면 의존 셀 `<f t="shared" si>` 가 고아가 된다). 수식 자리표시자를 바꾼 만큼
    수식 셀 수가 줄어드는 것은 검증 ②에서 허용 감소분으로 뺀다.
  · `--fx-actuals` — `변수` 환율 8행의 같은 기간 열(BS 새 기간과 동일: 세진·삼성重 2023Q4~2026Q2, 미포 2025Q1~Q3)에
    남아 있는 애널리스트 가정을 fx.json 실측으로 바꾼다. fx.json 이 `partial`(2026Q3~ 진행 중)이거나 없는 기간은
    건드리지 않아 이후 가정은 보존된다. 행 단위 배율은 교체 범위 앞 과거 실적 구간(2005Q1~)의 시트값/fx 중앙값 비율로
    정한다 — 레퍼런스 ` 원/100Y` 행은 라벨과 달리 원/엔(=fx JPY100KRW/100) 이어서 0.01 이 나온다. 배율을 못 정한
    행은 건너뛰고 보고한다. 교체 전후는 `fx.replaced[]`(was_value→now·scale) 에 남는다.
- 미포 출력이 `_2025Q3` 인 이유: HD현대미포는 2025Q4 부터 HD현대重에 합병돼 보고서가 없다(fin.quarters 마지막 =
  2025Q3). target 기본값이 fin 의 마지막 분기라 파일명도 2025Q3 이다(세진·삼성重은 2026Q2).
- calcChain.xml(세 원본 모두 있음 — 항목 수 == 수식 셀 수): `--overwrite-placeholders` 로 수식 자리표시자를 상수로 바꾼 셀의
  항목을 함께 뺀다(`calc_chain.pruned`). 수식이 없는 셀을 가리키는 항목이 남으면 Excel 이 "복구된 레코드: /xl/calcChain.xml"
  대화상자를 띄운다(ECMA-376 18.6.2 — 항목은 수식 셀만). 새로 쓴 수식 셀(--extend-formulas)은 항목이 없어도 된다(열 때 재구성).
- `--extend-formulas [plain|all|off]`(CLI 기본 켜짐 = all — 2026-10-08 오너 결정, `off` 로 끈다; `patch_file(extend=None)` API 기본은 꺼짐) — subQ 의 새 기간 열에서 **비어 있는** 셀을 같은 행의 소스 셀
  수식으로 채운다(사용자가 끌어 채우던 것). 소스 = 목표와 같은 분기 위치(1Q~4Q·연간)의 가장 최근 실적 열(3Q23 이 마지막이면
  4Q23←4Q22·2023←2022·1Q24←1Q23 …; 열 차이는 5의 배수라 1Q 열의 '직전 열=연간' 구조·연간 열의 SUM 범위가 유지된다).
  상대참조는 openpyxl Translator(토크나이저)로 열 차이만큼 치환하고 `$` 고정·이름정의(BS연결·SUBQH·U)·문자열은 그대로 —
  VLOOKUP 확정 행은 텍스트가 같다. `plain` 은 소스가 단순 `<f>` 인 셀만, `all` 은 공유수식 셀도 **앵커 텍스트를 앵커 기준으로
  치환한 단순 수식**으로 쓴다(그룹 자체는 손대지 않음 — 새 셀은 그룹 밖). 상수·배열수식·빈 소스·값이 있는 목표 셀은 건너뛴다.
  새 셀은 `<f>` 만(캐시값 없음, fullCalcOnLoad 로 계산). 보고: `extend`(연장 수·plain/shared·행 수·건너뜀·수식 목록).
- REF_DIR 기본값: `<레포>/jem_data/kship_models/reference` 가 없으면 `~/phalanx/jem_data/kship_models/reference`
  (레포 밖 작업 트리에서도 통합 테스트가 조용히 건너뛰지 않게). `KSHIP_REFERENCE_DIR` 가 항상 우선.

검증(`--verify`, tests/test_kship_xlsx_patch.py 가 같은 함수를 부른다):
 ① 패치본 zip 무결(testzip) ② openpyxl(도형 우회) 재오픈 → 시트 수·수식 셀 개수 원본과 동일(교체한 수식
 자리표시자 수만큼 감소 허용) ③ 패치 셀 목록·개수 ④ 원본 셀 전부 동일(허용 예외만 다름 — 교체 셀은 반드시 새
 기간 열 안이어야 하며 보고의 replaced 목록을 그대로 믿지 않는다) ⑤ subQ `매출액(수익)` 확정 행의
 `VLOOKUP($D, BS연결, MATCH(SUBQH, BS연결H, 0), 0)` 사슬을 파이썬으로 흉내내어 새 분기 값 == fin
 ⑥ calcChain.xml 전 항목이 패치본의 수식 셀을 가리킨다(고아 항목 0, 줄어든 수 == 교체한 자리표시자 수)
 ⑦ --extend-formulas 를 켰으면: 쓴 셀이 재오픈 시 전부 수식이고 텍스트가 보고와 같다 · 수식 수 증가분 == 연장 수(②) ·
 `VLOOKUP($D,<BS표>,MATCH(<헤더>,<BS표>H,0),0)[/U]` 모양 셀은 사슬 흉내 값(패치된 BS 셀 ÷ U) == fin ÷ U.

실행:
  python3 kship_xlsx_patch.py --stock 010140 --verify
  python3 kship_xlsx_patch.py --all --verify --report assets/xlsx_patch_report.json
  python3 kship_xlsx_patch.py --all --verify --overwrite-placeholders --fx-actuals --today 2026-09-30
  python3 kship_xlsx_patch.py --all --verify --is-convention 3m   # 예전 방식(항상 3개월 열)
  python3 kship_xlsx_patch.py --stock 010140 --verify --overwrite-placeholders --fx-actuals --is-convention 3m \
      --extend-formulas off --out /tmp/shi_plain.xlsx                 # 수식 연장 끄기(기본은 켜짐)
표준 라이브러리 + openpyxl(검증·--extend-formulas 의 수식 치환) 만 쓴다. 네트워크 없음.
"""
import argparse
import datetime
import json
import os
import re
import sys
import zipfile
from xml.sax.saxutils import escape, unescape

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
FIN_DIR = os.path.join(ASSETS, "fin")
PHALANX_ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
REF_DIR_CANDIDATES = (os.path.join(PHALANX_ROOT, "jem_data", "kship_models", "reference"),
                      os.path.expanduser("~/phalanx/jem_data/kship_models/reference"))


def default_ref_dir(candidates=REF_DIR_CANDIDATES):
    """레퍼런스 폴더 기본값 — 후보 중 처음 존재하는 것, 없으면 첫 후보(레포 기준). 환경변수가 있으면 그것."""
    env = os.environ.get("KSHIP_REFERENCE_DIR")
    if env:
        return env
    for c in candidates:
        if os.path.isdir(c):
            return c
    return candidates[0]


REF_DIR = default_ref_dir()
CALC_CHAIN = "xl/calcChain.xml"

# 레퍼런스 파일 → 패치 대상 시트(시트명 → (종목, 범위)). 시트명은 workbook.xml 로 파일을 찾는다.
FILES = {
    "075580": {"name": "세진중공업",
               "sheets": {"BS연결": ("075580", "cons"), "BS별도": ("075580", "sep"),
                          "BS일승(별도)": ("333430", "sep"), "BS동방선기(별도)": ("099410", "sep")}},
    "010140": {"name": "삼성중공업",
               "sheets": {"BS연결": ("010140", "cons"), "BS별도": ("010140", "sep")}},
    "010620": {"name": "HD현대미포",
               "sheets": {"BS연결": ("010620", "cons"), "BS별도": ("010620", "sep")}},
}

# `변수` 시트 환율 행 라벨(열 C) → fx.json quarters/annual 키. 레퍼런스 세 파일 공통(앞 공백 포함).
FX_ROWS = [(" 원/달러(평균)", "USDKRW_avg"), (" 원/달러(기말)", "USDKRW_end"),
           (" 원/100Y(평균)", "JPY100KRW_avg"), (" 원/100Y(기말)", "JPY100KRW_end"),
           (" W/Euro(평균)", "EURKRW_avg"), (" W/Euro(기말)", "EURKRW_end"),
           (" 원/중국원(평균)", "CNYKRW_avg"), (" 원/중국원(기말)", "CNYKRW_end")]

# FnGuide 주식수 계정(백만주) ← fin.shares[q] (주 단위). 우선주 = 발행 − 보통주.
SHARE_MAP = {"기말발행주식수(백만주)": "issued", "보통주기말발행주식수": "common_issued",
             "보통주기말자기주식수": "treasury"}
# 시가총액(십억원) ← prices.history_quarterly[q] × 보통주 발행주식수. 히스토리 없으면 비움.
MCAP_MAP = {"보통주시가총액(기말, 십억원)": "close_end", "보통주시가총액(최고, 십억원)": "high",
            "보통주시가총액(최저, 십억원)": "low", "보통주시가총액(평균, 십억원)": "avg"}
DPS_ACCT = "수정DPS(보통주, 기말현금)"
# IS 분기 값 관행 — "ytd_diff": fin.is_ytd_diff 우선(FnGuide 누적차분), "3m": 항상 fin.is(보고서 3개월 열).
IS_CONVENTIONS = ("ytd_diff", "3m")
IS_CONVENTION_DEFAULT = "ytd_diff"
RESTATED_TOL = 1.0          # 백만원 — is 와 is_ytd_diff 가 이보다 크게 다르면 재작성 셀(kship_fin 의 restated 문턱과 같다)

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
CELL_RE = re.compile(r'<c r="([A-Z]+)(\d+)"((?:\s+[\w:]+="[^"]*")*)\s*(?:/>|>(.*?)</c>)', re.S)
SI_RE = re.compile(r"<si>(.*?)</si>", re.S)
T_RE = re.compile(r"<t(?:\s[^>]*)?>([^<]*)</t>", re.S)


# ── 열 문자 ────────────────────────────────────────────────

def col_idx(letters):
    """'A'→1, 'CS'→97."""
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def col_letters(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


# ── 기간 ──────────────────────────────────────────────────

def q_parse(q):
    """'2026Q2' → (2026, 2)."""
    m = re.fullmatch(r"(\d{4})Q([1-4])", q)
    if not m:
        raise ValueError("분기키 형식 오류: %r" % q)
    return int(m.group(1)), int(m.group(2))


def period_seq(last, target_q):
    """마지막 실적 기간 다음부터 target 분기까지의 기간 열 순서.

    last: ('Q', y, n) 또는 ('A', y). 반환 원소: ('Q', y, n) | ('A', y). Q4 뒤에는 연간이 온다.
    """
    ty, tn = q_parse(target_q)
    out = []
    if last[0] == "A":
        y, n = last[1] + 1, 1
    else:
        y, n = last[1], last[2]
        if n == 4:
            out.append(("A", y))
            y, n = y + 1, 1
        else:
            n += 1
    while (y, n) <= (ty, tn):
        out.append(("Q", y, n))
        if n == 4:
            out.append(("A", y))
            y, n = y + 1, 1
        else:
            n += 1
    return out


def period_key(p):
    """('Q',2026,2)→'2026Q2', ('A',2025)→'2025A'."""
    return "%dQ%d" % (p[1], p[2]) if p[0] == "Q" else "%dA" % p[1]


def row1_label(p):
    """행1(BS연결H) 라벨: 분기 '2Q26'(문자열), 연간 2026(숫자)."""
    return "%dQ%02d" % (p[2], p[1] % 100) if p[0] == "Q" else str(p[1])


def row4_label(p):
    """행4 기간 라벨: 분기 숫자 'YYYY.MM', 연간 문자열 'YYYY.12A'."""
    return ("%d.%02d" % (p[1], p[2] * 3), False) if p[0] == "Q" else ("%d.12A" % p[1], True)


def parse_row4_label(text, is_str):
    """행4 라벨 → ('Q',y,n) | ('A',y) | None."""
    if is_str:
        m = re.fullmatch(r"(\d{4})\.12A", text.strip())
        return ("A", int(m.group(1))) if m else None
    try:
        v = float(text)
    except ValueError:
        return None
    y = int(v)
    mm = int(round((v - y) * 100))
    if 2000 <= y <= 2100 and mm in (3, 6, 9, 12):
        return ("Q", y, mm // 3)
    return None


# ── sharedStrings ─────────────────────────────────────────

class SharedStrings:
    """xl/sharedStrings.xml — 기존 항목은 바이트 그대로, 새 문자열은 끝에 append."""

    def __init__(self, xml):
        self.xml = xml
        self.texts = [unescape("".join(T_RE.findall(si))) for si in SI_RE.findall(xml)]
        self.n_orig = len(self.texts)
        self.index = {}
        for i, t in enumerate(self.texts):
            self.index.setdefault(t, i)
        self.added = []
        self.ref_delta = 0            # t="s" 참조 셀 증감(count 속성)

    def get(self, i):
        return self.texts[i]

    def find(self, text):
        return self.index.get(text)

    def add(self, text):
        """없으면 추가하고 인덱스를 돌려준다. 참조 1건 추가로 센다."""
        self.ref_delta += 1
        i = self.index.get(text)
        if i is not None:
            return i
        i = len(self.texts)
        self.texts.append(text)
        self.index[text] = i
        self.added.append(text)
        return i

    def release(self):
        """기존 t="s" 셀 하나를 다른 문자열로 바꿨을 때 참조 1건 감소."""
        self.ref_delta -= 1

    def render(self):
        if not self.added and self.ref_delta == 0:
            return self.xml
        xml = self.xml

        def _attr(m):
            head = m.group(0)
            head = re.sub(r'\bcount="(\d+)"', lambda k: 'count="%d"' % (int(k.group(1)) + self.ref_delta), head, count=1)
            head = re.sub(r'uniqueCount="(\d+)"', lambda k: 'uniqueCount="%d"' % (int(k.group(1)) + len(self.added)), head, count=1)
            return head
        xml = re.sub(r"<sst[^>]*>", _attr, xml, count=1)
        parts = []
        for t in self.added:
            sp = ' xml:space="preserve"' if t != t.strip() else ""
            parts.append("<si><t%s>%s</t></si>" % (sp, escape(t)))
        i = xml.rfind("</sst>")
        return xml[:i] + "".join(parts) + xml[i:]


# ── 시트 XML 조각 ─────────────────────────────────────────

def sheet_paths(wb_xml, rels_xml):
    """시트명 → zip 멤버 경로. workbook.xml sheet name→r:id→rels Target."""
    rid2t = {m.group(1): m.group(2)
             for m in re.finditer(r'<Relationship\s[^>]*?Id="(rId\d+)"[^>]*?Target="([^"]+)"', rels_xml)}
    rid2t.update({m.group(2): m.group(1)
                  for m in re.finditer(r'<Relationship\s[^>]*?Target="([^"]+)"[^>]*?Id="(rId\d+)"', rels_xml)})
    out = {}
    for m in re.finditer(r"<sheet\s([^>]*)/>", wb_xml):
        attrs = m.group(1)
        nm = re.search(r'name="([^"]*)"', attrs)
        rid = re.search(r'r:id="(rId\d+)"', attrs)
        if nm and rid and rid.group(1) in rid2t:
            t = rid2t[rid.group(1)]
            out[unescape(nm.group(1))] = t if t.startswith("xl/") else "xl/" + t.lstrip("/")
    return out


def defined_names(wb_xml):
    """이름정의(전역만) → 참조 문자열. 예: '종가' → \"'TP_PE PB'!$B$1\"."""
    out = {}
    for m in re.finditer(r"<definedName\s([^>]*)>([^<]*)</definedName>", wb_xml):
        attrs, ref = m.group(1), unescape(m.group(2))
        if "localSheetId=" in attrs:
            continue
        nm = re.search(r'name="([^"]*)"', attrs)
        if nm:
            out[unescape(nm.group(1))] = ref
    return out


def split_ref(ref):
    """\"'TP_PE PB'!$B$1\" → ('TP_PE PB', 'B', 1). 범위면 첫 셀. 열 범위('BS연결!$C:$FF')면 행은 None."""
    m = re.fullmatch(r"'?(.*?)'?!\$?([A-Z]+)\$?(\d*)(?::.*)?", ref or "")
    if not m:
        return None
    return m.group(1).replace("''", "'"), m.group(2), (int(m.group(3)) if m.group(3) else None)


def find_row(xml, r):
    """행 r 의 (start, end, head, inner). 없으면 None. 자기닫힘 행도 처리."""
    m = re.search(r'<row r="%d"(?:\s[^>]*?)?(/>|>)' % r, xml)
    if not m:
        return None
    if m.group(1) == "/>":
        return m.start(), m.end(), m.group(0), ""
    end = xml.index("</row>", m.end())
    return m.start(), end + len("</row>"), m.group(0), xml[m.end():end]


def parse_cells(inner):
    """행 내부 → [{'col','ci','attrs','inner','raw','has_value','s'}] (열 순서)."""
    cells = []
    for m in CELL_RE.finditer(inner):
        col, attrs, body = m.group(1), m.group(3) or "", m.group(4) or ""
        s = re.search(r'\ss="(\d+)"', attrs)
        cells.append({"col": col, "ci": col_idx(col), "attrs": attrs, "inner": body, "raw": m.group(0),
                      "has_value": bool(re.search(r"<(v|f|is)\b", body)),
                      "s": s.group(1) if s else None,
                      "t": (re.search(r'\st="(\w+)"', attrs) or [None, None])[1]})
    return cells


class Formula:
    """수식 셀 값 표식 — rebuild_row 의 writes 값으로 주면 `<f>…</f>` 셀(캐시값 없음)이 된다(--extend-formulas)."""
    __slots__ = ("text",)

    def __init__(self, text):
        self.text = text

    def __repr__(self):
        return "Formula(%r)" % (self.text,)

    def __eq__(self, other):
        return isinstance(other, Formula) and other.text == self.text


def cell_xml(col, r, s, value, is_sst=False):
    """숫자 · sharedStrings 인덱스 · 수식(Formula) 셀."""
    sa = ' s="%s"' % s if s is not None else ""
    if isinstance(value, Formula):
        return '<c r="%s%d"%s><f>%s</f></c>' % (col, r, sa, escape(value.text))
    if is_sst:
        return '<c r="%s%d"%s t="s"><v>%d</v></c>' % (col, r, sa, value)
    return '<c r="%s%d"%s><v>%s</v></c>' % (col, r, sa, fmt_num(value))


def fmt_num(v):
    """XML 숫자 표기 — 정수는 정수로, 실수는 최단 왕복 표기."""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int):
        return str(v)
    f = float(v)
    if f != f or f in (float("inf"), float("-inf")):
        raise ValueError("숫자 아님: %r" % (v,))
    if f.is_integer() and abs(f) < 1e15:
        return str(int(f))
    return repr(f)


def rebuild_row(xml, r, writes, overwrite_ok=None):
    """행 r 에 writes({colLetter: (value, is_sst)}) 를 끼워 넣은 새 XML 과 결과 목록.

    기존 값 셀은 절대 바꾸지 않는다(conflict). 비어 있던 자리 셀은 그 s 를, 없는 셀은 직전 셀 s 를 쓴다.
    예외는 `overwrite_ok(ref)` 가 True 를 주는 셀 — 애널리스트가 미발표 분기에 넣어 둔 가이던스·잠정치
    자리표시자(--overwrite-placeholders)만 실적으로 바꾸고, 바꾼 셀은 replaced 로 따로 돌려준다(원문 raw 보존).
    반환: (new_xml, written[(col, s)], conflicts[(col, raw)], replaced[(col, raw)])
    """
    found = find_row(xml, r)
    if not found:
        raise KeyError("행 %d 없음" % r)
    start, end, head, inner = found
    cells = parse_cells(inner)
    by_col = {c["ci"]: c for c in cells}
    written, conflicts, replaced = [], [], []
    new_cells = {}
    for col, (value, is_sst) in writes.items():
        ci = col_idx(col)
        ex = by_col.get(ci)
        if ex and ex["has_value"]:
            if overwrite_ok and overwrite_ok("%s%d" % (col, r), ex):
                replaced.append((col, ex["raw"]))
                new_cells[ci] = cell_xml(col, r, ex["s"], value, is_sst)
                written.append((col, ex["s"]))
            else:
                conflicts.append((col, ex["raw"]))
            continue
        if ex:
            s = ex["s"]
        else:
            s = None
            for c in reversed(cells):
                if c["ci"] < ci and c["s"] is not None:
                    s = c["s"]
                    break
        new_cells[ci] = cell_xml(col, r, s, value, is_sst)
        written.append((col, s))
    if not new_cells:
        return xml, written, conflicts, replaced
    # 셀 사이에 셀 아닌 텍스트가 있으면(정상 파일엔 없음) 순서를 보장할 수 없으니 실패한다.
    if CELL_RE.sub("", inner).strip():
        raise ValueError("행 %d: 셀 밖 XML 이 있어 재조립 불가" % r)
    # 열 순서로 재조립 — 기존 셀 raw 는 바이트 그대로, 빈 자리 셀만 교체.
    merged = []
    pending = sorted(new_cells)
    pi = 0
    for c in cells:
        while pi < len(pending) and pending[pi] < c["ci"]:
            merged.append(new_cells[pending[pi]])
            pi += 1
        if pi < len(pending) and pending[pi] == c["ci"]:
            merged.append(new_cells[c["ci"]])
            pi += 1
        else:
            merged.append(c["raw"])
    while pi < len(pending):
        merged.append(new_cells[pending[pi]])
        pi += 1
    if head.endswith("/>"):
        head = head[:-2] + ">"
    new_row = head + "".join(merged) + "</row>"
    return xml[:start] + new_row + xml[end:], written, conflicts, replaced


F_TAG_RE = re.compile(r"<f\b([^>]*?)/?>")


def f_attrs(body):
    """셀 내부 XML 의 첫 `<f …>` 속성 dict. 수식 없으면 None. 속성 순서·`ca="1"`(휘발성) 유무에 무관."""
    m = F_TAG_RE.search(body or "")
    if not m:
        return None
    return dict(re.findall(r'([\w:]+)="([^"]*)"', m.group(1)))


def shared_formula_groups(xml):
    """시트 안 공유수식 그룹 {si: set(셀 ref)} — 앵커의 ref 범위 전체 + 의존 셀.

    레퍼런스에는 `<f t="shared" ref=".." ca="1" si="0">`(ref 와 si 사이에 ca)·`<f t="shared" ca="1" si="0"/>` 처럼
    속성 순서가 다른 변형이 수천 개 있어(세진 482·삼성重 2,332·미포 2,839 의존 셀) 정규식으로 순서를 고정하면 놓친다.
    셀 단위로 `<f>` 속성을 파싱한다. ref 가 단일 셀("CS143")인 앵커도 처리.
    """
    groups = {}
    for m in CELL_RE.finditer(xml):
        body = m.group(4) or ""
        if "<f" not in body:
            continue
        at = f_attrs(body)
        if not at or at.get("t") != "shared" or "si" not in at:
            continue
        cell = m.group(1) + m.group(2)
        refs = groups.setdefault(at["si"], set())
        refs.add(cell)
        ref = at.get("ref")
        if ref:
            a, _, b = ref.partition(":")
            b = b or a
            c1, r1 = re.match(r"([A-Z]+)(\d+)", a).groups()
            c2, r2 = re.match(r"([A-Z]+)(\d+)", b).groups()
            for ci in range(col_idx(c1), col_idx(c2) + 1):
                for rr in range(int(r1), int(r2) + 1):
                    refs.add("%s%d" % (col_letters(ci), rr))
    return groups


def shared_formula_orphans(xml):
    """앵커(ref 있는 `t="shared"`) 없이 의존 셀만 있는 si 의 셀 목록 — 정상 시트는 []."""
    anchors, deps = set(), {}
    for m in CELL_RE.finditer(xml):
        body = m.group(4) or ""
        if "<f" not in body:
            continue
        at = f_attrs(body)
        if not at or at.get("t") != "shared" or "si" not in at:
            continue
        if at.get("ref"):
            anchors.add(at["si"])
        else:
            deps.setdefault(at["si"], []).append(m.group(1) + m.group(2))
    return sorted(c for si, cells in deps.items() if si not in anchors for c in cells)


def make_overwrite_ok(xml, planned):
    """--overwrite-placeholders 판정 함수. planned = 이번에 값을 쓸 셀 ref 집합.

    기존 값 셀은 (a) 새 기간 열 안이고 (b) 공유수식 그룹에 속하면 그 그룹의 모든 셀이 함께 바뀔 때만
    바꾼다 — 앵커만 지우면 의존 셀의 `<f t="shared" si>` 가 고아가 되어 Excel 이 복구 대화상자를 띄운다.
    """
    groups = shared_formula_groups(xml)
    cell2si = {c: si for si, cells in groups.items() for c in cells}

    def ok(ref, ex):
        si = cell2si.get(ref)
        if si is None:
            return True
        return groups[si] <= planned
    return ok


def replace_cell(xml, r, col, new_cell_xml):
    """행 r 의 기존 셀 col 을 통째로 교체(명시 허용 셀 전용). 없으면 KeyError."""
    found = find_row(xml, r)
    if not found:
        raise KeyError("행 %d 없음" % r)
    start, end, head, inner = found
    m = re.search(r'<c r="%s%d"(?:\s[^>]*)?(?:/>|>.*?</c>)' % (col, r), inner, re.S)
    if not m:
        raise KeyError("셀 %s%d 없음" % (col, r))
    new_inner = inner[:m.start()] + new_cell_xml + inner[m.end():]
    return xml[:start] + head + new_inner + "</row>" + xml[end:]


def header_columns(xml, sst, r=1):
    """행 r 라벨 → 열 문자. 문자열은 sharedStrings 로 풀고, 숫자는 정수 문자열('2023')."""
    found = find_row(xml, r)
    out = {}
    if not found:
        return out
    for c in parse_cells(found[3]):
        v = re.search(r"<v>([^<]*)</v>", c["inner"])
        if not v:
            continue
        if c["t"] == "s":
            lab = sst.get(int(v.group(1)))
        else:
            try:
                f = float(v.group(1))
                lab = str(int(f)) if f.is_integer() else v.group(1)
            except ValueError:
                lab = v.group(1)
        out.setdefault(lab, c["col"])
    return out


def last_period(xml, sst, r=4, first_col="D"):
    """행4 의 마지막 기간 라벨 → ('Q',y,n)|('A',y) 와 그 열."""
    found = find_row(xml, r)
    if not found:
        return None, None
    last, col = None, None
    for c in parse_cells(found[3]):
        if c["ci"] < col_idx(first_col):
            continue
        v = re.search(r"<v>([^<]*)</v>", c["inner"])
        if not v:
            continue
        txt = sst.get(int(v.group(1))) if c["t"] == "s" else v.group(1)
        p = parse_row4_label(txt, c["t"] == "s")
        if p:
            last, col = p, c["col"]
    return last, col


def account_rows(xml, sst, first_row=5):
    """[(row, code, name)] — 열 C(AccountNAME) 가 문자열인 행만."""
    out = []
    for m in re.finditer(r'<row r="(\d+)"(?:\s[^>]*)?>(.*?)</row>', xml, re.S):
        r = int(m.group(1))
        if r < first_row:
            continue
        cells = {c["col"]: c for c in parse_cells(m.group(2))}
        c = cells.get("C")
        if not c or c["t"] != "s":
            continue
        v = re.search(r"<v>(\d+)</v>", c["inner"])
        if not v:
            continue
        name = sst.get(int(v.group(1)))
        b = cells.get("B")
        code = None
        if b and b["t"] == "s":
            bv = re.search(r"<v>(\d+)</v>", b["inner"])
            code = sst.get(int(bv.group(1))) if bv else None
        out.append((r, code, name))
    return out


# ── fin → 값 ─────────────────────────────────────────────

def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _sum4(tbl, y, acct):
    vals = [_num((tbl.get("%dQ%d" % (y, n)) or {}).get(acct)) for n in (1, 2, 3, 4)]
    if any(v is None for v in vals):
        return None
    return round(sum(vals), 6)


def value_for(fin, scope, acct, p, prices=None, is_convention=IS_CONVENTION_DEFAULT):
    """fin json → 계정 acct 의 기간 p 값(백만원 등 FnGuide 단위). 없으면 (None, None).

    반환 (값, 출처표) — 출처표: 'bs'|'is_ytd_diff'|'is'|'cf'|'is_ytd'|'cf_ytd'|'is_sum4'|'cf_sum4'|'shares'|'dividend'|'mcap'.
    is_convention: "ytd_diff"(기본) 이면 분기 IS 는 `is_ytd_diff[q][acct]` 를 먼저 보고 없으면 `is`; "3m" 이면 항상 `is`.
    연간(YYYY.12A)은 관행과 무관하게 `is_ytd[Q4]`(없으면 `is` 4분기 합).
    """
    if is_convention not in IS_CONVENTIONS:
        raise ValueError("is_convention 은 %s 중 하나: %r" % ("|".join(IS_CONVENTIONS), is_convention))
    sect = fin.get(scope) or {}
    if p[0] == "Q":
        qk = period_key(p)
        for tbl in ("bs", "is", "cf"):
            if tbl == "is" and is_convention == "ytd_diff":
                v = _num(((sect.get("is_ytd_diff") or {}).get(qk) or {}).get(acct))
                if v is not None:
                    return v, "is_ytd_diff"
            v = _num(((sect.get(tbl) or {}).get(qk) or {}).get(acct))
            if v is not None:
                return v, tbl
    else:
        qk = "%dQ4" % p[1]
        v = _num(((sect.get("bs") or {}).get(qk) or {}).get(acct))
        if v is not None:
            return v, "bs"
        for tbl in ("is", "cf"):
            v = _num(((sect.get(tbl + "_ytd") or {}).get(qk) or {}).get(acct))
            if v is not None:
                return v, tbl + "_ytd"
            v = _sum4(sect.get(tbl) or {}, p[1], acct)
            if v is not None:
                return v, tbl + "_sum4"
    # 주식수(백만주) — 연결/별도 공통(회사 단위).
    sh = (fin.get("shares") or {}).get(qk) or {}
    if acct in SHARE_MAP:
        v = _num(sh.get(SHARE_MAP[acct]))
        return (v / 1e6, "shares") if v is not None else (None, None)
    if acct == "우선주기말발행주식수":
        a, b = _num(sh.get("issued")), _num(sh.get("common_issued"))
        return ((a - b) / 1e6, "shares") if a is not None and b is not None else (None, None)
    if acct == DPS_ACCT:
        v = _num(((fin.get("dividend") or {}).get(qk) or {}).get("dps_common"))
        return (v, "dividend") if v is not None else (None, None)
    if acct in MCAP_MAP and prices:
        hist = ((prices.get("rows") or {}).get(fin.get("stock")) or {}).get("history_quarterly") or {}
        px = _num((hist.get(qk) or {}).get(MCAP_MAP[acct]))
        n = _num(sh.get("common_issued"))
        if px is not None and n is not None:
            return round(px * n / 1e9, 3), "mcap"
    return None, None


# ── 패치 본체 ────────────────────────────────────────────

class Workbook:
    """zip 멤버를 메모리에 들고 일부만 바꾼다. 나머지는 바이트 그대로 다시 쓴다."""

    def __init__(self, path):
        self.path = path
        with zipfile.ZipFile(path) as z:
            self.infos = z.infolist()
            self.data = {i.filename: z.read(i.filename) for i in self.infos}
        self.modified = {}
        self.wb_xml = self.data["xl/workbook.xml"].decode("utf-8")
        self.rels_xml = self.data["xl/_rels/workbook.xml.rels"].decode("utf-8")
        self.sheets = sheet_paths(self.wb_xml, self.rels_xml)
        self.names = defined_names(self.wb_xml)
        self.sst = SharedStrings(self.data["xl/sharedStrings.xml"].decode("utf-8"))

    def sheet_xml(self, name):
        p = self.sheets[name]
        if p in self.modified:
            return self.modified[p]
        return self.data[p].decode("utf-8")

    def set_sheet(self, name, xml):
        self.modified[self.sheets[name]] = xml

    def set_calc_full(self):
        """<calcPr …/> 에 fullCalcOnLoad="1" 을 붙인다(이미 있으면 그대로)."""
        m = re.search(r"<calcPr\s[^>]*/>", self.wb_xml)
        if not m:
            self.wb_xml = self.wb_xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
            return True
        tag = m.group(0)
        if "fullCalcOnLoad=" in tag:
            if re.search(r'fullCalcOnLoad="(1|true)"', tag):
                return False
            new = re.sub(r'fullCalcOnLoad="[^"]*"', 'fullCalcOnLoad="1"', tag)
        else:
            new = tag[:-2] + ' fullCalcOnLoad="1"/>'
        self.wb_xml = self.wb_xml[:m.start()] + new + self.wb_xml[m.end():]
        return True

    def prune_calc_chain(self, removed_by_sheet):
        """수식 → 상수로 바뀐 셀({시트명: set(ref)})을 calcChain.xml 에서 뺀다(없으면 no-op). 보고 dict."""
        sids = sheet_ids(self.wb_xml)
        removed = {}
        for sheet, cells in removed_by_sheet.items():
            if cells and sheet in sids:
                removed.setdefault(sids[sheet], set()).update(cells)
        rep = {"present": CALC_CHAIN in self.data, "entries": 0,
               "requested": sum(len(v) for v in removed.values()), "pruned": [], "missing": []}
        if not rep["present"] or not removed:
            return rep
        xml = self.data[CALC_CHAIN].decode("utf-8")
        rep["entries"] = len(re.findall(r"<c\s", xml))
        new, pruned = prune_calc_chain(xml, removed)
        name_of = {v: k for k, v in sids.items()}
        rep["pruned"] = ["%s!%s" % (name_of.get(i, i), r) for i, r in pruned]
        done = {(i, r) for i, r in pruned}
        rep["missing"] = ["%s!%s" % (name_of.get(i, i), r) for i, cells in removed.items() for r in sorted(cells) if (i, r) not in done]
        if pruned:
            self.modified[CALC_CHAIN] = new
        return rep

    def write(self, out_path):
        assert os.path.abspath(out_path) != os.path.abspath(self.path), "원본 덮어쓰기 금지"
        assert not os.path.basename(out_path).endswith("_orig.xlsx"), "원본 덮어쓰기 금지"
        self.modified["xl/workbook.xml"] = self.wb_xml
        self.modified["xl/sharedStrings.xml"] = self.sst.render()
        tmp = out_path + ".tmp"
        with zipfile.ZipFile(tmp, "w") as z:
            for info in self.infos:
                body = self.modified.get(info.filename, self.data[info.filename])
                if isinstance(body, str):
                    body = body.encode("utf-8")
                zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
                zi.compress_type = info.compress_type
                zi.external_attr = info.external_attr
                zi.create_system = info.create_system
                z.writestr(zi, body)
        os.replace(tmp, out_path)


def patch_bs_sheet(wb, sheet, fin, scope, target_q, stamp, prices=None, overwrite=False,
                   is_convention=IS_CONVENTION_DEFAULT):
    """BS 시트 하나: 행4 라벨 + 계정 행 값 + 행3 스탬프. 보고 dict 반환.

    overwrite=True(--overwrite-placeholders): 새 기간 열에 이미 값이 있는 셀 — 애널리스트가 미발표 분기에
    넣어 둔 가이던스·잠정치(삼성重 4Q23 `CT69-SUM(CP69:CR69)`, 미포 1Q25 잠정 1,183,800 등) — 를 fin 실적으로
    바꾸고 replaced 에 원문을 남긴다. 공유수식 그룹은 전 셀이 함께 바뀔 때만(make_overwrite_ok).
    is_convention(value_for 참조): 분기 IS 를 `is_ytd_diff` 로 채운 셀 중 `is`(3개월 열) 와 RESTATED_TOL 초과로 다른
    것은 restated_cells 에 {cell, acct, period, is, ytd_diff, diff} 로 남긴다 — 실제로 쓴 셀만(충돌로 보존된 셀 제외).
    """
    xml = wb.sheet_xml(sheet)
    sst = wb.sst
    rep = {"sheet": sheet, "member": wb.sheets[sheet], "stock": fin.get("stock"), "scope": scope,
           "new_periods": [], "columns": {}, "cells_value": 0, "cells_label": 0, "cells_stamp": 0,
           "cells_blank": 0, "written": [], "conflicts": [], "replaced": [], "missing_accounts": [],
           "missing_quarters": [], "unmatched_fin_keys": [], "sources": {},
           "is_convention": is_convention, "is_ytd_diff_available": bool((fin.get(scope) or {}).get("is_ytd_diff")),
           "restated_cells": []}
    last, last_col = last_period(xml, sst)
    if last is None:
        rep["error"] = "행4 기간 라벨을 못 읽음"
        return rep
    rep["last_actual"] = period_key(last)
    rep["last_actual_col"] = last_col
    periods = period_seq(last, target_q)
    if not periods:
        rep["error"] = "추가할 기간 없음(이미 %s 까지)" % period_key(last)
        return rep
    hdr = header_columns(xml, sst, 1)
    cols = []
    prev = col_idx(last_col)
    for p in periods:
        col = hdr.get(row1_label(p))
        if col is None:
            rep["error"] = "행1 에 라벨 %r 없음" % row1_label(p)
            return rep
        if col_idx(col) != prev + 1:
            rep["error"] = "기간 열 불연속: %s → %s(%s)" % (col_letters(prev), col, row1_label(p))
            return rep
        prev = col_idx(col)
        cols.append((p, col))
    rep["new_periods"] = [period_key(p) for p, _ in cols]
    rep["columns"] = {period_key(p): c for p, c in cols}
    fin_qs = set(fin.get("quarters") or [])
    for p, _ in cols:
        if p[0] == "Q" and period_key(p) not in fin_qs:
            rep["missing_quarters"].append(period_key(p))

    # 행4 라벨
    writes = {}
    for p, col in cols:
        lab, is_str = row4_label(p)
        writes[col] = (sst.add(lab), True) if is_str else (float(lab), False)
    xml, w, conf, _ = rebuild_row(xml, 4, writes)
    rep["cells_label"] += len(w)
    rep["written"] += ["%s4" % c for c, _ in w]
    rep["conflicts"] += [{"cell": "%s4" % c, "kept": raw[:80]} for c, raw in conf]

    # 계정 행 값 — 1차: 쓸 값을 전부 모은다(덮어쓰기 판정에 전체 집합이 필요), 2차: 행별 재조립.
    seen_accts = set()
    plan = []                                   # (row, name, writes)
    restated = []                               # is_ytd_diff 로 채웠는데 is(3개월 열) 와 다른 셀 — 쓴 뒤 written 으로 거른다
    is_3m = (fin.get(scope) or {}).get("is") or {}
    for r, code, name in account_rows(xml, sst):
        if name.strip() == "":
            continue
        seen_accts.add(name)
        writes, any_val = {}, False
        for p, col in cols:
            v, src = value_for(fin, scope, name, p, prices, is_convention)
            if v is None:
                rep["cells_blank"] += 1
                continue
            writes[col] = (v, False)
            any_val = True
            rep["sources"][src] = rep["sources"].get(src, 0) + 1
            if src == "is_ytd_diff":
                v3 = _num((is_3m.get(period_key(p)) or {}).get(name))
                if v3 is not None and abs(v - v3) > RESTATED_TOL:
                    restated.append({"cell": "%s%d" % (col, r), "acct": name, "period": period_key(p),
                                     "is": v3, "ytd_diff": v, "diff": round(v - v3, 2)})
        if not any_val:
            rep["missing_accounts"].append(name)
            continue
        plan.append((r, name, writes))
    planned = {"%s%d" % (c, r) for r, _, ws in plan for c in ws}
    overwrite_ok = make_overwrite_ok(xml, planned) if overwrite else None
    for r, name, writes in plan:
        xml, w, conf, repl = rebuild_row(xml, r, writes, overwrite_ok)
        rep["cells_value"] += len(w)
        rep["written"] += ["%s%d" % (c, r) for c, _ in w]
        rep["conflicts"] += [{"cell": "%s%d" % (c, r), "acct": name, "kept": raw[:80]} for c, raw in conf]
        rep["replaced"] += [{"cell": "%s%d" % (c, r), "acct": name, "was": raw[:160],
                             "now": writes[c][0]} for c, raw in repl]
        for _, raw in repl:                      # 문자열 셀을 숫자로 바꿨으면 sst 참조 카운트도 맞춘다
            if re.search(r'\st="s"', raw):
                sst.release()
    written_set = set(rep["written"])
    rep["restated_cells"] = [c for c in restated if c["cell"] in written_set]
    # fin 에 있는데 시트에 행이 없는 계정(참고)
    sect = fin.get(scope) or {}
    fin_keys = set()
    for tbl in ("bs", "is", "cf"):
        for q, d in (sect.get(tbl) or {}).items():
            fin_keys.update(k for k, v in d.items() if v is not None)
    rep["unmatched_fin_keys"] = sorted(fin_keys - seen_accts)

    # 행3 UPDATE 스탬프(열 C) — 명시 허용 교체
    found = find_row(xml, 3)
    if found:
        c3 = {c["col"]: c for c in parse_cells(found[3])}.get("C")
        if c3 is not None:
            if c3["t"] == "s":
                sst.release()
            idx = sst.add(stamp)
            xml = replace_cell(xml, 3, "C", cell_xml("C", 3, c3["s"], idx, True))
            rep["cells_stamp"] = 1
            rep["written"].append("C3")
            rep["stamp"] = stamp
    wb.set_sheet(sheet, xml)
    return rep


FX_SCALE_CANDIDATES = (1.0, 0.01, 100.0)


def _median(xs):
    xs = sorted(xs)
    n = len(xs)
    return xs[n // 2] if n % 2 else (xs[n // 2 - 1] + xs[n // 2]) / 2


def fx_row_scale(cells, hdr, fx, key, before_quarters):
    """환율 행 하나의 단위 배율 — 과거 실적 구간 시트값 ÷ fx.json 실측 의 중앙값.

    cells: {열: 셀} (parse_cells 결과), hdr: 행1 라벨→열, before_quarters: 비교에 쓸 분기 키(교체 범위 앞).
    후보 {1, 0.01, 100} 중 중앙값 비율이 ±5% 안에 드는 것을 고른다. 레퍼런스 ` 원/100Y` 행은 라벨과 달리
    원/엔 이어서 0.01 이 나오고, ` W/Euro`·` 원/중국원` 은 2020~ 가정이 실측과 멀어도 2005~2019 실측이 중앙값을 지킨다.
    반환 {"scale": 배율|None, "median_ratio", "n"}.
    """
    ratios = []
    for qk in before_quarters:
        y, n = q_parse(qk)
        col = hdr.get(row1_label(("Q", y, n)))
        c = cells.get(col) if col else None
        if not c:
            continue
        v = re.search(r"<v>([^<]*)</v>", c["inner"])
        fv = _num(((fx.get("quarters") or {}).get(qk) or {}).get(key))
        if not v or not fv:
            continue
        try:
            sv = float(v.group(1))
        except ValueError:
            continue
        if sv:
            ratios.append(sv / fv)
    if not ratios:
        return {"scale": None, "median_ratio": None, "n": 0}
    med = _median(ratios)
    for cand in FX_SCALE_CANDIDATES:
        if abs(med / cand - 1) <= 0.05:
            return {"scale": cand, "median_ratio": round(med, 5), "n": len(ratios)}
    return {"scale": None, "median_ratio": round(med, 5), "n": len(ratios)}


def patch_fx_sheet(wb, fx, target_q, sheet="변수", actual_periods=None):
    """`변수` 환율 행: 비어 있는 분기/연간 열만 fx.json 으로 채운다(기존 가정 상수는 보존·보고).

    actual_periods(--fx-actuals): 이 기간 키들('2023Q4', '2023A' …, BS 새 기간과 동일)의 열은 값이 있어도 fx.json
    실측으로 바꾼다. fx.json 에 없거나 `partial`(2026Q3~ 진행 중)인 기간은 건드리지 않는다(이후 가정 보존).
    행 단위 배율은 fx_row_scale 로 정하고 못 정한 행은 건너뛰어 skipped_rows 에 적는다. 바꾼 셀은 replaced 에
    원문(was)·전값(was_value)·새값(now)·배율(scale) 을 남긴다. 공유수식 그룹은 make_overwrite_ok 규칙 그대로.
    """
    rep = {"sheet": sheet, "member": None, "filled": 0, "kept": 0, "rows_found": 0, "written": [], "missing_fx": 0,
           "replaced": [], "columns": {}, "scale": {}, "skipped_rows": [], "partial_skipped": []}
    if sheet not in wb.sheets:
        rep["error"] = "시트 없음"
        return rep
    rep["member"] = wb.sheets[sheet]
    xml = wb.sheet_xml(sheet)
    sst = wb.sst
    hdr = header_columns(xml, sst, 1)
    ty, tn = q_parse(target_q)
    labels = []
    y, n = 2005, 1
    while (y, n) <= (ty, tn):
        labels.append((("Q", y, n), row1_label(("Q", y, n))))
        if n == 4:
            labels.append((("A", y), str(y)))
            y, n = y + 1, 1
        else:
            n += 1

    def fx_src(p):
        return (fx.get("quarters") or {}).get(period_key(p)) if p[0] == "Q" else (fx.get("annual") or {}).get(str(p[1]))

    # 실측 교체 대상 기간 → 열. partial 이거나 fx 에 없으면 제외(가정 보존).
    actual_cols = {}
    first_actual = None
    for pk in sorted(actual_periods or []):
        p = ("A", int(pk[:-1])) if pk.endswith("A") else ("Q",) + q_parse(pk)
        src = fx_src(p)
        col = hdr.get(row1_label(p))
        if col is None or not src:
            continue
        if src.get("partial"):
            rep["partial_skipped"].append(pk)
            continue
        actual_cols[col] = pk
        if p[0] == "Q" and (first_actual is None or (p[1], p[2]) < first_actual):
            first_actual = (p[1], p[2])
    rep["columns"] = {pk: col for col, pk in actual_cols.items()}
    before = [period_key(p) for p, _ in labels if p[0] == "Q" and first_actual and (p[1], p[2]) < first_actual]

    # 라벨 행 찾기: 열 C 의 문자열
    rows = {}
    for m in re.finditer(r'<row r="(\d+)"(?:\s[^>]*)?>(.*?)</row>', xml, re.S):
        cells = {c["col"]: c for c in parse_cells(m.group(2))}
        c = cells.get("C")
        if c and c["t"] == "s":
            v = re.search(r"<v>(\d+)</v>", c["inner"])
            if v:
                rows.setdefault(sst.get(int(v.group(1))), (int(m.group(1)), cells))
    plan, allowed_all = [], set()
    for label, key in FX_ROWS:
        found = rows.get(label)
        if found is None:
            continue
        r, cells = found
        rep["rows_found"] += 1
        row_actual, scale = {}, 1.0
        if actual_cols:
            sc = fx_row_scale(cells, hdr, fx, key, before)
            rep["scale"][label.strip()] = sc
            if sc["scale"] is None:
                rep["skipped_rows"].append(dict(row=label.strip(), why="단위 배율 미확정(후보 1·0.01·100 ±5% 밖)", **sc))
            else:
                scale = sc["scale"]
                row_actual = dict(actual_cols)
        writes = {}
        for p, lab in labels:
            col = hdr.get(lab)
            if col is None:
                continue
            src = fx_src(p) or {}
            v = _num(src.get(key))
            if v is None:
                rep["missing_fx"] += 1
                continue
            if src.get("partial") and col not in row_actual:
                continue                           # 진행 중 분기는 빈칸도 채우지 않는다
            if col in row_actual:
                writes[col] = (round(v * scale, 6), False)
                allowed_all.add("%s%d" % (col, r))
            else:
                writes[col] = (v, False)
        plan.append((r, label, scale, writes))
    ok = None
    if allowed_all:
        grp_ok = make_overwrite_ok(xml, allowed_all)
        ok = lambda ref, ex: ref in allowed_all and grp_ok(ref, ex)   # noqa: E731
    for r, label, scale, writes in plan:
        xml, w, conf, repl = rebuild_row(xml, r, writes, ok)
        rep["filled"] += len(w) - len(repl)
        rep["kept"] += len(conf)
        rep["written"] += ["%s%d" % (c, r) for c, _ in w]
        for c, raw in repl:
            if re.search(r'\st="s"', raw):
                sst.release()
            wv = re.search(r"<v>([^<]*)</v>", raw)
            try:
                was_value = float(wv.group(1)) if wv else None
            except ValueError:
                was_value = wv.group(1)
            rep["replaced"].append({"cell": "%s%d" % (c, r), "row": label.strip(), "period": actual_cols.get(c),
                                    "was": raw[:160], "was_value": was_value, "now": writes[c][0], "scale": scale})
    wb.set_sheet(sheet, xml)
    return rep


def patch_price(wb, price_row, as_of):
    """이름정의 `종가` 셀에 종가, 그 왼쪽 라벨 셀에 '종가(MM월 DD일)'. 명시 허용 교체 두 곳."""
    rep = {"written": [], "overwritten": []}
    ref = wb.names.get("종가")
    if not ref:
        rep["error"] = "이름정의 종가 없음"
        return rep
    sp = split_ref(ref)
    if not sp or sp[0] not in wb.sheets:
        rep["error"] = "종가 참조 해석 실패: %s" % ref
        return rep
    sheet, col, r = sp
    close = _num(price_row.get("close"))
    if close is None:
        rep["error"] = "close 없음"
        return rep
    xml = wb.sheet_xml(sheet)
    found = find_row(xml, r)
    cells = {c["col"]: c for c in parse_cells(found[3])} if found else {}
    tgt = cells.get(col)
    if tgt is None:
        rep["error"] = "종가 셀 %s%d 없음" % (col, r)
        return rep
    if tgt["t"] == "s":
        wb.sst.release()
    xml = replace_cell(xml, r, col, cell_xml(col, r, tgt["s"], close))
    rep["written"].append("%s!%s%d" % (sheet, col, r))
    rep["overwritten"].append(tgt["raw"][:80])
    # 라벨: 왼쪽 셀이 '종가(…)' 문자열이면 날짜 갱신
    if col_idx(col) > 1:
        lc = col_letters(col_idx(col) - 1)
        lab = cells.get(lc)
        if lab and lab["t"] == "s":
            v = re.search(r"<v>(\d+)</v>", lab["inner"])
            if v and wb.sst.get(int(v.group(1))).lstrip().startswith("종가"):
                d = datetime.date.fromisoformat(as_of[:4] + "-" + as_of[4:6] + "-" + as_of[6:8]) if re.fullmatch(r"\d{8}", as_of) else datetime.date.fromisoformat(as_of)
                wb.sst.release()
                idx = wb.sst.add("종가(%02d월 %02d일)" % (d.month, d.day))
                xml = replace_cell(xml, r, lc, cell_xml(lc, r, lab["s"], idx, True))
                rep["written"].append("%s!%s%d" % (sheet, lc, r))
                rep["overwritten"].append(lab["raw"][:80])
    wb.set_sheet(sheet, xml)
    rep["close"] = close
    rep["as_of"] = as_of
    return rep


# ── 수식 연장(--extend-formulas) ──────────────────────────

UNESCAPE_MAP = {"&quot;": '"', "&apos;": "'"}
EXTEND_MODES = ("plain", "all")
EXTEND_CLI_MODES = EXTEND_MODES + ("off",)
EXTEND_CLI_DEFAULT = "all"          # 2026-10-08 오너 결정 — 운영 실행(CLI)은 수식 연장 켬. patch_file(extend=None) 의 API 기본은 꺼짐 그대로
EXTEND_SHEET = "subQ"


def f_text(body):
    """셀 내부 XML 의 `<f …>텍스트</f>` → 수식 문자열(XML 이스케이프 해제). 자기닫힘(`<f t="shared" si="3"/>`)·없음이면 None."""
    m = re.search(r"<f\b[^>]*>(.*?)</f>", body or "", re.S)
    return unescape(m.group(1), UNESCAPE_MAP) if m else None


def translate_formula(text, origin, dest):
    """수식 text(origin 셀 기준) 를 dest 셀로 옮길 때의 상대참조 치환 — Excel 채우기(끌기)와 같다.

    openpyxl 의 Translator(토크나이저 기반) 를 쓴다: A1 참조·범위·시트 한정 참조만 열/행 차이만큼 옮기고 `$` 고정 참조·
    이름정의(BS연결·SUBQH·U…)·문자열·함수명(LOG10)은 건드리지 않는다. 치환 실패(열 < A 등)는 ValueError.
    """
    from openpyxl.formula.translate import Translator
    try:
        out = Translator("=" + text, origin=origin).translate_formula(dest)
    except Exception as e:                                # TranslatorError·TokenizerError 등
        raise ValueError("수식 치환 실패 %s→%s (%s): %s" % (origin, dest, e.__class__.__name__, e))
    return out[1:] if out.startswith("=") else out


def source_period(p, last):
    """연장 소스 기간 — 목표 p 와 같은 분기 위치(1Q~4Q·연간)의 가장 최근 실적 기간(BS 마지막 실적 last 이하).

    같은 위치에서 가져와야 1Q 열(직전 열이 연간)·연간 열(SUM 범위) 의 상대참조 구조가 유지된다(열 차이는 항상 5의 배수).
    last=('Q',2023,3): 4Q→2022Q4 · 연간→2022 · 1Q~3Q→2023 / last=('A',2024): 연간→2024 · nQ→2024Qn.
    """
    if last[0] == "A":
        ly, ln = last[1], 4
    else:
        ly, ln = last[1], last[2]
    if p[0] == "A":
        return ("A", ly if last[0] == "A" else ly - 1)
    return ("Q", ly if p[2] <= ln else ly - 1, p[2])


def shared_anchors(xml):
    """{si: (앵커 셀 ref, 수식 텍스트)} — 공유수식 그룹의 앵커(ref 있는 `t="shared"`)."""
    out = {}
    for m in CELL_RE.finditer(xml):
        body = m.group(4) or ""
        if "<f" not in body:
            continue
        at = f_attrs(body)
        if at and at.get("t") == "shared" and at.get("ref") and "si" in at:
            out[at["si"]] = (m.group(1) + m.group(2), f_text(body) or "")
    return out


def extend_formulas(wb, last, periods, mode="all", sheet=EXTEND_SHEET, header_row=1):
    """`sheet`(subQ) 의 새 기간 열에서 **비어 있는** 셀을 같은 행 소스 셀의 수식으로 채운다 — 사용자가 끌어 채우던 것.

    소스 = 목표와 같은 분기 위치의 가장 최근 실적 열(source_period). 소스가 단순 `<f>` 면 그 텍스트를, 공유수식(앵커/의존)
    이면 앵커 텍스트를 앵커 기준으로(mode="all" 일 때만; "plain" 이면 건너뜀) 목표 셀로 치환해 `<f>` 로 쓴다(캐시값 없음).
    상수·배열수식·빈 소스는 건너뛴다. 값이 있는 목표 셀은 바꾸지 않는다(kept_nonempty). 공유수식 그룹은 손대지 않는다.
    periods: BS 새 기간(('Q',y,n)|('A',y)) 목록, last: BS 마지막 실적 기간.
    """
    if mode not in EXTEND_MODES:
        raise ValueError("extend mode 는 %s 중 하나: %r" % ("|".join(EXTEND_MODES), mode))
    rep = {"sheet": sheet, "member": None, "mode": mode, "header_row": header_row, "last_actual": period_key(last),
           "columns": {}, "sources": {}, "extended": 0, "extended_plain": 0, "extended_shared": 0, "rows": 0,
           "kept_nonempty": 0, "skipped_const": 0, "skipped_array": 0, "skipped_shared": 0, "no_source": 0,
           "errors": [], "written": [], "cells": []}
    if sheet not in wb.sheets:
        rep["error"] = "시트 없음"
        return rep
    rep["member"] = wb.sheets[sheet]
    xml = wb.sheet_xml(sheet)
    hdr = header_columns(xml, wb.sst, header_row)
    pairs = []
    for p in periods:
        sp = source_period(p, last)
        tcol, scol = hdr.get(row1_label(p)), hdr.get(row1_label(sp))
        if tcol is None or scol is None:
            rep["errors"].append("행%d 라벨 없음: 목표 %r→%s · 소스 %r→%s" % (header_row, row1_label(p), tcol, row1_label(sp), scol))
            continue
        pairs.append((p, tcol, scol))
        rep["columns"][period_key(p)] = tcol
        rep["sources"][period_key(p)] = "%s@%s" % (period_key(sp), scol)
    if not pairs:
        rep["error"] = "연장할 기간 열 없음"
        return rep
    anchors = shared_anchors(xml) if mode == "all" else {}
    plan = []
    for m in re.finditer(r'<row r="(\d+)"(?:\s[^>]*)?>(.*?)</row>', xml, re.S):
        r = int(m.group(1))
        if r <= header_row:
            continue
        cells = {c["col"]: c for c in parse_cells(m.group(2))}
        writes = {}
        for p, tcol, scol in pairs:
            t = cells.get(tcol)
            if t and t["has_value"]:
                rep["kept_nonempty"] += 1
                continue
            src = cells.get(scol)
            if not src or not src["has_value"]:
                rep["no_source"] += 1
                continue
            if "<f" not in src["inner"]:
                rep["skipped_const"] += 1
                continue
            at = f_attrs(src["inner"]) or {}
            if at.get("t") in ("array", "dataTable"):
                rep["skipped_array"] += 1
                continue
            if at.get("t") == "shared":
                if mode != "all" or at.get("si") not in anchors:
                    rep["skipped_shared"] += 1
                    continue
                origin, text, kind = anchors[at["si"]][0], anchors[at["si"]][1], "shared"
            else:
                origin, text, kind = "%s%d" % (scol, r), f_text(src["inner"]), "plain"
            if not text:
                rep["skipped_const"] += 1
                continue
            dest = "%s%d" % (tcol, r)
            try:
                new = translate_formula(text, origin, dest)
            except ValueError as e:
                rep["errors"].append(str(e))
                continue
            writes[tcol] = (Formula(new), False)
            rep["cells"].append({"cell": dest, "period": period_key(p), "source": "%s%d" % (scol, r), "kind": kind, "formula": new})
            rep["extended_plain" if kind == "plain" else "extended_shared"] += 1
        if writes:
            plan.append((r, writes))
    for r, writes in plan:
        xml, w, conf, _ = rebuild_row(xml, r, writes)
        rep["extended"] += len(w)
        rep["written"] += ["%s%d" % (c, r) for c, _ in w]
        rep["rows"] += 1
    wb.set_sheet(sheet, xml)
    return rep


# ── calcChain ────────────────────────────────────────────

def sheet_ids(wb_xml):
    """시트명 → sheetId(calcChain 항목의 `i` 가 가리키는 값; 시트 순번이 아니다 — 세진 BS연결 = 20)."""
    out = {}
    for m in re.finditer(r"<sheet\s([^>]*)/>", wb_xml):
        nm = re.search(r'name="([^"]*)"', m.group(1))
        sid = re.search(r'sheetId="(\d+)"', m.group(1))
        if nm and sid:
            out[unescape(nm.group(1))] = sid.group(1)
    return out


def calc_chain_entries(cc_xml):
    """[(sheetId, ref, attrs)] — `i` 생략 항목은 직전 항목의 sheetId 를 따른다(ECMA-376 18.6.2)."""
    out, cur = [], None
    for m in re.finditer(r"<c\s+([^>]*?)/>", cc_xml):
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        cur = attrs.get("i", cur)
        out.append((cur, attrs.get("r"), attrs))
    return out


def prune_calc_chain(cc_xml, removed):
    """calcChain.xml 에서 수식이 사라진 셀의 항목을 뺀다. removed = {sheetId: set(ref)}. 반환 (새 xml, [(sheetId, ref)]).

    항목 `<c r="CS69" i="13" l="1"/>` 의 `i` 는 생략되면 직전 항목을 따르고 `l`(새 의존 수준 시작)은 수준 첫 셀의 표식이라,
    지운 항목이 명시한 `i`·`l` 은 바로 다음 항목이 생략했을 때 넘겨준다. 나머지 바이트는 그대로.
    """
    out, pruned, pos, cur, carry = [], [], 0, None, {}
    for m in re.finditer(r"<c\s+([^>]*?)/>", cc_xml):
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        cur = attrs.get("i", cur)
        out.append(cc_xml[pos:m.start()])
        pos = m.end()
        if cur is not None and attrs.get("r") in removed.get(cur, ()):
            pruned.append((cur, attrs.get("r")))
            for k in ("i", "l"):
                if k in attrs:
                    carry[k] = attrs[k]
            continue
        if carry:
            for k, v in carry.items():
                attrs.setdefault(k, v)
            carry = {}
            order = ["r", "i", "l", "s", "a", "t"]
            keys = [k for k in order if k in attrs] + [k for k in attrs if k not in order]
            out.append("<c " + " ".join('%s="%s"' % (k, attrs[k]) for k in keys) + "/>")
            continue
        out.append(m.group(0))
    out.append(cc_xml[pos:])
    return "".join(out), pruned


def load_json(path):
    if not path or not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def patch_file(stock, fins, target_q, today, out_path=None, fx=None, prices=None, ref_dir=None,
               overwrite=False, fx_actuals=False, is_convention=IS_CONVENTION_DEFAULT, extend=None):
    """레퍼런스 한 파일 패치. fins: {종목: fin json}. 패치한 시트가 없으면 출력하지 않는다.

    overwrite=--overwrite-placeholders, fx_actuals=--fx-actuals(`변수` 환율 실측 교체 범위 = BS 새 기간의 합집합),
    is_convention=--is-convention(분기 IS 값 관행, value_for 참조), extend=--extend-formulas(None|"plain"|"all" —
    subQ 빈 새 기간 셀 수식 연장, extend_formulas 참조). 수식 자리표시자를 상수로 바꾼 셀은 calcChain.xml 에서도 뺀다.
    """
    ref_dir = ref_dir or REF_DIR
    spec = FILES[stock]
    orig = os.path.join(ref_dir, "%s_%s_subQ_orig.xlsx" % (spec["name"], stock))
    out_path = out_path or os.path.join(ref_dir, "%s_%s_subQ_%s.xlsx" % (spec["name"], stock, target_q))
    rep = {"stock": stock, "name": spec["name"], "orig": orig, "out": out_path, "target": target_q,
           "is_convention": is_convention, "extend_mode": extend,
           "sheets": [], "skipped_sheets": [], "fx": None, "price": None, "extend": None, "calc_chain": None,
           "calcPr_changed": False, "sst_added": [], "written": False}
    if not os.path.exists(orig):
        rep["error"] = "원본 없음: %s" % orig
        return rep
    wb = Workbook(orig)
    stamp = "UPDATE: " + datetime.date.fromisoformat(today).strftime("%y-%m-%d")
    for sheet, (st, scope) in spec["sheets"].items():
        if sheet not in wb.sheets:
            rep["skipped_sheets"].append({"sheet": sheet, "why": "시트 없음"})
            continue
        fin = fins.get(st)
        if not fin:
            rep["skipped_sheets"].append({"sheet": sheet, "why": "fin 없음(%s)" % st})
            continue
        if not (fin.get(scope) or {}):
            rep["skipped_sheets"].append({"sheet": sheet, "why": "fin.%s 비어 있음(%s)" % (scope, st)})
            continue
        rep["sheets"].append(patch_bs_sheet(wb, sheet, fin, scope, target_q, stamp, prices, overwrite=overwrite,
                                            is_convention=is_convention))
    if not [s for s in rep["sheets"] if not s.get("error")]:
        rep["error"] = "패치한 시트 없음 — 출력하지 않음"
        return rep
    if fx:
        actual = None
        if fx_actuals:
            actual = set()
            for s in rep["sheets"]:
                if not s.get("error"):
                    actual.update(s["new_periods"])
        rep["fx"] = patch_fx_sheet(wb, fx, target_q, actual_periods=actual)
    if prices:
        row = (prices.get("rows") or {}).get(stock)
        if row:
            rep["price"] = patch_price(wb, row, row.get("as_of") or prices.get("as_of") or today)
        else:
            rep["price"] = {"error": "prices.rows 에 %s 없음" % stock}
    if extend:
        first = next(s for s in rep["sheets"] if not s.get("error"))        # BS연결 — subQ 가 끌어오는 표
        last = ("A", int(first["last_actual"][:-1])) if first["last_actual"].endswith("A") else ("Q",) + q_parse(first["last_actual"])
        periods = [("A", int(pk[:-1])) if pk.endswith("A") else ("Q",) + q_parse(pk) for pk in first["new_periods"]]
        rep["extend"] = extend_formulas(wb, last, periods, mode=extend)
    # 수식 자리표시자 → 상수 로 바뀐 셀은 calcChain 항목도 뺀다(남으면 Excel 복구 대화상자)
    removed = {}
    for s_ in rep["sheets"]:
        cells = {c["cell"] for c in s_.get("replaced", []) if "<f" in c.get("was", "")}
        if cells:
            removed.setdefault(s_["sheet"], set()).update(cells)
    if rep.get("fx"):
        cells = {c["cell"] for c in rep["fx"].get("replaced", []) if "<f" in c.get("was", "")}
        if cells:
            removed.setdefault(rep["fx"]["sheet"], set()).update(cells)
    rep["calc_chain"] = wb.prune_calc_chain(removed)
    rep["calcPr_changed"] = wb.set_calc_full()
    rep["sst_added"] = list(wb.sst.added)
    wb.write(out_path)
    rep["written"] = True
    return rep


# ── 검증 ────────────────────────────────────────────────

def load_openpyxl():
    """도형(find_images) 때문에 깨지는 원본을 읽기 위한 몽키패치. 검증 전용."""
    import warnings
    import openpyxl
    import openpyxl.reader.excel as rx
    rx.find_images = lambda *a, **k: ([], [])
    warnings.filterwarnings("ignore", category=UserWarning)
    return openpyxl


def verify_zip(out_path):
    with zipfile.ZipFile(out_path) as z:
        bad = z.testzip()
    return {"ok": bad is None, "bad_member": bad}


def _formula_counts(openpyxl, path):
    wb = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    out = {}
    for ws in wb.worksheets:
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    n += 1
        out[ws.title] = n
    return wb.sheetnames, out


def verify_reopen(orig, out_path, rep=None):
    """② 몽키패치 openpyxl 재오픈 → 시트 목록·시트별 수식 셀 수 동일.

    --overwrite-placeholders 로 수식 자리표시자를 값으로 바꾼 시트는 그 개수만큼 수식이 줄어드는 것이
    정상이다 — rep.sheets[].replaced 의 `was` 에 `<f` 가 든 셀 수를 시트별 허용 감소분으로 뺀다.
    --extend-formulas 로 subQ 에 쓴 수식 수(rep.extend.extended)만큼은 그 시트에서 늘어나야 한다(정확히 그만큼).
    """
    op = load_openpyxl()
    names_o, f_o = _formula_counts(op, orig)
    names_p, f_p = _formula_counts(op, out_path)
    allowed_drop = {}
    for s in (rep or {}).get("sheets", []) + [x for x in [(rep or {}).get("fx")] if x]:
        n = sum(1 for c in s.get("replaced", []) if "<f" in c.get("was", ""))
        if n:
            allowed_drop[s["sheet"]] = allowed_drop.get(s["sheet"], 0) + n
    allowed_add = {}
    ext = (rep or {}).get("extend")
    if ext and not ext.get("error") and ext.get("extended"):
        allowed_add[ext["sheet"]] = ext["extended"]
    expected = {k: v - allowed_drop.get(k, 0) + allowed_add.get(k, 0) for k, v in f_o.items()}
    return {"ok": names_o == names_p and expected == f_p, "sheets_orig": len(names_o), "sheets_patched": len(names_p),
            "formula_cells_orig": sum(f_o.values()), "formula_cells_patched": sum(f_p.values()),
            "formula_placeholders_replaced": sum(allowed_drop.values()), "formula_cells_extended": sum(allowed_add.values()),
            "diff": {k: (expected.get(k), f_p.get(k)) for k in set(expected) | set(f_p) if expected.get(k) != f_p.get(k)}}


def verify_unchanged(orig, out_path, rep):
    """④ 원본 셀 전부 동일 — 허용 예외(행3 C3 스탬프·종가 두 셀·비어 있던 자리 셀)만 다를 수 있다."""
    zo, zp = zipfile.ZipFile(orig), zipfile.ZipFile(out_path)
    res = {"ok": True, "members_orig": len(zo.namelist()), "members_patched": len(zp.namelist()),
           "identical_members": 0, "changed_members": [], "changed_existing_cells": [], "added_cells": 0,
           "replaced_empty_cells": 0, "problems": []}
    if zo.namelist() != zp.namelist():
        res["ok"] = False
        res["problems"].append("멤버 목록 다름")
    allowed_replace = {}                   # member → set(ref)
    for s in rep.get("sheets", []):
        allowed_replace.setdefault(s["member"], set()).add("C3")
        # --overwrite-placeholders 로 바꾼 자리표시자 셀(원문은 rep.sheets[].replaced 에 있다)
        for c in s.get("replaced", []):
            allowed_replace[s["member"]].add(c["cell"])
    fxr = rep.get("fx") or {}
    if fxr.get("member"):
        for c in fxr.get("replaced", []):   # --fx-actuals 로 바꾼 `변수` 환율 셀
            allowed_replace.setdefault(fxr["member"], set()).add(c["cell"])
    # 보고의 replaced 목록을 그대로 믿지 않는다 — 교체 셀은 반드시 그 시트의 새 기간(실측 교체) 열 안이어야 한다.
    for s in rep.get("sheets", []) + ([fxr] if fxr.get("member") else []):
        cols = set((s.get("columns") or {}).values())
        for c in s.get("replaced", []):
            if re.match(r"[A-Z]+", c["cell"]).group(0) not in cols:
                res["ok"] = False
                res["problems"].append("%s: 교체 셀 %s 가 새 기간 열 밖" % (s["sheet"], c["cell"]))
    pr = rep.get("price") or {}
    wb_xml = zo.read("xl/workbook.xml").decode("utf-8")
    sheets = sheet_paths(wb_xml, zo.read("xl/_rels/workbook.xml.rels").decode("utf-8"))
    for w in pr.get("written", []):
        sh, ref = w.rsplit("!", 1)
        allowed_replace.setdefault(sheets.get(sh), set()).add(ref)
    for name in zo.namelist():
        a, b = zo.read(name), zp.read(name)
        if a == b:
            res["identical_members"] += 1
            continue
        res["changed_members"].append(name)
        if name == "xl/workbook.xml":
            ta = re.sub(r"<calcPr[^>]*/>", "", a.decode("utf-8"))
            tb = re.sub(r"<calcPr[^>]*/>", "", b.decode("utf-8"))
            if ta != tb:
                res["ok"] = False
                res["problems"].append("workbook.xml 이 calcPr 외에도 다름")
            continue
        if name == "xl/sharedStrings.xml":
            sa, sb = SI_RE.findall(a.decode("utf-8")), SI_RE.findall(b.decode("utf-8"))
            if sb[:len(sa)] != sa:
                res["ok"] = False
                res["problems"].append("sharedStrings 기존 항목이 바뀜")
            continue
        if name == CALC_CHAIN:
            # 수식 자리표시자 → 상수 셀의 항목만 빠질 수 있다(⑥ 이 고아 항목 0 을 따로 본다). 추가·그 밖의 제거는 문제.
            ea = {(i, r) for i, r, _ in calc_chain_entries(a.decode("utf-8"))}
            eb = {(i, r) for i, r, _ in calc_chain_entries(b.decode("utf-8"))}
            sids = sheet_ids(wb_xml)
            allowed = set()
            for s in rep.get("sheets", []) + ([fxr] if fxr.get("member") else []):
                for c in s.get("replaced", []):
                    if "<f" in c.get("was", "") and s["sheet"] in sids:
                        allowed.add((sids[s["sheet"]], c["cell"]))
            if eb - ea:
                res["ok"] = False
                res["problems"].append("calcChain 항목 추가 %s" % sorted(eb - ea)[:5])
            if (ea - eb) - allowed:
                res["ok"] = False
                res["problems"].append("calcChain 허용 밖 항목 제거 %s" % sorted((ea - eb) - allowed)[:5])
            res["calc_chain_pruned"] = len(ea - eb)
            continue
        if not name.startswith("xl/worksheets/sheet"):
            res["ok"] = False
            res["problems"].append("예상 밖 멤버 변경: %s" % name)
            continue
        ca = {m.group(1) + m.group(2): m.group(0) for m in CELL_RE.finditer(a.decode("utf-8"))}
        cb = {m.group(1) + m.group(2): m.group(0) for m in CELL_RE.finditer(b.decode("utf-8"))}
        for ref, raw in ca.items():
            if ref not in cb:
                res["ok"] = False
                res["problems"].append("%s: 셀 %s 사라짐" % (name, ref))
                continue
            if cb[ref] == raw:
                continue
            had_value = bool(re.search(r"<(v|f|is)\b", raw))
            if not had_value:
                res["replaced_empty_cells"] += 1
                continue
            if ref in allowed_replace.get(name, set()):
                res["changed_existing_cells"].append("%s!%s" % (name, ref))
                continue
            res["ok"] = False
            res["problems"].append("%s: 기존 값 셀 %s 변경 %r → %r" % (name, ref, raw[:60], cb[ref][:60]))
        res["added_cells"] += len(set(cb) - set(ca))
        # 셀 외 텍스트(행 머리·sheetData 밖)도 동일해야 한다.
        na = CELL_RE.sub("", a.decode("utf-8"))
        nb = CELL_RE.sub("", b.decode("utf-8"))
        if na != nb:
            res["ok"] = False
            res["problems"].append("%s: 셀 외 XML 변경" % name)
        # 공유수식 고아: 앵커 없이 `<f t="shared" si>` 의존 셀만 남으면 Excel 이 복구 대화상자를 띄운다.
        orphans = shared_formula_orphans(b.decode("utf-8"))
        if orphans and set(orphans) - set(shared_formula_orphans(a.decode("utf-8"))):
            res["ok"] = False
            res["problems"].append("%s: 공유수식 고아 %d셀 %s" % (name, len(orphans), orphans[:6]))
    return res


def verify_chain(orig, out_path, fins, rep, acct="매출액(수익)"):
    """⑤ subQ 확정 행 `VLOOKUP($D, <BS표>, MATCH(<헤더>, <BS표>H, 0), 0)[/U]` 를 흉내내어 새 분기 값 == fin.

    - <BS표>·<BS표>H·헤더(SUBQH/YQ)·U 는 workbook.xml 이름정의로 푼다.
    - MATCH 는 subQ 행1(캐시값) 라벨 ↔ BS 시트 행1 라벨, VLOOKUP 은 BS 시트 열 C == $D 값.
    """
    op = load_openpyxl()
    wbf = op.load_workbook(orig, data_only=False, keep_links=False)       # 수식(원본 = 패치본, subQ 불변)
    wbv = op.load_workbook(out_path, data_only=True, keep_links=False)    # 캐시값 + 패치된 상수
    with zipfile.ZipFile(orig) as z:
        names = defined_names(z.read("xl/workbook.xml").decode("utf-8"))
    res = {"ok": True, "checked": 0, "rows": [], "problems": []}
    if "subQ" not in wbf.sheetnames:
        res["ok"] = False
        res["problems"].append("subQ 시트 없음")
        return res
    ws = wbf["subQ"]
    wsv = wbv["subQ"]
    pat = re.compile(r"VLOOKUP\(\$D(\d+),(\w+),MATCH\((\w+),(\w+),0\),0\)(/(\w+))?")
    # 표 이름 → 시트명
    def sheet_of(tbl):
        ref = names.get(tbl)
        sp = split_ref(ref) if ref else None
        return sp[0] if sp else None
    targets = {s["sheet"]: s for s in rep.get("sheets", []) if not s.get("error")}
    # 헤더 이름(SUBQH/YQ) 은 subQ!$1:$1 → 행 번호
    def header_row(nm):
        ref = names.get(nm, "")
        m = re.search(r"!\$(\d+):\$(\d+)$", ref)
        return int(m.group(1)) if m else 1
    seen = set()
    for r in range(1, min(ws.max_row, 1400) + 1):
        if ws.cell(r, 4).value != acct:
            continue
        # 그 행에서 VLOOKUP 수식이 있는 마지막 열의 수식 패턴을 쓴다.
        formula = None
        for c in range(ws.max_column, 4, -1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "VLOOKUP" in v:
                m = pat.search(v.replace(" ", ""))
                if m and int(m.group(1)) == r:
                    formula = m
                    break
        if not formula:
            continue
        tbl, hdr_nm, hdr_tbl, div_nm = formula.group(2), formula.group(3), formula.group(4), formula.group(6)
        bs_sheet = sheet_of(tbl)
        if bs_sheet not in targets or (bs_sheet, tbl) in seen:
            continue
        seen.add((bs_sheet, tbl))
        srep = targets[bs_sheet]
        div = 1.0
        if div_nm:
            dv = names.get(div_nm)
            try:
                div = float(dv)
            except (TypeError, ValueError):
                res["problems"].append("이름정의 %s 를 숫자로 못 풂: %r" % (div_nm, dv))
                res["ok"] = False
                continue
        hrow = header_row(hdr_nm)
        sub_labels = {}
        for c in range(4, wsv.max_column + 1):
            lv = wsv.cell(hrow, c).value
            if lv is not None:
                sub_labels.setdefault(str(int(lv)) if isinstance(lv, (int, float)) and float(lv).is_integer() else str(lv), c)
        bsv = wbv[bs_sheet]
        bs_labels = {}
        for c in range(3, bsv.max_column + 1):
            lv = bsv.cell(1, c).value
            if lv is not None:
                bs_labels.setdefault(str(int(lv)) if isinstance(lv, (int, float)) and float(lv).is_integer() else str(lv), c)
        acct_row = None
        for rr in range(1, bsv.max_row + 1):
            if bsv.cell(rr, 3).value == acct:
                acct_row = rr
                break
        if acct_row is None:
            res["ok"] = False
            res["problems"].append("%s 열C 에 %s 없음" % (bs_sheet, acct))
            continue
        fin = fins.get(srep["stock"]) or {}
        row_res = {"subQ_row": r, "table": tbl, "bs_sheet": bs_sheet, "divisor": div, "acct_row": acct_row, "checks": []}
        conflict_cells = {c["cell"] for c in srep.get("conflicts", [])}
        for pk in srep["new_periods"]:
            p = ("A", int(pk[:-1])) if pk.endswith("A") else ("Q",) + q_parse(pk)
            expect, src = value_for(fin, srep["scope"], acct, p,
                                    is_convention=srep.get("is_convention", IS_CONVENTION_DEFAULT))
            col_letter = srep["columns"][pk]
            if expect is None or ("%s%d" % (col_letter, acct_row)) in conflict_cells:
                continue
            lab = row1_label(p)
            if lab not in sub_labels:
                res["ok"] = False
                res["problems"].append("subQ 행%d 에 라벨 %r 없음" % (hrow, lab))
                continue
            bc = bs_labels.get(lab)
            if bc is None:
                res["ok"] = False
                res["problems"].append("%s 행1 에 라벨 %r 없음" % (bs_sheet, lab))
                continue
            got = bsv.cell(acct_row, bc).value
            ok = got is not None and abs(float(got) / div - expect / div) < 1e-6
            row_res["checks"].append({"period": pk, "label": lab, "bs_col": col_letters(bc), "subQ_col": col_letters(sub_labels[lab]),
                                      "value_bs": got, "expected_fin": expect, "src": src, "ok": ok})
            res["checked"] += 1
            if not ok:
                res["ok"] = False
                res["problems"].append("%s!%s%d=%r ≠ fin %r (%s)" % (bs_sheet, col_letters(bc), acct_row, got, expect, pk))
        res["rows"].append(row_res)
    if res["checked"] == 0:
        res["ok"] = False
        res["problems"].append("검증한 셀 0 — 수식 행 또는 fin 값 없음")
    return res


def _formula_cells_by_member(z, member):
    """zip 멤버(시트 XML)에서 `<f` 가 있는 셀 ref 집합."""
    return {m.group(1) + m.group(2) for m in CELL_RE.finditer(z.read(member).decode("utf-8")) if "<f" in (m.group(4) or "")}


def verify_calc_chain(orig, out_path, rep=None):
    """⑥ calcChain.xml 의 모든 항목이 패치본에서 실제 수식 셀을 가리킨다(고아 항목 0). calcChain 이 없는 파일은 ok.

    원본 대비 줄어든 항목 수 == 교체한 수식 자리표시자 수(rep.sheets/fx 의 replaced 중 `<f`). 보고의 pruned 를 믿지 않고
    패치본 시트 XML 을 직접 읽어 판정한다 — --overwrite-placeholders 가 자리표시자를 상수로 바꾼 뒤 항목이 남아 있으면 Excel 이
    "복구된 레코드: /xl/calcChain.xml" 을 띄우므로 이 검사는 그 결함을 바로 잡는다.
    """
    zp = zipfile.ZipFile(out_path)
    res = {"ok": True, "present": CALC_CHAIN in zp.namelist(), "entries_orig": 0, "entries_patched": 0,
           "expected_drop": 0, "stale": [], "problems": []}
    if not res["present"]:
        return res
    wb_xml = zp.read("xl/workbook.xml").decode("utf-8")
    sheets = sheet_paths(wb_xml, zp.read("xl/_rels/workbook.xml.rels").decode("utf-8"))
    member_by_sid = {sid: sheets[name] for name, sid in sheet_ids(wb_xml).items() if name in sheets}
    fcells = {}
    for sid, ref, _ in calc_chain_entries(zp.read(CALC_CHAIN).decode("utf-8")):
        res["entries_patched"] += 1
        if sid not in fcells:
            fcells[sid] = _formula_cells_by_member(zp, member_by_sid[sid]) if sid in member_by_sid else set()
        if ref not in fcells[sid]:
            res["stale"].append("%s!%s" % (sid, ref))
    zo = zipfile.ZipFile(orig)
    if CALC_CHAIN in zo.namelist():
        res["entries_orig"] = len(calc_chain_entries(zo.read(CALC_CHAIN).decode("utf-8")))
    for s in (rep or {}).get("sheets", []) + [x for x in [(rep or {}).get("fx")] if x]:
        res["expected_drop"] += sum(1 for c in s.get("replaced", []) if "<f" in c.get("was", ""))
    if res["stale"]:
        res["ok"] = False
        res["problems"].append("수식 없는 셀을 가리키는 calcChain 항목 %d: %s" % (len(res["stale"]), res["stale"][:6]))
    if res["entries_orig"] - res["entries_patched"] != res["expected_drop"]:
        res["ok"] = False
        res["problems"].append("calcChain 항목 감소 %d ≠ 교체한 수식 자리표시자 %d"
                               % (res["entries_orig"] - res["entries_patched"], res["expected_drop"]))
    return res


VLOOKUP_RE = re.compile(r"VLOOKUP\(\$D(\d+),(\w+),MATCH\((\w+),(\w+),0\),0\)(?:/(\w+))?")


def _divisor(names, div):
    """VLOOKUP 뒤 `/U`·`/10` 의 나눗수 — 이름정의(U=100) 또는 숫자 리터럴. 못 풀면 None."""
    if not div:
        return 1.0
    try:
        return float(div)
    except ValueError:
        pass
    try:
        return float(names.get(div))
    except (TypeError, ValueError):
        return None


def verify_extend(orig, out_path, fins, rep, prices=None):
    """⑦ --extend-formulas 로 쓴 셀: (a) 재오픈 시 전부 수식이고 텍스트가 보고와 같다 (b) 순수 `VLOOKUP($D,<BS표>,MATCH(<헤더>,
    <BS표>H,0),0)[/U]` 모양은 사슬을 흉내내어 값(패치된 BS 셀 ÷ U) == fin ÷ U(시가총액 행은 prices 로 — 패치와 같은 입력).
    그 밖의 모양(REF-REF 등)은 not_checkable. BS 셀이 비어 있고 fin 도 없는 계정(unresolved) 은 세지 않는다(Excel 은 0)."""
    ext = rep.get("extend") or {}
    res = {"ok": True, "cells": ext.get("extended", 0), "formula_ok": 0, "formula_bad": [], "checked": 0, "matched": 0,
           "mismatch": [], "unresolved": 0, "not_checkable": 0, "problems": []}
    if not ext or ext.get("error") or not ext.get("cells"):
        res["ok"] = not ext or not ext.get("error")
        if ext.get("error"):
            res["problems"].append(ext["error"])
        return res
    op = load_openpyxl()
    wbf = op.load_workbook(out_path, data_only=False, keep_links=False)
    wbv = op.load_workbook(out_path, data_only=True, keep_links=False)
    with zipfile.ZipFile(orig) as z:
        names = defined_names(z.read("xl/workbook.xml").decode("utf-8"))
    ws, wsv = wbf[ext["sheet"]], wbv[ext["sheet"]]
    hrow = ext.get("header_row", 1)
    sub_labels = {}
    for c in range(1, wsv.max_column + 1):
        lv = wsv.cell(hrow, c).value
        if lv is not None:
            sub_labels[c] = str(int(lv)) if isinstance(lv, (int, float)) and float(lv).is_integer() else str(lv)
    targets = {s["sheet"]: s for s in rep.get("sheets", []) if not s.get("error")}
    bs_cache = {}

    def bs_info(sheet):
        if sheet not in bs_cache:
            bsv = wbv[sheet]
            labels, rows = {}, {}
            for c in range(3, bsv.max_column + 1):
                lv = bsv.cell(1, c).value
                if lv is not None:
                    labels.setdefault(str(int(lv)) if isinstance(lv, (int, float)) and float(lv).is_integer() else str(lv), c)
            for r in range(1, bsv.max_row + 1):
                nm = bsv.cell(r, 3).value
                if isinstance(nm, str):
                    rows.setdefault(nm, r)
            bs_cache[sheet] = (bsv, labels, rows)
        return bs_cache[sheet]
    for c in ext["cells"]:
        col, r = re.match(r"([A-Z]+)(\d+)", c["cell"]).groups()
        r = int(r)
        got_f = ws.cell(r, col_idx(col)).value
        if got_f != "=" + c["formula"]:
            res["formula_bad"].append({"cell": c["cell"], "expected": c["formula"], "got": got_f})
            continue
        res["formula_ok"] += 1
        m = VLOOKUP_RE.fullmatch(c["formula"].replace(" ", ""))
        if not m or int(m.group(1)) != r:
            res["not_checkable"] += 1
            continue
        tbl, div = m.group(2), _divisor(names, m.group(5))
        sp = split_ref(names.get(tbl) or "")
        bs_sheet = sp[0] if sp else None
        if bs_sheet not in targets or div is None:
            res["not_checkable"] += 1
            continue
        srep = targets[bs_sheet]
        acct = wsv.cell(r, 4).value
        lab = sub_labels.get(col_idx(col))
        bsv, labels, rows = bs_info(bs_sheet)
        if not isinstance(acct, str) or lab not in labels or acct not in rows:
            res["unresolved"] += 1
            continue
        got = bsv.cell(rows[acct], labels[lab]).value
        p = ("A", int(c["period"][:-1])) if c["period"].endswith("A") else ("Q",) + q_parse(c["period"])
        expect, src = value_for(fins.get(srep["stock"]) or {}, srep["scope"], acct, p, prices,
                                is_convention=srep.get("is_convention", IS_CONVENTION_DEFAULT))
        if expect is None and got is None:
            res["unresolved"] += 1
            continue
        res["checked"] += 1
        if got is not None and expect is not None and abs(float(got) / div - expect / div) < 1e-6:
            res["matched"] += 1
        else:
            res["mismatch"].append({"cell": c["cell"], "acct": acct, "period": c["period"], "bs": "%s!%s%d" % (bs_sheet, col_letters(labels[lab]), rows[acct]),
                                    "value_bs": got, "expected_fin": expect, "src": src})
    if res["formula_bad"]:
        res["ok"] = False
        res["problems"].append("재오픈 수식 불일치 %d: %s" % (len(res["formula_bad"]), res["formula_bad"][:3]))
    if res["mismatch"]:
        res["ok"] = False
        res["problems"].append("사슬 흉내 값 ≠ fin %d: %s" % (len(res["mismatch"]), res["mismatch"][:3]))
    return res


def verify_all(rep, fins, prices=None):
    """검증 ①~⑥(+⑦ 연장 시)을 모아 rep['verify'] 에 넣는다. prices 는 ⑦ 의 시가총액 행 기대값에만 쓴다."""
    orig, out = rep["orig"], rep["out"]
    v = {"1_zip": verify_zip(out), "2_reopen": verify_reopen(orig, out, rep)}
    cells = []
    for s in rep["sheets"]:
        cells += ["%s!%s" % (s["sheet"], c) for c in s.get("written", [])]
    if rep.get("fx"):
        cells += ["변수!%s" % c for c in rep["fx"].get("written", [])]
    if rep.get("price"):
        cells += rep["price"].get("written", [])
    if rep.get("extend") and not rep["extend"].get("error"):
        cells += ["%s!%s" % (rep["extend"]["sheet"], c) for c in rep["extend"].get("written", [])]
    v["3_patched_cells"] = {"count": len(cells), "sample": cells[:12] + (["…"] if len(cells) > 12 else [])}
    v["4_unchanged"] = verify_unchanged(orig, out, rep)
    v["5_chain"] = verify_chain(orig, out, fins, rep)
    v["6_calc_chain"] = verify_calc_chain(orig, out, rep)
    keys = ["1_zip", "2_reopen", "4_unchanged", "5_chain", "6_calc_chain"]
    if rep.get("extend") is not None:
        v["7_extend"] = verify_extend(orig, out, fins, rep, prices)
        keys.append("7_extend")
    v["all_ok"] = all(v[k]["ok"] for k in keys)
    rep["verify"] = v
    return v


# ── CLI ──────────────────────────────────────────────────

def _fin_path(stock, fin_dir):
    return os.path.join(fin_dir, "%s.json" % stock)


def summarize(rep):
    lines = ["[%s %s] → %s" % (rep["stock"], rep["name"], rep["out"] if rep.get("written") else "(미출력) " + rep.get("error", ""))]
    for s in rep.get("sheets", []):
        if s.get("error"):
            lines.append("  %s: 오류 %s" % (s["sheet"], s["error"]))
            continue
        lines.append("  %s(%s %s): 마지막실적 %s@%s → 새 기간 %d열 %s..%s | 값셀 %d·라벨 %d·스탬프 %d·빈칸 %d·충돌 %d·교체 %d | 빈 계정 %d·fin 없는 분기 %d"
                     % (s["sheet"], s["stock"], s["scope"], s["last_actual"], s["last_actual_col"], len(s["new_periods"]),
                        s["new_periods"][0] if s["new_periods"] else "-", s["new_periods"][-1] if s["new_periods"] else "-",
                        s["cells_value"], s["cells_label"], s["cells_stamp"], s["cells_blank"], len(s["conflicts"]),
                        len(s.get("replaced", [])), len(s["missing_accounts"]), len(s["missing_quarters"])))
        srcs = s.get("sources") or {}
        lines.append("    IS 관행 %s(%s): 출처 is_ytd_diff %d·is %d·bs %d·cf %d · 재작성 셀 %d"
                     % (s.get("is_convention"), "fin.is_ytd_diff 있음" if s.get("is_ytd_diff_available") else "fin.is_ytd_diff 없음 → is",
                        srcs.get("is_ytd_diff", 0), srcs.get("is", 0), srcs.get("bs", 0), srcs.get("cf", 0), len(s.get("restated_cells", []))))
        for c in s.get("restated_cells", [])[:6]:
            lines.append("    ≠ %s %s %s: is %s → ytd_diff %s (차 %s)" % (c["cell"], c["acct"], c["period"], c["is"], c["ytd_diff"], c["diff"]))
        if len(s.get("restated_cells", [])) > 6:
            lines.append("    ≠ … 외 %d셀" % (len(s["restated_cells"]) - 6))
        for c in s.get("replaced", []):
            lines.append("    ↻ %s %s: %s → %s" % (c["cell"], c["acct"], re.sub(r"<[^>]+>", " ", c["was"]).strip()[:70], c["now"]))
    for s in rep.get("skipped_sheets", []):
        lines.append("  %s: 건너뜀 — %s" % (s["sheet"], s["why"]))
    if rep.get("fx"):
        f = rep["fx"]
        lines.append("  변수 환율: 채움 %d·기존값 보존 %d·실측 교체 %d·fx 없음 %d (행 %d)%s"
                     % (f.get("filled", 0), f.get("kept", 0), len(f.get("replaced", [])), f.get("missing_fx", 0), f.get("rows_found", 0),
                        " · partial 보존 %s" % ",".join(f["partial_skipped"]) if f.get("partial_skipped") else ""))
        byrow = {}
        for c in f.get("replaced", []):
            byrow.setdefault(c["row"], []).append(c)
        for row, cs in byrow.items():
            sc = (f.get("scale") or {}).get(row, {})
            lines.append("    ↻ %s: %d셀 실측 교체(배율 %s·중앙값비 %s·n=%s) %s..%s 예 %s %s %s → %s"
                         % (row, len(cs), sc.get("scale"), sc.get("median_ratio"), sc.get("n"), cs[0]["period"], cs[-1]["period"],
                            cs[0]["cell"], cs[0]["period"], cs[0]["was_value"], cs[0]["now"]))
        for sk in f.get("skipped_rows", []):
            lines.append("    ! %s 건너뜀 — %s (중앙값비 %s·n=%s)" % (sk["row"], sk["why"], sk.get("median_ratio"), sk.get("n")))
    if rep.get("price"):
        p = rep["price"]
        lines.append("  종가: %s" % (p.get("error") or "%s → %s (%s)" % (p.get("written"), p.get("close"), p.get("as_of"))))
    if rep.get("extend"):
        e = rep["extend"]
        if e.get("error"):
            lines.append("  수식 연장(%s): 오류 %s" % (e.get("mode"), e["error"]))
        else:
            lines.append("  수식 연장 %s(%s, 마지막실적 %s): %d셀(단순 %d·공유앵커 %d) %d행 | 값 있는 목표 보존 %d·상수 소스 %d·배열 %d·공유 건너뜀 %d·치환 오류 %d"
                         % (e["sheet"], e["mode"], e["last_actual"], e["extended"], e["extended_plain"], e["extended_shared"], e["rows"],
                            e["kept_nonempty"], e["skipped_const"], e["skipped_array"], e["skipped_shared"], len(e["errors"])))
            lines.append("    소스: " + " ".join("%s←%s" % (k, v) for k, v in list(e["sources"].items())[:5]) + (" …" if len(e["sources"]) > 5 else ""))
            for c in e["cells"][:3]:
                lines.append("    + %s(%s) ← %s: %s" % (c["cell"], c["kind"], c["source"], c["formula"][:70]))
            for er in e["errors"][:3]:
                lines.append("    ! " + er)
    if rep.get("calc_chain") and rep["calc_chain"].get("present"):
        cc = rep["calc_chain"]
        lines.append("  calcChain: 항목 %d · 수식→상수 교체 셀 항목 제거 %d%s" % (cc["entries"], len(cc["pruned"]),
                     (" · 못 찾음 %s" % cc["missing"]) if cc.get("missing") else ""))
    if rep.get("verify"):
        v = rep["verify"]
        lines.append("  검증: ①zip %s ②재오픈 %s(시트 %d/%d·수식 %d/%d%s) ③패치셀 %d ④무변경 %s(교체허용 %d·빈자리 %d·추가 %d) ⑤사슬 %s(%d건) ⑥calcChain %s(항목 %d→%d·고아 %d)%s → %s"
                     % ("OK" if v["1_zip"]["ok"] else "FAIL", "OK" if v["2_reopen"]["ok"] else "FAIL",
                        v["2_reopen"]["sheets_orig"], v["2_reopen"]["sheets_patched"], v["2_reopen"]["formula_cells_orig"], v["2_reopen"]["formula_cells_patched"],
                        (" 연장 +%d" % v["2_reopen"]["formula_cells_extended"]) if v["2_reopen"].get("formula_cells_extended") else "",
                        v["3_patched_cells"]["count"], "OK" if v["4_unchanged"]["ok"] else "FAIL",
                        len(v["4_unchanged"]["changed_existing_cells"]), v["4_unchanged"]["replaced_empty_cells"], v["4_unchanged"]["added_cells"],
                        "OK" if v["5_chain"]["ok"] else "FAIL", v["5_chain"]["checked"],
                        ("OK" if v["6_calc_chain"]["ok"] else "FAIL") if v["6_calc_chain"]["present"] else "없음",
                        v["6_calc_chain"]["entries_orig"], v["6_calc_chain"]["entries_patched"], len(v["6_calc_chain"]["stale"]),
                        (" ⑦연장 %s(수식 %d·사슬 %d/%d 일치·미검증 %d·미해결 %d)" % ("OK" if v["7_extend"]["ok"] else "FAIL", v["7_extend"]["formula_ok"],
                                                                     v["7_extend"]["matched"], v["7_extend"]["checked"], v["7_extend"]["not_checkable"], v["7_extend"]["unresolved"]))
                        if "7_extend" in v else "",
                        "ALL OK" if v["all_ok"] else "FAIL"))
        probs = v["4_unchanged"]["problems"][:5] + v["5_chain"]["problems"][:5] + v["6_calc_chain"]["problems"][:3]
        if "7_extend" in v:
            probs += v["7_extend"]["problems"][:3]
        for p in probs:
            lines.append("    ! " + str(p))
    return "\n".join(lines)


def main(argv=None):
    ap = argparse.ArgumentParser(description="레퍼런스 subQ xlsx 의 BS 시트를 fin json 으로 제자리 패치")
    ap.add_argument("--stock", action="append", help="075580|010140|010620 (반복 가능)")
    ap.add_argument("--all", action="store_true", help="세 파일 전부")
    ap.add_argument("--fin-dir", default=FIN_DIR, help="fin json 폴더(<stock>.json)")
    ap.add_argument("--fin", action="append", default=[], help="종목=경로 로 fin 을 직접 지정(픽스처 개발용)")
    ap.add_argument("--fx", default=os.path.join(ASSETS, "fx.json"))
    ap.add_argument("--prices", default=os.path.join(ASSETS, "prices.json"))
    ap.add_argument("--target", help="마지막 기간 분기키(기본: fin.quarters 의 마지막)")
    ap.add_argument("--today", default=datetime.date.today().isoformat())
    ap.add_argument("--ref-dir", default=REF_DIR)
    ap.add_argument("--out", help="출력 경로(--stock 하나일 때만)")
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--report", help="보고 JSON 저장 경로")
    ap.add_argument("--overwrite-placeholders", action="store_true",
                    help="새 기간 열에 애널리스트가 넣어 둔 가이던스·잠정치 셀을 fin 실적으로 바꾼다(원문은 보고에 남김)")
    ap.add_argument("--fx-actuals", action="store_true",
                    help="`변수` 환율 8행의 새 기간 열(BS 와 동일)에 남은 가정을 fx.json 실측으로 바꾼다(partial 분기는 보존, 전후는 보고에)")
    ap.add_argument("--is-convention", choices=IS_CONVENTIONS, default=IS_CONVENTION_DEFAULT,
                    help="분기 IS 값: ytd_diff=fin.is_ytd_diff 우선(FnGuide 누적차분, 기본) · 3m=항상 fin.is(보고서 3개월 열)")
    ap.add_argument("--extend-formulas", nargs="?", const="all", choices=EXTEND_CLI_MODES, default=EXTEND_CLI_DEFAULT,
                    help="subQ 새 기간 열의 빈 셀을 같은 분기위치 최근 실적 열 수식으로 채운다(상대참조 치환; 기본 켜짐 = all). "
                         "plain=단순 수식 소스만 · all=공유수식 앵커 텍스트도 · off=끄기(예전 동작)")
    a = ap.parse_args(argv)
    stocks = list(FILES) if a.all else (a.stock or [])
    if not stocks:
        ap.error("--stock 또는 --all")
    if a.out and len(stocks) != 1:
        ap.error("--out 은 --stock 하나일 때만")
    overrides = dict(s.split("=", 1) for s in a.fin)
    fx = load_json(a.fx)
    prices = load_json(a.prices)
    reports = []
    for stock in stocks:
        need = sorted({st for st, _ in FILES[stock]["sheets"].values()})
        fins = {}
        for st in need:
            fin = load_json(overrides.get(st) or _fin_path(st, a.fin_dir))
            if fin:
                fins[st] = fin
        target = a.target
        if not target:
            qs = sorted(q for f in fins.values() for q in (f.get("quarters") or []))
            target = qs[-1] if qs else None
        if not target:
            reports.append({"stock": stock, "name": FILES[stock]["name"], "error": "fin 없음 — target 결정 불가", "sheets": [], "skipped_sheets": []})
            print(summarize(reports[-1]))
            continue
        rep = patch_file(stock, fins, target, a.today, out_path=a.out, fx=fx, prices=prices, ref_dir=a.ref_dir,
                         overwrite=a.overwrite_placeholders, fx_actuals=a.fx_actuals, is_convention=a.is_convention,
                         extend=None if a.extend_formulas == "off" else a.extend_formulas)
        if a.verify and rep.get("written"):
            verify_all(rep, fins, prices)
        rep["fx_available"] = fx is not None
        rep["prices_available"] = prices is not None
        reports.append(rep)
        print(summarize(rep))
    if a.report:
        os.makedirs(os.path.dirname(os.path.abspath(a.report)), exist_ok=True)
        with open(a.report, "w", encoding="utf-8") as f:
            json.dump(reports, f, ensure_ascii=False, indent=1)
            f.write("\n")
    return 0 if all(r.get("written") and (not a.verify or r.get("verify", {}).get("all_ok")) for r in reports) else 1


if __name__ == "__main__":
    sys.exit(main())
