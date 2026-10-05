#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""selfcheck_models — 한국조선 실적 모델 파이프라인 **전 산출물 교차 정합 검사기**(V10 reproducibility, 2026-10-05).

다른 레인의 산출물을 읽기만 한다(고치지 않는다). 한 산출물이 다른 산출물·입력과 같은 숫자를 말하는지 전수로 대조해
표로 보고하고, 실패가 있으면 종료 코드 1 을 돌려준다. 성공 문구가 아니라 검사 건수·실패 건수·실패 셀 목록으로 말한다.

검사 그룹(스펙 MODEL_SPEC §2 스키마 · MODEL.md §3 계산 사슬):
  fin↔model      모델 actual 셀 = fin 값 ÷ 100 — 셀의 src 문구를 **파서로 풀어** 전수 재계산(fin.cons.is.X, 파생 ÷·+·−, 잔차, 부호 반전,
                 종속사 합, sls.reconcile, 토스 분기 평균 ÷ TTM EPS …). 거꾸로 fin 에 있는 값이 모델에 빠졌는지(coverage)도 본다.
  sls↔model      조선사 매출조선(실적) = sls.reconcile 부문매출, 매출조선(추정) = SLS 해양 원화(signed_by_origin) + 원장 밖 잔고 소진
                 + 신규수주 매출(패널 base), 매출액 = 매출조선 + 매출기타, OP조선신규 = 매출조선신규 × OPM 경로, sls 내부 항등식
                 (marine = signed_by_origin + post_origin), 타겟 OPM 경로 = sls.target_opm, 지주 = HD현대重 × 비율, 시나리오 FY 합.
  model↔summary  summary.json 행 전수(status·last_actual·driver·fy 9항목·per/pbr_now·backtest·warnings_n·scenarios) = 모델 json, counts·n·origin·정렬.
  model↔xlsx     README 스탬프, BS연결/BS별도 시트 = fin(백만원, 분기·연간 전수), subQ 확정 셀 VLOOKUP 을 직접 추적해 모델값과 대조,
                 추정 열은 kship_model_xlsx.verify 의 수식 에뮬레이터(레인 코드 재사용, 표기)로 재계산, 문서 속성 날짜 = built_at.
  model↔page     회사 페이지 섹션(id=kship-model-<stock>) data-* 속성·KPI 타일(FY26E~28E 매출/OP/OPM/EPS, PER/PBR/종가)·분기 손익표
                 전 셀을 **정규식으로 숫자만 뽑아** 모델과 대조(렌더러 포매터를 import 하지 않음), xlsx 다운로드 KB = 실제 파일 크기.
  summary↔hub    models.html 행 수·종목 집합 = 모집단, 각 행의 data-v(FY 매출/OP/OPM/EPS·PER·PBR)·상태 라벨·신규수주 = summary.
  집합           universe(57) ↔ 회사 폴더(56 = universe − 지주) ↔ 페이지(섹션 유무) ↔ models/fin/xlsx(58 = universe ∪ fin 전용 010620) ↔ sls(6) ↔ prices.
  prices↔model   valuation.price(종가·as_of) = prices.json, per_now/pbr_now/적정가치 재계산(내부 정합).
  fx↔model       assumptions.fx 추정 분기 = fx.json(quarters 완결 > forward).
  contracts↔sls  sls 계약 rcp 집합 = contracts.json(건너뛴 건·대체된 건 포함) — 원장이 sls 보다 새로우면 여기서 드러난다.
  panel↔model    매출조선신규 = forecast_panel base new_order_revenue ÷ 100.

종류(kind): consistency — 같은 세대의 산출물끼리는 항상 맞아야 한다(실패 = 결함 또는 일부만 재생성한 stale) ·
            freshness — 입력(prices·fx·contracts·panel)이 산출물보다 새로울 수 있다(일일 갱신은 시세·환율·계약만 받고 모델은 주 1회 —
            MODEL.md §8). 기본 종료 코드는 consistency 실패만 본다(--strict 면 freshness 도).

    python3 selfcheck_models.py                       # 레포 산출물 전수(58사) — 표 + 실패 상세
    python3 selfcheck_models.py --stocks 010140 075580 --json out.json
    python3 selfcheck_models.py --root /tmp/copy/argus/kship --no-emulate   # 다른 트리(재생성 사본) 검사

표준 라이브러리 + openpyxl. 네트워크 없음. 다른 레인 파일은 읽기만 한다.
"""
import argparse
import collections
import datetime
import gzip
import json
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))

# ── 스펙 §2-5 행 ↔ fin 계정(ROW_DEFS 의 거울 — 레인 코드를 import 하지 않고 여기 박아 둔다. 레인이 매핑을 바꾸면 여기서 드러난다)
CANON_MAP = [
    ("매출액", "is", "매출액(수익)"), ("매출원가", "is", "매출원가"), ("매출총이익", "is", "매출총이익"), ("판관비", "is", "판관비"),
    ("기타영업손익", "is", "기타영업손익"), ("영업이익", "is", "영업이익"), ("금융손익", "is", "금융손익"),
    ("기타영업외손익", "is", "기타영업외손익"), ("지분법손익", "is", "종속기업,공동지배기업및관계기업관련손익"),
    ("세전이익", "is", "법인세비용차감전계속사업이익"), ("법인세비용", "is", "법인세비용"), ("당기순이익", "is", "당기순이익"),
    ("지배주주순이익", "is", "(지배주주지분)당기순이익"), ("중단사업이익", "is", "중단사업이익"), ("감가상각비", "cf", "감가상각비"),
    ("자산총계", "bs", "자산총계"), ("부채총계", "bs", "부채총계"), ("자본총계", "bs", "자본총계"), ("지배주주지분", "bs", "지배주주지분"),
    ("총차입금", "bs", "총차입금"), ("순차입금", "bs", "순차입금"), ("이자발생자산", "bs", "이자발생자산"), ("현금및현금성자산", "bs", "현금및현금성자산"),
    ("CAPEX", "cf", "CAPEX"), ("영업활동현금흐름", "cf", "영업활동으로인한현금흐름"),
]
UNIT_DIV = 100.0
YARD_DRIVER = "sls_marine_plus_uncovered_backlog_runoff"
HOLDING, HOLDING_CORE = "009540", "329180"
SEJIN_SUBS = ("333430", "099410")
FY_EST = ("2026", "2027", "2028")
STATUS_KO = {"full": "완성", "partial": "부분", "no_fin": "재무 없음"}
# 문서화된 집합 예외: 지주는 폴더를 만들지 않는다(MODEL.md §11) · HD현대미포는 합병 소멸로 universe 밖, fin 만 있어 모델은 만든다(kship_model Ctx.population)
ALLOW_NO_FOLDER = {"009540"}
ALLOW_FIN_ONLY = {"010620"}
MINUS = "−"          # U+2212 — 모델 src 의 뺄셈 기호(하이픈 아님)


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _load(path):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def _read(path):
    try:
        with open(path, encoding="utf-8") as f:
            return f.read()
    except OSError:
        return None


def q_year(q):
    return int(q[:4])


def q_add(q, n):
    y, k = int(q[:4]), int(q[5])
    i = y * 4 + (k - 1) + n
    return "%dQ%d" % (i // 4, i % 4 + 1)


def qlabel_to_q(label):
    """xlsx 행1 라벨 '2Q26' → '2026Q2'. 연간 열은 int 연도."""
    if isinstance(label, int):
        return str(label)
    m = re.fullmatch(r"([1-4])Q(\d{2})", str(label or ""))
    return "20%sQ%s" % (m.group(2), m.group(1)) if m else None


# ── 숫자 표기 역파싱(렌더러 포매터와 독립 — 화면 문자열에서 숫자를 되뽑는다) ─────────────────────

def parse_num(s):
    """'127,235' → 127235.0 · '-76' → -76 · '1.7%' → 1.7 · '14.3배' → 14.3 · '—' → None."""
    if s is None:
        return None
    t = s.strip().replace(",", "").replace("배", "").replace("%", "").replace("원", "").replace("억", "").strip()
    if t in ("", "—", "-"):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def tol_a(v):
    """fmt_a 역허용치: |v|<10 은 소수 1자리, 그 밖은 정수 반올림."""
    return 0.05 + 1e-9 if abs(v) < 10 else 0.5 + 1e-9


def close(a, b, tol):
    return _num(a) and _num(b) and abs(a - b) <= tol


# ── 보고 구조 ───────────────────────────────────────────────

class Check:
    def __init__(self, key, label, kind):
        self.key, self.label, self.kind = key, label, kind      # kind: consistency | freshness | info
        self.n = 0
        self.fails = []
        self.notes = []

    def ok(self, n=1):
        self.n += n

    def fail(self, msg):
        self.n += 1
        self.fails.append(msg)

    def note(self, msg):
        if msg not in self.notes:
            self.notes.append(msg)

    def as_dict(self):
        return {"key": self.key, "label": self.label, "kind": self.kind, "checked": self.n, "failed": len(self.fails),
                "fails": self.fails, "notes": self.notes}


class Report:
    def __init__(self):
        self.checks = collections.OrderedDict()

    def c(self, key, label, kind="consistency"):
        if key not in self.checks:
            self.checks[key] = Check(key, label, kind)
        return self.checks[key]

    def failed(self, kinds=("consistency",)):
        return sum(len(c.fails) for c in self.checks.values() if c.kind in kinds)

    def as_dict(self):
        return {"checks": [c.as_dict() for c in self.checks.values()],
                "failed_consistency": self.failed(("consistency",)), "failed_freshness": self.failed(("freshness",))}

    def table(self, max_detail=8):
        w = max(len(c.label) for c in self.checks.values()) if self.checks else 10
        lines = ["%-*s  %-11s  %7s  %6s  %s" % (w, "검사", "종류", "검사수", "실패", "비고"),
                 "-" * (w + 40)]
        for c in self.checks.values():
            lines.append("%-*s  %-11s  %7d  %6d  %s" % (w, c.label, c.kind, c.n, len(c.fails), " · ".join(c.notes)[:120]))
        lines.append("-" * (w + 40))
        lines.append("consistency 실패 %d · freshness 실패 %d" % (self.failed(("consistency",)), self.failed(("freshness",))))
        for c in self.checks.values():
            if c.fails:
                lines.append("")
                lines.append("[%s] %s — 실패 %d (처음 %d)" % (c.kind, c.label, len(c.fails), min(max_detail, len(c.fails))))
                for f in c.fails[:max_detail]:
                    lines.append("  · " + f)
        return "\n".join(lines)


# ── 트리(경로) ──────────────────────────────────────────────

class Tree:
    """argus/kship 한 그루 — 레포 또는 재생성 사본. 모든 경로를 여기서 푼다."""

    def __init__(self, root):
        self.root = os.path.abspath(root)
        self.tools = os.path.join(self.root, "tools")
        self.assets = os.path.join(self.tools, "assets")
        self.models_dir = os.path.join(self.assets, "models")
        self.fin_dir = os.path.join(self.assets, "fin")
        self.sls_dir = os.path.join(self.assets, "sls")
        self.xlsx_dir = os.path.join(self.root, "models")
        self._cache = {}

    def model(self, stock):
        return self._json(os.path.join(self.models_dir, stock + ".json"))

    def fin(self, stock):
        return self._json(os.path.join(self.fin_dir, stock + ".json"))

    def sls(self, stock):
        return self._json(os.path.join(self.sls_dir, stock + ".json"))

    def _json(self, path):
        if path not in self._cache:
            self._cache[path] = _load(path)
        return self._cache[path]

    @property
    def summary(self):
        return self._json(os.path.join(self.models_dir, "summary.json"))

    @property
    def prices(self):
        return self._json(os.path.join(self.assets, "prices.json")) or {}

    @property
    def fx(self):
        return self._json(os.path.join(self.assets, "fx.json")) or {}

    @property
    def universe(self):
        return (self._json(os.path.join(self.assets, "universe.json")) or {}).get("rows") or []

    @property
    def suppliers(self):
        return self._json(os.path.join(self.assets, "suppliers.json")) or {}

    @property
    def contracts(self):
        return (self._json(os.path.join(self.assets, "contracts.json")) or {}).get("rows") or []

    @property
    def panel(self):
        if "panel" not in self._cache:
            p = os.path.join(self.assets, "forecast_panel.json.gz")
            try:
                with gzip.open(p, "rt", encoding="utf-8") as f:
                    d = json.load(f)
                self._cache["panel"] = {c["stock"]: c for c in d.get("companies") or [] if c.get("stock")}
            except (OSError, ValueError):
                self._cache["panel"] = {}
        return self._cache["panel"]

    def model_stocks(self):
        if not os.path.isdir(self.models_dir):
            return []
        return sorted(f[:6] for f in os.listdir(self.models_dir) if re.fullmatch(r"\d{6}\.json", f))

    def folders(self):
        return sorted(d for d in os.listdir(self.root) if re.fullmatch(r"\d{6}", d) and os.path.isdir(os.path.join(self.root, d)))

    def page(self, stock):
        return _read(os.path.join(self.root, stock, "index.html"))

    def xlsx_path(self, stock):
        return os.path.join(self.xlsx_dir, "%s_model.xlsx" % stock)


def rows_of(model):
    return {r["key"]: r for r in (model or {}).get("rows") or []}


def cell_v(rm, key, q, kind="q"):
    c = ((rm.get(key) or {}).get(kind) or {}).get(q)
    return c.get("v") if isinstance(c, dict) and _num(c.get("v")) else None


def fin_scope_of(fin, q):
    """분기별 사용 스코프(MODEL.md §3-5 별도 보충 규칙): 연결 손익이 있으면 cons, 그 분기 연결 매출이 없고 별도 매출이 있으면 sep."""
    cons_is = ((fin.get("cons") or {}).get("is") or {})
    if not cons_is:
        return "sep"
    if _num((cons_is.get(q) or {}).get("매출액(수익)")):
        return "cons"
    if _num((((fin.get("sep") or {}).get("is") or {}).get(q) or {}).get("매출액(수익)")):
        return "sep"
    return "cons"


def fin_val(fin, scope, kind, q, acct):
    v = ((((fin or {}).get(scope) or {}).get(kind) or {}).get(q) or {}).get(acct)
    return v if _num(v) else None


def fin_shares(fin, q):
    s = ((fin or {}).get("shares") or {}).get(q) or {}
    v = s.get("common_outstanding") or s.get("outstanding")
    return v if _num(v) and v > 0 else None


# ── ① fin↔model: actual 셀 src 파서 ─────────────────────────

class Unverifiable(Exception):
    pass


class Mislabel(Exception):
    """값은 파생(당기순이익 − 비지배)인데 src 가 fin face 계정으로 적힌 셀 — expected 를 들고 간다."""

    def __init__(self, msg, expected):
        super().__init__(msg)
        self.expected = expected


class SrcEval:
    """셀 src 문구 → 기대값. 문법은 모델 산출물의 실제 src 전집(2026-10-05 전수 58사)에서 뽑았다."""

    def __init__(self, tree, stock, model, fin):
        self.t, self.stock, self.m, self.fin = tree, stock, model, fin
        self.rm = rows_of(model)
        self.row_key = None

    def row(self, key, q):
        v = cell_v(self.rm, key, q)
        if v is None:
            raise Unverifiable("모델 행 %s %s 없음" % (key, q))
        return v

    def base(self, key, q):
        """행 key 의 **반올림 전** 값(억원). 모델은 파생 행(OPM·EPS·BPS·EBITDA·잔차)을 반올림 전 fin 값으로 계산하므로,
        actual 셀의 src 가 단일 fin 토큰이면 fin 값/100 을, 매출조선(실적)이면 sls.reconcile 값을, 그 밖은 셀 값(r2)을 돌려준다."""
        cell = ((self.rm.get(key) or {}).get("q") or {}).get(q) or {}
        v = cell.get("v")
        if not _num(v):
            raise Unverifiable("모델 행 %s %s 없음" % (key, q))
        if cell.get("kind") != "actual":
            return v
        core = (cell.get("src") or "").split(" (연결 손익 없는 분기")[0].split(" — ")[0].replace("(파생)", "")
        # 지배주주순이익 파생 src(2026-10-05 모델: 'fin.X.is.당기순이익 − fin.X.is.(비지배주주지분)당기순이익(파생)') — 반올림 전 값으로 재계산
        md = re.fullmatch(r"fin\.(cons|sep)\.is\.당기순이익 %s fin\.(cons|sep)\.is\.\(비지배주주지분\)당기순이익" % MINUS, core)
        if md and key == "지배주주순이익":
            ni, nci = fin_val(self.fin, md.group(1), "is", q, "당기순이익"), fin_val(self.fin, md.group(2), "is", q, "(비지배주주지분)당기순이익")
            if ni is not None and nci is not None:
                return (ni - nci) / UNIT_DIV
        m = re.fullmatch(r"fin\.(cons|sep)\.(is|bs|cf)\.(\S+)", core)
        if m:
            fv = fin_val(self.fin, m.group(1), m.group(2), q, m.group(3))
            if fv is not None:
                return fv / UNIT_DIV
            if key == "지배주주순이익":            # face 누락 → 당기순이익 − 비지배(모델 _actual_series 파생 규칙)
                ni, nci = fin_val(self.fin, m.group(1), "is", q, "당기순이익"), fin_val(self.fin, m.group(1), "is", q, "(비지배주주지분)당기순이익")
                if ni is not None and nci is not None:
                    return (ni - nci) / UNIT_DIV
        if key == "매출조선" and core.startswith("sls.reconcile"):
            rv = (((self.t.sls(self.stock) or {}).get("reconcile") or {}).get(q) or {}).get("reported_segment_rev_m")
            if _num(rv):
                return rv / UNIT_DIV
        return v

    def fin_tok(self, tok, q):
        neg = tok.endswith("¬")
        tok = tok.rstrip("¬")
        m = re.fullmatch(r"fin\.(cons|sep)\.(is|bs|cf)\.(.+)", tok)
        if not m:
            raise Unverifiable("fin 토큰 해석 불가 %r" % tok)
        v = fin_val(self.fin, m.group(1), m.group(2), q, m.group(3))
        if v is None:
            if m.group(3) == "(지배주주지분)당기순이익" and self.row_key == "지배주주순이익":
                ni, nci = fin_val(self.fin, m.group(1), "is", q, "당기순이익"), fin_val(self.fin, m.group(1), "is", q, "(비지배주주지분)당기순이익")
                if ni is not None and nci is not None:
                    raise Mislabel("fin face 에 (지배주주지분)당기순이익 %s 없음 → 당기순이익 − 비지배 파생값인데 src 는 fin face 로 표기" % q,
                                   (ni - nci) / UNIT_DIV)
            raise Unverifiable("fin.%s.%s.%s %s 값 없음" % (m.group(1), m.group(2), m.group(3), q))
        v = v / UNIT_DIV
        return -v if neg else v

    def eval(self, src, q, row_key=None):
        self.row_key = row_key
        s = (src or "").strip()
        if not s:
            raise Unverifiable("src 없음")
        # 특수 패턴(파생·외부 산출물) — 파생 행은 반올림 전 값(base)으로 다시 계산한다
        if s == "영업이익 ÷ 매출액(파생)":
            return self.base("영업이익", q) / self.base("매출액", q), "r4"
        if s in ("지배주주순이익 ÷ 유통주식수(파생)", "지배주주지분 ÷ 유통주식수(파생)"):
            key = "지배주주순이익" if s.startswith("지배주주순이익") else "지배주주지분"
            sh_cell = (self.rm.get("주식수") or {}).get("q", {}).get(q) or {}
            sh = fin_shares(self.fin, q) if sh_cell.get("kind") == "actual" else (sh_cell.get("v") * 1e6 if _num(sh_cell.get("v")) else None)
            if not sh:
                raise Unverifiable("유통주식수 %s 없음" % q)
            return self.base(key, q) * 1e8 / sh, "r1"
        if s == "영업이익 + 감가상각비(CF)":
            return self.base("영업이익", q) + self.base("감가상각비", q), "r2"
        m = re.fullmatch(r"= ([^\s(]+)\(.*\)", s)          # '= 매출액(부문 미분리)' · '= EPS(일회성 의심 없음)' — 셀 값을 그대로 복사한 행
        if m:
            if m.group(1) not in self.rm:
                raise Unverifiable("복사 원본 행 %s 없음" % m.group(1))
            return self.row(m.group(1), q), "exact"
        if s in ("토스 일봉 분기 평균 ÷ TTM EPS", "토스 일봉 분기 평균 ÷ BPS"):
            hist = ((self.t.prices.get("rows") or {}).get(self.stock) or {}).get("history_quarterly") or {}
            avg = (hist.get(q) or {}).get("avg")
            if not _num(avg):
                raise Unverifiable("prices.history_quarterly %s avg 없음" % q)
            if s.endswith("BPS"):
                return avg / self.row("BPS", q), "r2"
            eps = [cell_v(self.rm, "EPS", q_add(q, -i)) for i in range(4)]
            if any(e is None for e in eps):
                raise Unverifiable("TTM EPS 4분기 미완 %s" % q)
            return avg / sum(eps), "r2"
        if s.startswith("sls.reconcile.reported_segment_rev_m"):
            sls = self.t.sls(self.stock) or {}
            v = ((sls.get("reconcile") or {}).get(q) or {}).get("reported_segment_rev_m")
            if not _num(v):
                raise Unverifiable("sls.reconcile[%s].reported_segment_rev_m 없음" % q)
            return v / UNIT_DIV, "r2"
        if s == "연결 매출 − 조선 부문(파생)":
            return max(self.base("매출액", q) - self.base("매출조선", q), 0.0), "r2"
        m = re.fullmatch(r"일승·동방선기 fin (매출|영업이익) 합", s)
        if m:
            key = "매출액" if m.group(1) == "매출" else "영업이익"
            tot = 0.0
            for sub in SEJIN_SUBS:
                sm = self.t.model(sub)
                v = cell_v(rows_of(sm), key, q) if sm else None
                if v is None:                       # 종속사 모델에 없으면 fin 으로
                    f = self.t.fin(sub) or {}
                    v = fin_val(f, fin_scope_of(f, q), "is", q, "매출액(수익)" if key == "매출액" else "영업이익")
                    v = v / UNIT_DIV if v is not None else None
                if v is None:
                    raise Unverifiable("종속사 %s %s %s 없음" % (sub, key, q))
                tot += v
            return tot, "r2x2"
        if s == "연결 − 별도 − 종속사(베트남 종속·내부거래 잔차)":
            # 연결·별도는 반올림 전 fin, 종속사 합은 종속사 모델 행(r2) 합 — 모델도 같은 입력을 쓴다
            return self.base("매출액", q) - self.base("매출조선기자재", q) - (cell_v(self.rm, "매출종속사", q) or 0.0), "r2x2"
        if s == "연결 OP − 별도 OP − 종속사 OP(잔차)":
            return self.base("영업이익", q) - self.base("OP조선기자재", q) - (cell_v(self.rm, "OP종속사", q) or 0.0), "r2x2"
        if s == "fin.shares.common_outstanding":
            sh = fin_shares(self.fin, q)
            if not sh:
                raise Unverifiable("fin.shares %s 없음" % q)
            return sh / 1e6, "r3"
        if s.startswith("fin.dividend.dps_common"):
            d = ((self.fin.get("dividend") or {}).get(q) or {}).get("dps_common")
            if not _num(d):
                raise Unverifiable("fin.dividend[%s].dps_common 없음" % q)
            return float(d), "exact"
        # 일반 식: fin 토큰·모델 행 토큰을 ' − ' / ' + ' 로 잇는다. 꼬리 주석은 잘라낸다.
        core = s.split(" (연결 손익 없는 분기")[0].split(" — ")[0]
        core = core.replace("(음수 저장 → 부호 반전)", "¬").replace("(잔차)", "").replace("(파생)", "")
        parts = re.split(r"\s([%s+])\s" % MINUS, core)
        total, sign = 0.0, 1.0
        for i, p in enumerate(parts):
            if i % 2 == 1:
                sign = -1.0 if p == MINUS else 1.0
                continue
            p = p.strip()
            if p.startswith("fin."):
                v = self.fin_tok(p, q)
            elif p in self.rm:
                v = self.row(p, q)
            else:
                raise Unverifiable("토큰 해석 불가 %r (src %r)" % (p, src))
            total += sign * v
        return total, "r2" if len(parts) > 1 else "r2"


TOL = {"r2": 0.005 + 1e-7, "r2x2": 0.0101, "r2x3": 0.0151, "r4": 0.00005 + 1e-9, "r1": 0.05 + 1e-7, "r3": 0.0005 + 1e-9, "exact": 1e-9}


def check_fin_model(rep, tree, stocks):
    c = rep.c("fin_model", "fin↔model actual 셀(src 전수 재계산)")
    cu = rep.c("fin_model_src", "fin↔model 추적 불가 src")
    cl = rep.c("fin_model_mislabel", "fin↔model 파생값인데 src 가 fin face(지배주주순이익 = 당기순이익 − 비지배)")
    ch = rep.c("prices_hist_model", "prices.history_quarterly↔model PER/PBR 실적 셀(토스 분기 평균)", "freshness")
    cc = rep.c("fin_coverage", "fin→model coverage(fin 값이 모델에 빠짐)")
    srcs_unv = collections.Counter()
    for st in stocks:
        m = tree.model(st)
        if not m:
            continue
        fin = tree.fin(st)
        if not fin:
            if m.get("status") != "no_fin":
                c.fail("%s fin json 없음(status %s)" % (st, m.get("status")))
            continue
        ev = SrcEval(tree, st, m, fin)
        rm = ev.rm
        for key, r in rm.items():
            for q, cell in (r.get("q") or {}).items():
                if not isinstance(cell, dict) or cell.get("kind") != "actual":
                    continue
                # PER/PBR 실적 셀은 prices.json 이력(입력)에서 오므로 신선도 그룹으로 — 시세를 다시 받으면 모델보다 새로울 수 있다
                tgt = ch if str(cell.get("src") or "").startswith("토스 일봉") else c
                try:
                    exp, cls = ev.eval(cell.get("src"), q, row_key=key)
                except Mislabel as e:
                    v = cell.get("v")
                    cl.fail("%s %s %s: %s (모델 %s · 파생 재계산 %.2f%s)" % (
                        st, key, q, e, v, e.expected, "" if close(v, e.expected, TOL["r2"]) else " — 값도 다름"))
                    continue
                except Unverifiable as e:
                    if tgt is ch:
                        ch.fail("%s %s %s: %s" % (st, key, q, e))
                    else:
                        cu.fail("%s %s %s: %s" % (st, key, q, e))
                        srcs_unv[re.sub(r"\d{4}Q\d", "<q>", str(cell.get("src")))] += 1
                    continue
                v = cell.get("v")
                if not _num(v) or abs(v - exp) > TOL[cls]:
                    tgt.fail("%s %s %s: 모델 %s vs 재계산 %.4f (src %s)" % (st, key, q, v, exp, cell.get("src")))
                else:
                    tgt.ok()
        # coverage: fin 에 있는 값은 모델 actual 셀로 나와야 한다(스코프 규칙 포함)
        la = (m.get("periods") or {}).get("last_actual") or ""
        for q in fin.get("quarters") or []:
            if la and q > la:
                continue
            sc = fin_scope_of(fin, q)
            for key, kind, acct in CANON_MAP:
                fv = fin_val(fin, sc, kind, q, acct)
                if fv is None:
                    continue
                cell = ((rm.get(key) or {}).get("q") or {}).get(q)
                if not isinstance(cell, dict) or cell.get("kind") != "actual":
                    cc.fail("%s %s %s: fin.%s.%s 값 %.2f 있는데 모델 actual 셀 없음(셀 %s)" % (st, key, q, sc, kind, fv, (cell or {}).get("kind")))
                elif not close(cell.get("v"), fv / UNIT_DIV, 0.0051):
                    cc.fail("%s %s %s: 모델 %s vs fin %.4f" % (st, key, q, cell.get("v"), fv / UNIT_DIV))
                else:
                    cc.ok()
    if srcs_unv:
        cu.note("패턴 %d종: %s" % (len(srcs_unv), "; ".join("%s×%d" % kv for kv in srcs_unv.most_common(4))))


# ── ② sls↔model ────────────────────────────────────────────

def _fq(model):
    la = (model.get("periods") or {}).get("last_actual") or ""
    return [q for q in (model.get("periods") or {}).get("quarters") or [] if q > la]


def _seg_driver(model):
    seg = (model.get("segments") or [{}])[0] if model.get("segments") else {}
    return seg.get("driver") or {}, seg.get("opm_path") or {}


def check_sls_model(rep, tree, stocks):
    c = rep.c("sls_model", "sls↔model(조선사 매출조선·항등식·OPM 경로·지주·시나리오)")
    ci = rep.c("sls_internal", "sls 내부 항등식(marine = signed_by_origin + post_origin)")
    for st in stocks:
        m, sls = tree.model(st), tree.sls(st)
        if not m or not sls:
            continue
        rm = rows_of(m)
        drv, opm_path = _seg_driver(m)
        fq = _fq(m)
        bq = sls.get("by_quarter") or {}
        for q, b in bq.items():
            if "marine_hedged_krw_m_post_origin" in b:
                tot, a, p = b.get("marine_hedged_krw_m"), b.get("marine_hedged_krw_m_signed_by_origin"), b.get("marine_hedged_krw_m_post_origin")
                if not close(tot, (a or 0) + (p or 0), 0.2):
                    ci.fail("%s %s marine %s ≠ signed %s + post %s" % (st, q, tot, a, p))
                else:
                    ci.ok()
        if m.get("role") == "holding":
            core = tree.model(HOLDING_CORE)
            ratio = drv.get("ratio_used")
            if core and _num(ratio) and drv.get("type") == "subsidiary_yard_scaled":
                crm = rows_of(core)
                for q in fq:
                    cv, hv = cell_v(crm, "매출액", q), cell_v(rm, "매출액", q)
                    if cv is None or hv is None:
                        continue
                    if not close(hv, cv * ratio, abs(cv) * 0.00006 + 0.01):      # ratio_used 는 4자리 반올림
                        c.fail("%s %s 매출액 %s ≠ HD현대重 %s × %s" % (st, q, hv, cv, ratio))
                    else:
                        c.ok()
            continue
        if drv.get("type") != YARD_DRIVER:
            c.note("%s 드라이버 %s — 선표 검사 생략" % (st, drv.get("type")))
            continue
        # 실적: 매출조선 = reconcile 부문매출 / 100 (fin↔model 그룹에서도 보지만 sls 쪽 키로 다시)
        for q, r in (sls.get("reconcile") or {}).items():
            v = cell_v(rm, "매출조선", q)
            if _num(r.get("reported_segment_rev_m")) and v is not None:
                if close(v, r["reported_segment_rev_m"] / UNIT_DIV, 0.0051):
                    c.ok()
                else:
                    c.fail("%s %s 매출조선(실적) %s ≠ sls.reconcile %.2f" % (st, q, v, r["reported_segment_rev_m"] / UNIT_DIV))
        # 추정: 매출조선 = SLS(signed_by_origin | marine) + 잔고 소진 + 신규
        inc = bool(m.get("new_orders_included"))
        rate, remaining = drv.get("runoff_per_q"), drv.get("uncovered_backlog_at_origin")
        if not (_num(rate) and _num(remaining)):
            c.fail("%s driver.runoff_per_q/uncovered_backlog_at_origin 없음" % st)
            continue
        for q in fq:
            b = bq.get(q) or {}
            ex = b.get("marine_hedged_krw_m_signed_by_origin") if inc else b.get("marine_hedged_krw_m")
            if ex is None:
                ex = b.get("marine_hedged_krw_m") or 0.0
            r = min(rate, remaining)
            remaining -= r
            new = cell_v(rm, "매출조선신규", q) or 0.0
            exp = ex / UNIT_DIV + r + new
            v = cell_v(rm, "매출조선", q)
            if v is None:
                c.fail("%s %s 매출조선 추정 셀 없음" % (st, q))
            elif not close(v, exp, 0.06):
                c.fail("%s %s 매출조선 %s ≠ SLS %.2f + 소진 %.2f + 신규 %.2f = %.2f" % (st, q, v, ex / UNIT_DIV, r, new, exp))
            else:
                c.ok()
            tot, oth = cell_v(rm, "매출액", q), cell_v(rm, "매출기타", q)
            if tot is not None and v is not None and oth is not None:
                if close(tot, v + oth, 0.02):
                    c.ok()
                else:
                    c.fail("%s %s 매출액 %s ≠ 매출조선 %s + 매출기타 %s" % (st, q, tot, v, oth))
            op_path = opm_path.get(q)
            if _num(op_path):
                for rk, ok_ in (("OP조선신규", "매출조선신규"), ("OP조선", "매출조선"), ("OP기타", "매출기타")):
                    ov, rv = cell_v(rm, rk, q), cell_v(rm, ok_, q)
                    if ov is None or rv is None:
                        continue
                    if close(ov, rv * op_path, abs(rv) * 0.00006 + 0.011):
                        c.ok()
                    else:
                        c.fail("%s %s %s %s ≠ %s %s × OPM %s" % (st, q, rk, ov, ok_, rv, op_path))
                opm_v = cell_v(rm, "OPM", q)
                if opm_v is not None and not close(opm_v, op_path, 0.00011):
                    c.fail("%s %s OPM 행 %s ≠ 세그먼트 opm_path %s" % (st, q, opm_v, op_path))
            # 타겟 OPM 경로 = sls.target_opm
            tp = (drv.get("target_path") or {})
            so = (sls.get("target_opm") or {}).get(q) or {}
            if _num((tp.get("cohort_raw") or {}).get(q)) and _num(so.get("opm")):
                if close(tp["cohort_raw"][q], so["opm"], 0.00011) and close((tp.get("graded_share") or {}).get(q), so.get("graded_share"), 0.00011):
                    c.ok()
                else:
                    c.fail("%s %s 타겟 OPM %s/비중 %s ≠ sls.target_opm %s/%s" % (st, q, tp["cohort_raw"][q], (tp.get("graded_share") or {}).get(q), so.get("opm"), so.get("graded_share")))
        # 시나리오 FY: 매출 = 행 매출(base) − base 신규 + 시나리오 신규(분기 합)
        sc = m.get("scenarios") if isinstance(m.get("scenarios"), dict) else {}
        base = (sc.get("base") or {}).get("quarterly") or {}
        for scn in ("existing_only", "conservative", "optimistic", "base"):
            blk = sc.get(scn) or {}
            for y, a in (blk.get("annual") or {}).items():
                ks = ["%sQ%d" % (y, k) for k in range(1, 5)]
                exp = 0.0
                ok_all = True
                for k in ks:
                    qd = (blk.get("quarterly") or {}).get(k)
                    if qd:
                        exp += qd["rev"]
                    else:
                        v = cell_v(rm, "매출액", k)
                        if v is None or (((rm.get("매출액") or {}).get("q") or {}).get(k) or {}).get("kind") != "actual":
                            ok_all = False
                            break
                        exp += v
                if not ok_all:
                    continue
                if close(a.get("rev"), exp, 0.03):
                    c.ok()
                else:
                    c.fail("%s 시나리오 %s FY%s 매출 %s ≠ 분기 합 %.2f" % (st, scn, y, a.get("rev"), exp))
            for q, qd in (blk.get("quarterly") or {}).items():
                nb = (base.get(q) or {}).get("new_order_revenue") or 0.0
                rv = cell_v(rm, "매출액", q)
                if rv is None:
                    continue
                exp = rv - nb + (qd.get("new_order_revenue") or 0.0)
                if close(qd.get("rev"), exp, 0.02):
                    c.ok()
                else:
                    c.fail("%s 시나리오 %s %s 매출 %s ≠ 행 %s − base 신규 %s + 신규 %s" % (st, scn, q, qd.get("rev"), rv, nb, qd.get("new_order_revenue")))


# ── ③ model↔summary ────────────────────────────────────────

def check_summary(rep, tree, stocks):
    c = rep.c("summary", "model↔summary.json(행 전수)")
    s = tree.summary
    if not s:
        c.fail("summary.json 없음")
        return
    rows = s.get("rows") or []
    by = {r["stock"]: r for r in rows}
    if s.get("n") != len(rows):
        c.fail("n %s ≠ rows %d" % (s.get("n"), len(rows)))
    if [r["stock"] for r in rows] != sorted(r["stock"] for r in rows):
        c.fail("rows 가 종목코드 순이 아님")
    cnt = collections.Counter(r.get("status") for r in rows)
    for k in ("full", "partial", "no_fin"):
        if (s.get("counts") or {}).get(k) != cnt.get(k, 0):
            c.fail("counts.%s %s ≠ 실제 %d" % (k, (s.get("counts") or {}).get(k), cnt.get(k, 0)))
    la_max = max((r.get("last_actual") or "" for r in rows), default="")
    if s.get("origin") != (la_max or None):
        c.fail("origin %s ≠ max(last_actual) %s" % (s.get("origin"), la_max))
    ms = set(tree.model_stocks())
    if set(by) != ms:
        c.fail("summary 종목 집합 ≠ 모델 json 집합: summary만 %s · json만 %s" % (sorted(set(by) - ms), sorted(ms - set(by))))
    c.ok()
    for st in stocks:
        m, r = tree.model(st), by.get(st)
        if not m or not r:
            if m and not r:
                c.fail("%s summary 행 없음" % st)
            continue
        rm = rows_of(m)
        la = (m.get("periods") or {}).get("last_actual")
        pairs = [("name", m.get("name")), ("role", m.get("role")), ("status", m.get("status")), ("last_actual", la),
                 ("fin_quarters", (m.get("quality") or {}).get("fin_quarters")), ("driver", m.get("driver_type")),
                 ("driver_fallback", m.get("driver_fallback")), ("per_now", (m.get("valuation") or {}).get("per_now")),
                 ("pbr_now", (m.get("valuation") or {}).get("pbr_now")), ("new_orders_included", m.get("new_orders_included")),
                 ("identities_ok", (m.get("quality") or {}).get("identities_ok")),
                 ("warnings_n", len((m.get("quality") or {}).get("warnings") or []))]
        for k, v in pairs:
            if r.get(k) != v:
                c.fail("%s summary.%s %r ≠ 모델 %r" % (st, k, r.get(k), v))
            else:
                c.ok()
        bt = m.get("backtest") or {}
        for k in ("revenue_wape_pct", "op_wape_pct", "n", "driver_at_freeze"):
            if (r.get("backtest") or {}).get(k) != bt.get(k):
                c.fail("%s summary.backtest.%s %r ≠ 모델 %r" % (st, k, (r.get("backtest") or {}).get(k), bt.get(k)))
            else:
                c.ok()
        if la:
            y0 = q_year(la)
            exp_tags = ["%d%s" % (y, "E" if y >= y0 else "A") for y in range(y0 - 1, 2029)]
            if sorted(r.get("fy") or {}) != sorted(exp_tags):
                c.fail("%s fy 태그 %s ≠ 기대 %s" % (st, sorted(r.get("fy") or {}), exp_tags))
            for tag, f in (r.get("fy") or {}).items():
                y = tag[:4]
                for k, key in (("rev", "매출액"), ("op", "영업이익"), ("opm", "OPM"), ("ni_ctrl", "지배주주순이익"), ("eps", "EPS"),
                               ("bps", "BPS"), ("per", "PER"), ("pbr", "PBR"), ("dps", "DPS")):
                    mv = cell_v(rm, key, y, "a")
                    if f.get(k) != mv:
                        c.fail("%s fy.%s.%s %r ≠ 모델 %s.a[%s] %r" % (st, tag, k, f.get(k), key, y, mv))
                    else:
                        c.ok()
        sc = m.get("scenarios") if isinstance(m.get("scenarios"), dict) else {}
        if r.get("scenarios"):
            for tag, blk in r["scenarios"].items():
                for scn, vals in blk.items():
                    ann = ((sc.get(scn) or {}).get("annual") or {}).get(tag[:4]) or {}
                    if vals.get("rev") != ann.get("rev") or vals.get("op") != ann.get("op"):
                        c.fail("%s summary.scenarios.%s.%s %s ≠ 모델 %s" % (st, tag, scn, vals, {"rev": ann.get("rev"), "op": ann.get("op")}))
                    else:
                        c.ok()
        elif sc.get("base"):
            c.fail("%s 모델에 scenarios 있는데 summary.scenarios 없음" % st)


# ── ④ model↔xlsx ───────────────────────────────────────────

def check_xlsx(rep, tree, stocks, emulate=True):
    c = rep.c("xlsx", "model↔xlsx(README 스탬프·BS 시트=fin·subQ VLOOKUP 직접 추적)")
    ce = rep.c("xlsx_emul", "model↔xlsx 추정 열 재계산(kship_model_xlsx.verify 에뮬레이터 — 레인 코드 재사용)", "consistency" if emulate else "info")
    try:
        from openpyxl import load_workbook
    except ImportError:
        c.fail("openpyxl 없음")
        return
    X = None
    if emulate:
        # 에뮬레이터는 이 파일 옆(레인 코드)에서 가져온다 — sys.path 를 건드리지 않는다(검사 대상 트리의 모듈로 테스트 프로세스가 오염되지 않게).
        try:
            import kship_model_xlsx as X          # noqa: N812 — verify(path, model) 만 쓴다
            if not callable(getattr(X, "verify", None)):
                raise AttributeError("kship_model_xlsx.verify 없음(레인 코드 편집 중?)")
        except Exception as e:                   # noqa: BLE001
            ce.note("에뮬레이터 사용 불가(%s) — 추정 열 재계산 생략" % e)
            ce.kind = "info"
            X = None
    for st in stocks:
        m = tree.model(st)
        path = tree.xlsx_path(st)
        if not m:
            continue
        if not os.path.isfile(path):
            c.fail("%s xlsx 없음" % st)
            continue
        fin = tree.fin(st) or {}
        try:
            wb = load_workbook(path, data_only=False)
        except Exception as e:                   # noqa: BLE001
            c.fail("%s xlsx 열기 실패 %s" % (st, e))
            continue
        la = (m.get("periods") or {}).get("last_actual") or ""
        # README 스탬프
        ws = wb["README"] if "README" in wb.sheetnames else None
        if ws is None:
            c.fail("%s README 시트 없음" % st)
        else:
            a1, a2 = str(ws.cell(1, 1).value or ""), str(ws.cell(2, 1).value or "")
            if "%s(%s)" % (m.get("name"), st) not in a1:
                c.fail("%s README A1 %r 에 회사명(종목) 없음" % (st, a1))
            exp2 = "기준 분기 %s · 마지막 실적 %s · 생성 %s" % (m.get("origin"), la, m.get("built_at"))
            if a2 != exp2:
                c.fail("%s README A2 %r ≠ %r" % (st, a2, exp2))
            else:
                c.ok()
        # 문서 속성 날짜 = built_at(결정론 스탬프)
        bd = str(m.get("built_at") or "")[:10]
        mod = wb.properties.modified
        if not (isinstance(mod, datetime.datetime) and mod.date().isoformat() == bd):
            c.fail("%s 문서 속성 modified %s ≠ built_at %s" % (st, mod, bd))
        else:
            c.ok()
        # BS 시트 = fin(백만원). 행1 라벨 → 분기/연도, C열 계정, E열부터 값
        bs_vals = {}
        for sheet, scope in (("BS연결", "cons"), ("BS별도", "sep")):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            cols = {}
            for col in range(5, ws.max_column + 1):
                k = qlabel_to_q(ws.cell(1, col).value)
                if k:
                    cols[col] = k
            vals = {}
            fs = fin.get(scope) or {}
            nchk, nbad = 0, 0
            for r in range(5, ws.max_row + 1):
                acct = ws.cell(r, 3).value
                if not acct:
                    continue
                d = vals.setdefault(acct, {})
                stmt = next((s for s in ("bs", "is", "cf") if any(acct in (fs.get(s) or {}).get(q, {}) for q in (fs.get(s) or {}))), None)
                for col, k in cols.items():
                    v = ws.cell(r, col).value
                    if v is not None:
                        d[k] = v
                    if not fin or stmt is None:
                        continue
                    if "Q" in k:
                        e = ((fs.get(stmt) or {}).get(k) or {}).get(acct)
                    else:                                        # 연간: BS 는 Q4, 흐름은 4분기 합
                        qs = ["%sQ%d" % (k, i) for i in range(1, 5)]
                        if stmt == "bs":
                            e = ((fs.get("bs") or {}).get(k + "Q4") or {}).get(acct)
                        else:
                            parts = [((fs.get(stmt) or {}).get(q) or {}).get(acct) for q in qs]
                            e = round(sum(parts), 2) if all(_num(p) for p in parts) else None
                    e = e if _num(e) else None
                    nchk += 1
                    if (v is None) != (e is None) or (v is not None and not close(v, e, 0.011)):
                        nbad += 1
                        if nbad <= 3:
                            c.fail("%s %s!%s %s: 시트 %s ≠ fin %s" % (st, sheet, acct, k, v, e))
            c.ok(nchk - nbad)
            bs_vals[sheet] = vals
        # subQ 확정 셀 — VLOOKUP 을 직접 추적(에뮬레이터 없이)
        if "subQ" in wb.sheetnames:
            ws = wb["subQ"]
            rm = rows_of(m)
            cols = {col: qlabel_to_q(ws.cell(1, col).value) for col in range(5, ws.max_column + 1)}
            for r in range(5, ws.max_row + 1):
                key = ws.cell(r, 1).value
                if not key or key not in rm:
                    continue
                acct = ws.cell(r, 4).value
                for col, k in cols.items():
                    if not k or "Q" not in k or k > la:
                        continue
                    f = ws.cell(r, col).value
                    mv = cell_v(rm, key, k)
                    if isinstance(f, str) and f.startswith("=IFERROR(VLOOKUP("):
                        mm = re.search(r"VLOOKUP\(\$D(\d+),(BS[^,]+),MATCH\(([A-Z]+)\$1,", f)
                        if not mm or int(mm.group(1)) != r:
                            c.fail("%s subQ!%s%d VLOOKUP 행 참조 어긋남 %r" % (st, ws.cell(r, col).column_letter, r, f))
                            continue
                        got = (bs_vals.get(mm.group(2)) or {}).get(acct, {}).get(k)
                        if got is None:
                            c.fail("%s subQ %s %s: %s 에 %s 값 없음(VLOOKUP 0 표시 위험)" % (st, key, k, mm.group(2), acct))
                        elif mv is None:
                            c.fail("%s subQ %s %s: 시트 VLOOKUP %s 인데 모델 셀 없음" % (st, key, k, got))
                        elif not close(got / UNIT_DIV, mv, 0.0051):
                            c.fail("%s subQ %s %s: VLOOKUP→%s/100 ≠ 모델 %s" % (st, key, k, got, mv))
                        else:
                            c.ok()
                    elif isinstance(f, (int, float)):
                        if mv is None or abs(f - mv) > 1e-9:
                            c.fail("%s subQ %s %s: 직접값 %s ≠ 모델 %s" % (st, key, k, f, mv))
                        else:
                            c.ok()
        if X is not None:
            try:
                v = X.verify(path, m)
                bad = [ch for ch in v.get("checks") or [] if not ch.get("ok")]
                mis = (v.get("estimate_recalc") or {}).get("mismatch") or []
                if bad or mis or not v.get("size_ok", True):
                    ce.fail("%s verify: VLOOKUP 실패 %d · 추정 재계산 불일치 %d(%s) · size_ok %s" % (
                        st, len(bad), len(mis), "; ".join("%s %s 모델 %s vs %s" % (x["key"], x["q"], x["model"], x["emulated"]) for x in mis[:3]), v.get("size_ok")))
                else:
                    ce.ok((v.get("estimate_recalc") or {}).get("match") or 0)
            except Exception as e:               # noqa: BLE001
                ce.fail("%s verify 예외 %s: %s" % (st, type(e).__name__, e))


# ── ⑤ model↔page ───────────────────────────────────────────

def _section_html(page, stock):
    i = page.find('id="kship-model-%s"' % stock)
    if i < 0:
        return None
    j = page.find("<div", i - 200) if i >= 200 else 0
    start = page.rfind("<div", 0, i)
    return page[start if start >= 0 else i:]


def check_pages(rep, tree, stocks):
    c = rep.c("page", "model↔회사 페이지(data-* 속성·KPI·분기 손익표 전 셀·xlsx KB)")
    cm = rep.c("page_markup", "회사 페이지 섹션 마크업 인식(정규식) — 실패면 렌더러 변경 가능", "consistency")
    cp = rep.c("page_prices", "회사 페이지 TTM PER/PBR 표기 ↔ prices.json", "freshness")
    summary_by = {r["stock"]: r for r in (tree.summary or {}).get("rows") or []}
    for st in stocks:
        m = tree.model(st)
        page = tree.page(st)
        if not m or page is None:
            continue
        sec = _section_html(page, st)
        if sec is None:
            c.fail("%s 페이지에 섹션 id=kship-model-%s 없음" % (st, st))
            continue
        if page.count('id="kship-model-%s"' % st) != 1:
            c.fail("%s 섹션 id 가 %d번" % (st, page.count('id="kship-model-%s"' % st)))
        rm = rows_of(m)
        # data-* 속성
        attrs = dict(re.findall(r'data-model-([a-z-]+)="([^"]*)"', sec[:2000]))
        exp_status = (summary_by.get(st) or {}).get("status") or m.get("status")
        inc = m.get("new_orders_included")
        exp_attr = {"status": exp_status, "origin": m.get("origin"), "driver": m.get("driver_type"),
                    "new-orders": "included" if inc is True else ("excluded" if inc is False else "n/a")}
        for k, v in exp_attr.items():
            if attrs.get(k) != v:
                c.fail("%s data-model-%s %r ≠ %r" % (st, k, attrs.get(k), v))
            else:
                c.ok()
        # KPI 타일
        tiles = re.findall(r"<b>([^<]*)<small>억</small></b>\s*<span>FY(\d{4})E 매출 · (추정|실적)</span>\s*<i>OP ([^<·]+)억 · OPM ([^<·]+) · EPS ([^<]+?)원</i>", sec)
        if len(tiles) != 3:
            cm.fail("%s KPI 타일 %d개 인식(기대 3)" % (st, len(tiles)))
        else:
            cm.ok()
            for rev_s, y, _, op_s, opm_s, eps_s in tiles:
                rev, op, eps = cell_v(rm, "매출액", y, "a"), cell_v(rm, "영업이익", y, "a"), cell_v(rm, "EPS", y, "a")
                opm = (op / rev) if (rev and op is not None) else None
                for nm, shown, exp, tol in (("매출", rev_s, rev, tol_a(rev) if rev is not None else 0), ("OP", op_s, op, tol_a(op) if op is not None else 0),
                                            ("OPM%", opm_s, (opm * 100) if opm is not None else None, 0.05 + 1e-9), ("EPS", eps_s, eps, 0.5 + 1e-9)):
                    got = parse_num(shown)
                    if (got is None) != (exp is None) or (exp is not None and abs(got - exp) > tol):
                        c.fail("%s KPI FY%sE %s 화면 %r ≠ 모델 %s" % (st, y, nm, shown, exp))
                    else:
                        c.ok()
        val = m.get("valuation") or {}
        mper = re.search(r"<b>([^<]*)</b>\s*<span>현재 PER\(12M fwd EPS ([^)]*?)원\)</span>\s*<i>([^<]*)</i>", sec)
        mpbr = re.search(r"<b>([^<]*)</b>\s*<span>현재 PBR\(최근 BPS ([^)]*?)원\)</span>\s*<i>종가 ([^<·]+?)원 · ([^<]*)</i>", sec)
        if not mper or not mpbr:
            cm.fail("%s PER/PBR 타일 인식 실패(PER %s · PBR %s)" % (st, bool(mper), bool(mpbr)))
        else:
            cm.ok()
            checks = [("per_now", mper.group(1), val.get("per_now"), 0.05), ("eps_fwd12m", mper.group(2), val.get("eps_fwd12m"), 0.5),
                      ("pbr_now", mpbr.group(1), val.get("pbr_now"), 0.05), ("bps_latest", mpbr.group(2), val.get("bps_latest"), 0.5),
                      ("close", mpbr.group(3), (val.get("price") or {}).get("close"), 0.5)]
            for nm, shown, exp, tol in checks:
                got = parse_num(shown)
                if (got is None) != (exp is None) or (exp is not None and abs(got - exp) > tol + 1e-9):
                    c.fail("%s KPI %s 화면 %r ≠ 모델 %s" % (st, nm, shown, exp))
                else:
                    c.ok()
            if (mpbr.group(4) or "").strip() != str((val.get("price") or {}).get("as_of") or ""):
                c.fail("%s KPI 종가 as_of 화면 %r ≠ 모델 %r" % (st, mpbr.group(4), (val.get("price") or {}).get("as_of")))
            # TTM 표기 ↔ prices.json(페이지 생성 시점의 prices)
            pr = (tree.prices.get("rows") or {}).get(st) or {}
            mt = re.search(r"TTM ([^ /]+) / ([^(]+)\(aikstockdata\)", mper.group(3))
            if mt and (_num(pr.get("pe_ttm")) or _num(pr.get("pb"))):
                for shown, exp in ((mt.group(1), pr.get("pe_ttm")), (mt.group(2), pr.get("pb"))):
                    got = parse_num(shown)
                    if (got is None) != (not _num(exp)) or (_num(exp) and abs(got - exp) > 0.05 + 1e-9):
                        cp.fail("%s 페이지 TTM %r ≠ prices.json %s" % (st, shown, exp))
                    else:
                        cp.ok()
        # 분기 손익표: 헤더 분기 + 각 행 셀
        tbl = re.search(r"분기 손익.*?<table class=\"pnl\">\s*<thead><tr>(.*?)</tr></thead><tbody>(.*?)</tbody></table>", sec, re.S)
        if not tbl:
            cm.fail("%s 분기 손익표 인식 실패" % st)
        else:
            qs = re.findall(r"<th[^>]*>(\d{4}Q\d)<br>(?:A|E)</th>", tbl.group(1))
            la = (m.get("periods") or {}).get("last_actual") or ""
            all_q = (m.get("periods") or {}).get("quarters") or []
            exp_q = [q for q in all_q if q <= la][-8:] + [q for q in all_q if q > la][:10]
            if qs != exp_q:
                cm.fail("%s 손익표 분기 머리 %s ≠ 기대 %s" % (st, qs[:3] + ["…"], exp_q[:3] + ["…"]))
            label_to_key = {}
            for r in m.get("rows") or []:
                label_to_key.setdefault(r.get("label") or r["key"], r["key"])
            nrow = 0
            for th_attr, label, unit, cells in re.findall(r'<tr[^>]*><th class="rowh" scope="row"([^>]*)>([^<]*)</th><td class="l mut">([^<]*)</td>(.*?)</tr>', tbl.group(2), re.S):
                full = re.search(r'title="([^"]*)"', th_attr)
                key = label_to_key.get(label) or label_to_key.get(full.group(1) if full else "") or (label if label in rm else None)
                if key not in rm:
                    continue
                tds = re.findall(r"<td[^>]*>([^<]*)</td>", cells)
                if len(tds) != len(qs):
                    cm.fail("%s 손익표 %s 셀 %d ≠ 분기 %d" % (st, key, len(tds), len(qs)))
                    continue
                nrow += 1
                for q, shown in zip(qs, tds):
                    mv = cell_v(rm, key, q)
                    got = parse_num(shown)
                    if mv is None:
                        if got is None:
                            c.ok()
                        else:
                            c.fail("%s 손익표 %s %s 화면 %r 인데 모델 셀 없음" % (st, key, q, shown))
                        continue
                    if unit == "%":
                        exp, tol = mv * 100, 0.05 + 1e-9
                    elif unit == "원":
                        exp, tol = mv, 0.5 + 1e-9
                    else:
                        exp, tol = mv, tol_a(mv)
                    if got is None or abs(got - exp) > tol:
                        c.fail("%s 손익표 %s %s 화면 %r ≠ 모델 %s" % (st, key, q, shown, mv))
                    else:
                        c.ok()
            if nrow < 5:
                cm.fail("%s 손익표에서 모델 행과 짝지은 행 %d개(기대 ≥5)" % (st, nrow))
            else:
                cm.ok()
        # xlsx 다운로드 KB = 실제 크기
        md = re.search(r'href="\.\./models/%s_model\.xlsx" download>[^<]*<span class="mut">(\d+) KB</span>' % st, sec)
        xp = tree.xlsx_path(st)
        if os.path.isfile(xp):
            if not md:
                c.fail("%s 다운로드 링크 없음(xlsx 는 있음)" % st)
            elif abs(int(md.group(1)) - round(os.path.getsize(xp) / 1024)) > 1:
                c.fail("%s 다운로드 %s KB ≠ 실제 xlsx %.0f KB(페이지가 xlsx 보다 오래됨)" % (st, md.group(1), os.path.getsize(xp) / 1024))
            else:
                c.ok()
        elif md:
            c.fail("%s 다운로드 링크 있는데 xlsx 파일 없음" % st)


# ── ⑥ summary↔hub ──────────────────────────────────────────

def population(tree):
    """허브 모집단 = universe ∪ suppliers.cos ∪ 회사 폴더(렌더 확인 출력 제외)."""
    s = {r["stock"] for r in tree.universe}
    s |= {c["stock"] for c in tree.suppliers.get("cos") or []}
    for d in tree.folders():
        pg = tree.page(d)
        if pg is not None and "실적 모델(렌더 확인)" not in pg[:4000]:
            s.add(d)
    return sorted(s)


def check_hub(rep, tree):
    c = rep.c("hub", "summary↔models.html(행 집합·data-v·상태·신규수주·머리 KPI)")
    html = _read(os.path.join(tree.root, "models.html"))
    s = tree.summary or {}
    by = {r["stock"]: r for r in s.get("rows") or []}
    if html is None:
        c.fail("models.html 없음")
        return
    pop = population(tree)
    body = re.search(r'<table id="mtab"[^>]*>.*?<tbody>(.*?)</tbody>', html, re.S)
    if not body:
        c.fail("허브 표 tbody 인식 실패")
        return
    # 머리글 열 이름 → 번호. 열이 추가·이동돼도(폴백 사유 열 등) 이름으로 찾는다.
    head = re.search(r'<table id="mtab"[^>]*>\s*<thead>\s*<tr>(.*?)</tr>', html, re.S)
    ths = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", t)).strip() for t in re.findall(r"<th[^>]*>(.*?)</th>", head.group(1), re.S)] if head else []
    col = {name: i for i, name in enumerate(ths)}
    need = ["회사", "종목코드", "드라이버", "신규 수주", "PER 현재", "PBR 현재", "상태", "xlsx"] + \
           ["FY%sE %s" % (y[2:], k) for y in FY_EST for k in ("매출(억)", "OP(억)", "OPM", "EPS(원)")]
    missing = [n for n in need if n not in col]
    if missing:
        c.fail("허브 머리글에 열 없음 %s (머리글 %s)" % (missing, ths))
        return

    def td_attr(td, name):
        m = re.search(r'%s="([^"]*)"' % name, td)
        return m.group(1) if m else None

    def td_text(td):
        return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", td)).strip()

    trs = re.findall(r"<tr>(.*?)</tr>", body.group(1), re.S)
    seen = []
    for tr in trs:
        tds = re.findall(r"<td[^>]*>.*?</td>", tr, re.S)
        if len(tds) < 2:
            continue
        st = td_text(tds[col["종목코드"]])
        seen.append(st)
        r = by.get(st)
        if not r:
            if "모델 없음" in tr:
                c.ok()
            else:
                c.fail("%s 허브 행에 모델 값이 있는데 summary 에 없음" % st)
            continue
        if len(tds) != len(ths):
            c.fail("%s 허브 행 셀 %d ≠ 머리글 %d" % (st, len(tds), len(ths)))
            continue
        st_txt = td_text(tds[col["상태"]])
        if st_txt != STATUS_KO.get(r.get("status"), r.get("status")):
            c.fail("%s 상태 라벨 %r ≠ summary %s" % (st, st_txt, r.get("status")))
        else:
            c.ok()
        exp_no = {True: "2", False: "1"}.get(r.get("new_orders_included"), "0")
        if td_attr(tds[col["신규 수주"]], "data-v") != exp_no:
            c.fail("%s 신규수주 data-v %r ≠ %r" % (st, td_attr(tds[col["신규 수주"]], "data-v"), exp_no))
        else:
            c.ok()
        for y in FY_EST:
            f = (r.get("fy") or {}).get(y + "E") or {}
            rev, op, eps = f.get("rev"), f.get("op"), f.get("eps")
            opm = (op / rev * 100) if (rev and op is not None) else None
            for nm, exp, tol in (("매출(억)", rev, 1e-6), ("OP(억)", op, 1e-6), ("OPM", opm, 0.0051), ("EPS(원)", eps, 1e-6)):
                shown = td_attr(tds[col["FY%sE %s" % (y[2:], nm)]], "data-v")
                got = parse_num(shown)
                if (got is None) != (exp is None) or (exp is not None and abs(got - exp) > tol):
                    c.fail("%s FY%sE %s data-v %r ≠ summary %s" % (st, y, nm, shown, exp))
                else:
                    c.ok()
        for nm, exp in (("PER 현재", r.get("per_now")), ("PBR 현재", r.get("pbr_now"))):
            shown = td_attr(tds[col[nm]], "data-v")
            got = parse_num(shown)
            if (got is None) != (exp is None) or (exp is not None and abs(got - exp) > 1e-6):
                c.fail("%s %s data-v %r ≠ summary %s" % (st, nm, shown, exp))
            else:
                c.ok()
        drv = td_attr(tds[col["드라이버"]], "title")
        if drv != (r.get("driver") or ""):
            c.fail("%s 드라이버 title %r ≠ summary %r" % (st, drv, r.get("driver")))
        else:
            c.ok()
        has_sec = ('id="kship-model-%s"' % st) in (tree.page(st) or "")
        if has_sec != (('href="%s/index.html#kship-model-%s"' % (st, st)) in tr):
            c.fail("%s 앵커 링크 %s ≠ 페이지 섹션 유무 %s" % (st, not has_sec, has_sec))
        else:
            c.ok()
        if os.path.isfile(tree.xlsx_path(st)) != ('href="models/%s_model.xlsx"' % st in tr):
            c.fail("%s xlsx 링크 ≠ 파일 유무" % st)
        else:
            c.ok()
    # 행 집합: 모집단(universe ∪ suppliers ∪ 폴더)은 전부 있어야 하고, 그 밖의 행은 모델 json 이 있는 회사(fin 전용 010620)만 허용한다.
    # 10-05 L7 이 fin 전용 회사를 표에 넣기 시작했다(57 → 58행) — 두 쪽 다 받되 모집단 밖 행은 비고로 남긴다.
    ms = set(tree.model_stocks())
    if set(pop) - set(seen) or set(seen) - set(pop) - ms or seen != sorted(seen) or len(seen) != len(set(seen)):
        c.fail("허브 행 %d vs 모집단 %d: 모집단만 %s · 허브만(모델 없음) %s · 정렬/중복 %s" % (
            len(seen), len(pop), sorted(set(pop) - set(seen)), sorted(set(seen) - set(pop) - ms), seen != sorted(seen) or len(seen) != len(set(seen))))
    else:
        c.ok()
        if set(seen) - set(pop):
            c.note("모집단 밖 행(모델만 있음) %s" % sorted(set(seen) - set(pop)))
    # 머리 KPI — 분자(모델 생성)는 표 행 중 모델 있는 회사 수, 분모는 모집단 또는 표 행 수(L7 정의가 바뀌는 중이라 둘 다 받고 비고)
    mk = re.search(r"<b>(\d+)<small>/ (\d+)</small></b><span>모델 생성 · 모집단", html)
    built = sum(1 for st in seen if tree.model(st))
    if not mk or int(mk.group(1)) != built or int(mk.group(2)) not in (len(pop), len(seen)):
        c.fail("허브 머리 '모델 생성/모집단' %s ≠ %d/%d(또는 %d)" % (mk.groups() if mk else None, built, len(pop), len(seen)))
    else:
        c.ok()
        if int(mk.group(2)) != int(mk.group(1)) and int(mk.group(2)) < built:
            c.note("머리 '모델 생성 %s / 모집단 %s' — 분자가 분모보다 큼(fin 전용 회사 포함)" % mk.groups())
    mk = re.search(r"<b>(\d+)</b><span>완성\(full\) · 부분 (\d+) · 재무 없음 (\d+) · 페이지 섹션 없음 (\d+)", html)
    cnt = collections.Counter((by.get(st) or {}).get("status") for st in seen if tree.model(st))
    nosec = sum(1 for st in seen if tree.model(st) and tree.page(st) is not None and ('id="kship-model-%s"' % st) not in tree.page(st))
    exp = (cnt.get("full", 0), cnt.get("partial", 0), cnt.get("no_fin", 0), nosec)
    if not mk or tuple(int(x) for x in mk.groups()) != exp:
        c.fail("허브 머리 상태 집계 %s ≠ 표 행 기준 %s" % (mk.groups() if mk else None, exp))
    else:
        c.ok()
    mk = re.search(r"<b>([^<]*)</b><span>최근 생성", html)
    latest = max((str(tree.model(st).get("built_at") or "") for st in seen if tree.model(st)), default="")
    if not mk or mk.group(1) != latest:
        c.fail("허브 '최근 생성' %r ≠ 모델 built_at 최대 %r" % (mk.group(1) if mk else None, latest))
    else:
        c.ok()


# ── ⑦ 집합 ─────────────────────────────────────────────────

def check_sets(rep, tree):
    c = rep.c("sets", "집합 universe↔폴더↔페이지↔models/fin/xlsx↔sls↔prices")
    uni = [r["stock"] for r in tree.universe]
    us = set(uni)
    if len(uni) != len(us) or any(not re.fullmatch(r"\d{6}", s) for s in uni):
        c.fail("universe 종목 중복/형식 오류")
    folders = set(tree.folders())
    if folders - us:
        c.fail("universe 밖 폴더 %s" % sorted(folders - us))
    miss = us - folders - ALLOW_NO_FOLDER
    if miss:
        c.fail("universe 회사인데 폴더 없음 %s" % sorted(miss))
    for st in sorted(folders):
        pg = tree.page(st)
        if pg is None:
            c.fail("%s 폴더에 index.html 없음" % st)
        elif "실적 모델(렌더 확인)" in pg[:4000]:
            c.fail("%s 폴더의 index.html 이 렌더 확인 출력" % st)
        elif tree.model(st) and ('id="kship-model-%s"' % st) not in pg:
            c.fail("%s 모델 있는데 페이지에 섹션 없음(페이지 재생성 누락)" % st)
        else:
            c.ok()
    ms = set(tree.model_stocks())
    fins = {f[:6] for f in os.listdir(tree.fin_dir) if re.fullmatch(r"\d{6}\.json", f)} if os.path.isdir(tree.fin_dir) else set()
    xl = {f[:6] for f in os.listdir(tree.xlsx_dir) if re.fullmatch(r"\d{6}_model\.xlsx", f)} if os.path.isdir(tree.xlsx_dir) else set()
    extra = ms - us - ALLOW_FIN_ONLY
    if extra:
        c.fail("universe 밖 모델(허용 목록 밖) %s" % sorted(extra))
    if us - ms:
        c.fail("universe 회사인데 모델 없음 %s" % sorted(us - ms))
    if fins != ms:
        c.fail("fin 집합 ≠ 모델 집합: fin만 %s · 모델만 %s" % (sorted(fins - ms), sorted(ms - fins)))
    if xl != ms:
        c.fail("xlsx 집합 ≠ 모델 집합: xlsx만 %s · 모델만 %s" % (sorted(xl - ms), sorted(ms - xl)))
    yards = {r["stock"] for r in tree.universe if r.get("role") in ("yard", "holding")}
    sls = {f[:6] for f in os.listdir(tree.sls_dir) if re.fullmatch(r"\d{6}\.json", f)} if os.path.isdir(tree.sls_dir) else set()
    if sls != yards:
        c.fail("sls 집합 %s ≠ 조선사+지주 %s" % (sorted(sls), sorted(yards)))
    pr = set((tree.prices.get("rows") or {}).keys())
    if us - pr:
        c.fail("universe 회사인데 prices 행 없음 %s" % sorted(us - pr))
    c.ok()
    c.note("universe %d · 폴더 %d · 모델 %d · fin %d · xlsx %d · sls %d · prices %d" % (len(us), len(folders), len(ms), len(fins), len(xl), len(sls), len(pr)))


# ── ⑧ prices↔model · ⑨ fx↔model · ⑩ contracts↔sls · ⑪ panel↔model ───────────

def check_inputs(rep, tree, stocks):
    cp = rep.c("prices_model", "prices.json↔model.valuation.price(종가·as_of)", "freshness")
    cv = rep.c("valuation", "valuation 내부 재계산(per_now·pbr_now·적정가치·fwd EPS)")
    cf = rep.c("fx_model", "fx.json↔model.assumptions.fx(추정 분기)", "freshness")
    cc = rep.c("contracts_sls", "contracts.json↔sls 계약 집합", "freshness")
    cn = rep.c("panel_model", "forecast_panel↔model 매출조선신규(base)", "freshness")
    prices = tree.prices.get("rows") or {}
    fx = tree.fx
    for st in stocks:
        m = tree.model(st)
        if not m:
            continue
        rm = rows_of(m)
        val = m.get("valuation") or {}
        pr = prices.get(st)
        vp = val.get("price") or {}
        if pr:
            if vp.get("close") != pr.get("close") or str(vp.get("as_of")) != str(pr.get("as_of")):
                cp.fail("%s 모델 종가 %s(%s) ≠ prices %s(%s)" % (st, vp.get("close"), vp.get("as_of"), pr.get("close"), pr.get("as_of")))
            else:
                cp.ok()
        elif vp.get("close") is not None:
            cp.fail("%s prices 행 없는데 모델 종가 %s" % (st, vp.get("close")))
        # 내부 재계산
        close_v, eps_f, bps = vp.get("close"), val.get("eps_fwd12m"), val.get("bps_latest")
        qs = val.get("eps_fwd_quarters") or []
        if qs:
            parts = [cell_v(rm, "EPS", q) for q in qs]
            if all(p is not None for p in parts):
                if not close(eps_f, round(sum(parts), 1), 0.051):
                    cv.fail("%s eps_fwd12m %s ≠ Σ EPS %s = %.1f" % (st, eps_f, qs, sum(parts)))
                else:
                    cv.ok()
        la = (m.get("periods") or {}).get("last_actual")
        if la and bps is not None and cell_v(rm, "BPS", la) is not None and not close(bps, cell_v(rm, "BPS", la), 0.051):
            cv.fail("%s bps_latest %s ≠ BPS[%s] %s" % (st, bps, la, cell_v(rm, "BPS", la)))
        if _num(close_v):
            if _num(eps_f) and eps_f > 0:
                if not close(val.get("per_now"), close_v / eps_f, 0.0051):
                    cv.fail("%s per_now %s ≠ %s/%s" % (st, val.get("per_now"), close_v, eps_f))
                else:
                    cv.ok()
            if _num(bps) and bps > 0:
                if not close(val.get("pbr_now"), close_v / bps, 0.0051):
                    cv.fail("%s pbr_now %s ≠ %s/%s" % (st, val.get("pbr_now"), close_v, bps))
                else:
                    cv.ok()
        band, fv = val.get("per_band") or {}, val.get("fair_value_per") or {}
        if _num(eps_f) and eps_f > 0:
            for k in ("lo", "mid", "hi"):
                if _num(band.get(k)) and fv.get(k) != round(band[k] * eps_f, -2):
                    cv.fail("%s fair_value_per.%s %s ≠ round(%s×%s,-2)" % (st, k, fv.get(k), band[k], eps_f))
                else:
                    cv.ok()
        if _num(val.get("fair_pbr")) and _num(bps) and val.get("fair_value_pbr") != round(bps * val["fair_pbr"], -2):
            cv.fail("%s fair_value_pbr %s ≠ round(%s×%s,-2)" % (st, val.get("fair_value_pbr"), bps, val["fair_pbr"]))
        # fx
        afx = (m.get("assumptions") or {}).get("fx") or {}
        for key, byq in afx.items():
            if not isinstance(byq, dict):
                continue
            for q, v in byq.items():
                qd = (fx.get("quarters") or {}).get(q) or {}
                fd = (fx.get("forward") or {}).get(q) or {}
                exp = qd.get(key) if (_num(qd.get(key)) and not qd.get("partial")) else (fd.get(key) if _num(fd.get(key)) else qd.get(key))
                if not _num(exp):
                    continue
                if not close(v, exp, 0.0051):
                    cf.fail("%s fx %s %s 모델 %s ≠ fx.json %s" % (st, key, q, v, exp))
                else:
                    cf.ok()
        # contracts ↔ sls
        sls = tree.sls(st)
        if sls:
            mine = {r["rcp"] for r in tree.contracts if r.get("stock") == st}
            known = {c.get("rcp") for c in sls.get("contracts") or []} | {x.get("rcp") for x in (sls.get("skipped") or []) if isinstance(x, dict)} | set(sls.get("dropped_superseded") or [])
            if mine - known:
                cc.fail("%s contracts.json 에만 있는 계약 %d건 %s — sls 가 원장보다 오래됨" % (st, len(mine - known), sorted(mine - known)[:4]))
            else:
                cc.ok()
            if known - mine - {None}:
                cc.fail("%s sls 에만 있는 계약 %s" % (st, sorted(x for x in known - mine if x)[:4]))
            po = sls.get("post_origin") or {}
            d0 = (_seg_driver(m)[0].get("post_origin_excluded") or {})
            if po.get("n") is not None and d0 and d0.get("n") != po.get("n"):
                cc.fail("%s origin 이후 계약 건수 모델 %s ≠ sls %s" % (st, d0.get("n"), po.get("n")))
        # panel ↔ 매출조선신규
        if m.get("new_orders_included") and m.get("role") == "yard":
            base = ((tree.panel.get(st) or {}).get("scenarios") or {}).get("base") or {}
            pq = {x.get("quarter"): x.get("new_order_revenue") for x in base.get("quarterly") or []}
            for q in _fq(m):
                v = cell_v(rm, "매출조선신규", q)
                exp = pq.get(q)
                if v is None or not _num(exp):
                    continue
                if not close(v, exp / UNIT_DIV, 0.0051):
                    cn.fail("%s %s 매출조선신규 %s ≠ 패널 base %.2f" % (st, q, v, exp / UNIT_DIV))
                else:
                    cn.ok()


# ── 실행 ────────────────────────────────────────────────────

def run(root, stocks=None, xlsx=True, emulate=True):
    tree = Tree(root)
    rep = Report()
    all_stocks = tree.model_stocks()
    stocks = [s for s in all_stocks if not stocks or s in set(stocks)]
    check_fin_model(rep, tree, stocks)
    check_sls_model(rep, tree, stocks)
    check_summary(rep, tree, stocks)
    if xlsx:
        check_xlsx(rep, tree, stocks, emulate=emulate)
    check_pages(rep, tree, stocks)
    check_hub(rep, tree)
    check_sets(rep, tree)
    check_inputs(rep, tree, stocks)
    return rep


def main(argv=None):
    ap = argparse.ArgumentParser(description="한국조선 실적 모델 전 산출물 교차 정합 검사기(읽기 전용)")
    ap.add_argument("--root", default=os.path.dirname(HERE), help="argus/kship 디렉터리(기본: 이 파일의 상위)")
    ap.add_argument("--stocks", nargs="*", help="대상 종목(기본 전부)")
    ap.add_argument("--no-xlsx", action="store_true", help="xlsx 검사 생략")
    ap.add_argument("--no-emulate", action="store_true", help="kship_model_xlsx.verify 에뮬레이터(레인 코드) 생략")
    ap.add_argument("--json", help="보고 JSON 출력 경로")
    ap.add_argument("--max-detail", type=int, default=8, help="그룹별 실패 상세 출력 수")
    ap.add_argument("--strict", action="store_true", help="freshness 실패도 종료 코드 1")
    a = ap.parse_args(argv)
    rep = run(a.root, a.stocks, xlsx=not a.no_xlsx, emulate=not a.no_emulate)
    print("root %s · 종목 %d" % (os.path.abspath(a.root), len(a.stocks) if a.stocks else len(Tree(a.root).model_stocks())))
    print(rep.table(a.max_detail))
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(rep.as_dict(), f, ensure_ascii=False, indent=1)
    bad = rep.failed(("consistency", "freshness") if a.strict else ("consistency",))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
