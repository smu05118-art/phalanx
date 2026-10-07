#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""kship_model_xlsx — 모델 json(MODEL_SPEC 2-5) → 애널리스트 subQ 형 xlsx(MODEL_SPEC 2-6).

레퍼런스(사용자 subQ 모델 3개)의 골격을 그대로 따른다.
  README   시트 안내·단위·색 규약·가정 요약·연결 지도·출처·한계 — 모델 산출값이며 목표주가·추천이 아님
  변수      가정 시계열 — 환율 8행(fx.json; 추정 분기 노란 셀) · 추정 드라이버(모델값에서 역산, 노란 셀) ·
           스칼라 가정(세율·판관비율·지배주주비중·이자율·종가·PER 밴드·COE·외화 순노출·헤지비율 — 노란 셀. 추정 전 분기의
           드라이버 셀이 `=$D$n` 으로 참조하므로 하나를 바꾸면 전 분기가 따라간다)
  BS연결    FnGuide 계정명 × 기간 덤프(백만원). fin json(2-1)이 있으면 전 계정, 없으면 모델 행이 매핑되는 계정만
  BS별도    같은 구조(별도). fin 없으면 헤더만
  subQ     분기 서브모델(억원). 확정치는 `=IFERROR(VLOOKUP($D행,BS연결,MATCH(E$1,BS연결H,0),0)/100,"")` 진짜 수식
           (외환손익·파생상품손익은 BS 계정 여럿의 합·차, 기타금융손익은 잔차 항등식, OPM·EPS·BPS·PER·PBR 은 파생 수식),
           추정치는 `변수` 가정 셀을 참조하는 수식 — 매출 = 전년동기 × (1+성장률), 법인세 = MAX(세전,0) × 세율,
           이자손익 = (전분기 이자발생자산 × 자산이자율 − 전분기 총차입금 × 부채이자율)/4, 환관련손익 = 외화 순노출 × Δ원/달러(기말)/100,
           조선사 매출조선 = SLS 원화 합/100 + 원장 밖 잔고 소진 + 매출조선신규, 금융손익·세전·순이익·자산·EBITDA 는 항등식
  SLS      선표(sls json 2-4 가 있을 때만) — 요약 행(값) + 계약별 스케줄(백만$) × (헤지비율 × 헤지환율 + (1−헤지비율) × 건조시점 환율
           [변수 원/달러(평균)]) → 원화(백만원) 수식 블록 → 합계가 subQ 매출조선에 연결(환율 forward 를 바꾸면 매출이 재계산)
  시나리오   신규수주 보수/기준/낙관(forecast_panel, 모델 scenarios 가 있을 때만) — 매출액·영업이익 = subQ 기준값 ± (시나리오 신규 − base 신규) × OPM
  분기·연간예상  subQ 를 이름으로 VLOOKUP 한 보고서 표
  TP       PER 밴드·PBR(ROE/COE) 적정가치 구간 — 모델 산출값, 목표주가·추천 아님
  외화      환율 QoQ · 외화 순노출 × Δ기말환율 → 환관련손익 추정

헤더 규약: 행1 분기 라벨('4Q21','1Q22',… 연간은 2022 숫자), 행4 'YYYY.MM'/'YYYY.12A', 데이터는 E열부터,
4분기 뒤 연간 1열. 이름정의 SUBQH·BS연결H·BS별도H·변수H(행1), BS연결·BS별도(계정명 열 C 부터 전체), 환율.
추정 열은 연노랑 배경, 실적/추정 경계는 굵은 테두리, 숫자 서식 #,##0. 추정 셀 메모 = 모델 json 의 basis.
노란 셀 = 가정 입력, 연녹 셀 = 스칼라 가정을 참조하는 수식(숫자로 덮어쓰면 그 분기만 바뀐다).
확정 셀은 BS 시트에 그 기간 값이 있고 모델값과 맞을 때만 VLOOKUP(빈 셀 VLOOKUP 은 0 으로 보이고, 비용 부호가 반전된 회사는
BS 값 ≠ 모델값), 없으면 모델값 직접 + 메모, 둘 다 없으면 빈 칸.
결정론: 같은 모델·fin·fx·prices·sls 면 바이트 동일 — BS 행3 UPDATE 스탬프·문서 속성·zip 엔트리 시각 전부 모델 built_at 날짜로 고정.

원칙: 출처 없는 숫자를 만들지 않는다 — 추정 드라이버 값은 전부 모델 json 의 v 에서 역산(implied)하고 그 사실을 셀 메모에 적는다.
스칼라(세율·판관비율·지배주주비중·이자율)는 추정 전 분기를 한 값으로 재현할 수 있을 때만(±0.05억) 연결하고, 아니면 분기별 역산값으로
남기고 README 에 '연결 안 됨' 으로 적는다. 선표 연결은 계약별 재구성이 sls 저장값과 맞을 때만(스케줄 소수 3자리 반올림 허용).
검증(--verify)은 추정 전 분기·실적 수식 셀을 자체 계산기로 재계산해 모델값과 대조하고, 세율·OPM·환율 forward·이자율·신규수주 셀을
실제로 바꿔 종속 셀이 기대만큼 움직이는지(sensitivity) 확인한다. LibreOffice/Excel 은 없다 — Excel 실개봉은 미확인.

    python3 kship_model_xlsx.py --model assets/models/010140.json --verify
    python3 kship_model_xlsx.py --all --verify        # assets/models/*.json → ../models/<stock>_model.xlsx (요약표 출력)
    python3 kship_model_xlsx.py --sample --verify     # tests/fixtures/model_sample.json 으로 개발·점검
"""
import argparse
import datetime
import json
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from kship_lib import ASSETS, KSHIP, opm_table_text

MODELS_DIR = os.path.join(ASSETS, "models")
FIN_DIR = os.path.join(ASSETS, "fin")
SLS_DIR = os.path.join(ASSETS, "sls")
OUT_DIR = os.path.join(KSHIP, "models")
SAMPLE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tests", "fixtures", "model_sample.json")

FIRST_COL = 5            # E — 데이터 첫 열(모든 시계열 시트 공통)
HEADER_ROW, DATE_ROW, DATA_ROW = 1, 4, 5
SIZE_LIMIT = 20 * 1024 * 1024
TOL_EOK = 0.05           # 억원 행 재계산 허용(모델 셀 0.01 반올림 × 연쇄)
TOL_SLS_M = 30.0         # 선표 재구성 허용(백만원) — 계약 스케줄이 소수 3자리라 Σ(0.0005×환율) 만큼 어긋난다

# 모델 행 key → FnGuide 계정명(BS 시트 VLOOKUP 키). 없는 key 는 수식 파생 또는 값 직접 기입.
ROW_ACCOUNT = {
    "매출액": "매출액(수익)", "매출원가": "매출원가", "매출총이익": "매출총이익", "판관비": "판관비",
    "영업이익": "영업이익", "기타영업손익": "기타영업손익", "금융손익": "금융손익", "금융수익": "금융수익",
    "금융비용": "금융비용", "이자수익": "이자수익", "이자비용": "이자비용", "이자손익": "이자손익",
    "기타영업외손익": "기타영업외손익", "지분법손익": "종속기업,공동지배기업및관계기업관련손익",
    "세전이익": "법인세비용차감전계속사업이익", "법인세비용": "법인세비용", "당기순이익": "당기순이익",
    "지배주주순이익": "(지배주주지분)당기순이익", "비지배주주순이익": "(비지배주주지분)당기순이익",
    "자산총계": "자산총계", "부채총계": "부채총계", "자본총계": "자본총계", "지배주주지분": "지배주주지분",
    "비지배주주지분": "비지배주주지분", "총차입금": "총차입금", "순차입금": "순차입금",
    "현금및현금성자산": "현금및현금성자산", "단기금융자산": "단기금융자산", "이자발생자산": "이자발생자산",
    "감가상각비": "감가상각비", "CAPEX": "CAPEX", "재고자산": "재고자산", "충당부채": "충당부채",
    "매출채권": "매출채권및기타채권", "매입채무": "매입채무및기타채무", "선수금": "선수금",
    "자본금": "자본금", "이익잉여금": "이익잉여금", "영업활동현금흐름": "영업활동으로인한현금흐름",
    "중단사업이익": "중단사업이익",
}
# 실적 구간 복합 계정 — BS 시트 계정 여럿의 합·차(T4 영업외 세부 행). 전 계정이 그 분기에 있을 때만 수식
COMPOSITE_ACCOUNTS = {
    "외환손익": [("+", "외환차익"), ("-", "외환차손"), ("+", "외화환산이익"), ("-", "외화환산손실")],
    "파생상품손익": [("+", "파생상품이익"), ("-", "파생상품손실")],
}
FIN_DETAIL_KEYS = ("이자손익", "외환손익", "파생상품손익", "기타금융손익")     # 금융손익 = Σ 세부(기타금융손익은 잔차)
# 시점(BS) 성격 — 연간 열은 4Q 값, 추정 기본 드라이버는 QoQ
STOCK_KEYS = {"자산총계", "부채총계", "자본총계", "지배주주지분", "비지배주주지분", "총차입금", "순차입금",
              "현금및현금성자산", "단기금융자산", "이자발생자산", "재고자산", "충당부채", "매출채권", "매입채무",
              "선수금", "자본금", "이익잉여금", "BPS", "주식수"}
STOCK_GROUPS = {"재무상태", "BS", "자본"}
# 부호가 흔들려 비율 역산이 무의미한 행 — 절대값 가정(노란 셀에 모델값 그대로)
ABS_KEYS = {"금융손익", "기타영업외손익", "기타영업손익", "지분법손익", "이자손익", "일회성", "환관련손익", "파생상품손익",
            "외환손익", "기타금융손익", "DPS", "감가상각비", "중단사업이익"}
# 억원이 아닌 행
PER_SHARE_KEYS = {"EPS", "BPS", "DPS"}
# 드라이버 행(ratio) ↔ 변수 3절 스칼라 — 추정 전 분기를 한 값으로 재현할 수 있으면 드라이버 셀이 스칼라를 참조한다
SCALAR_LINK = {"판관비": "판관비율", "지배주주순이익": "지배주주비중"}

FX_ROWS = [("USDKRW_avg", " 원/달러(평균)"), ("USDKRW_end", " 원/달러(기말)"),
           ("JPY100KRW_avg", " 원/100Y(평균)"), ("JPY100KRW_end", " 원/100Y(기말)"),
           ("EURKRW_avg", " W/Euro(평균)"), ("EURKRW_end", " W/Euro(기말)"),
           ("CNYKRW_avg", " 원/중국원(평균)"), ("CNYKRW_end", " 원/중국원(기말)")]
PRICE_ROWS = [("close_end", "종가(기말)"), ("high", "종가(최고)"), ("low", "종가(최저)"), ("avg", "종가(평균)")]
SCN_CASES = ("existing_only", "conservative", "base", "optimistic")
SCN_LABEL = {"existing_only": "기존 잔고만(신규수주 0)", "conservative": "보수", "base": "기준(base — subQ 행에 포함)", "optimistic": "낙관"}

# ── 스타일 ───────────────────────────────────────────────────
F_BOLD = Font(bold=True)
F_HEAD = Font(bold=True, color="FFFFFF")
F_NOTE = Font(italic=True, color="666666", size=9)
F_KEY = Font(color="1F4E79")
FILL_HEAD = PatternFill("solid", fgColor="44546A")
FILL_EST = PatternFill("solid", fgColor="FFF9DB")       # 추정 열 연노랑
FILL_ANNUAL = PatternFill("solid", fgColor="EDEDED")    # 연간 열 회색
FILL_INPUT = PatternFill("solid", fgColor="FFF2AB")     # 사용자 가정 입력 셀
FILL_LINK = PatternFill("solid", fgColor="E2EFDA")      # 스칼라 가정을 참조하는 수식 셀(연녹) — 숫자로 덮어쓰면 그 분기만 바뀐다
THIN, MED = Side(style="thin", color="999999"), Side(style="medium", color="000000")
NF_INT, NF_1, NF_PCT, NF_FX, NF_PCT2 = "#,##0", "#,##0.0", "0.0%", "#,##0.00", "0.00%"


# ── 기간 헬퍼 ───────────────────────────────────────────────

def qlabel(q):
    """'2021Q4' → '4Q21' (레퍼런스 행1 규약)."""
    return "%sQ%s" % (q[5], q[2:4])


def qdate(q):
    """'2021Q4' → '2021.12' (레퍼런스 행4 규약, FnGuide 열키)."""
    return "%s.%02d" % (q[:4], int(q[5]) * 3)


def build_columns(quarters, annual):
    """분기 뒤에 그 해의 연간 열을 끼운 열 배치. 각 항목 {kind,key,label,date,col}."""
    annual = [str(a) for a in (annual or [])]
    cols, c = [], FIRST_COL
    years_seen = []
    for q in quarters:
        cols.append({"kind": "q", "key": q, "label": qlabel(q), "date": qdate(q), "col": c})
        c += 1
        y = q[:4]
        if q.endswith("Q4") or q == quarters[-1]:
            if y in annual and y not in years_seen:
                cols.append({"kind": "a", "key": y, "label": int(y), "date": y + ".12A", "col": c})
                years_seen.append(y)
                c += 1
    for y in annual:              # 분기가 하나도 없는 해(드묾)는 뒤에 붙인다
        if y not in years_seen:
            cols.append({"kind": "a", "key": y, "label": int(y), "date": y + ".12A", "col": c})
            years_seen.append(y)
            c += 1
    return cols


class Layout:
    """열 배치 + 실적/추정 경계. 모든 시계열 시트가 같은 열 위치를 쓴다(변수!E5 ↔ subQ!E5)."""

    def __init__(self, model):
        p = model["periods"]
        self.quarters = list(p["quarters"])
        self.annual = [str(a) for a in p.get("annual", [])]
        self.last_actual = p["last_actual"]
        self.cols = build_columns(self.quarters, self.annual)
        self.by_key = {c["key"]: c for c in self.cols}
        self.last_col = self.cols[-1]["col"] if self.cols else FIRST_COL
        self.qcols = [c for c in self.cols if c["kind"] == "q"]
        self.acols = [c for c in self.cols if c["kind"] == "a"]
        self.boundary_col = self.by_key[self.last_actual]["col"] if self.last_actual in self.by_key else None
        self.last_year = int(self.last_actual[:4])
        self.est_quarters = [q for q in self.quarters if q > self.last_actual]

    def col_of(self, key):
        return self.by_key[key]["col"]

    def letter(self, key):
        return get_column_letter(self.col_of(key))

    def is_est(self, c):
        """추정 음영 대상: last_actual 이후 분기, last_actual 연도 이후(포함) 연간."""
        if c["kind"] == "q":
            return c["key"] > self.last_actual
        return int(c["key"]) >= self.last_year

    def prev_q(self, q):
        i = self.quarters.index(q)
        return self.quarters[i - 1] if i > 0 else None

    def prev_year_q(self, q):
        py = "%d%s" % (int(q[:4]) - 1, q[4:])
        return py if py in self.by_key else None

    def year_quarters(self, y):
        return [q for q in self.quarters if q[:4] == y]


# ── 모델 접근 ───────────────────────────────────────────────

def cell_of(row, q):
    return (row.get("q") or {}).get(q) or {}


def val_of(row, q):
    return cell_of(row, q).get("v")


def is_stock(row):
    return row["key"] in STOCK_KEYS or row.get("group") in STOCK_GROUPS


def bs_sheet_for(row, q=None):
    src = (cell_of(row, q).get("src") if q else None) or row.get("src") or ""
    return "BS별도" if ".sep." in str(src) or str(src).startswith("sep") else "BS연결"


def tol_for(key, unit, v, ttm_eps=None, bps=None, base=None):
    """재계산 허용오차 — 억원 0.05(또는 1e-4 상대), 주당(원) 0.5(또는 1e-3 상대: 주식수 백만주 3자리 반올림 증폭),
    비율(%) 6e-5(모델 4자리 반올림) + 0.011/|분모 매출액|(분자·분모 셀 0.01억 반올림),
    배 0.006(2자리) + 분모 셀(TTM EPS·BPS)의 자체 허용오차가 배수로 전파되는 상대항, 백만주 0.001."""
    if key in PER_SHARE_KEYS or unit == "원":
        return max(0.5, abs(v) * 1e-3)
    if unit == "%":
        return 6e-5 + (0.011 / abs(base) if base else 0.0)
    if unit == "배":
        if key == "PER" and ttm_eps:
            return 0.006 + abs(v) * (4 * 0.5) / abs(ttm_eps)            # EPS 셀 4개 × 각 허용 0.5원
        if key == "PBR" and bps:
            return 0.006 + abs(v) * max(0.5, abs(bps) * 1e-3) / abs(bps)  # BPS 셀 허용오차 전파
        return 0.006
    if unit == "백만주":
        return 0.001
    return max(TOL_EOK, abs(v) * 1e-4)


def tol_args(rows_by_key, quarters, q):
    """tol_for 의 보조 인자 — 그 분기의 모델 TTM EPS·BPS·매출액."""
    return {"ttm_eps": ttm_eps_of(rows_by_key, quarters, q),
            "bps": val_of(rows_by_key["BPS"], q) if "BPS" in rows_by_key else None,
            "base": val_of(rows_by_key["매출액"], q) if "매출액" in rows_by_key else None}


def ttm_eps_of(rows_by_key, quarters, q):
    """모델 EPS(v) 4분기 합 — PER 허용오차 계산용. 하나라도 없으면 None."""
    if "EPS" not in rows_by_key or q not in quarters:
        return None
    i = quarters.index(q)
    if i < 3:
        return None
    vals = [val_of(rows_by_key["EPS"], qq) for qq in quarters[i - 3:i + 1]]
    return sum(vals) if all(x is not None for x in vals) else None


# ── 추정 드라이버 계획 ───────────────────────────────────────

def _est_cells(row):
    return {q: c for q, c in (row.get("q") or {}).items() if c.get("kind") == "estimate" and c.get("v") is not None}


def _consistent(by, target, expr, tol=TOL_EOK):
    """target 의 추정 셀마다 Σ sign×part == target(±tol). 추정 셀이 없거나 part 가 빠진 분기가 있으면 False — 항등식 채택 조건."""
    if target not in by:
        return False
    qs = _est_cells(by[target])
    if not qs:
        return False
    for q, c in qs.items():
        s = 0.0
        for sign, k in expr:
            v = val_of(by[k], q) if k in by else None
            if v is None:
                return False
            s += v if sign == "+" else -v
        if abs(s - c["v"]) > tol:
            return False
    return True


def plan_drivers(rows):
    """행마다 추정 수식 종류를 정한다. 반환 {key: {"type", ...}}.
    identity(다른 subQ 행 합·차) · divide(비율) · per_share · valuation(종가/TTM EPS·종가/BPS) · tax(MAX(세전,0)×세율) ·
    interest(전분기 잔액 × 이자율) · ratio_parent(부모 마진 공유) 은 가정 셀이 없고, yoy·qoq·ratio·abs·roll 은 `변수` 드라이버 행을 하나씩 가진다.
    세그먼트가 둘 이상이고 합이 맞으면 매출액·영업이익은 세그먼트 합(레퍼런스 사업부 빌드업), 단일 '전사' 세그먼트면 전사 = 매출액.
    Builder 가 데이터를 보고 tax·interest 를 ratio·abs 로 내리거나 매출<seg> 를 sls_link·환관련손익을 fxpnl 로 올릴 수 있다."""
    keys = [r["key"] for r in rows]
    by = {r["key"]: r for r in rows}
    has = by.__contains__
    segs = [k for k in keys if k.startswith("매출") and k not in ("매출액", "매출원가", "매출총이익", "매출채권", "매출전사") and not k.endswith("신규")]
    multi = len(segs) >= 2 and has("매출액") and _consistent(by, "매출액", [("+", k) for k in segs])
    op_segs = ["OP" + k[2:] for k in segs]
    multi_op = multi and has("영업이익") and all(has(k) for k in op_segs) and _consistent(by, "영업이익", [("+", k) for k in op_segs])
    plan = {}
    for r in rows:
        k = r["key"]
        if k == "매출액":
            d = ({"type": "identity", "expr": [("+", s) for s in segs], "note": "매출액 = " + " + ".join(segs)} if multi
                 else {"type": "yoy", "label": "매출액 YoY"})
        elif k == "매출전사" and has("매출액") and _consistent(by, k, [("+", "매출액")]):
            d = {"type": "identity", "expr": [("+", "매출액")], "note": "매출전사 = 매출액(단일 부문)"}
        elif k == "OP전사" and has("영업이익") and _consistent(by, k, [("+", "영업이익")]):
            d = {"type": "identity", "expr": [("+", "영업이익")], "note": "OP전사 = 영업이익(단일 부문)"}
        elif k == "영업이익" and has("매출액"):
            d = ({"type": "identity", "expr": [("+", s) for s in op_segs], "note": "영업이익 = " + " + ".join(op_segs)} if multi_op
                 else {"type": "ratio", "base": "매출액", "label": "영업이익률(OPM)"})
        elif k == "OPM" and has("영업이익") and has("매출액"):
            d = {"type": "divide", "num": "영업이익", "den": "매출액", "note": "OPM = 영업이익 / 매출액"}
        elif k == "판관비" and has("매출액"):
            d = {"type": "ratio", "base": "매출액", "label": "판관비율(/매출)"}
        elif k == "매출총이익" and has("영업이익") and has("판관비"):
            d = {"type": "identity", "expr": [("+", "영업이익"), ("+", "판관비")] + ([("-", "기타영업손익")] if has("기타영업손익") else []),
                 "note": "매출총이익 = 영업이익 + 판관비" + (" − 기타영업손익" if has("기타영업손익") else "")}
        elif k == "매출원가" and has("매출액") and has("매출총이익"):
            d = {"type": "identity", "expr": [("+", "매출액"), ("-", "매출총이익")], "note": "매출원가 = 매출액 − 매출총이익"}
        elif k == "금융손익":
            parts = [x for x in FIN_DETAIL_KEYS if has(x)]
            d = ({"type": "identity", "expr": [("+", x) for x in parts], "note": "금융손익 = " + " + ".join(parts) + "(T4 세부 행 항등식)"}
                 if parts and _consistent(by, k, [("+", x) for x in parts]) else {"type": "abs", "label": "금융손익 (가정, 억원)"})
        elif k == "이자손익" and has("이자발생자산"):
            d = {"type": "interest", "ia": "이자발생자산", "debt": "총차입금" if has("총차입금") else None,
                 "note": "이자손익 = (전분기 이자발생자산 × 자산이자율 − 전분기 총차입금 × 부채이자율) / 4"}
        elif k == "세전이익" and has("영업이익"):
            # 모델(kship_model.py) basis 와 같은 순서 — 환관련손익(yard 외화 순노출 × Δ환율) 을 빼면 법인세→NI→EPS→BPS 까지 연쇄로 어긋난다
            parts = [("+", "영업이익")] + [("+", x) for x in ("금융손익", "기타영업외손익", "지분법손익", "환관련손익") if has(x)]
            d = {"type": "identity", "expr": parts, "note": "세전이익 = " + " + ".join(x for _, x in parts)}
        elif k == "법인세비용" and has("세전이익"):
            d = {"type": "tax", "base": "세전이익", "label": "법인세율(/세전) — 법인세 = MAX(세전,0) × 세율"}
        elif k == "당기순이익" and has("세전이익") and has("법인세비용"):
            d = {"type": "identity", "expr": [("+", "세전이익"), ("-", "법인세비용")], "note": "당기순이익 = 세전이익 − 법인세비용"}
        elif k == "지배주주순이익" and has("당기순이익"):
            d = {"type": "ratio", "base": "당기순이익", "label": "지배주주 비중(/당기순이익)"}
        elif k == "EPS" and has("지배주주순이익") and has("주식수"):
            d = {"type": "per_share", "num": "지배주주순이익", "den": "주식수", "note": "EPS = 지배주주순이익(억원)×100 / 유통주식수(백만주)"}
        elif k == "BPS" and has("지배주주지분") and has("주식수"):
            d = {"type": "per_share", "num": "지배주주지분", "den": "주식수", "note": "BPS = 지배주주지분(억원)×100 / 유통주식수(백만주)"}
        elif k == "PER" and has("EPS"):
            d = {"type": "valuation", "den": "ttm_eps", "note": "PER = 종가(추정 구간: 변수 종가 · 실적 구간: 분기 평균 종가) / TTM EPS(4분기 합)"}
        elif k == "PBR" and has("BPS"):
            d = {"type": "valuation", "den": "BPS", "note": "PBR = 종가(추정 구간: 변수 종가 · 실적 구간: 분기 평균 종가) / BPS"}
        elif k == "EBITDA" and has("영업이익") and has("감가상각비"):
            d = {"type": "identity", "expr": [("+", "영업이익"), ("+", "감가상각비")], "note": "EBITDA = 영업이익 + 감가상각비"}
        elif k == "자본총계" and has("당기순이익"):
            d = {"type": "roll", "flow": "당기순이익", "label": "자본총계 기타변동(배당 등, 억원)"}
        elif k == "지배주주지분" and has("지배주주순이익"):
            d = {"type": "roll", "flow": "지배주주순이익", "label": "지배주주지분 기타변동(배당 등, 억원)"}
        elif k == "자산총계" and has("부채총계") and has("자본총계"):
            d = {"type": "identity", "expr": [("+", "부채총계"), ("+", "자본총계")], "note": "자산총계 = 부채총계 + 자본총계"}
        elif k.startswith("매출") and k.endswith("신규"):
            d = {"type": "abs", "label": "%s (forecast_panel base 신규수주 매출, 억원 — 가정)" % k}
        elif k.startswith("OP") and k.endswith("신규") and has("매출" + k[2:]):
            d = {"type": "ratio_parent", "base": "매출" + k[2:], "parent": "OP" + k[2:-2], "note": "%s = 매출%s × %s 마진(같은 타겟 OPM)" % (k, k[2:], "OP" + k[2:-2])}
        elif k in segs:
            if k == "매출연결조정" and has("매출종속사"):
                d = {"type": "ratio", "base": "매출종속사", "label": "매출연결조정 비율(/매출종속사 — 내부거래·잔차)"}
            elif multi:
                d = {"type": "yoy", "label": "%s YoY" % k}
            else:
                d = {"type": "ratio", "base": "매출액", "label": "%s 비중(/매출액)" % k}
        elif k.startswith("OP") and has("매출" + k[2:]):
            d = {"type": "ratio", "base": "매출" + k[2:], "label": "%s 마진(/매출%s)" % (k, k[2:])}
        elif k in ABS_KEYS:
            d = {"type": "abs", "label": "%s (가정, %s)" % (k, r.get("unit") or "억원")}
        elif is_stock(r):
            d = {"type": "qoq", "label": "%s QoQ" % k}
        else:
            d = {"type": "yoy", "label": "%s YoY" % k}
        plan[k] = d
    for k, d in plan.items():            # 부모 마진 행이 ratio 가 아니면(식별식 등) 자기 비율 드라이버로
        if d["type"] == "ratio_parent" and plan.get(d["parent"], {}).get("type") != "ratio":
            plan[k] = {"type": "ratio", "base": d["base"], "label": "%s 마진(/%s)" % (k, d["base"])}
    return plan


def implied_param(row, q, drv, rows_by_key, lay, base_of=None):
    """추정 셀의 가정값을 모델 v 에서 역산. 기준값이 없으면 None(→ 값 직접 기입 폴백).
    base_of(row, q) 는 기준 셀이 시트에서 실제로 갖게 될 값 — 확정 구간 VLOOKUP 이면 BS 값/100(반올림 전).
    모델 v 는 0.01 억원 반올림이라 기준이 0 에 가까운 행(순차입금 부호 전환 등)은 QoQ 역산이 크게 흔들린다.
    기준도 값도 0 이면 0(수식 prev×(1+0)=0 이 모델과 같다 — 총차입금 0 회사·정의상 0 행)."""
    base_of = base_of or val_of
    v = val_of(row, q)
    if v is None:
        return None
    t = drv["type"]
    if t == "abs":
        return v
    if t in ("yoy", "qoq"):
        pq = lay.prev_year_q(q) if t == "yoy" else lay.prev_q(q)
        b = base_of(row, pq) if pq else None
        if b:
            return v / b - 1
        return 0.0 if (b == 0 and v == 0) else None
    if t in ("ratio", "tax"):
        b = val_of(rows_by_key[drv["base"]], q)
        if b:
            return v / b
        return 0.0 if (b == 0 and v == 0) else None
    if t == "roll":
        pq = lay.prev_q(q)
        b = base_of(row, pq) if pq else None
        f = val_of(rows_by_key[drv["flow"]], q)
        return (b + f - v) if (b is not None and f is not None) else None
    return None


def _fit_two(pts):
    """z = a·x − b·y 의 최소자승 (x, y). pts = [(a, b, z)]. b 가 전부 0 이면 y = 0. 특이하면 None."""
    saa = sum(a * a for a, _, _ in pts)
    sab = sum(a * b for a, b, _ in pts)
    sbb = sum(b * b for _, b, _ in pts)
    saz = sum(a * z for a, _, z in pts)
    sbz = sum(b * z for _, b, z in pts)
    if sbb == 0:
        return (saz / saa, 0.0) if saa > 0 else None
    det = saa * sbb - sab * sab
    if abs(det) <= 1e-12 * (saa * sbb):
        return None
    x = (saz * sbb - sab * sbz) / det
    y_neg = (saa * sbz - sab * saz) / det          # 계수 of b (= −y)
    return x, -y_neg


# ── 워크북 조립 ─────────────────────────────────────────────

class Builder:
    def __init__(self, model, fin=None, sls=None, fx=None, prices=None):
        self.m = model
        self.fin = fin
        self.sls = sls or model.get("sls")
        self.fx = fx
        self.prices = prices
        self.lay = Layout(model)
        self.rows = list(model["rows"])
        self.rows_by_key = {r["key"]: r for r in self.rows}
        self.plan = plan_drivers(self.rows)
        self.subq_row = {}       # key → subQ 행 번호
        self.drv_row = {}        # key → 변수 드라이버 행 번호
        self.scalar_row = {}     # 스칼라 가정 이름 → 변수 행 번호
        self.fx_row = {}         # 환율 metric → 변수 행 번호
        self.price_row = {}      # 종가 metric → 변수 행 번호
        self.links = {}          # 연결 판정(세율·판관비율·지배주주비중·이자율·환관련손익·선표) — README·verify 가 읽는다
        self.fit = {}            # 스칼라 역산값
        self.sls_link = None     # 선표 연결 정보(재구성 원화·포함 usd·잔고 소진)
        self.sls_geo = None
        # BS 값은 수식 판정(VLOOKUP 쓸지)·역산 기준·이자율 적합에 먼저 필요하다
        self.bs_vals = {"BS연결": self._bs_values("cons"), "BS별도": self._bs_values("sep")}
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.stats = {}
        self._resolve_links()

    # ── 공통 ──
    def _header(self, ws, name_label, cols=None, unit_note=None):
        """행1 라벨('4Q21'/2021 숫자)·행3 인덱스·행4 'YYYY.MM'/'YYYY.12A'. cols 의 col 번호를 그대로 쓴다."""
        cols = cols if cols is not None else self.lay.cols
        ws.cell(HEADER_ROW, 3, name_label).font = F_NOTE
        if unit_note:
            ws.cell(HEADER_ROW, 2, unit_note).font = F_NOTE
        for i, c in enumerate(cols):
            col = c["col"]
            h = ws.cell(HEADER_ROW, col, c["label"])
            h.font = F_HEAD
            h.fill = FILL_HEAD
            h.alignment = Alignment(horizontal="center")
            ws.cell(3, col, i + 1).font = F_NOTE
            d = ws.cell(DATE_ROW, col, c["date"])
            d.font = F_NOTE
            d.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = 11
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 9
        ws.column_dimensions["C"].width = 26
        ws.column_dimensions["D"].width = 30
        ws.freeze_panes = ws.cell(DATA_ROW, FIRST_COL)

    def _shade(self, ws, last_row, cols=None):
        """추정 열 연노랑·연간 열 회색·경계 굵은 테두리(행1~last_row)."""
        cols = cols if cols is not None else self.lay.cols
        for c in cols:
            fill = FILL_EST if self.lay.is_est(c) else (FILL_ANNUAL if c["kind"] == "a" else None)
            for r in range(DATA_ROW, last_row + 1):
                cell = ws.cell(r, c["col"])
                if fill is not None and cell.fill.fgColor.rgb in (None, "00000000"):
                    cell.fill = fill
                if c["kind"] == "a":
                    cell.font = F_BOLD
        if self.lay.boundary_col:
            for r in range(HEADER_ROW, last_row + 1):
                cell = ws.cell(r, self.lay.boundary_col)
                cell.border = Border(right=MED, left=cell.border.left, top=cell.border.top, bottom=cell.border.bottom)

    def _name(self, name, ref):
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)

    def built_date(self):
        """모델 built_at 의 날짜 — BS 시트 UPDATE 스탬프·문서 속성·zip 시각을 이 날짜로 고정(같은 입력 → 같은 바이트)."""
        m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(self.m.get("built_at") or ""))
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else datetime.date(2026, 1, 1)

    def _bs_get(self, sheet, acct, key):
        """BS 시트에 (계정, 열키) 값이 있으면 백만원 값, 없으면 None."""
        return ((self.bs_vals.get(sheet) or {}).get(acct) or {}).get(key)

    def _sheet_base(self, row, q):
        """역산 기준값 = 그 셀이 시트에서 실제로 계산할 값. 확정 구간 VLOOKUP 셀이면 BS 값/100(반올림 전), 아니면 모델 v."""
        acct = ROW_ACCOUNT.get(row["key"])
        mc = cell_of(row, q)
        if q <= self.lay.last_actual and acct and (mc.get("kind") == "actual" or not mc):
            v = self._bs_get(bs_sheet_for(row, q), acct, q)
            if v is not None and (mc.get("v") is None or abs(v / 100 - mc["v"]) <= 0.011 + abs(mc["v"]) * 1e-6):
                return v / 100
        return val_of(row, q)

    def _v(self, k, q):
        return val_of(self.rows_by_key[k], q) if k in self.rows_by_key else None

    # ── 환율 값(변수 시트가 갖게 될 값) ──
    def _fx_value(self, metric, c):
        """(값, 메모) — 모델 assumptions.fx(추정 가정) → fx.json quarters(실측) → fx.json forward(가정) → annual."""
        fx_model = ((self.m.get("assumptions") or {}).get("fx") or {})
        fxq = (self.fx or {}).get("quarters") or {}
        fxf = (self.fx or {}).get("forward") or {}
        fxa = (self.fx or {}).get("annual") or {}
        if c["kind"] == "q":
            q = c["key"]
            if isinstance(fx_model.get(metric), dict) and fx_model[metric].get(q) is not None:
                return fx_model[metric][q], "모델 assumptions.fx — " + str(fx_model.get("basis") or "")
            if q in fxq and fxq[q].get(metric) is not None:
                return fxq[q][metric], ("fx.json quarters(partial — 분기 진행 중 실측 평균)" if fxq[q].get("partial") else None)
            if isinstance(fxf.get(q), dict) and fxf[q].get(metric) is not None:
                return fxf[q][metric], "fx.json forward(%s) — 가정" % fxf.get("method", "")
            return None, None
        a = fxa.get(c["key"]) or {}
        return a.get(metric), None

    def _fx_series(self, metric, q):
        return self._fx_value(metric, self.lay.by_key[q])[0] if q in self.lay.by_key else None

    # ── 연결 판정(스칼라·이자율·환관련·선표) ──
    def _resolve_links(self):
        lay, asm, v = self.lay, (self.m.get("assumptions") or {}), self._v
        fq = lay.est_quarters
        # 세율 — 법인세 = MAX(세전,0) × 세율 을 추정 전 분기에서 맞추는 세율 하나(Σ법인세 / Σmax(세전,0))가 있으면 연결
        d = self.plan.get("법인세비용")
        if d and d["type"] == "tax":
            pts = [(max(v("세전이익", q), 0.0), v("법인세비용", q)) for q in fq if v("세전이익", q) is not None and v("법인세비용", q) is not None]
            sp = sum(a for a, _ in pts)
            rate = (sum(b for _, b in pts) / sp) if sp > 0 else None
            err = max(abs(a * rate - b) for a, b in pts) if rate is not None else None
            if rate is not None and err <= TOL_EOK:
                self.fit["세율"] = rate
                d["scalar"] = "세율"
                self.links["세율"] = {"linked": True, "fit": rate, "model_assumption": asm.get("tax_rate"), "n": len(pts), "max_err": err}
            else:
                self.plan["법인세비용"] = {"type": "ratio", "base": "세전이익", "label": "법인세율(/세전)"}
                self.links["세율"] = {"linked": False, "n": len(pts), "max_err": err,
                                    "reason": "추정 법인세 없음" if not pts else "MAX(세전,0)×단일 세율로 재현 불가(지주 자회사 합산 등) → 분기별 역산"}
        # 판관비율·지배주주비중 — ratio 드라이버를 스칼라 하나로(Σ분자/Σ분모 가 전 분기 ±0.05억)
        for k, sc in SCALAR_LINK.items():
            d = self.plan.get(k)
            if not d or d["type"] != "ratio":
                continue
            pts = [(v(d["base"], q), v(k, q)) for q in fq if v(d["base"], q) is not None and v(k, q) is not None]
            sb = sum(a for a, _ in pts)
            ratio = (sum(b for _, b in pts) / sb) if sb else None
            err = max(abs(a * ratio - b) for a, b in pts) if ratio is not None else None
            if ratio is not None and err <= TOL_EOK:
                self.fit[sc] = ratio
                d["scalar"] = sc
                self.links[sc] = {"linked": True, "fit": ratio, "n": len(pts), "max_err": err}
            else:
                self.links[sc] = {"linked": False, "n": len(pts), "max_err": err, "reason": "추정 없음" if not pts else "단일 비율로 재현 불가 → 분기별 역산"}
        # 이자율 — 이자손익(q) = (이자발생자산(q−1) × r_a − 총차입금(q−1) × r_d) / 4, 두 변수 최소자승(모델 assumptions 는 4자리 반올림이라 쓰지 않음)
        d = self.plan.get("이자손익")
        if d and d["type"] == "interest":
            pts = []
            for q in fq:
                iv = v("이자손익", q)
                pq = lay.prev_q(q)
                if iv is None or pq is None:
                    continue
                ia = self._sheet_base(self.rows_by_key[d["ia"]], pq)
                de = self._sheet_base(self.rows_by_key[d["debt"]], pq) if d["debt"] else 0.0
                pts.append((ia or 0.0, de or 0.0, 4.0 * iv))
            fit, ok = None, False
            if pts:
                cands = [_fit_two(pts)]
                saa, sbb = sum(a * a for a, _, _ in pts), sum(b * b for _, b, _ in pts)
                cands.append((sum(a * z for a, _, z in pts) / saa, 0.0) if saa > 0 else None)            # 자산 쪽만(모델 부채 금리 '없음' = 0)
                cands.append((0.0, -sum(b * z for _, b, z in pts) / sbb) if sbb > 0 else None)           # 부채 쪽만
                for c in cands:
                    if c is not None and c[0] >= 0 and c[1] >= 0 and max(abs((a * c[0] - b * c[1]) / 4 - z / 4) for a, b, z in pts) <= TOL_EOK:
                        fit, ok = c, True
                        break
            if ok:
                self.fit["이자율_자산"] = fit[0]
                self.fit["이자율_부채"] = fit[1] if d["debt"] else None
                self.links["이자율"] = {"linked": True, "fit_asset": fit[0], "fit_debt": self.fit["이자율_부채"], "n": len(pts),
                                      "model_assumption": [asm.get("interest_rate_asset"), asm.get("interest_rate_debt")]}
            else:
                self.plan["이자손익"] = {"type": "abs", "label": "이자손익 (가정, 억원)"}
                self.links["이자율"] = {"linked": False, "n": len(pts), "reason": "추정 이자손익 없음" if not pts else "전분기 잔액 × 이자율로 재현 불가 → 분기별 가정값"}
        # 환관련손익 — 외화 순노출 × Δ원/달러(기말)/100 이 변수 시트 환율 값으로 재현되면 연결
        if "환관련손익" in self.rows_by_key:
            e = (self.m.get("fx_pnl") or {}).get("net_usd_exposure_m")
            errs, ok = [], e is not None
            for q in fq:
                fv = v("환관련손익", q)
                if fv is None:
                    continue
                pq = lay.prev_q(q)
                e1, e0 = self._fx_series("USDKRW_end", q), (self._fx_series("USDKRW_end", pq) if pq else None)
                if e1 is None or e0 is None:
                    ok = False
                    break
                errs.append(abs(e * (e1 - e0) / 100.0 - fv))
            if ok and errs and max(errs) <= TOL_EOK:
                self.plan["환관련손익"] = {"type": "fxpnl", "exp": "외화순노출_백만$", "note": "환관련손익 = 외화 순노출(백만$) × Δ원/달러(기말) / 100"}
                self.links["환관련손익"] = {"linked": True, "exposure_usd_m": e, "n": len(errs), "max_err": max(errs)}
            else:
                self.links["환관련손익"] = {"linked": False, "exposure_usd_m": e, "reason": "노출 없음" if e is None else "환율 열 없음·재현 불가"}
        self._resolve_sls_link()

    @staticmethod
    def _sls_include(c, incl_new):
        """모델이 매출조선에 쓰는 계약 — 해양(type≠OTHER) · counted · (신규수주 패널 포함 모델이면 origin 분기말까지 체결분만)."""
        return bool(c.get("counted", True)) and c.get("type") != "OTHER" and (bool(c.get("signed_by_origin", True)) or not incl_new)

    def _sls_geometry(self):
        """SLS 시트 배치 — subQ(매출조선 수식)와 SLS 시트가 같은 행·열 번호를 쓰도록 미리 계산한다."""
        if self.sls_geo is not None:
            return self.sls_geo
        s = self.sls or {}
        byq, byy = s.get("by_quarter") or {}, s.get("by_year") or {}
        qs = sorted(k for k in byq if re.fullmatch(r"\d{4}Q[1-4]", k))
        ys = sorted(k for k in byy if re.fullmatch(r"\d{4}", k))
        cols = [{"kind": "q", "key": q, "label": qlabel(q), "date": qdate(q), "col": FIRST_COL + i} for i, q in enumerate(qs)]
        cols += [{"kind": "a", "key": y, "label": int(y), "date": y + ".12A", "col": FIRST_COL + len(qs) + i} for i, y in enumerate(ys)]
        types, cohorts = [], []
        for d in list(byq.values()) + list(byy.values()):
            for t in (d.get("by_type") or {}):
                if t not in types:
                    types.append(t)
            for t in (d.get("by_cohort") or {}):
                if t not in cohorts:
                    cohorts.append(t)
        specs = [("usd_m", "매출인식 합계(백만$)", NF_1), ("marine_usd_m", "해양(type≠OTHER) 합계(백만$)", NF_1)] + \
                [("type:" + t, "선종 " + t, NF_1) for t in types] + [("cohort:" + t, "코호트 " + t, NF_1) for t in cohorts] + \
                [("hedge_ratio", "헤지비율", NF_PCT), ("hedge_rate", "헤지환율(가중)", NF_FX), ("spot_assumed", "건조시점 환율(sls 가정)", NF_FX),
                 ("applied_rate", "적용환율", NF_FX), ("hedged_krw_m", "헤지적용 원화매출(백만원)", NF_INT),
                 ("marine_hedged_krw_m", "해양 원화(백만원)", NF_INT), ("marine_hedged_krw_m_signed_by_origin", "해양 원화 — origin 분기말까지 체결분(백만원)", NF_INT),
                 ("marine_hedged_krw_m_post_origin", "해양 원화 — origin 이후 체결분(백만원)", NF_INT), ("target_opm", "타겟 OPM(코호트)", NF_PCT)]
        cts = [c for c in (s.get("contracts") or []) if c.get("counted", True)]
        g = {"cols": cols, "qcol": {c["key"]: c["col"] for c in cols if c["kind"] == "q"}, "specs": specs,
             "spec_first": DATA_ROW, "spec_last": DATA_ROW + len(specs) - 1, "contracts": cts, "n": len(cts)}
        r = g["spec_last"] + 2
        g.update(title_row=r, h_row=r + 1, mode_row=r + 2, spot_row=r + 3, after_row=r + 4, usd_head=r + 6, usd_first=r + 7)
        g["usd_last"] = g["usd_first"] + max(len(cts), 1) - 1
        g["krw_head"] = g["usd_last"] + 2
        g["krw_first"] = g["krw_head"] + 1
        g["krw_last"] = g["krw_first"] + max(len(cts), 1) - 1
        g["usd_incl_row"] = g["krw_last"] + 1
        g["krw_link_row"] = g["krw_last"] + 2
        g["ledger_row"] = g["krw_link_row"] + 2
        pc = (cols[-1]["col"] if cols else FIRST_COL) + 2
        g["pcol"] = {"incl": pc, "rate": pc + 1, "factor": pc + 2, "signed": pc + 3, "by_origin": pc + 4, "cohort": pc + 5, "amt_usd": pc + 6, "end": pc + 7}
        self.sls_geo = g
        return g

    def _resolve_sls_link(self):
        """조선사: 계약별 스케줄 × 헤지 환율 재구성이 sls by_quarter 저장값과 맞으면(±TOL_SLS_M 백만원) subQ 매출<seg> 를 SLS 시트 합계에 연결.
        매출<seg>(q) = SLS 원화/100 + 원장 밖 잔고 소진(변수 노란 셀, 역산) + 매출<seg>신규(변수 노란 셀). 건조시점 환율은 변수 원/달러(평균)."""
        s = self.sls
        if not s or not s.get("contracts") or not s.get("by_quarter"):
            return
        seg = next((sg for sg in (self.m.get("segments") or []) if str((sg.get("driver") or {}).get("type") or "").startswith("sls")), None)
        if not seg:
            return
        row_key = "매출" + seg["key"]
        if row_key not in self.rows_by_key:
            return
        new_key = row_key + "신규" if (row_key + "신규") in self.rows_by_key else None
        incl_new = bool(self.m.get("new_orders_included"))
        series_key = "marine_hedged_krw_m_signed_by_origin" if incl_new else "marine_hedged_krw_m"
        hedge = s.get("hedge") or {}
        h = hedge.get("hedge_ratio")
        if h is None:
            h = next((d.get("hedge_ratio") for d in s["by_quarter"].values() if d.get("hedge_ratio") is not None), None)
        hrate = hedge.get("hedge_rate")               # measured(정기보고서 평균 약정환율, 한화오션) — 전 계약 공통; None 이면 계약별 수주시점 환율
        cap = s.get("backlog_cap") or {}
        factor = cap.get("factor", 1.0) if s.get("backlog_cap_applied") else 1.0
        origin = s.get("origin") or self.lay.last_actual
        g = self._sls_geometry()
        if h is None or not g["contracts"]:
            self.links["선표"] = {"linked": False, "reason": "헤지비율·계약 없음"}
            return
        recon, usd_incl, runoff, diffs = {}, {}, {}, []
        for q in self.lay.est_quarters:
            mv = self._v(row_key, q)
            if mv is None:
                continue
            stored = (s["by_quarter"].get(q) or {}).get(series_key)
            spot = self._fx_series("USDKRW_avg", q)
            if q not in g["qcol"] or stored is None or spot is None:
                self.links["선표"] = {"linked": False, "reason": "%s — SLS 열·저장값·환율 중 없음" % q}
                return
            after = 1 if q > origin else 0
            k_sum = u_sum = 0.0
            for c in g["contracts"]:
                sched = (c.get("schedule") or {}).get(q)
                if not sched or not self._sls_include(c, incl_new):
                    continue
                rate = hrate if hrate is not None else c.get("fx_at_sign")
                if rate is None:
                    self.links["선표"] = {"linked": False, "reason": "계약 %s 환율 없음" % c.get("rcp")}
                    return
                eff = sched * (1 + ((factor if c.get("backlog_cap_applied") else 1.0) - 1) * after)
                k_sum += eff * (h * rate + (1 - h) * spot)
                u_sum += eff
            diffs.append(abs(k_sum - stored))
            recon[q], usd_incl[q] = k_sum, u_sum
            runoff[q] = mv - k_sum / 100.0 - ((self._v(new_key, q) or 0.0) if new_key else 0.0)
        if not recon:
            return
        mx = max(diffs)
        if mx > max(TOL_SLS_M, 1e-3 * max(abs(x) for x in recon.values())):
            self.links["선표"] = {"linked": False, "reason": "계약별 재구성 ≠ sls %s(최대 %.1f 백만원)" % (series_key, mx), "max_diff_m": mx}
            return
        self.sls_link = {"row": row_key, "new": new_key, "seg": seg["key"], "h": h, "hedge_rate": hrate, "factor": factor, "origin": origin,
                         "incl_new": incl_new, "series_key": series_key, "recon": recon, "usd_incl": usd_incl, "runoff": runoff,
                         "runoff_per_q_model": (seg.get("driver") or {}).get("runoff_per_q")}
        self.plan[row_key] = {"type": "sls_link", "new": new_key,
                              "label": "원장 밖 잔고 소진(억원) — %s = SLS 원화/100 + 이 행%s" % (row_key, (" + " + new_key) if new_key else "")}
        self.links["선표"] = {"linked": True, "series_key": series_key, "hedge_ratio": h, "hedge_rate_mode": "약정환율(measured)" if hrate is not None else "계약별 수주시점 환율",
                            "contracts": len(g["contracts"]), "included": sum(1 for c in g["contracts"] if self._sls_include(c, incl_new)),
                            "max_diff_m": mx, "n": len(recon), "backlog_cap_factor": factor}

    # ── README ──
    def sheet_readme(self):
        ws = self.wb.create_sheet("README")
        m = self.m
        asm = m.get("assumptions") or {}
        fit, links = self.fit, self.links

        def pct(x, d=2):
            return ("%%.%df%%%%" % d) % (x * 100) if isinstance(x, (int, float)) else "—"
        linked = [k for k, v in links.items() if v.get("linked")]
        unlinked = ["%s(%s)" % (k, v.get("reason") or "") for k, v in links.items() if not v.get("linked")]
        lines = [
            ("%s(%s) 실적 모델 — ARGUS 한국조선" % (m.get("name", ""), m.get("stock", "")), F_BOLD),
            ("기준 분기 %s · 마지막 실적 %s · 생성 %s" % (m.get("origin", ""), m["periods"]["last_actual"], m.get("built_at", "")), None),
            ("모든 수치는 모델 산출값(DART 공개 재무제표 + 우리 추정)이다 — 목표주가·투자 추천 아님. 참고용 · 투자조언 아님.", F_BOLD),
            ("", None),
            ("시트", F_BOLD),
            ("변수      가정 시계열. 환율 8행(ECB via frankfurter — L2 fx.json; 추정 분기 = 노란 가정), 추정 드라이버(노란 셀 = 모델값에서 역산한 가정, 바꾸면 재계산; 연녹 셀 = 3절 스칼라 참조 수식), 3절 스칼라 가정(세율·판관비율·지배주주비중·이자율·종가·PER 밴드·COE·외화 순노출·헤지비율)", None),
            ("BS연결/BS별도  FnGuide 계정명 × 기간 값 덤프(백만원). 행1 분기라벨, 행4 YYYY.MM / YYYY.12A. 이름정의 BS연결H·BS연결(계정명 C열부터)", None),
            ("subQ      분기 서브모델(억원 = 백만원/100). 확정 행: =IFERROR(VLOOKUP($D행,BS연결,MATCH(E$1,BS연결H,0),0)/100,\"\")  추정 행: 변수 시트 가정 셀 참조 수식. 외환손익·파생상품손익은 BS 계정 합·차, 기타금융손익은 잔차, OPM·EPS·BPS·PER·PBR 은 파생 수식", None),
            ("SLS       선표(조선사만, sls json 있을 때): 요약 행(sls 저장값) + 계약별 스케줄(백만$) × (헤지비율 × 헤지환율 + (1−헤지비율) × 건조시점 환율[변수 원/달러(평균)]) → 원화 수식 블록 → 합계 행이 subQ 매출조선에 연결", None),
            ("시나리오   신규수주 보수/기준/낙관(forecast_panel 패널, 모델에 scenarios 가 있을 때만): 매출액·영업이익 = subQ 기준값 ± (시나리오 신규 − base 신규) × OPM. base 만 subQ 행에 포함, 나머지는 합산 안 함", None),
            ("분기/연간예상  subQ 를 이름(A열 key)으로 VLOOKUP 한 보고서 표(억원)", None),
            ("TP        PER 밴드 × 12M fwd EPS, PBR = ROE/COE × BPS — 모델 산출값 구간, 목표주가·추천 아님", None),
            ("외화      환율 QoQ, 외화 순노출 × Δ기말환율 → 환관련손익 추정", None),
            ("", None),
            ("규약", F_BOLD),
            ("· 데이터는 E열부터, 4분기 뒤 연간 1열(라벨은 숫자 연도). 실적/추정 경계 = 굵은 세로선, 추정 열 = 연노랑, 연간 열 = 회색, 가정 입력 = 노란 셀, 스칼라 참조 수식 = 연녹 셀(숫자로 덮어쓰면 그 분기만 바뀜)", None),
            ("· 숫자 서식 #,##0(억원) · EPS/BPS 원 · 주식수 백만주 · 환율 원 · 선표 백만$/백만원", None),
            ("· 추정 셀 메모 = 모델 json 의 basis. 변수 시트 노란 셀 메모 = 역산 방법. 출처 없는 숫자는 없다 — 가정이면 가정이라고 적었다", None),
            ("· BS 시트에 값이 없거나 모델값과 다른(비용 부호 반전 등) 기간의 확정 셀은 모델값 직접 + 메모. fin(L1) 수집 후 재생성", None),
            ("", None),
            ("가정 요약 — 변수 시트 3절 노란 셀을 바꾸면 subQ·분기·연간예상·TP·시나리오가 재계산된다(연결된 항목만)", F_BOLD),
            ("· 세율 %s — %s" % (pct(fit.get("세율", asm.get("tax_rate"))), ("법인세 = MAX(세전,0) × 세율, 추정 전 분기 연결(Σ법인세/Σmax(세전,0) 역산, 재현 오차 ≤ 0.05억). 모델 가정 %s · 근거: %s" % (pct(asm.get("tax_rate")), asm.get("tax_basis") or "")) if links.get("세율", {}).get("linked") else "연결 안 됨(%s) — 드라이버 행 분기별 역산값 사용" % links.get("세율", {}).get("reason", "")), None),
            ("· 판관비율 %s — %s" % (pct(fit.get("판관비율", asm.get("sga_ratio"))), "판관비 = 매출액 × 판관비율, 추정 전 분기 연결" if links.get("판관비율", {}).get("linked") else "연결 안 됨(%s)" % links.get("판관비율", {}).get("reason", "")), None),
            ("· 지배주주 비중 %s — %s" % (pct(fit.get("지배주주비중"), 1), "지배주주순이익 = 당기순이익 × 비중, 추정 전 분기 연결" if links.get("지배주주비중", {}).get("linked") else "연결 안 됨(%s)" % links.get("지배주주비중", {}).get("reason", "")), None),
            ("· 이자율 자산 %s / 부채 %s — %s" % (pct(fit.get("이자율_자산")), pct(fit.get("이자율_부채")), ("이자손익 = (전분기 이자발생자산 × 자산이자율 − 전분기 총차입금 × 부채이자율)/4, 최소자승 역산(모델 assumptions %s 는 4자리 반올림이라 그대로 쓰면 ±1억 어긋남) · 출처: %s" % (links["이자율"].get("model_assumption"), json.dumps(asm.get("interest_rate_source") or {}, ensure_ascii=False))) if links.get("이자율", {}).get("linked") else "연결 안 됨(%s)" % links.get("이자율", {}).get("reason", "")), None),
            ("· 환율 — 변수 1절 8행. 실적 분기 = ECB 분기 평균/기말(fx.json), 추정 분기(노란) = %s. 원/달러(평균) → SLS 건조시점 환율 → 매출조선(조선사, 선표 연결 시) ; 원/달러(기말) → %s" % (
                ((asm.get("fx") or {}).get("basis") or "fx.json forward"),
                ("환관련손익 = 외화 순노출 %s백만$ × Δ기말/100 (연결)" % (self.m.get("fx_pnl") or {}).get("net_usd_exposure_m")) if links.get("환관련손익", {}).get("linked")
                else ("환관련손익 행 없음(%s)" % ((self.m.get("fx_pnl") or {}).get("note") or "외화 순노출 미공시") if "환관련손익" not in self.rows_by_key
                      else "환관련손익 연결 안 됨: %s" % links.get("환관련손익", {}).get("reason", ""))), None),
        ]
        if self.sls:
            lk = links.get("선표") or {}
            hedge = (self.sls.get("hedge") or {})
            lines.append(("· 선표(SLS) — 헤지비율 %s(%s) · 헤지환율 %s · 코호트 모드 %s · 잔고 캡 %s · %s" % (
                pct(hedge.get("hedge_ratio"), 1), hedge.get("kind") or "", lk.get("hedge_rate_mode") or ("약정환율(measured)" if hedge.get("hedge_rate") else "계약별 수주시점 환율"),
                self.sls.get("cohort_mode"), ("적용 × %.4f" % lk.get("backlog_cap_factor", 1.0)) if self.sls.get("backlog_cap_applied") else "미적용",
                ("subQ %s = SLS 합계/100 + 원장 밖 잔고 소진(변수 drv:%s, 역산) + %s — 계약 %d건 중 포함 %d건(%s), 재구성 최대차 %.1f 백만원" % (
                    self.sls_link["row"], self.sls_link["row"], self.sls_link["new"] or "0", lk.get("contracts", 0), lk.get("included", 0), lk.get("series_key"), lk.get("max_diff_m", 0.0)))
                if lk.get("linked") else "subQ 매출조선은 SLS 에 연결되지 않음(%s)" % lk.get("reason", "sls 없음")), None))
        if self.m.get("scenarios"):
            no = self.m.get("new_orders") or {}
            lines.append(("· 신규수주 — forecast_panel %s 시나리오를 매출조선신규 행에 포함(calibrated=%s, status %s — 통계 흐름, 공시 계약 개별 반영 아님). 보수/낙관/기존만 은 시나리오 시트에만(합산 안 함)" % (
                no.get("scenario_in_rows", "base"), no.get("calibrated"), no.get("panel_status")), None))
        elif self.m.get("new_orders_included") is False:
            lines.append(("· 신규수주 — 패널 값 없음 → 2026Q2 잔고 + 공시 수주 소진분만(2028 감소는 이 한계)", None))
        lines += [
            ("· 연결된 가정: %s · 연결 안 된 가정(모델값·분기별 역산 그대로): %s" % (", ".join(linked) or "없음", ", ".join(unlinked) or "없음"), None),
            ("· 밸류에이션 — PER 밴드 %s배(%s) · COE %s · 12M fwd EPS 가중 0.5(가정). 모델 산출값 구간이며 추천이 아니다" % (
                "/".join(str((m.get("valuation") or {}).get("per_band", {}).get(k)) for k in ("lo", "mid", "hi")), (m.get("valuation") or {}).get("per_band", {}).get("basis", ""), pct((m.get("valuation") or {}).get("coe"), 1)), None),
            ("", None),
            ("출처: DART 정기보고서(연결/별도 재무제표·주석), FnGuide 계정명 규약, ECB 환율(api.frankfurter.app), 종가 네이버 정규장(aikstockdata 금융위 확정종가 대조·출처표기·비영리), 척당 계약 공시(contracts.json), forecast_panel(신규수주 패널)", None),
            ("재현: cd argus/kship/tools && python3 kship_model_xlsx.py --model assets/models/%s.json --verify   "
             "(모델 json 부터: kship_model.py --build --all --xlsx · xlsx 58사 일괄: kship_model_xlsx.py --all --verify)" % m.get("stock", "<stock>"), None),
            ("검증(--verify): 추정 전 분기·실적 수식 셀을 자체 계산기로 재계산해 모델값과 대조(억원 ±0.05), 세율·OPM·환율 forward·이자율·신규수주 셀을 바꿔 종속 셀이 기대만큼 움직이는지 확인. Excel 실개봉은 미확인(LibreOffice/Excel 없음)", None),
            ("한계: 모델 quality.warnings — " + " / ".join((m.get("quality") or {}).get("warnings") or ["(없음)"]), None),
            ("백테스트: " + json.dumps(m.get("backtest") or {}, ensure_ascii=False), None),
        ]
        for i, (t, f) in enumerate(lines, 1):
            c = ws.cell(i, 1, t)
            if f:
                c.font = f
        ws.column_dimensions["A"].width = 160
        self.stats["README"] = (len(lines), 1)

    # ── 변수 ──
    def sheet_vars(self):
        ws = self.wb.create_sheet("변수")
        lay = self.lay
        self._header(ws, "이름: 변수H")
        ws.cell(2, 1, "가정 시계열 — 노란 셀은 추정 가정(모델값에서 역산), 연녹 셀은 3절 스칼라 참조 수식. 바꾸면 subQ·분기·연간예상·TP·시나리오가 재계산됩니다.").font = F_NOTE
        r = DATA_ROW
        # 1. 환율
        ws.cell(r, 2, "1. FX").font = F_BOLD
        ws.cell(r, 3, "이름: 환율").font = F_NOTE
        fx_first = r + 1
        for metric, label in FX_ROWS:
            r += 1
            self.fx_row[metric] = r
            ws.cell(r, 1, metric).font = F_KEY
            ws.cell(r, 3, label)
            for c in lay.cols:
                v, note = self._fx_value(metric, c)
                if v is None:
                    continue
                cell = ws.cell(r, c["col"], v)
                cell.number_format = NF_FX
                if lay.is_est(c):
                    cell.fill = FILL_INPUT
                if note:
                    cell.comment = Comment(note, "kship")
        fx_last = r
        # 1b. 종가 4행(prices.json history_quarterly) — 레퍼런스 PER 4종(기말/고/저/평균)용. 평균은 subQ 실적 PER/PBR 수식이 참조
        if self.prices and self.prices.get("history_quarterly"):
            r += 1
            ws.cell(r, 2, "1b. 종가").font = F_BOLD
            hist = self.prices["history_quarterly"]
            for metric, label in PRICE_ROWS:
                r += 1
                self.price_row[metric] = r
                ws.cell(r, 1, "price_" + metric).font = F_KEY
                ws.cell(r, 3, label)
                for c in lay.qcols:
                    h = hist.get(c["key"]) or {}
                    if h.get(metric) is not None:
                        cell = ws.cell(r, c["col"], h[metric])
                        cell.number_format = NF_INT
                        if h.get("partial"):
                            cell.comment = Comment("분기 진행 중(partial) — aikstockdata/토스 일봉", "kship")
        # 2. 추정 드라이버
        r += 2
        ws.cell(r, 2, "2. 추정 드라이버").font = F_BOLD
        ws.cell(r, 3, "실적 열 = subQ 실현값(수식), 추정 열 = 가정(노란 = 모델값 역산, 연녹 = 3절 스칼라 참조)").font = F_NOTE
        self.vars_ws = ws
        self.drv_first = r + 1
        for row in self.rows:
            d = self.plan[row["key"]]
            if d["type"] in ("identity", "per_share", "divide", "valuation", "interest", "fxpnl", "ratio_parent"):
                continue
            r += 1
            self.drv_row[row["key"]] = r
            ws.cell(r, 1, "drv:" + row["key"]).font = F_KEY
            ws.cell(r, 2, d["type"]).font = F_NOTE
            ws.cell(r, 3, d["label"])
        self.drv_last = r
        # 3. 밸류에이션·스칼라 가정 — 노란 셀. 연결된 것은 추정 전 분기 드라이버 셀(연녹)이 `=$D$n` 으로 참조한다
        r += 2
        ws.cell(r, 2, "3. 밸류에이션·스칼라").font = F_BOLD
        ws.cell(r, 3, "노란 셀 = 가정. '연결' 표시는 추정 전 분기 수식이 이 셀을 참조한다는 뜻").font = F_NOTE
        val = self.m.get("valuation") or {}
        price = val.get("price") or {}
        band = val.get("per_band") or {}
        asm = self.m.get("assumptions") or {}
        fit, links = self.fit, self.links
        hedge = (self.sls or {}).get("hedge") or {}

        def lk(name, formula_txt, model_v=None):
            L = links.get(name) or {}
            if L.get("linked"):
                return "연결 — %s(역산, 재현 오차 ≤ %.2f억%s)" % (formula_txt, TOL_EOK, (" · 모델 assumptions %s" % model_v) if model_v is not None else "")
            return "참고(연결 안 됨: %s) — 모델 assumptions 값" % (L.get("reason") or "해당 없음")
        scalars = [
            ("종가", price.get("close"), "원 — %s" % (price.get("source") or price.get("src") or "prices.json"), NF_INT),
            ("기준일", price.get("as_of"), "종가 기준일", None),
            ("COE", val.get("coe"), "자기자본비용(가정) — TP 적정 PBR = ROE/COE", NF_PCT),
            ("PER_lo", band.get("lo"), "PER 밴드 하단(배) — %s" % (band.get("basis") or ""), NF_1),
            ("PER_mid", band.get("mid"), "PER 밴드 중단(배)", NF_1),
            ("PER_hi", band.get("hi"), "PER 밴드 상단(배)", NF_1),
            ("EPS_Y1가중", 0.5, "12M fwd EPS = FY(last_actual연도)E × w + FY+1E × (1−w) — 기준 분기 2Q 이면 0.5(가정)", NF_PCT),
            ("세율", fit.get("세율", asm.get("tax_rate")), lk("세율", "법인세 = MAX(세전,0) × 세율", asm.get("tax_rate")) + (" · " + str(asm.get("tax_basis") or "")), NF_PCT2),
            ("판관비율", fit.get("판관비율", asm.get("sga_ratio")), lk("판관비율", "판관비 = 매출액 × 판관비율", asm.get("sga_ratio")), NF_PCT2),
            ("지배주주비중", fit.get("지배주주비중", (1 - asm["minority_share"]) if isinstance(asm.get("minority_share"), (int, float)) else None),
             lk("지배주주비중", "지배주주순이익 = 당기순이익 × 비중"), NF_PCT2),
            ("이자율_자산", fit.get("이자율_자산", asm.get("interest_rate_asset")), lk("이자율", "이자손익 = (전분기 이자발생자산 × 자산이자율 − 전분기 총차입금 × 부채이자율)/4", asm.get("interest_rate_asset")), NF_PCT2),
            ("이자율_부채", fit.get("이자율_부채", asm.get("interest_rate_debt")), lk("이자율", "같은 수식의 부채 쪽(총차입금 행이 없으면 0)", asm.get("interest_rate_debt")), NF_PCT2),
            ("배당성향", asm.get("payout"), "배당성향(참고 — 모델은 직전 DPS 유지·Q2 지급 가정)", NF_PCT),
            ("외화순노출_백만$", (self.m.get("fx_pnl") or {}).get("net_usd_exposure_m"),
             ((self.m.get("fx_pnl") or {}).get("method") or "") + (" — 연결: subQ 환관련손익 = 노출 × Δ원/달러(기말)/100" if links.get("환관련손익", {}).get("linked") else " — " + ((self.m.get("fx_pnl") or {}).get("note") or "")), NF_1),
            ("헤지비율", hedge.get("hedge_ratio"), ("SLS 원화 = Σ 스케줄 × (헤지비율 × 헤지환율 + (1−헤지비율) × 건조시점 환율[원/달러(평균)]) — %s · %s" % (hedge.get("kind") or "", hedge.get("basis") or "")) if self.sls else "sls 없음", NF_PCT),
            ("헤지환율_약정", hedge.get("hedge_rate"), "정기보고서 평균 약정환율(measured) — 비어 있으면 계약별 수주시점 환율(fx_at_sign)을 쓴다" if self.sls else "sls 없음", NF_FX),
        ]
        for name, v, note, nf in scalars:
            r += 1
            self.scalar_row[name] = r
            ws.cell(r, 1, "sc:" + name).font = F_KEY
            ws.cell(r, 3, name)
            cell = ws.cell(r, 4, v)
            cell.fill = FILL_INPUT
            if nf:
                cell.number_format = nf
            ws.cell(r, 5, note).font = F_NOTE
        self.vars_last = r
        last = get_column_letter(lay.last_col)
        self._name("변수H", "변수!$C$%d:$%s$%d" % (HEADER_ROW, last, HEADER_ROW))
        self._name("환율", "변수!$C$%d:$%s$%d" % (fx_first, last, fx_last))
        self.stats["변수"] = (r, lay.last_col)

    def _yoy_param(self, row, q):
        """yoy 드라이버 셀의 (가정값, 기준 분기, 모드). 전년동기가 없거나 0 이면(상장 직후·기간 짧은 회사) 전분기 대비로 내려간다 — 셀을 죽은 값으로 두지 않기 위해."""
        lay = self.lay
        pq = lay.prev_year_q(q)
        b = self._sheet_base(row, pq) if pq else None
        v = val_of(row, q)
        if v is None:
            return None, None, None
        if b:
            return v / b - 1, pq, "yoy"
        if b == 0 and v == 0:
            return 0.0, pq, "yoy"
        pq2 = lay.prev_q(q)
        b2 = self._sheet_base(row, pq2) if pq2 else None
        if b2:
            return v / b2 - 1, pq2, "qoq_fallback"
        return None, None, None

    def fill_driver_cells(self):
        """subQ 행 번호가 정해진 뒤 — 드라이버 행의 실적 열(실현값 수식)·추정 열(역산 가정값 또는 스칼라 참조)을 채운다.
        roll(자본총계·지배주주지분 '기타변동') 드라이버는 2차 패스에서 시트가 실제로 계산하는 흐름값(당기순이익 수식 결과)으로 역산한다 —
        세율·이자율 스칼라 적합의 분기별 오차(≤0.05억)가 자본 롤 체인에 10분기 쌓여 자산총계·BPS 가 허용오차를 넘는 것을 막는다."""
        ws = self.vars_ws
        lay = self.lay
        geo = self._sls_geometry() if self.sls_link else None
        for row in self.rows:
            k = row["key"]
            if k not in self.drv_row:
                continue
            d, r, sr = self.plan[k], self.drv_row[k], self.subq_row[k]
            nf = NF_PCT if d["type"] in ("yoy", "qoq", "ratio", "tax") else NF_1
            for c in lay.cols:
                L = get_column_letter(c["col"])
                cell = ws.cell(r, c["col"])
                cell.number_format = nf
                if c["kind"] == "q" and lay.is_est(c):
                    q = c["key"]
                    if d["type"] == "roll":
                        continue                      # 2차 패스
                    if d["type"] == "sls_link":
                        p = self.sls_link["runoff"].get(q)
                        if p is None:
                            continue
                        cell.value = round(p, 6)
                        cell.fill = FILL_INPUT
                        cell.comment = Comment("역산: 모델 %s − SLS 재구성 원화/100%s (모델 runoff_per_q %s; 계약 스케줄 소수 3자리 반올림 차 포함)\n%s" % (
                            k, (" − " + d["new"]) if d.get("new") else "", self.sls_link.get("runoff_per_q_model"), cell_of(row, q).get("basis") or ""), "kship")
                        continue
                    if d.get("scalar") and d["scalar"] in self.scalar_row:
                        if val_of(row, q) is None:
                            continue
                        cell.value = "=$D$%d" % self.scalar_row[d["scalar"]]
                        cell.fill = FILL_LINK
                        cell.comment = Comment("3절 스칼라 '%s' 참조 — 숫자로 덮어쓰면 이 분기만 바뀐다\n%s" % (d["scalar"], cell_of(row, q).get("basis") or ""), "kship")
                        continue
                    if d["type"] == "yoy":
                        p, _, mode = self._yoy_param(row, q)
                    else:
                        p, mode = implied_param(row, q, d, self.rows_by_key, lay, self._sheet_base), d["type"]
                    if p is None:
                        continue
                    cell.value = round(p, 9)
                    cell.fill = FILL_INPUT
                    cell.comment = Comment("모델값 역산(implied): %s\n%s" % (
                        {"yoy": "v / 전년동기 − 1", "qoq_fallback": "v / 전분기 − 1 (전년동기 없음 → 전분기 대비)", "qoq": "v / 전분기 − 1",
                         "ratio": "v / %s" % d.get("base"), "tax": "v / %s" % d.get("base"), "abs": "모델 v 그대로"}[mode],
                        cell_of(row, q).get("basis") or ""), "kship")
                    continue
                # 실적(또는 연간) 열 — subQ 실현값
                f = None
                if d["type"] in ("ratio", "tax"):
                    br = self.subq_row.get(d["base"])
                    if br:
                        f = '=IFERROR(subQ!%s%d/subQ!%s%d,"")' % (L, sr, L, br)
                elif d["type"] == "yoy" and c["kind"] == "q":
                    pq = lay.prev_year_q(c["key"])
                    if pq:
                        f = '=IFERROR(subQ!%s%d/subQ!%s%d-1,"")' % (L, sr, lay.letter(pq), sr)
                elif d["type"] == "qoq" and c["kind"] == "q":
                    pq = lay.prev_q(c["key"])
                    if pq:
                        f = '=IFERROR(subQ!%s%d/subQ!%s%d-1,"")' % (L, sr, lay.letter(pq), sr)
                elif d["type"] == "abs":
                    f = "=subQ!%s%d" % (L, sr)
                elif d["type"] == "roll" and c["kind"] == "q":
                    pq = lay.prev_q(c["key"])
                    fr = self.subq_row.get(d["flow"])
                    if pq and fr:
                        f = '=IFERROR(subQ!%s%d+subQ!%s%d-subQ!%s%d,"")' % (lay.letter(pq), sr, L, fr, L, sr)
                elif d["type"] == "sls_link" and c["kind"] == "q" and geo and c["key"] in geo["qcol"]:
                    f = '=IFERROR(subQ!%s%d-SLS!%s%d/100,"")' % (L, sr, get_column_letter(geo["qcol"][c["key"]]), geo["krw_link_row"])
                if f:
                    cell.value = f
        # 2차 패스 — roll 드라이버: 전분기 시트값 + 흐름 시트값(수식 평가) − 모델 v. 분기 순서대로(전분기 드라이버가 먼저 있어야 전분기 값이 나온다)
        roll_rows = [row for row in self.rows if row["key"] in self.drv_row and self.plan[row["key"]]["type"] == "roll"]
        for q in lay.est_quarters:
            em = Emulator(self.wb) if roll_rows else None
            L = lay.letter(q)
            pq = lay.prev_q(q)
            for row in roll_rows:
                k = row["key"]
                d, r, sr = self.plan[k], self.drv_row[k], self.subq_row[k]
                fr = self.subq_row.get(d["flow"])
                v = val_of(row, q)
                cell = ws.cell(r, lay.col_of(q))
                if v is None or not fr or pq is None:
                    continue
                try:
                    b = em.value("subQ", "%s%d" % (lay.letter(pq), sr))
                    f = em.value("subQ", "%s%d" % (L, fr))
                except Exception:  # noqa: BLE001
                    b = f = None
                if isinstance(b, (int, float)) and isinstance(f, (int, float)) and not isinstance(b, bool) and not isinstance(f, bool):
                    p, how = b + f - v, "전분기(시트값) + %s(시트 수식값) − v" % d["flow"]
                else:
                    p, how = implied_param(row, q, d, self.rows_by_key, lay, self._sheet_base), "전분기 + %s − v (모델값 기준 — 시트값 평가 불가)" % d["flow"]
                if p is None:
                    continue
                cell.value = round(p, 9)
                cell.fill = FILL_INPUT
                cell.comment = Comment("모델값 역산(implied): %s\n%s" % (how, cell_of(row, q).get("basis") or ""), "kship")
        self._shade(ws, self.vars_last)

    # ── BS연결 / BS별도 ──
    def _bs_values(self, scope):
        """{계정명: {열키: 값(백만원)}} — fin json 우선, 없으면 모델 실적 행(억원×100)."""
        lay = self.lay
        out = {}
        fin = (self.fin or {}).get(scope)
        if fin:
            for stmt in ("bs", "is", "cf"):
                data = fin.get(stmt) or {}
                accts = []
                for q in lay.quarters:
                    for a in (data.get(q) or {}):
                        if a not in accts:
                            accts.append(a)
                for a in accts:
                    d = out.setdefault(a, {})
                    for q in lay.quarters:
                        v = (data.get(q) or {}).get(a)
                        if v is not None:
                            d[q] = v
                    for y in lay.annual:
                        qs = lay.year_quarters(y)
                        if stmt == "bs":
                            q4 = y + "Q4"
                            if q4 in d:
                                d[y] = d[q4]
                        elif len(qs) == 4 and all(q in d for q in qs):
                            d[y] = round(sum(d[q] for q in qs), 2)
            return out
        if scope != "cons":
            return out
        for row in self.rows:
            acct = ROW_ACCOUNT.get(row["key"])
            if not acct or bs_sheet_for(row) != "BS연결":
                continue
            d = out.setdefault(acct, {})
            for q in lay.quarters:
                c = cell_of(row, q)
                if c.get("kind") == "actual" and c.get("v") is not None:
                    d[q] = round(c["v"] * 100, 2)
            for y, a in (row.get("a") or {}).items():
                if str(y) in lay.annual and a.get("kind") == "actual" and a.get("v") is not None:
                    d[str(y)] = round(a["v"] * 100, 2)
        return out

    def sheet_bs(self, scope):
        name = "BS연결" if scope == "cons" else "BS별도"
        ws = self.wb.create_sheet(name)
        lay = self.lay
        self._header(ws, "이름: %sH" % name, unit_note="(백만원)")
        ws.cell(2, 2, "연결" if scope == "cons" else "별도")
        ws.cell(3, 2, "%s(%s)" % (self.m.get("name", ""), "연" if scope == "cons" else "별"))
        stamp = self.built_date().strftime("%y-%m-%d")      # 레퍼런스 행3 규약 'UPDATE: yy-mm-dd' — 실행일이 아니라 모델 built_at(결정론)
        ws.cell(3, 3, "UPDATE: " + stamp).font = F_NOTE
        ws.cell(DATE_ROW, 2, "Account Code").font = F_NOTE
        ws.cell(DATE_ROW, 3, "AccountNAME").font = F_NOTE
        values = self.bs_vals[name]
        src = "fin json(kship_fin.py)" if (self.fin or {}).get(scope) else (
            "모델 실적 행(억원×100) — fin json 없음, 매핑 계정만" if scope == "cons" else "별도 fin 없음 — L1 kship_fin.py 산출 후 재생성")
        ws.cell(2, 4, src).font = F_NOTE
        r = DATA_ROW - 1
        for acct, d in values.items():
            r += 1
            ws.cell(r, 1, 1)
            ws.cell(r, 3, acct)
            for c in lay.cols:
                v = d.get(c["key"])
                if v is not None:
                    cell = ws.cell(r, c["col"], v)
                    cell.number_format = NF_INT
        last_row = max(r, DATA_ROW)
        last = get_column_letter(lay.last_col)
        self._name(name + "H", "%s!$C$%d:$%s$%d" % (name, HEADER_ROW, last, HEADER_ROW))
        self._name(name, "%s!$C$%d:$%s$%d" % (name, DATA_ROW, last, last_row))
        self._shade(ws, last_row)
        self.stats[name] = (last_row, lay.last_col)
        return values

    # ── subQ ──
    def sheet_subq(self):
        ws = self.wb.create_sheet("subQ")
        lay = self.lay
        self._header(ws, "subQ 분기 서브모델(억원, U=100)")
        ws.cell(HEADER_ROW, 1, "YQ").font = F_BOLD
        ws.cell(2, 1, "A=이름(key) B=그룹 C=단위 D=FnGuide 계정명(VLOOKUP 키). 확정=VLOOKUP(BS), 추정=변수 가정 참조").font = F_NOTE
        r = DATA_ROW - 1
        for row in self.rows:
            r += 1
            self.subq_row[row["key"]] = r
        for row in self.rows:
            r = self.subq_row[row["key"]]
            ws.cell(r, 1, row["key"]).font = F_KEY
            ws.cell(r, 2, row.get("group") or "")
            ws.cell(r, 3, row.get("unit") or "억원")
            acct = ROW_ACCOUNT.get(row["key"])
            ws.cell(r, 4, acct or row.get("label") or row["key"])
            unit = row.get("unit") or "억원"
            nf = NF_1 if row["key"] == "주식수" else (NF_PCT if unit == "%" else (NF_FX if unit == "배" else NF_INT))
            for c in lay.cols:
                cell = ws.cell(r, c["col"])
                cell.number_format = nf
                if c["kind"] == "q":
                    self._subq_quarter(ws, row, c, acct)
                else:
                    self._subq_annual(ws, row, c, acct)
        last_row = r
        # 비율 행(참고) — 이름으로 분기·연간예상에서 끌어감. 모델에 같은 이름의 행(OPM)이 있으면 만들지 않는다(VLOOKUP 키 중복 금지)
        ratios = [("OPM", "영업이익", "매출액"), ("NPM", "지배주주순이익", "매출액"), ("매출원가율", "매출원가", "매출액"), ("판관비율", "판관비", "매출액")]
        r += 1
        ws.cell(r, 2, "비율(참고)").font = F_NOTE
        for name, num, den in ratios:
            if num not in self.subq_row or den not in self.subq_row or name in self.subq_row:
                continue
            r += 1
            self.subq_row[name] = r
            ws.cell(r, 1, name).font = F_KEY
            ws.cell(r, 3, "%")
            ws.cell(r, 4, "%s / %s" % (num, den))
            for c in lay.cols:
                L = get_column_letter(c["col"])
                cell = ws.cell(r, c["col"], '=IFERROR(%s%d/%s%d,"")' % (L, self.subq_row[num], L, self.subq_row[den]))
                cell.number_format = NF_PCT
        last_row = r
        last = get_column_letter(lay.last_col)
        self._name("SUBQH", "subQ!$A$%d:$%s$%d" % (HEADER_ROW, last, HEADER_ROW))
        self._name("subQ", "subQ!$A$%d:$%s$%d" % (DATA_ROW, last, last_row))
        self._shade(ws, last_row)
        self.stats["subQ"] = (last_row, lay.last_col)

    def _vlookup(self, row, q_or_y, acct, L):
        sheet = bs_sheet_for(row, q_or_y if isinstance(q_or_y, str) and "Q" in q_or_y else None)
        return '=IFERROR(VLOOKUP($D%d,%s,MATCH(%s$1,%sH,0),0)/100,"")' % (self.subq_row[row["key"]], sheet, L, sheet)

    @staticmethod
    def _composite_vlookup(sheet, comp, L):
        parts = []
        for sign, a in comp:
            parts.append('%sVLOOKUP("%s",%s,MATCH(%s$1,%sH,0),0)' % (sign if parts or sign == "-" else "", a, sheet, L, sheet))
        return '=IFERROR((%s)/100,"")' % "".join(parts)

    def _identity_formula(self, drv, L):
        parts = []
        for sign, k in drv["expr"]:
            if k not in self.subq_row:
                return None
            parts.append("%s%s%d" % (sign if parts or sign == "-" else "", L, self.subq_row[k]))
        return '=IFERROR(%s,"")' % "".join(parts)

    def _eps4(self, q):
        """TTM EPS 셀 4개(해당 분기 포함 직전 4분기) — 연간 열이 끼어 있어 범위 SUM 대신 셀을 나열한다. 부족하면 None."""
        qs = self.lay.quarters
        i = qs.index(q)
        if i < 3 or "EPS" not in self.subq_row:
            return None
        return "+".join("%s%d" % (self.lay.letter(qq), self.subq_row["EPS"]) for qq in qs[i - 3:i + 1])

    def _valuation(self, row, q, L, est):
        """(수식, 파이썬 재계산값) — PER = 종가/TTM EPS, PBR = 종가/BPS. 추정 구간은 변수 종가, 실적 구간은 분기 평균 종가(변수 1b). 못 만들면 (None, None)."""
        key = row["key"]
        close = ((self.m.get("valuation") or {}).get("price") or {}).get("close")
        if est:
            if close is None or "종가" not in self.scalar_row:
                return None, None
            num_f, num_v = "변수!$D$%d" % self.scalar_row["종가"], close
        else:
            pr = self.price_row.get("avg")
            pv = (((self.prices or {}).get("history_quarterly") or {}).get(q) or {}).get("avg") if pr else None
            if pr is None or pv is None:
                return None, None
            num_f, num_v = "변수!%s%d" % (L, pr), pv
        if key == "PER":
            cells = self._eps4(q)
            qs = self.lay.quarters
            i = qs.index(q)
            eps = [self._v("EPS", qq) for qq in qs[i - 3:i + 1]] if i >= 3 else []
            if not cells or any(e is None for e in eps) or sum(eps) == 0:
                return None, None
            return '=IFERROR(%s/(%s),"")' % (num_f, cells), num_v / sum(eps)
        if key == "PBR":
            b = self._v("BPS", q)
            if "BPS" not in self.subq_row or not b:
                return None, None
            return '=IFERROR(%s/%s%d,"")' % (num_f, L, self.subq_row["BPS"]), num_v / b
        return None, None

    def _subq_quarter(self, ws, row, c, acct):
        lay = self.lay
        q, L, r = c["key"], get_column_letter(c["col"]), self.subq_row[row["key"]]
        cell = ws.cell(r, c["col"])
        mc = cell_of(row, q)
        drv = self.plan[row["key"]]
        v = mc.get("v")
        unit = row.get("unit") or "억원"
        if q <= lay.last_actual:
            # 확정 구간 — BS 시트에 값이 있고 모델값과 맞는 실측은 VLOOKUP, 복합 계정은 합·차 VLOOKUP, 파생 행은 수식, 아니면 모델값 직접(메모), 둘 다 없으면 빈 칸(0 표시 금지)
            if mc.get("kind") == "actual" or not mc:
                sheet = bs_sheet_for(row, q)
                bsv = self._bs_get(sheet, acct, q) if acct else None
                if bsv is not None:
                    if v is None or abs(bsv / 100 - v) <= 0.011 + abs(v) * 1e-6:       # 모델 v 는 0.01억 반올림 — 그 밖의 차이는 BS 값이 모델 입력이 아니라는 뜻
                        cell.value = self._vlookup(row, q, acct, L)
                    else:
                        cell.value = v
                        cell.comment = Comment("실측(모델값 직접) — %s %s %s 값 %.2f ≠ 모델값(비용 부호 반전·재분류 등; src %s)" % (
                            sheet, acct, qdate(q), bsv / 100, mc.get("src") or ""), "kship")
                    return
                comp = COMPOSITE_ACCOUNTS.get(row["key"])
                if comp and all(self._bs_get(sheet, a, q) is not None for _, a in comp):
                    calc = sum((1 if s == "+" else -1) * self._bs_get(sheet, a, q) for s, a in comp) / 100
                    if v is None or abs(calc - v) <= 0.011 + abs(v) * 1e-6:
                        cell.value = self._composite_vlookup(sheet, comp, L)
                        cell.comment = Comment("BS 계정 합·차: " + " ".join(s + a for s, a in comp), "kship")
                        return
                f = self._derived_actual(row, q, L)
                if f:
                    cell.value = f
                    return
                if v is not None:
                    cell.value = v
                    cell.comment = Comment("실측(%s — 모델값 직접) — %s" % (
                        ("%s 시트에 %s %s 값 없음" % (sheet, acct, qdate(q))) if acct else "계정 매핑 없음(파생·부문·시세 행)", mc.get("src") or ""), "kship")
            else:
                if v is not None:
                    cell.value = v
                    cell.fill = FILL_EST
                    cell.comment = Comment("추정(실적 구간 갭필): " + str(mc.get("basis") or ""), "kship")
            return
        # 추정 구간 — 가정 셀 참조 수식
        f = self._estimate_formula(row, q, L, drv)
        t = drv["type"]
        if f:
            cell.value = f
            if mc.get("basis") or mc.get("kind"):
                cell.comment = Comment("추정: %s\n[%s]" % (mc.get("basis") or drv.get("note") or "", t), "kship")
        elif v is not None:
            cell.value = v
            cell.comment = Comment("추정(가정 셀 참조 불가 — 기준값 없음, 모델값 직접): " + str(mc.get("basis") or ""), "kship")

    def _derived_actual(self, row, q, L):
        """실적 구간 파생 행 수식 — EPS·BPS(per_share), OPM(divide), 기타금융손익(잔차 항등식), PER·PBR(분기 평균 종가 / TTM EPS·BPS)."""
        key, drv, v = row["key"], self.plan[row["key"]], val_of(row, q)
        if v is None:
            return None
        if drv["type"] == "per_share":
            return '=IFERROR(%s%d*100/%s%d,"")' % (L, self.subq_row[drv["num"]], L, self.subq_row[drv["den"]])
        if drv["type"] == "divide" and self._v(drv["num"], q) is not None and self._v(drv["den"], q):
            return '=IFERROR(%s%d/%s%d,"")' % (L, self.subq_row[drv["num"]], L, self.subq_row[drv["den"]])
        if key == "기타금융손익" and all(k in self.subq_row and self._v(k, q) is not None for k in ("금융손익", "이자손익", "외환손익", "파생상품손익")):
            calc = self._v("금융손익", q) - self._v("이자손익", q) - self._v("외환손익", q) - self._v("파생상품손익", q)
            if abs(calc - v) <= 0.011:
                return '=IFERROR(%s%d-%s%d-%s%d-%s%d,"")' % (L, self.subq_row["금융손익"], L, self.subq_row["이자손익"], L, self.subq_row["외환손익"], L, self.subq_row["파생상품손익"])
        if drv["type"] == "valuation":
            f, calc = self._valuation(row, q, L, est=False)
            if f and calc is not None and abs(calc - v) <= tol_for(key, row.get("unit") or "배", v, **tol_args(self.rows_by_key, self.lay.quarters, q)):
                return f
        return None

    def _estimate_formula(self, row, q, L, drv):
        lay, t, r = self.lay, drv["type"], self.subq_row[row["key"]]
        if t == "identity":
            return self._identity_formula(drv, L)
        if t == "per_share":
            return '=IFERROR(%s%d*100/%s%d,"")' % (L, self.subq_row[drv["num"]], L, self.subq_row[drv["den"]])
        if t == "divide":
            return '=IFERROR(%s%d/%s%d,"")' % (L, self.subq_row[drv["num"]], L, self.subq_row[drv["den"]])
        if t == "valuation":
            v = val_of(row, q)
            f, calc = self._valuation(row, q, L, est=True)
            return f if (f and v is not None and calc is not None and abs(calc - v) <= tol_for(row["key"], row.get("unit") or "배", v, **tol_args(self.rows_by_key, lay.quarters, q))) else None
        if t == "tax":
            dr = self.drv_row.get(row["key"])
            return '=IFERROR(MAX(%s%d,0)*변수!%s%d,"")' % (L, self.subq_row[drv["base"]], L, dr) if (dr and val_of(row, q) is not None) else None
        if t == "interest":
            pq = lay.prev_q(q)
            if pq is None or val_of(row, q) is None or "이자율_자산" not in self.scalar_row:
                return None
            Lp = lay.letter(pq)
            f = "%s%d*변수!$D$%d" % (Lp, self.subq_row[drv["ia"]], self.scalar_row["이자율_자산"])
            if drv.get("debt") and self.fit.get("이자율_부채") is not None:
                f += "-%s%d*변수!$D$%d" % (Lp, self.subq_row[drv["debt"]], self.scalar_row["이자율_부채"])
            return '=IFERROR((%s)/4,"")' % f
        if t == "fxpnl":
            pq = lay.prev_q(q)
            er = self.fx_row.get("USDKRW_end")
            if pq is None or er is None or val_of(row, q) is None:
                return None
            return '=IFERROR(변수!$D$%d*(변수!%s%d-변수!%s%d)/100,"")' % (self.scalar_row[drv["exp"]], L, er, lay.letter(pq), er)
        if t == "sls_link":
            g = self._sls_geometry()
            dr = self.drv_row.get(row["key"])
            if q not in g["qcol"] or not dr or q not in self.sls_link["runoff"]:
                return None
            f = "SLS!%s%d/100+변수!%s%d" % (get_column_letter(g["qcol"][q]), g["krw_link_row"], L, dr)
            if drv.get("new"):
                f += "+%s%d" % (L, self.subq_row[drv["new"]])
            return '=IFERROR(%s,"")' % f
        if t == "ratio_parent":
            pdr = self.drv_row.get(drv["parent"])
            return '=IFERROR(%s%d*변수!%s%d,"")' % (L, self.subq_row[drv["base"]], L, pdr) if (pdr and val_of(row, q) is not None) else None
        dr = self.drv_row.get(row["key"])
        if not dr:
            return None
        if t == "yoy":
            p, ref_q, _ = self._yoy_param(row, q)
            return '=IFERROR(%s%d*(1+변수!%s%d),"")' % (lay.letter(ref_q), r, L, dr) if p is not None else None
        p = implied_param(row, q, drv, self.rows_by_key, lay, self._sheet_base)
        if p is None:
            return None
        if t == "qoq":
            return '=IFERROR(%s%d*(1+변수!%s%d),"")' % (lay.letter(lay.prev_q(q)), r, L, dr)
        if t == "ratio":
            return '=IFERROR(%s%d*변수!%s%d,"")' % (L, self.subq_row[drv["base"]], L, dr)
        if t == "abs":
            return "=변수!%s%d" % (L, dr)
        if t == "roll":
            return '=IFERROR(%s%d+%s%d-변수!%s%d,"")' % (lay.letter(lay.prev_q(q)), r, L, self.subq_row[drv["flow"]], L, dr)
        return None

    def _subq_annual(self, ws, row, c, acct):
        lay = self.lay
        y, L, r = c["key"], get_column_letter(c["col"]), self.subq_row[row["key"]]
        cell = ws.cell(r, c["col"])
        a = (row.get("a") or {}).get(y) or {}
        qs = lay.year_quarters(y)
        stock = is_stock(row)
        t = self.plan[row["key"]]["type"]
        if a.get("kind") == "actual" and acct and not lay.is_est(c):
            if self._bs_get(bs_sheet_for(row), acct, y) is not None:
                bsv = self._bs_get(bs_sheet_for(row), acct, y) / 100
                if a.get("v") is None or abs(bsv - a["v"]) <= 0.011 + abs(a["v"]) * 1e-6:
                    cell.value = self._vlookup(row, y, acct, L)
                    return
            if len(qs) < 4 and a.get("v") is not None:       # BS 에 연간 값이 없고 분기도 미완 → 모델 연간값 직접
                cell.value = a["v"]
                cell.comment = Comment("연간 실측(BS 시트에 %s 없음 — 모델 a 값 직접)" % (y + ".12A"), "kship")
                return
        if len(qs) == 4:
            if stock or (t == "per_share" and row["key"] == "BPS") or (t == "valuation"):
                cell.value = "=%s%d" % (lay.letter(qs[-1]), r)          # 시점·배수 행은 4Q 값
            elif t == "divide":
                cell.value = '=IFERROR(%s%d/%s%d,"")' % (L, self.subq_row[self.plan[row["key"]]["num"]], L, self.subq_row[self.plan[row["key"]]["den"]])
            else:
                cell.value = "=SUM(%s%d:%s%d)" % (lay.letter(qs[0]), r, lay.letter(qs[-1]), r)
            return
        if (stock or t == "valuation") and qs:
            cell.value = "=%s%d" % (lay.letter(qs[-1]), r)
        elif a.get("v") is not None:
            cell.value = a["v"]
            cell.comment = Comment("연간(%s) — 분기 미완 연도, 모델 a 값 직접" % a.get("kind"), "kship")

    # ── SLS ──
    def sheet_sls(self):
        s = self.sls
        if not s:
            return
        ws = self.wb.create_sheet("SLS")
        g = self._sls_geometry()
        byq, byy, cols = s.get("by_quarter") or {}, s.get("by_year") or {}, g["cols"]
        self._header(ws, "이름: SLS_H · 선표 매출인식(백만$) — sls json origin %s" % s.get("origin", ""), cols=cols)
        ws.cell(2, 1, opm_table_text(s) + " — 회사별 캘리브레이션은 target_opm.basis 참조 · "
                      + "코호트 모드 %s(대안 %s) · 잔고 캡 %s" % (s.get("cohort_mode"), s.get("cohort_mode_alt"),
                                                        ("적용 × %.4f" % (s.get("backlog_cap") or {}).get("factor", 1.0)) if s.get("backlog_cap_applied") else "미적용")).font = F_NOTE
        topm = s.get("target_opm") or {}
        r = DATA_ROW - 1
        for key, label, nf in g["specs"]:
            r += 1
            ws.cell(r, 1, key).font = F_KEY
            ws.cell(r, 4, label)
            for c in cols:
                d = (byq if c["kind"] == "q" else byy).get(c["key"]) or {}
                if key.startswith("type:"):
                    v = (d.get("by_type") or {}).get(key[5:])
                elif key.startswith("cohort:"):
                    v = (d.get("by_cohort") or {}).get(key[7:])
                elif key == "target_opm":
                    v = (topm.get(c["key"]) or {}).get("opm")
                else:
                    v = d.get(key)
                if v is not None:
                    cell = ws.cell(r, c["col"], v)
                    cell.number_format = nf
        # 선표 연결 블록 — 계약별 스케줄(값) × 헤지 환율(수식) → 원화. 합계 행이 subQ 매출조선 수식의 피참조
        cts = g["contracts"]
        pc = g["pcol"]
        lk = self.sls_link
        incl_new = bool(self.m.get("new_orders_included"))
        hedge = s.get("hedge") or {}
        ws.cell(g["title_row"], 1, "선표 연결(수식) — 원화(백만원) = 스케줄(백만$) × 포함 × (1 + (캡배율 − 1) × origin 이후) × (헤지비율 × 헤지환율 + (1 − 헤지비율) × 건조시점 환율)").font = F_BOLD
        ws.cell(g["title_row"], 10, ("subQ %s = 합계/100 + 변수 잔고 소진 + %s" % (lk["row"], lk["new"] or "0")) if lk else
                "subQ 매출조선에 연결되지 않음(%s)" % (self.links.get("선표") or {}).get("reason", "")).font = F_NOTE
        ws.cell(g["h_row"], 1, "헤지비율").font = F_KEY
        ws.cell(g["h_row"], 4, "=변수!$D$%d" % self.scalar_row["헤지비율"]).number_format = NF_PCT
        ws.cell(g["h_row"], 5, "변수 3절 '헤지비율' 참조 — %s" % (hedge.get("basis") or "")).font = F_NOTE
        ws.cell(g["mode_row"], 1, "헤지환율 방식").font = F_KEY
        ws.cell(g["mode_row"], 4, "정기보고서 평균 약정환율(measured) %.2f — 전 계약 공통" % hedge["hedge_rate"] if hedge.get("hedge_rate") else "계약별 수주시점 환율(fx_at_sign)")
        ws.cell(g["spot_row"], 1, "spot_krw").font = F_KEY
        ws.cell(g["spot_row"], 4, "건조시점 원/달러(평균) — 변수 환율 행 참조, 변수 범위 밖 분기는 sls spot_assumed")
        ws.cell(g["after_row"], 1, "after_origin").font = F_KEY
        ws.cell(g["after_row"], 4, "origin(%s) 이후 분기 = 1 (잔고 캡 배율 적용 구간)" % (s.get("origin") or ""))
        origin = s.get("origin") or self.lay.last_actual
        for c in cols:
            if c["kind"] != "q":
                continue
            L = get_column_letter(c["col"])
            fb = (byq.get(c["key"]) or {}).get("spot_assumed")
            cell = ws.cell(g["spot_row"], c["col"], '=IFERROR(VLOOKUP(" 원/달러(평균)",환율,MATCH(%s$1,변수H,0),0),%s)' % (L, fb if fb is not None else '""'))
            cell.number_format = NF_FX
            ws.cell(g["after_row"], c["col"], 1 if c["key"] > origin else 0)
        heads = {"incl": "포함(1/0)", "rate": "헤지환율", "factor": "캡배율", "signed": "체결일", "by_origin": "origin이내", "cohort": "코호트", "amt_usd": "계약액(백만$)", "end": "종료"}
        for blk, title in (("usd", "① 계약별 매출인식 스케줄(백만$, sls 저장값)"), ("krw", "② 계약별 원화(백만원, 수식)")):
            hr = g[blk + "_head"]
            ws.cell(hr, 1, title).font = F_BOLD
            for i, h in enumerate(("rcp", "type", "ships", "name")):
                ws.cell(hr, 1 + i, h if i else title).font = F_BOLD
            if blk == "usd":
                for k, col in pc.items():
                    ws.cell(hr, col, heads[k]).font = F_BOLD
                    ws.column_dimensions[get_column_letter(col)].width = 12
        H = "$D$%d" % g["h_row"]
        for i, ct in enumerate(cts):
            ur, kr = g["usd_first"] + i, g["krw_first"] + i
            for rr in (ur, kr):
                ws.cell(rr, 1, ct.get("rcp"))
                ws.cell(rr, 2, ct.get("type"))
                ws.cell(rr, 3, ct.get("ships"))
                ws.cell(rr, 4, ct.get("name"))
            incl = ws.cell(ur, pc["incl"], 1 if self._sls_include(ct, incl_new) else 0)
            incl.fill = FILL_INPUT
            incl.comment = Comment("모델 포함 규칙: 해양(type≠OTHER)%s. 0 으로 바꾸면 그 계약을 매출에서 뺀다" % (" · origin 분기말까지 체결분(이후 체결은 패널 신규수주에 포함으로 봄)" if incl_new else ""), "kship")
            rate = ws.cell(ur, pc["rate"], hedge.get("hedge_rate") if hedge.get("hedge_rate") is not None else ct.get("fx_at_sign"))
            rate.number_format = NF_FX
            rate.fill = FILL_INPUT
            fac = ws.cell(ur, pc["factor"], (s.get("backlog_cap") or {}).get("factor", 1.0) if (s.get("backlog_cap_applied") and ct.get("backlog_cap_applied")) else 1.0)
            fac.number_format = "0.0000"
            ws.cell(ur, pc["signed"], ct.get("signed") or ct.get("start"))
            ws.cell(ur, pc["by_origin"], 1 if ct.get("signed_by_origin", True) else 0)
            ws.cell(ur, pc["cohort"], ct.get("cohort"))
            ws.cell(ur, pc["amt_usd"], ct.get("amt_usd_m")).number_format = NF_1
            ws.cell(ur, pc["end"], ct.get("end"))
            sched = ct.get("schedule") or {}
            for c in cols:
                if c["kind"] != "q":
                    continue
                L = get_column_letter(c["col"])
                v = sched.get(c["key"])
                if v is not None:
                    ws.cell(ur, c["col"], v).number_format = NF_1
                    kcell = ws.cell(kr, c["col"], '=IFERROR(%s%d*$%s%d*(1+($%s%d-1)*%s$%d)*(%s*$%s%d+(1-%s)*%s$%d),"")' % (
                        L, ur, get_column_letter(pc["incl"]), ur, get_column_letter(pc["factor"]), ur, L, g["after_row"],
                        H, get_column_letter(pc["rate"]), ur, H, L, g["spot_row"]))
                    kcell.number_format = NF_INT
        ws.cell(g["usd_incl_row"], 1, "usd_incl").font = F_KEY
        ws.cell(g["usd_incl_row"], 4, "포함 계약 스케줄 합(백만$, 캡 배율 전)")
        ws.cell(g["krw_link_row"], 1, "krw_link").font = F_KEY
        ws.cell(g["krw_link_row"], 4, "원화 합계(백만원) — subQ 연결 행").font = F_BOLD
        PI = get_column_letter(pc["incl"])
        for c in cols:
            if c["kind"] != "q":
                continue
            L = get_column_letter(c["col"])
            if cts:
                ws.cell(g["usd_incl_row"], c["col"], "=SUMPRODUCT(%s%d:%s%d,$%s$%d:$%s$%d)" % (L, g["usd_first"], L, g["usd_last"], PI, g["usd_first"], PI, g["usd_last"])).number_format = NF_1
                ws.cell(g["krw_link_row"], c["col"], "=SUM(%s%d:%s%d)" % (L, g["krw_first"], L, g["krw_last"])).number_format = NF_INT
        # 계약 원장(원문 필드) · 경고
        r = g["ledger_row"]
        ws.cell(r, 1, "계약 원장(contracts.json 기반 — counted 아닌 계약 포함)").font = F_BOLD
        heads2 = ["rcp", "type", "ships", "amt_krw_m", "amt_usd_m", "fx_at_sign", "signed", "start", "end", "cohort", "cohort_alt", "curve", "counted", "signed_by_origin", "estimated_series"]
        r += 1
        for i, h in enumerate(heads2):
            ws.cell(r, 1 + i, h).font = F_BOLD
        for ct in s.get("contracts") or []:
            r += 1
            for i, h in enumerate(heads2):
                v = ct.get(h)
                ws.cell(r, 1 + i, v if not isinstance(v, (dict, list)) else json.dumps(v, ensure_ascii=False))
        for w in s.get("warnings") or []:
            r += 1
            ws.cell(r, 1, "⚠ " + w).font = F_NOTE
        if cols:
            self._name("SLS_H", "SLS!$A$1:$%s$1" % get_column_letter(cols[-1]["col"]))
        self.stats["SLS"] = (r, (max(pc.values()) if cols else FIRST_COL))

    # ── 시나리오(결정 ⓓ — 신규수주 보수/기준/낙관) ──
    def _opm_ref(self, L):
        """시나리오 OP 델타에 쓰는 OPM 셀 — 영업이익 ratio 드라이버, 아니면 신규수주 행이 붙은 세그먼트의 OP 마진 드라이버(모델: 신규분에도 같은 타겟)."""
        if self.plan.get("영업이익", {}).get("type") == "ratio" and "영업이익" in self.drv_row:
            return "변수!%s%d" % (L, self.drv_row["영업이익"])
        for k in self.rows_by_key:
            if k.startswith("매출") and k.endswith("신규"):
                op = "OP" + k[2:-2]
                if self.plan.get(op, {}).get("type") == "ratio" and op in self.drv_row:
                    return "변수!%s%d" % (L, self.drv_row[op])
        return None

    def sheet_scenarios(self):
        sc = self.m.get("scenarios")
        if not sc or not any(k in sc for k in SCN_CASES):
            return
        ws = self.wb.create_sheet("시나리오")
        lay = self.lay
        meta = sc.get("meta") or {}
        fq = [q for q in lay.est_quarters if any(q in ((sc.get(k) or {}).get("quarterly") or {}) for k in SCN_CASES)]
        years = [str(y) for y in (meta.get("fiscal_years") or sorted({q[:4] for q in fq}))]
        cols = build_columns(fq, years)
        self._header(ws, "신규수주 시나리오(억원) — subQ 기준값 ± (시나리오 신규 − base 신규) × OPM", cols=cols)
        ws.cell(2, 1, "출처: %s · calibrated=%s · 행(subQ)에 포함된 시나리오: %s. 보수/낙관/기존만 은 합산하지 않는다. 모델 산출값 — 목표주가·추천 아님%s" % (
            meta.get("source") or "forecast_panel", meta.get("calibrated"), meta.get("in_rows") or "base",
            (" · ※ " + str(meta.get("fallback_note") or "covered_scope_new_revenue 폴백(저신뢰)")) if meta.get("fallback") else "")).font = F_NOTE
        ws.cell(3, 1, str(meta.get("note") or "")).font = F_NOTE
        new_key = next((k for k in self.rows_by_key if k.startswith("매출") and k.endswith("신규")), None)
        r_rev, r_op = self.subq_row.get("매출액"), self.subq_row.get("영업이익")
        base_q = (sc.get("base") or {}).get("quarterly") or {}
        r = DATA_ROW
        base_new_row = None
        self.scn_rows = {}
        order = [k for k in SCN_CASES if k in sc]
        if "base" in order:                                    # base 행이 먼저 있어야 다른 케이스 수식이 참조할 수 있다
            order.remove("base")
            order.insert(0, "base")
        for case in order:
            blk = sc.get(case) or {}
            qd, ann = blk.get("quarterly") or {}, blk.get("annual") or {}
            ws.cell(r, 1, case).font = F_BOLD
            ws.cell(r, 4, "%s%s" % (SCN_LABEL.get(case, case), " — 행에 포함" if blk.get("in_rows") else "")).font = F_BOLD
            r += 1
            rn, rr, ro, rm = r, r + 1, r + 2, r + 3
            self.scn_rows[case] = {"new": rn, "rev": rr, "op": ro, "opm": rm}
            for rr_, key, label, nf in ((rn, "신규수주매출", "신규수주 매출(억원) — %s" % ("subQ %s 참조" % new_key if (case == "base" and new_key) else "패널 값(노란, 가정)"), NF_INT),
                                        (rr, "매출액", "매출액 = subQ 매출액 − base 신규 + 이 시나리오 신규", NF_INT),
                                        (ro, "영업이익", "영업이익 = subQ 영업이익 + (이 시나리오 신규 − base 신규) × OPM(변수 드라이버)", NF_INT),
                                        (rm, "OPM", "OPM", NF_PCT)):
                ws.cell(rr_, 1, "scn:%s:%s" % (case, key)).font = F_KEY
                ws.cell(rr_, 3, "%" if key == "OPM" else "억원")
                ws.cell(rr_, 4, label)
                for c in cols:
                    ws.cell(rr_, c["col"]).number_format = nf
            for c in cols:
                L = get_column_letter(c["col"])
                if c["kind"] == "q":
                    q = c["key"]
                    Lq = lay.letter(q)
                    cell_new = ws.cell(rn, c["col"])
                    if case == "base" and new_key and val_of(self.rows_by_key[new_key], q) is not None:
                        cell_new.value = "=subQ!%s%d" % (Lq, self.subq_row[new_key])
                        cell_new.fill = FILL_LINK
                    else:
                        nv = (qd.get(q) or {}).get("new_order_revenue")
                        if nv is None and case == "existing_only":
                            nv = 0.0
                        if nv is not None:
                            cell_new.value = nv
                            cell_new.fill = FILL_INPUT
                            cell_new.comment = Comment("forecast_panel %s %s(억원)%s — %s" % (case, meta.get("source_field") or "new_order_revenue", " · 폴백·저신뢰" if meta.get("fallback") else "", meta.get("source") or ""), "kship")
                    if q not in qd:
                        continue
                    bn = "%s%d" % (L, base_new_row if base_new_row else rn)
                    if r_rev:
                        ws.cell(rr, c["col"], '=IFERROR(subQ!%s%d-%s+%s%d,"")' % (Lq, r_rev, bn, L, rn))
                    opm = self._opm_ref(Lq)
                    if r_op and opm:
                        ws.cell(ro, c["col"], '=IFERROR(subQ!%s%d+(%s%d-%s)*%s,"")' % (Lq, r_op, L, rn, bn, opm))
                    elif (qd.get(q) or {}).get("op") is not None:
                        ws.cell(ro, c["col"], qd[q]["op"]).comment = Comment("OPM 드라이버 없음 — 모델 scenarios 값 직접", "kship")
                    ws.cell(rm, c["col"], '=IFERROR(%s%d/%s%d,"")' % (L, ro, L, rr))
                else:
                    y = c["key"]
                    est_qs = [q for q in fq if q[:4] == y]
                    act_qs = [q for q in lay.year_quarters(y) if q <= lay.last_actual]
                    if len(est_qs) + len(act_qs) < 4 or not est_qs:
                        continue
                    L1, L2 = get_column_letter(cols[[cc["key"] for cc in cols].index(est_qs[0])]["col"]), get_column_letter(cols[[cc["key"] for cc in cols].index(est_qs[-1])]["col"])
                    ws.cell(rn, c["col"], "=SUM(%s%d:%s%d)" % (L1, rn, L2, rn))
                    for row_i, sub_r in ((rr, r_rev), (ro, r_op)):
                        if not sub_r:
                            continue
                        acts = "".join("+subQ!%s%d" % (lay.letter(q), sub_r) for q in act_qs)
                        ws.cell(row_i, c["col"], '=IFERROR(SUM(%s%d:%s%d)%s,"")' % (L1, row_i, L2, row_i, acts))
                    ws.cell(rm, c["col"], '=IFERROR(%s%d/%s%d,"")' % (L, ro, L, rr))
            if case == "base":
                base_new_row = rn
            r += 5
        for c in cols:
            if lay.is_est(c):
                for rr_ in range(DATA_ROW, r):
                    cell = ws.cell(rr_, c["col"])
                    if cell.fill.fgColor.rgb in (None, "00000000"):
                        cell.fill = FILL_EST
        self.stats["시나리오"] = (r - 1, cols[-1]["col"] if cols else FIRST_COL)

    # ── 분기 / 연간예상 ──
    def _view_keys(self, view):
        v = (self.m.get("views") or {}).get(view)
        keys = None
        if isinstance(v, list) and v and isinstance(v[0], list) and all(isinstance(x, str) for x in v[0]):
            keys = [k for k in v[0] if k in self.subq_row]
        if not keys:
            keys = [r["key"] for r in self.rows]
        return keys + [k for k in ("OPM", "NPM") if k in self.subq_row and k not in keys]

    def sheet_view(self, name, kind):
        ws = self.wb.create_sheet(name)
        lay = self.lay
        cols = [dict(c, col=FIRST_COL + i) for i, c in enumerate(lay.qcols if kind == "q" else lay.acols)]
        self._header(ws, "이름: %sH — subQ 를 이름으로 VLOOKUP(억원)" % name, cols=cols)
        price = (self.m.get("valuation") or {}).get("price") or {}
        ws.cell(2, 3, "종가:")
        ws.cell(2, 4, "=변수!D%d" % self.scalar_row["종가"]).number_format = NF_INT
        ws.cell(3, 3, "기준일:")
        ws.cell(3, 4, price.get("as_of"))
        r = DATA_ROW - 1
        for k in self._view_keys(name):
            r += 1
            row = self.rows_by_key.get(k) or {"key": k, "label": k, "unit": "%"}
            ws.cell(r, 1, k).font = F_KEY
            ws.cell(r, 4, row.get("label") or k)
            unit = row.get("unit") or "억원"
            nf = NF_PCT if (k in ("OPM", "NPM") or unit == "%") else (NF_1 if k == "주식수" else (NF_FX if unit == "배" else NF_INT))
            for c in cols:
                L = get_column_letter(c["col"])
                cell = ws.cell(r, c["col"], '=IFERROR(VLOOKUP($A%d,subQ,MATCH(%s$1,SUBQH,0),0),"")' % (r, L))
                cell.number_format = nf
        # YoY 행(매출액·영업이익·지배주주순이익)
        for k in ("매출액", "영업이익", "지배주주순이익"):
            if k not in self.subq_row:
                continue
            r += 1
            ws.cell(r, 1, k + " YoY").font = F_KEY
            ws.cell(r, 4, k + " YoY")
            base_r = [rr for rr in range(DATA_ROW, r) if ws.cell(rr, 1).value == k][0]
            step = 4 if kind == "q" else 1
            for i, c in enumerate(cols):
                if i < step:
                    continue
                L, LP = get_column_letter(c["col"]), get_column_letter(cols[i - step]["col"])
                cell = ws.cell(r, c["col"], '=IFERROR(%s%d/%s%d-1,"")' % (L, base_r, LP, base_r))
                cell.number_format = NF_PCT
        # 음영(열 배치가 다르므로 직접)
        for c in cols:
            if lay.is_est(c):
                for rr in range(DATA_ROW, r + 1):
                    ws.cell(rr, c["col"]).fill = FILL_EST
            elif c["kind"] == "q" and c["key"] == lay.last_actual or (c["kind"] == "a" and int(c["key"]) == lay.last_year - 1):
                for rr in range(HEADER_ROW, r + 1):
                    cell = ws.cell(rr, c["col"])
                    cell.border = Border(right=MED)
        self.stats[name] = (r, cols[-1]["col"] if cols else FIRST_COL)

    # ── TP ──
    def sheet_tp(self):
        ws = self.wb.create_sheet("TP")
        lay = self.lay
        sc = self.scalar_row
        ws.cell(1, 2, "종가").font = F_BOLD
        ws.cell(1, 3, "=변수!D%d" % sc["종가"]).number_format = NF_INT
        ws.cell(2, 2, "기준일").font = F_BOLD
        ws.cell(2, 3, "=변수!D%d" % sc["기준일"])
        ws.cell(3, 2, "모델 산출값 — 목표주가·추천 아님. 밴드·COE 는 변수 시트 3절의 가정").font = F_NOTE
        y0 = lay.last_year
        years = [str(y) for y in (y0 - 1, y0, y0 + 1, y0 + 2) if str(y) in lay.annual]
        ws.cell(5, 3, "FY").font = F_BOLD
        for i, y in enumerate(years):
            c = ws.cell(5, 4 + i, int(y))
            c.font = F_HEAD
            c.fill = FILL_HEAD
            ws.cell(6, 4 + i, ("A" if int(y) < y0 else "E")).font = F_NOTE
        rows = [("EPS", "EPS(원)", NF_INT), ("BPS", "BPS(원)", NF_INT), ("지배주주순이익", "지배주주순이익(억원)", NF_INT), ("영업이익", "영업이익(억원)", NF_INT)]
        r = 6
        tp_row = {}
        for key, label, nf in rows:
            if key not in self.subq_row:
                continue
            r += 1
            tp_row[key] = r
            ws.cell(r, 3, label)
            for i, y in enumerate(years):
                L = get_column_letter(4 + i)
                ws.cell(r, 4 + i, '=IFERROR(VLOOKUP("%s",subQ,MATCH(%s$5,SUBQH,0),0),"")' % (key, L)).number_format = nf
        if "EPS" in tp_row and "BPS" in tp_row:
            r += 1
            ws.cell(r, 3, "ROE(EPS/기초BPS)")
            for i in range(1, len(years)):
                L, LP = get_column_letter(4 + i), get_column_letter(3 + i)
                ws.cell(r, 4 + i, '=IFERROR(%s%d/%s%d,"")' % (L, tp_row["EPS"], LP, tp_row["BPS"])).number_format = NF_PCT
        # 12M fwd EPS = FY(y0)E × w + FY(y0+1)E × (1−w)
        iy0 = years.index(str(y0)) if str(y0) in years else None
        r += 2
        ws.cell(r, 2, "PER").font = F_BOLD
        r += 1
        ws.cell(r, 3, "12M fwd EPS")
        if iy0 is not None and iy0 + 1 < len(years) and "EPS" in tp_row:
            L0, L1 = get_column_letter(4 + iy0), get_column_letter(5 + iy0)
            ws.cell(r, 4, '=IFERROR(%s%d*변수!D%d+%s%d*(1-변수!D%d),"")' % (L0, tp_row["EPS"], sc["EPS_Y1가중"], L1, tp_row["EPS"], sc["EPS_Y1가중"])).number_format = NF_INT
        eps_r = r
        for band in ("PER_lo", "PER_mid", "PER_hi"):
            r += 1
            ws.cell(r, 3, "적정가치 @%s" % band)
            ws.cell(r, 4, "=변수!D%d" % sc[band]).number_format = NF_1
            ws.cell(r, 5, '=IFERROR(ROUND($D$%d*D%d,-2),"")' % (eps_r, r)).number_format = NF_INT
            ws.cell(r, 6, '=IFERROR(E%d/$C$1-1,"")' % r).number_format = NF_PCT
            ws.cell(r, 7, "배 · 원 · 종가 대비").font = F_NOTE
        r += 1
        ws.cell(r, 3, "현재 PER(종가/12M fwd EPS)")
        ws.cell(r, 4, '=IFERROR($C$1/$D$%d,"")' % eps_r).number_format = NF_1
        r += 2
        ws.cell(r, 2, "PBR").font = F_BOLD
        r += 1
        ws.cell(r, 3, "최근 BPS(마지막 실적 분기)")
        ws.cell(r, 4, '=IFERROR(VLOOKUP("BPS",subQ,MATCH("%s",SUBQH,0),0),"")' % qlabel(lay.last_actual)).number_format = NF_INT
        bps_r = r
        r += 1
        ws.cell(r, 3, "ROE fwd(12M fwd EPS / BPS)")
        ws.cell(r, 4, '=IFERROR($D$%d/$D$%d,"")' % (eps_r, bps_r)).number_format = NF_PCT
        roe_r = r
        r += 1
        ws.cell(r, 3, "COE")
        ws.cell(r, 4, "=변수!D%d" % sc["COE"]).number_format = NF_PCT
        coe_r = r
        r += 1
        ws.cell(r, 3, "적정 PBR = ROE/COE")
        ws.cell(r, 4, '=IFERROR(D%d/D%d,"")' % (roe_r, coe_r)).number_format = NF_FX
        fpbr_r = r
        r += 1
        ws.cell(r, 3, "적정가치(PBR) = BPS × 적정PBR")
        ws.cell(r, 4, '=IFERROR(ROUND(D%d*D%d,-2),"")' % (bps_r, fpbr_r)).number_format = NF_INT
        ws.cell(r, 5, '=IFERROR(D%d/$C$1-1,"")' % r).number_format = NF_PCT
        r += 1
        ws.cell(r, 3, "현재 PBR(종가/BPS)")
        ws.cell(r, 4, '=IFERROR($C$1/D%d,"")' % bps_r).number_format = NF_FX
        val = self.m.get("valuation") or {}
        r += 2
        ws.cell(r, 2, "모델 json valuation(참고)").font = F_NOTE
        for k in ("eps_fwd12m", "bps_latest", "per_now", "pbr_now", "roe_fwd", "fair_pbr", "fair_value_pbr", "ev_ebitda", "note"):
            r += 1
            ws.cell(r, 3, k).font = F_NOTE
            v = val.get(k)
            ws.cell(r, 4, v if not isinstance(v, (dict, list)) else json.dumps(v, ensure_ascii=False))
        ws.column_dimensions["C"].width = 34
        self.stats["TP"] = (r, 7)

    # ── 외화 ──
    def sheet_fx(self):
        ws = self.wb.create_sheet("외화")
        lay = self.lay
        cols = [dict(c, col=FIRST_COL + i) for i, c in enumerate(lay.qcols)]
        self._header(ws, "이름: 외화QH — 환율 QoQ · 외화 순노출 × Δ기말환율", cols=cols)
        fxp = self.m.get("fx_pnl") or {}
        specs = [(" 원/달러(평균)", True), (" 원/달러(기말)", True), (" W/Euro(평균)", True), (" W/Euro(기말)", True), (" 원/100Y(기말)", True)]
        r = DATA_ROW - 1
        end_r = None
        for label, _ in specs:
            r += 1
            ws.cell(r, 4, label)
            for c in cols:
                L = get_column_letter(c["col"])
                ws.cell(r, c["col"], '=IFERROR(VLOOKUP($D%d,환율,MATCH(%s$1,변수H,0),0),"")' % (r, L)).number_format = NF_FX
            if label == " 원/달러(기말)":
                end_r = r
            r += 1
            ws.cell(r, 1, label.strip() + " QoQ").font = F_KEY
            for i, c in enumerate(cols[1:], 1):
                L, LP = get_column_letter(c["col"]), get_column_letter(cols[i - 1]["col"])
                ws.cell(r, c["col"], '=IFERROR(%s%d/%s%d-1,"")' % (L, r - 1, LP, r - 1)).number_format = NF_PCT
        r += 2
        ws.cell(r, 1, "외화순노출").font = F_KEY
        ws.cell(r, 3, "외화 순노출(백만$)")
        ws.cell(r, 4, "=변수!D%d" % self.scalar_row["외화순노출_백만$"]).number_format = NF_1
        ws.cell(r, 5, fxp.get("method") or "").font = F_NOTE
        exp_r = r
        r += 1
        ws.cell(r, 1, "환관련손익추정").font = F_KEY
        ws.cell(r, 3, "환관련손익 추정(억원) = 순노출 × Δ원/달러(기말) / 100%s" % (" — subQ 환관련손익 행과 같은 수식(변수 환율·노출 참조)" if self.plan.get("환관련손익", {}).get("type") == "fxpnl" else ""))
        for i, c in enumerate(cols[1:], 1):
            L, LP = get_column_letter(c["col"]), get_column_letter(cols[i - 1]["col"])
            ws.cell(r, c["col"], '=IFERROR($D$%d*(%s%d-%s%d)/100,"")' % (exp_r, L, end_r, LP, end_r)).number_format = NF_INT
        r += 1
        ws.cell(r, 1, "fx_pnl.by_q").font = F_KEY
        ws.cell(r, 3, "모델 fx_pnl.by_q(억원, 참고)")
        for c in cols:
            v = (fxp.get("by_q") or {}).get(c["key"])
            if isinstance(v, dict):
                v = v.get("v")
            if v is not None:
                ws.cell(r, c["col"], v).number_format = NF_INT
        if fxp.get("note"):
            r += 1
            ws.cell(r, 3, fxp["note"]).font = F_NOTE
        for c in cols:
            if lay.is_est(c):
                for rr in range(DATA_ROW, r + 1):
                    ws.cell(rr, c["col"]).fill = FILL_EST
        self.stats["외화"] = (r, cols[-1]["col"] if cols else FIRST_COL)

    def build(self):
        self.sheet_readme()
        self.sheet_vars()
        self.sheet_bs("cons")
        self.sheet_bs("sep")
        self.sheet_subq()
        self.sheet_sls()                 # fill_driver_cells 의 2차 패스(roll)가 시트 수식을 평가하므로 subQ 매출조선이 참조하는 SLS 시트가 먼저 있어야 한다
        self.fill_driver_cells()
        self.sheet_scenarios()
        self.sheet_view("분기", "q")
        self.sheet_view("연간예상", "a")
        self.sheet_tp()
        self.sheet_fx()
        # 결정론: 문서 속성 시각을 모델 built_at 날짜로 고정. openpyxl 은 save 때 modified 를 now 로 덮고 zip 엔트리 시각도
        # 실행 시각을 쓰므로 save_atomic 이 created 기준으로 다시 정규화한다(같은 입력 → 같은 바이트)
        d = self.built_date()
        dt = datetime.datetime(d.year, d.month, d.day)
        self.wb.properties.created = dt
        self.wb.properties.modified = dt
        self.wb.properties.creator = "kship_model_xlsx"
        self.wb.calculation.fullCalcOnLoad = True
        return self.wb


# ── 저장·검증 ───────────────────────────────────────────────

def _normalize_zip(src, dst, dt):
    """openpyxl 산출 zip 을 결정론적으로 다시 쓴다 — 모든 엔트리 시각 = dt, docProps/core.xml 의 dcterms:modified = created.
    (openpyxl writer 는 save 시 properties.modified 를 now 로 덮고 ZipInfo 시각에 실행 시각을 넣는다.)"""
    import zipfile
    stamp = (dt.year, dt.month, dt.day, 0, 0, 0)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                m = re.search(r"<dcterms:created[^>]*>([^<]*)</dcterms:created>", text)
                if m:
                    text = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", lambda mm: mm.group(1) + m.group(1) + mm.group(2), text)
                data = text.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=stamp)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            zout.writestr(zi, data)


def save_atomic(wb, path):
    d = os.path.dirname(os.path.abspath(path))
    os.makedirs(d, exist_ok=True)
    tmp = os.path.join(d, ".kship_tmp_" + os.path.basename(path))
    raw = tmp + ".raw"
    wb.save(raw)
    try:
        _normalize_zip(raw, tmp, wb.properties.created or datetime.datetime(2026, 1, 1))
    finally:
        os.remove(raw)
    os.replace(tmp, path)
    return os.path.getsize(path)


def load_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_optional(path):
    return load_json(path) if path and os.path.exists(path) else None


class XlError(Exception):
    """엑셀 오류값(#N/A·#DIV/0! …) — IFERROR 가 잡는다."""


class Emulator:
    """생성 파일 검증용 초소형 수식 계산기(LibreOffice 없음).
    지원: 숫자·문자열·셀/범위 참조(시트!·$)·이름정의·+ - * / ^ & 비교·IFERROR·IF·VLOOKUP(정확일치)·
    MATCH(0)·SUM·SUMPRODUCT·ROUND·ABS·MIN·MAX. 그 밖은 XlError."""

    TOK = re.compile(r"""\s*(?:(?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?)|(?P<str>"(?:[^"]|"")*")|
        (?P<ref>(?:'[^']+'|[^\s!(),:;+\-*/^&<>=]+)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)|
        (?P<name>[^\s!(),:;+\-*/^&<>="']+)|(?P<op><>|<=|>=|[-+*/^&<>=(),:]))""", re.X)

    def __init__(self, wb):
        self.wb = wb
        self.names = {n: dn.attr_text for n, dn in wb.defined_names.items()}
        self.memo = {}
        self.stack = set()

    # 참조 해석
    def _split_ref(self, ref, cur_sheet):
        if "!" in ref:
            sh, rng = ref.rsplit("!", 1)
            sh = sh.strip("'")
        else:
            sh, rng = cur_sheet, ref
        rng = rng.replace("$", "")
        return sh, rng

    def _range_cells(self, sheet, rng):
        if ":" in rng:
            a, b = rng.split(":")
            c0, r0 = self._coord(a)
            c1, r1 = self._coord(b)
            return [[(sheet, get_column_letter(c) + str(r)) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)]
        return [[(sheet, rng)]]

    @staticmethod
    def _coord(a):
        m = re.fullmatch(r"([A-Z]{1,3})(\d+)", a)
        col = 0
        for ch in m.group(1):
            col = col * 26 + ord(ch) - 64
        return col, int(m.group(2))

    def value(self, sheet, coord):
        key = (sheet, coord)
        if key in self.memo:
            return self.memo[key]
        if key in self.stack:
            raise XlError("circular %s!%s" % key)
        v = self.wb[sheet][coord].value
        if isinstance(v, str) and v.startswith("="):
            # 중첩 평가 — 바깥 파서 상태(toks·i·sheet)를 보존한다
            saved = (getattr(self, "toks", None), getattr(self, "i", 0), getattr(self, "sheet", None))
            self.stack.add(key)
            try:
                v = self.eval(v[1:], sheet)
            finally:
                self.stack.discard(key)
                self.toks, self.i, self.sheet = saved
        self.memo[key] = v
        return v

    # 파서(재귀 하강)
    def eval(self, text, sheet):
        self.toks = [(m.lastgroup, m.group(m.lastgroup)) for m in self.TOK.finditer(text) if m.lastgroup]
        self.i = 0
        self.sheet = sheet
        v = self._cmp()
        return v

    def _peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def _next(self):
        t = self.toks[self.i]
        self.i += 1
        return t

    def _cmp(self):
        a = self._concat()
        while self._peek()[0] == "op" and self._peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self._next()[1]
            b = self._concat()
            a = {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]
        return a

    def _concat(self):
        a = self._add()
        while self._peek() == ("op", "&"):
            self._next()
            a = str(a) + str(self._add())
        return a

    def _add(self):
        a = self._mul()
        while self._peek()[0] == "op" and self._peek()[1] in "+-":
            op = self._next()[1]
            b = self._mul()
            a = self._n(a) + self._n(b) if op == "+" else self._n(a) - self._n(b)
        return a

    def _mul(self):
        a = self._pow()
        while self._peek()[0] == "op" and self._peek()[1] in "*/":
            op = self._next()[1]
            b = self._pow()
            if op == "*":
                a = self._n(a) * self._n(b)
            else:
                if self._n(b) == 0:
                    raise XlError("#DIV/0!")
                a = self._n(a) / self._n(b)
        return a

    def _pow(self):
        a = self._unary()
        while self._peek() == ("op", "^"):
            self._next()
            a = self._n(a) ** self._n(self._unary())
        return a

    def _unary(self):
        if self._peek() == ("op", "-"):
            self._next()
            return -self._n(self._unary())
        if self._peek() == ("op", "+"):
            self._next()
            return self._unary()
        return self._atom()

    @staticmethod
    def _n(v):
        if v is None or v == "":
            return 0
        if isinstance(v, bool):
            return int(v)
        if isinstance(v, str):
            raise XlError("#VALUE!")
        return v

    def _atom(self):
        kind, tok = self._next()
        if kind == "num":
            return float(tok) if ("." in tok or "e" in tok.lower()) else int(tok)
        if kind == "str":
            return tok[1:-1].replace('""', '"')
        if kind == "op" and tok == "(":
            v = self._cmp()
            assert self._next() == ("op", ")")
            return v
        if kind == "ref":
            sh, rng = self._split_ref(tok, self.sheet)
            if ":" in rng:
                return ("range", sh, rng)
            return self.value(sh, rng)
        if kind == "name":
            if self._peek() == ("op", "("):
                self._next()
                args = []
                if self._peek() != ("op", ")"):
                    args.append(self._cmp_or_range())
                    while self._peek() == ("op", ","):
                        self._next()
                        args.append(self._cmp_or_range())
                assert self._next() == ("op", ")")
                return self._call(tok.upper(), args)
            if tok in self.names:
                sh, rng = self._split_ref(self.names[tok], self.sheet)
                return ("range", sh, rng) if ":" in rng else self.value(sh, rng)
            raise XlError("#NAME? " + tok)
        raise XlError("parse " + str(tok))

    def _cmp_or_range(self):
        # IFERROR 는 첫 인자 오류를 잡아야 하므로 지연 평가 대신 예외를 값으로 돌린다
        save = self.i
        try:
            return self._cmp()
        except XlError as e:
            # 인자 끝까지 토큰을 소비
            depth = 0
            self.i = save
            while self.i < len(self.toks):
                k, t = self.toks[self.i]
                if k == "op" and t == "(":
                    depth += 1
                elif k == "op" and t == ")":
                    if depth == 0:
                        break
                    depth -= 1
                elif k == "op" and t == "," and depth == 0:
                    break
                self.i += 1
            return e

    def _cells(self, arg):
        if isinstance(arg, tuple) and arg[0] == "range":
            return self._range_cells(arg[1], arg[2])
        raise XlError("#VALUE! range expected")

    def _call(self, fn, args):
        if fn == "IFERROR":
            return args[1] if isinstance(args[0], XlError) else args[0]
        for a in args:
            if isinstance(a, XlError):
                raise a
        if fn == "IF":
            return args[1] if args[0] else (args[2] if len(args) > 2 else False)
        if fn == "SUM":
            s = 0
            for a in args:
                if isinstance(a, tuple):
                    for rowc in self._cells(a):
                        for sh, co in rowc:
                            v = self.value(sh, co)
                            if isinstance(v, (int, float)) and not isinstance(v, bool):
                                s += v
                else:
                    s += self._n(a)
            return s
        if fn == "SUMPRODUCT":
            cells = [[c for rowc in self._cells(a) for c in rowc] for a in args]
            if len({len(x) for x in cells}) != 1:
                raise XlError("#VALUE! SUMPRODUCT shape")
            s = 0
            for tup in zip(*cells):
                p = 1
                for sh, co in tup:
                    v = self.value(sh, co)
                    p *= v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0
                s += p
            return s
        if fn == "ROUND":
            return round(self._n(args[0]), int(args[1]))
        if fn == "ABS":
            return abs(self._n(args[0]))
        if fn == "MIN":
            return min(self._n(a) for a in args)
        if fn == "MAX":
            return max(self._n(a) for a in args)
        if fn == "MATCH":
            look, rng = args[0], args[1]
            cells = self._cells(rng)
            flat = [c for rowc in cells for c in rowc]
            for i, (sh, co) in enumerate(flat, 1):
                if self._eq(self.value(sh, co), look):
                    return i
            raise XlError("#N/A MATCH %r" % (look,))
        if fn == "VLOOKUP":
            look, rng, idx = args[0], args[1], int(self._n(args[2]))
            for rowc in self._cells(rng):
                sh, co = rowc[0]
                if self._eq(self.value(sh, co), look):
                    sh2, co2 = rowc[idx - 1]
                    v = self.value(sh2, co2)
                    return 0 if v is None else v          # 엑셀: 빈 셀 → 0
            raise XlError("#N/A VLOOKUP %r" % (look,))
        raise XlError("#NAME? " + fn)

    @staticmethod
    def _eq(a, b):
        if isinstance(a, str) and isinstance(b, str):
            return a.strip().lower() == b.strip().lower()
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            return float(a) == float(b)
        return a == b


def _is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def _find_rows(ws, col=1, start=DATA_ROW):
    out = {}
    for r in range(start, ws.max_row + 1):
        k = ws.cell(r, col).value
        if isinstance(k, str) and k not in out:
            out[k] = r
    return out


def _probe(wb, changes, targets):
    """가정 셀을 바꾸고 종속 셀이 움직이는지 — (전, 후). 바꾼 셀은 원복한다."""
    em0 = Emulator(wb)
    before = [em0.value(sh, co) for sh, co in targets]
    saved = {(sh, co): wb[sh][co].value for sh, co in changes}
    for (sh, co), nv in changes.items():
        wb[sh][co].value = nv
    try:
        em1 = Emulator(wb)
        after = [em1.value(sh, co) for sh, co in targets]
    finally:
        for (sh, co), ov in saved.items():
            wb[sh][co].value = ov
    return before, after


def _num_or_none(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def verify(path, model, keys=("매출액", "영업이익", "지배주주순이익", "자산총계"), quarters_after=None):
    """재오픈 → 시트·이름정의·수식 개수 · VLOOKUP 흉내(마지막 실적 분기) · 추정 전 분기 재계산(quarters_after=None 이면 전부) ·
    실적 구간 수식 셀 재계산 · 추정 셀 수식/값 인구조사(links) · 시나리오 시트 대조 · 민감도(세율·OPM·환율·이자율·신규수주를 실제로 바꿔 종속 셀 확인)."""
    wb = load_workbook(path, data_only=False)
    rep = {"file": path, "size_bytes": os.path.getsize(path), "sheets": {}, "defined_names": {}, "formulas_total": 0, "checks": []}
    for ws in wb.worksheets:
        nf = sum(1 for row in ws.iter_rows() for c in row if _is_formula(c.value))
        ncomment = sum(1 for row in ws.iter_rows() for c in row if c.comment is not None)
        rep["sheets"][ws.title] = {"rows": ws.max_row, "cols": ws.max_column, "formulas": nf, "comments": ncomment}
        rep["formulas_total"] += nf
    for n, dn in wb.defined_names.items():
        rep["defined_names"][n] = dn.attr_text
    em = Emulator(wb)
    lay = Layout(model)
    ws = wb["subQ"]
    rows_by_key = {r["key"]: r for r in model["rows"]}
    subq_row = _find_rows(ws)
    la = lay.last_actual
    for k in keys:
        if k not in subq_row:
            continue
        r = subq_row[k]
        coord = "%s%d" % (lay.letter(la), r)
        f = ws[coord].value
        model_v = val_of(rows_by_key[k], la)
        chk = {"cell": "subQ!" + coord, "quarter": la, "formula": f, "model_v": model_v}
        try:
            got = em.value("subQ", coord)
            chk["emulated"] = got
            if model_v is None:
                chk["ok"] = got in (None, "")        # 모델에 값이 없는 확정 셀 — 시트도 비어 있어야 한다(VLOOKUP 0 표시 금지)
                chk["no_model_value"] = True
            else:
                chk["ok"] = isinstance(got, (int, float)) and abs(got - model_v) < 0.01
            m = re.search(r"VLOOKUP\(\$D(\d+),(BS[^,]+),MATCH\(([A-Z]+)\$1,(BS[^,]+)H,0\),0\)", f if isinstance(f, str) else "")
            if m:
                acct = ws.cell(int(m.group(1)), 4).value
                bs = wb[m.group(2)]
                hdr = {bs.cell(1, c).value: c for c in range(3, bs.max_column + 1)}
                col = hdr.get(ws.cell(1, ws[coord].column).value)
                acct_row = next((rr for rr in range(DATA_ROW, bs.max_row + 1) if bs.cell(rr, 3).value == acct), None)
                chk["vlookup_trace"] = {"acct": acct, "bs_sheet": m.group(2), "bs_col": get_column_letter(col) if col else None,
                                        "bs_row": acct_row, "bs_value_million": bs.cell(acct_row, col).value if (col and acct_row) else None}
        except Exception as e:  # noqa: BLE001
            chk["error"] = repr(e)
            chk["ok"] = False
        rep["checks"].append(chk)
    # 추정 열 재계산 — last_actual 다음 분기부터(기본 전부) + 추정 셀 인구조사(수식/값)
    est_q = lay.est_quarters if quarters_after is None else lay.est_quarters[:quarters_after]
    est = {"quarters": est_q, "rows": 0, "match": 0, "mismatch": []}
    links = {}
    for q in est_q:
        L = lay.letter(q)
        for k, r in subq_row.items():
            if k not in rows_by_key:
                continue
            mv = val_of(rows_by_key[k], q)
            if mv is None:
                continue
            est["rows"] += 1
            cell = ws["%s%d" % (L, r)]
            lk = links.setdefault(k, {"formula": 0, "value": 0})
            lk["formula" if _is_formula(cell.value) else "value"] += 1
            try:
                got = em.value("subQ", "%s%d" % (L, r))
            except Exception as e:  # noqa: BLE001
                got = repr(e)
            tol = tol_for(k, rows_by_key[k].get("unit") or "억원", mv, **tol_args(rows_by_key, lay.quarters, q))
            if isinstance(got, (int, float)) and abs(got - mv) <= tol:
                est["match"] += 1
            else:
                est["mismatch"].append({"key": k, "q": q, "model": mv, "emulated": got})
    rep["estimate_recalc"] = est
    rep["links"] = {"by_key": links, "formula_cells": sum(v["formula"] for v in links.values()), "value_cells": sum(v["value"] for v in links.values()),
                    "value_keys": sorted(k for k, v in links.items() if v["value"])}
    # 실적 구간 — 수식 셀(VLOOKUP·복합·파생)을 재계산해 모델값과 대조
    act = {"formula_cells": 0, "match": 0, "mismatch": []}
    for q in [q for q in lay.quarters if q <= la]:
        L = lay.letter(q)
        for k, r in subq_row.items():
            if k not in rows_by_key:
                continue
            mv = val_of(rows_by_key[k], q)
            cell = ws["%s%d" % (L, r)]
            if mv is None or not _is_formula(cell.value):
                continue
            act["formula_cells"] += 1
            try:
                got = em.value("subQ", "%s%d" % (L, r))
            except Exception as e:  # noqa: BLE001
                got = repr(e)
            tol = tol_for(k, rows_by_key[k].get("unit") or "억원", mv, **tol_args(rows_by_key, lay.quarters, q))
            if isinstance(got, (int, float)) and abs(got - mv) <= tol:
                act["match"] += 1
            else:
                act["mismatch"].append({"key": k, "q": q, "model": mv, "emulated": got})
    rep["actual_recalc"] = act
    # 시나리오 시트 — 모델 scenarios 와 대조(분기: 억원 규칙 ±0.05 또는 1e-4 상대, 연간 ±0.25: 분기 셀 반올림 합)
    sc = model.get("scenarios") or {}
    if "시나리오" in wb.sheetnames and sc:
        wsn = wb["시나리오"]
        srow = _find_rows(wsn)
        hdr = {wsn.cell(1, c).value: c for c in range(FIRST_COL, wsn.max_column + 1)}
        scn = {"cells": 0, "match": 0, "mismatch": []}
        for case in SCN_CASES:
            blk = sc.get(case) or {}
            for q, d in (blk.get("quarterly") or {}).items():
                col = hdr.get(qlabel(q))
                for key, mk in (("매출액", "rev"), ("영업이익", "op")):
                    r = srow.get("scn:%s:%s" % (case, key))
                    if not r or col is None or d.get(mk) is None:
                        continue
                    scn["cells"] += 1
                    try:
                        got = em.value("시나리오", "%s%d" % (get_column_letter(col), r))
                    except Exception as e:  # noqa: BLE001
                        got = repr(e)
                    if isinstance(got, (int, float)) and abs(got - d[mk]) <= tol_for(key, "억원", d[mk]):
                        scn["match"] += 1
                    else:
                        scn["mismatch"].append({"case": case, "q": q, "key": key, "model": d[mk], "emulated": got})
            for y, d in (blk.get("annual") or {}).items():
                col = hdr.get(int(y))
                for key, mk in (("매출액", "rev"), ("영업이익", "op")):
                    r = srow.get("scn:%s:%s" % (case, key))
                    if not r or col is None or d.get(mk) is None:
                        continue
                    scn["cells"] += 1
                    try:
                        got = em.value("시나리오", "%s%d" % (get_column_letter(col), r))
                    except Exception as e:  # noqa: BLE001
                        got = repr(e)
                    if isinstance(got, (int, float)) and abs(got - d[mk]) <= max(0.25, tol_for(key, "억원", d[mk])):
                        scn["match"] += 1
                    else:
                        scn["mismatch"].append({"case": case, "q": y, "key": key, "model": d[mk], "emulated": got})
        rep["scenarios"] = scn
    # 민감도 — 가정 셀을 실제로 바꿔 종속 셀이 기대만큼 움직이는지
    rep["sensitivity"] = _sensitivity(wb, model, lay, subq_row, em)
    rep["size_ok"] = rep["size_bytes"] <= SIZE_LIMIT
    rep["ok"] = (all(c.get("ok") for c in rep["checks"]) and not est["mismatch"] and not act["mismatch"]
                 and not (rep.get("scenarios") or {}).get("mismatch") and all(p["ok"] for p in rep["sensitivity"]) and rep["size_ok"])
    return rep


def _sensitivity(wb, model, lay, subq_row, em):
    """세율·OPM·환율(평균→선표, 기말→환관련)·이자율·신규수주 가정 셀을 바꾼 뒤 종속 셀 변화량이 기대값과 맞는지. 각 항목 {name, changed, target, before, after, expected_delta, ok}."""
    out = []
    vs = wb["변수"]
    vrow = _find_rows(vs)
    fq = lay.est_quarters
    if not fq:
        return out
    q1 = fq[0]
    q2 = fq[1] if len(fq) > 1 else fq[0]
    L1, L2 = lay.letter(q1), lay.letter(q2)

    def num(sh, co):
        try:
            return _num_or_none(em.value(sh, co))
        except Exception:  # noqa: BLE001
            return None

    def add(name, changes, target, expected, tol_abs=TOL_EOK, extra_targets=()):
        try:
            before, after = _probe(wb, changes, [target] + list(extra_targets))
        except Exception as e:  # noqa: BLE001
            out.append({"name": name, "ok": False, "error": repr(e)})
            return
        b, a = _num_or_none(before[0]), _num_or_none(after[0])
        delta = (a - b) if (a is not None and b is not None) else None
        ok = delta is not None and expected is not None and abs(delta - expected) <= max(tol_abs, abs(expected) * 1e-4)
        rec = {"name": name, "changed": ["%s!%s" % k for k in changes], "target": "%s!%s" % target, "before": b, "after": a,
               "delta": delta, "expected_delta": expected, "ok": ok}
        if extra_targets:
            rec["chain"] = [{"cell": "%s!%s" % t, "before": _num_or_none(bb), "after": _num_or_none(aa), "moved": (_num_or_none(bb) != _num_or_none(aa))}
                            for t, bb, aa in zip(extra_targets, before[1:], after[1:])]
            ok = ok and all(c["moved"] for c in rec["chain"])
            rec["ok"] = ok
        out.append(rec)
    # ① 세율 +1%p → 법인세(T+1) Δ = MAX(세전,0) × 0.01 ; 당기순이익·EPS 도 움직여야 한다
    if "sc:세율" in vrow and "drv:법인세비용" in vrow and "법인세비용" in subq_row and "세전이익" in subq_row:
        dcell = vs["%s%d" % (L1, vrow["drv:법인세비용"])].value
        if _is_formula(dcell) and "$D$%d" % vrow["sc:세율"] in dcell:
            pt = num("subQ", "%s%d" % (L1, subq_row["세전이익"]))
            cur = _num_or_none(vs["D%d" % vrow["sc:세율"]].value) or 0.0
            chain = [("subQ", "%s%d" % (L1, subq_row[k])) for k in ("당기순이익", "EPS") if k in subq_row]
            add("세율 +1%p → 법인세비용", {("변수", "D%d" % vrow["sc:세율"]): cur + 0.01}, ("subQ", "%s%d" % (L1, subq_row["법인세비용"])),
                max(pt or 0.0, 0.0) * 0.01, extra_targets=chain)
    # ② OPM +1%p → 영업이익(T+1) Δ = 기준 매출 × 0.01 ; 세전·지배NI·EPS 연쇄
    for drv, base in (("drv:영업이익", "매출액"), ("drv:OP조선", "매출조선"), ("drv:OP조선기자재", "매출조선기자재"), ("drv:OP전사", "매출전사")):
        if drv in vrow and base in subq_row and "영업이익" in subq_row and _num_or_none(vs["%s%d" % (L1, vrow[drv])].value) is not None:
            bv = num("subQ", "%s%d" % (L1, subq_row[base]))
            cur = vs["%s%d" % (L1, vrow[drv])].value
            chain = [("subQ", "%s%d" % (L1, subq_row[k])) for k in ("세전이익", "지배주주순이익", "EPS") if k in subq_row]
            add("OPM(%s) +1%%p → 영업이익" % drv[4:], {("변수", "%s%d" % (L1, vrow[drv])): cur + 0.01}, ("subQ", "%s%d" % (L1, subq_row["영업이익"])),
                (bv or 0.0) * 0.01, extra_targets=chain)
            break
    # ③ 원/달러(평균) T+2 +100원 → 매출조선(T+2) Δ = (1−h) × Σ포함 스케줄(캡 반영) × 100 / 100 ; 매출액 연쇄
    if "SLS" in wb.sheetnames and "USDKRW_avg" in vrow:
        sl = wb["SLS"]
        srow = _find_rows(sl)
        seg_key = next((k for k in subq_row if k.startswith("매출") and _is_formula(ws_val(wb, "subQ", "%s%d" % (L2, subq_row[k]))) and "SLS!" in ws_val(wb, "subQ", "%s%d" % (L2, subq_row[k]))), None)
        if seg_key and "krw_link" in srow and "spot_krw" in srow and "after_origin" in srow:
            hdr = {sl.cell(1, c).value: c for c in range(FIRST_COL, sl.max_column + 1)}
            col = hdr.get(qlabel(q2))
            h = num("SLS", "D%d" % srow["헤지비율"]) if "헤지비율" in srow else None
            if col and h is not None:
                # 블록 ① 행 범위: 'rcp' 머리글 다음 행부터 usd_incl 전까지 — 포함·캡배율 열은 머리글로 찾는다
                head_r = next((r for r in range(DATA_ROW, sl.max_row + 1) if sl.cell(r, 1).value and str(sl.cell(r, 1).value).startswith("① ")), None)
                pcol = {sl.cell(head_r, c).value: c for c in range(FIRST_COL, sl.max_column + 1) if sl.cell(head_r, c).value} if head_r else {}
                after = _num_or_none(sl.cell(srow["after_origin"], col).value) or 0
                expected = 0.0
                if head_r and "포함(1/0)" in pcol and "캡배율" in pcol:
                    for r in range(head_r + 1, srow["usd_incl"]):
                        if sl.cell(r, 1).value is None:
                            break
                        u = _num_or_none(sl.cell(r, col).value)
                        if u is None:
                            continue
                        incl = _num_or_none(sl.cell(r, pcol["포함(1/0)"]).value) or 0
                        fac = _num_or_none(sl.cell(r, pcol["캡배율"]).value) or 1.0
                        expected += u * incl * (1 + (fac - 1) * after)
                    expected = (1 - h) * expected * 100 / 100.0          # 백만$ × 100원 = 백만원 → 억원 /100
                    chain = [("subQ", "%s%d" % (L2, subq_row[k])) for k in ("매출액", "영업이익") if k in subq_row]
                    cur = _num_or_none(vs["%s%d" % (L2, vrow["USDKRW_avg"])].value) or 0.0
                    add("원/달러(평균) %s +100원 → %s(선표)" % (q2, seg_key), {("변수", "%s%d" % (L2, vrow["USDKRW_avg"])): cur + 100},
                        ("subQ", "%s%d" % (L2, subq_row[seg_key])), expected, extra_targets=chain)
    # ④ 원/달러(기말) T+1 +100원 → 환관련손익(T+1) Δ = 노출(백만$) × 100 / 100 ; 세전 연쇄
    if "환관련손익" in subq_row and "USDKRW_end" in vrow and "sc:외화순노출_백만$" in vrow:
        f = ws_val(wb, "subQ", "%s%d" % (L1, subq_row["환관련손익"]))
        if _is_formula(f) and "변수!$D$%d" % vrow["sc:외화순노출_백만$"] in f:
            e = _num_or_none(vs["D%d" % vrow["sc:외화순노출_백만$"]].value) or 0.0
            cur = _num_or_none(vs["%s%d" % (L1, vrow["USDKRW_end"])].value) or 0.0
            add("원/달러(기말) %s +100원 → 환관련손익" % q1, {("변수", "%s%d" % (L1, vrow["USDKRW_end"])): cur + 100},
                ("subQ", "%s%d" % (L1, subq_row["환관련손익"])), e * 100 / 100.0, extra_targets=[("subQ", "%s%d" % (L1, subq_row["세전이익"]))] if "세전이익" in subq_row else ())
    # ⑤ 자산이자율 +1%p → 이자손익(T+1) Δ = 이자발생자산(마지막 실적) × 0.01 / 4 ; 금융손익 연쇄
    if "sc:이자율_자산" in vrow and "이자손익" in subq_row and "이자발생자산" in subq_row:
        f = ws_val(wb, "subQ", "%s%d" % (L1, subq_row["이자손익"]))
        if _is_formula(f) and "$D$%d" % vrow["sc:이자율_자산"] in f:
            ia = num("subQ", "%s%d" % (lay.letter(lay.last_actual), subq_row["이자발생자산"]))
            cur = _num_or_none(vs["D%d" % vrow["sc:이자율_자산"]].value) or 0.0
            add("자산이자율 +1%p → 이자손익", {("변수", "D%d" % vrow["sc:이자율_자산"]): cur + 0.01}, ("subQ", "%s%d" % (L1, subq_row["이자손익"])),
                (ia or 0.0) * 0.01 / 4, extra_targets=[("subQ", "%s%d" % (L1, subq_row["금융손익"]))] if "금융손익" in subq_row else ())
    # ⑥ 신규수주 매출(T+2) +1,000억 → 매출액(T+2) Δ 1,000 ; 시나리오 base 매출액 연쇄
    new_key = next((k for k in subq_row if k.startswith("매출") and k.endswith("신규")), None)
    seg_f = ws_val(wb, "subQ", "%s%d" % (L2, subq_row["매출" + new_key[2:-2]])) if (new_key and ("매출" + new_key[2:-2]) in subq_row) else None
    if (new_key and "drv:" + new_key in vrow and "매출액" in subq_row and _num_or_none(vs["%s%d" % (L2, vrow["drv:" + new_key])].value) is not None
            and _is_formula(seg_f) and ("+%s%d" % (L2, subq_row[new_key])) in seg_f):     # 매출<seg> 가 신규 행을 더하는 구조(선표 연결)일 때만 의미 있는 프로브
        cur = vs["%s%d" % (L2, vrow["drv:" + new_key])].value
        chain = []
        if "시나리오" in wb.sheetnames:
            wsn = wb["시나리오"]
            srow = _find_rows(wsn)
            hdr = {wsn.cell(1, c).value: c for c in range(FIRST_COL, wsn.max_column + 1)}
            if "scn:base:매출액" in srow and hdr.get(qlabel(q2)):
                chain.append(("시나리오", "%s%d" % (get_column_letter(hdr[qlabel(q2)]), srow["scn:base:매출액"])))
        add("%s %s +1,000억 → 매출액" % (new_key, q2), {("변수", "%s%d" % (L2, vrow["drv:" + new_key])): cur + 1000}, ("subQ", "%s%d" % (L2, subq_row["매출액"])), 1000.0, extra_targets=chain)
    return out


def ws_val(wb, sheet, coord):
    return wb[sheet][coord].value


def build_one(model_path, out_path=None, fin_path=None, sls_path=None, fx_path=None, prices_path=None, do_verify=False):
    model = load_json(model_path)
    stock = model["stock"]
    fin = load_optional(fin_path or os.path.join(FIN_DIR, stock + ".json"))
    sls = load_optional(sls_path or os.path.join(SLS_DIR, stock + ".json"))
    fx = load_optional(fx_path or os.path.join(ASSETS, "fx.json"))
    prices_all = load_optional(prices_path or os.path.join(ASSETS, "prices.json"))
    prices = ((prices_all or {}).get("rows") or {}).get(stock)
    out_path = out_path or os.path.join(OUT_DIR, "%s_model.xlsx" % stock)
    b = Builder(model, fin=fin, sls=sls, fx=fx, prices=prices)
    wb = b.build()
    size = save_atomic(wb, out_path)
    rep = {"out": out_path, "size_bytes": size, "sheet_dims": b.stats, "fin": bool(fin), "sls": bool(sls), "fx": bool(fx), "prices": bool(prices),
           "links": b.links, "fit": b.fit}
    if size > SIZE_LIMIT:
        rep["error"] = "size > 20MB"
    if do_verify:
        rep["verify"] = verify(out_path, model)
        if not rep["verify"].get("ok"):
            rep["error"] = rep.get("error") or "verify failed"
    return rep


def _compact(rep):
    """--all 출력용 한 줄 요약."""
    v = rep.get("verify") or {}
    est, act, sc, sens = v.get("estimate_recalc") or {}, v.get("actual_recalc") or {}, v.get("scenarios") or {}, v.get("sensitivity") or []
    return {"out": os.path.basename(rep["out"]), "bytes": rep["size_bytes"], "sheets": len(v.get("sheets") or rep.get("sheet_dims") or {}),
            "formulas": v.get("formulas_total"), "links": {k: ("Y" if d.get("linked") else "n") for k, d in (rep.get("links") or {}).items()},
            "est": "%s/%s" % (est.get("match"), est.get("rows")), "est_value_cells": (v.get("links") or {}).get("value_cells"),
            "actual_formula": "%s/%s" % (act.get("match"), act.get("formula_cells")),
            "scenario": ("%s/%s" % (sc.get("match"), sc.get("cells"))) if sc else "-",
            "sensitivity": "%d/%d" % (sum(1 for p in sens if p.get("ok")), len(sens)), "ok": v.get("ok"), "error": rep.get("error")}


def main(argv=None):
    ap = argparse.ArgumentParser(description="모델 json → subQ 형 xlsx")
    ap.add_argument("--model", help="모델 json 경로(스펙 2-5)")
    ap.add_argument("--sample", action="store_true", help="tests/fixtures/model_sample.json 사용")
    ap.add_argument("--all", action="store_true", help="assets/models/*.json 전부")
    ap.add_argument("--out", help="출력 xlsx 경로(기본 argus/kship/models/<stock>_model.xlsx)")
    ap.add_argument("--fin", help="fin json(2-1) 경로 — 기본 assets/fin/<stock>.json 이 있으면 사용")
    ap.add_argument("--sls", help="sls json(2-4) 경로")
    ap.add_argument("--fx", help="fx.json 경로")
    ap.add_argument("--prices", help="prices.json 경로")
    ap.add_argument("--verify", action="store_true", help="재오픈 검증(수식·이름정의·VLOOKUP 흉내·추정 전 분기 재계산·시나리오·민감도)")
    ap.add_argument("--full-report", action="store_true", help="--all 에서도 회사별 전체 보고(기본은 한 줄 요약표)")
    a = ap.parse_args(argv)
    targets = []
    if a.sample:
        targets.append(SAMPLE)
    if a.model:
        targets.append(a.model)
    if a.all:
        targets += sorted(os.path.join(MODELS_DIR, f) for f in os.listdir(MODELS_DIR) if f.endswith(".json") and f != "summary.json") if os.path.isdir(MODELS_DIR) else []
    if not targets:
        ap.error("--model / --sample / --all 중 하나")
    reps = []
    for t in targets:
        rep = build_one(t, out_path=a.out if len(targets) == 1 else None, fin_path=a.fin, sls_path=a.sls, fx_path=a.fx,
                        prices_path=a.prices, do_verify=a.verify)
        reps.append(rep)
    if len(reps) > 1 and not a.full_report:
        for r in reps:
            print(json.dumps(_compact(r), ensure_ascii=False, default=str))
        bad = [r for r in reps if "error" in r]
        tot = {"files": len(reps), "errors": len(bad), "bytes": sum(r["size_bytes"] for r in reps)}
        if a.verify:
            tot.update(est_match=sum((r["verify"]["estimate_recalc"]["match"]) for r in reps), est_rows=sum((r["verify"]["estimate_recalc"]["rows"]) for r in reps),
                       est_value_cells=sum(r["verify"]["links"]["value_cells"] for r in reps),
                       actual_formula_match=sum(r["verify"]["actual_recalc"]["match"] for r in reps), actual_formula_cells=sum(r["verify"]["actual_recalc"]["formula_cells"] for r in reps),
                       sensitivity_ok=sum(sum(1 for p in r["verify"]["sensitivity"] if p.get("ok")) for r in reps), sensitivity_n=sum(len(r["verify"]["sensitivity"]) for r in reps),
                       scenario_match=sum((r["verify"].get("scenarios") or {}).get("match", 0) for r in reps), scenario_cells=sum((r["verify"].get("scenarios") or {}).get("cells", 0) for r in reps),
                       linked=dict(sorted(((k, sum(1 for r in reps if (r["links"].get(k) or {}).get("linked"))) for k in {kk for r in reps for kk in r["links"]}))))
        print("TOTAL " + json.dumps(tot, ensure_ascii=False))
    else:
        print(json.dumps(reps if len(reps) > 1 else reps[0], ensure_ascii=False, indent=1, default=str))
    return 0 if all("error" not in r for r in reps) else 1


if __name__ == "__main__":
    sys.exit(main())
