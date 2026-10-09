#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""kship_model_xlsx 계약 테스트 — 표본 모델 json(fixtures/model_sample.json) → xlsx → 재오픈 검증.

실행: cd argus/kship/tools && python3 -m unittest tests.test_kship_model_xlsx -v
확인하는 것: 시트 목록·헤더 규약(행1 '4Q21'/2021 숫자, 행4 '2021.12'/'2021.12A', E열 시작)·이름정의·
확정 행 VLOOKUP 수식과 파이썬 VLOOKUP 흉내(2026Q2 매출액 = BS연결 값/100)·추정 행 가정 셀 참조와 재계산 일치·
추정 열 연노랑·경계 굵은 테두리·메모(basis)·파일 크기.
V5 검증 회귀(2차): README 재현 명령·UPDATE 스탬프/문서속성/zip 시각 = built_at·두 번 생성 바이트 동일·
BS 에 없는 확정 셀은 모델값/빈 칸(VLOOKUP 0 금지)·역산 기준 = 시트 값·세전 항등식에 환관련손익.
V5 검증 회귀(5차, 2026-10-05): 추정 전 분기(T+1~T+10) 재계산·죽은 값 셀 0·스칼라 연결(세율·판관비율·지배주주비중·이자율 — 바꾸면 전 분기 재계산)·
법인세 = MAX(세전,0)×세율·금융손익 = Σ 세부 4행·EBITDA 항등식·외환손익 BS 복합 VLOOKUP·기타금융손익 잔차·BS≠모델(부호 반전)이면 모델값 직접·
OPM/PER/PBR 파생 수식·전년동기 없을 때 QoQ 폴백·VLOOKUP 키 중복 금지·시나리오 시트·선표 연결(환율 forward → 매출조선)·환관련손익 연결·민감도.
실데이터 테스트(TestRealModelLinks)는 assets/models·sls·fx·prices 가 있을 때만 돈다 — 구조·재계산 일치만 단언(수치 고정 없음).
"""
import copy
import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.dirname(HERE)
sys.path.insert(0, TOOLS)

from openpyxl import load_workbook                          # noqa: E402

import kship_model_xlsx as X                                # noqa: E402

SAMPLE = os.path.join(HERE, "fixtures", "model_sample.json")
REAL_MODEL = os.path.join(X.MODELS_DIR, "010140.json")
REAL_SLS = os.path.join(X.SLS_DIR, "010140.json")


def _load_sample():
    with open(SAMPLE, encoding="utf-8") as f:
        return json.load(f)


def _row_of(ws, key):
    for r in range(X.DATA_ROW, ws.max_row + 1):
        if ws.cell(r, 1).value == key:
            return r
    raise KeyError(key)


class TestPeriods(unittest.TestCase):
    def test_labels(self):
        self.assertEqual(X.qlabel("2021Q4"), "4Q21")
        self.assertEqual(X.qlabel("2026Q1"), "1Q26")
        self.assertEqual(X.qdate("2021Q4"), "2021.12")
        self.assertEqual(X.qdate("2026Q2"), "2026.06")

    def test_columns_annual_after_q4(self):
        cols = X.build_columns(["2021Q4", "2022Q1", "2022Q2", "2022Q3", "2022Q4", "2023Q1"], ["2021", "2022", "2023"])
        keys = [c["key"] for c in cols]
        self.assertEqual(keys, ["2021Q4", "2021", "2022Q1", "2022Q2", "2022Q3", "2022Q4", "2022", "2023Q1", "2023"])
        self.assertEqual(cols[0]["col"], 5)                       # E열 시작
        self.assertEqual(cols[1]["label"], 2021)                  # 연간 라벨은 숫자
        self.assertEqual(cols[1]["date"], "2021.12A")

    def test_tolerances(self):
        self.assertEqual(X.tol_for("매출액", "억원", 100.0), 0.05)
        self.assertEqual(X.tol_for("매출액", "억원", 100000.0), 10.0)
        self.assertEqual(X.tol_for("EPS", "원", 100.0), 0.5)
        self.assertAlmostEqual(X.tol_for("OPM", "%", 0.1, base=110.0), 6e-5 + 0.011 / 110.0)
        self.assertGreater(X.tol_for("PER", "배", 144.0, ttm_eps=200.0), 0.1)      # EPS 0.1원 반올림 × 4 가 배수로 증폭
        self.assertEqual(X.tol_for("PBR", "배", 3.0), 0.006)


class TestBuild(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.model = _load_sample()
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_")
        cls.out = os.path.join(cls.tmp, "sample_model.xlsx")
        # 다른 레인 산출(fin/sls/fx/prices)에 의존하지 않도록 없는 경로를 준다 — 표본만으로 만든다
        none = os.path.join(cls.tmp, "none.json")
        cls.rep = X.build_one(SAMPLE, out_path=cls.out, fin_path=none, sls_path=none, fx_path=none, prices_path=none, do_verify=True)
        cls.wb = load_workbook(cls.out, data_only=False)
        cls.lay = X.Layout(cls.model)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def _subq_row(self, key):
        return _row_of(self.wb["subQ"], key)

    def test_sheets_and_names(self):
        self.assertEqual(self.wb.sheetnames, ["README", "변수", "BS연결", "BS별도", "subQ", "분기", "연간예상", "TP", "외화"])
        names = set(self.wb.defined_names.keys())
        for n in ("SUBQH", "BS연결H", "BS별도H", "BS연결", "BS별도", "변수H", "환율", "subQ"):
            self.assertIn(n, names)
        self.assertTrue(self.wb.defined_names["BS연결"].attr_text.startswith("BS연결!$C$5:"))
        self.assertTrue(self.wb.defined_names["SUBQH"].attr_text.startswith("subQ!$A$1:"))

    def test_header_convention(self):
        for sn in ("subQ", "BS연결", "BS별도", "변수"):
            ws = self.wb[sn]
            self.assertEqual(ws.cell(1, 5).value, "4Q21", sn)
            self.assertEqual(ws.cell(1, 6).value, 2021, sn)             # 4분기 뒤 연간(숫자)
            self.assertEqual(ws.cell(1, 7).value, "1Q22", sn)
            self.assertEqual(ws.cell(4, 5).value, "2021.12", sn)
            self.assertEqual(ws.cell(4, 6).value, "2021.12A", sn)
            self.assertIsNone(ws.cell(1, 4).value if sn != "subQ" else None)   # D열은 데이터 아님
        self.assertEqual(self.wb["subQ"]["A1"].value, "YQ")
        self.assertEqual(self.wb["BS연결"]["B1"].value, "(백만원)")

    def test_actual_cell_is_vlookup_and_resolves(self):
        ws = self.wb["subQ"]
        r = self._subq_row("매출액")
        L = self.lay.letter("2026Q2")
        f = ws["%s%d" % (L, r)].value
        self.assertEqual(f, '=IFERROR(VLOOKUP($D%d,BS연결,MATCH(%s$1,BS연결H,0),0)/100,"")' % (r, L))
        self.assertEqual(ws.cell(r, 4).value, "매출액(수익)")
        chk = self.rep["verify"]["checks"][0]
        self.assertTrue(chk["ok"], chk)
        self.assertAlmostEqual(chk["emulated"], 32307.12, places=2)   # 3,230,712 백만원 / 100 (shi_2026Q2_cons)
        self.assertEqual(chk["vlookup_trace"]["bs_value_million"], 3230712)
        # 직접 흉내: BS연결 행1 라벨 → 열, C열 계정명 → 행
        bs = self.wb["BS연결"]
        col = next(c for c in range(5, bs.max_column + 1) if bs.cell(1, c).value == "2Q26")
        row = next(rr for rr in range(5, bs.max_row + 1) if bs.cell(rr, 3).value == "매출액(수익)")
        self.assertAlmostEqual(bs.cell(row, col).value / 100, 32307.12, places=2)

    def test_estimate_cell_references_vars_and_recomputes(self):
        ws = self.wb["subQ"]
        r = self._subq_row("매출액")
        L, LP = self.lay.letter("2026Q3"), self.lay.letter("2025Q3")
        f = ws["%s%d" % (L, r)].value
        self.assertTrue(f.startswith("=IFERROR(%s%d*(1+변수!%s" % (LP, r, L)), f)
        self.assertIsNotNone(ws["%s%d" % (L, r)].comment)
        self.assertIn("표본 가정", ws["%s%d" % (L, r)].comment.text)
        em = X.Emulator(self.wb)
        got = em.value("subQ", "%s%d" % (L, r))
        model_v = [x for x in self.model["rows"] if x["key"] == "매출액"][0]["q"]["2026Q3"]["v"]
        self.assertAlmostEqual(got, model_v, delta=0.05)
        # 가정 셀을 바꾸면 결과가 따라간다(수식이 진짜 가정 셀을 참조)
        vr = self.wb["변수"]
        drv = _row_of(vr, "drv:매출액")
        g = vr["%s%d" % (L, drv)].value
        self.assertIsInstance(g, float)
        vr["%s%d" % (L, drv)].value = g + 0.10
        em2 = X.Emulator(self.wb)
        self.assertAlmostEqual(em2.value("subQ", "%s%d" % (L, r)), got + 0.10 * em2.value("subQ", "%s%d" % (LP, r)), delta=0.05)
        vr["%s%d" % (L, drv)].value = g

    def test_estimate_recalc_all_rows_all_quarters(self):
        e = self.rep["verify"]["estimate_recalc"]
        self.assertEqual(e["quarters"], [q for q in self.model["periods"]["quarters"] if q > "2026Q2"])    # T+1~T+10 전부
        self.assertEqual(len(e["quarters"]), 10)
        self.assertGreater(e["rows"], 200)
        self.assertEqual(e["mismatch"], [], e["mismatch"])
        self.assertEqual(e["rows"], e["match"])
        # 죽은 값 셀(수식 아닌 추정 셀) 0
        self.assertEqual(self.rep["verify"]["links"]["value_cells"], 0, self.rep["verify"]["links"]["value_keys"])
        self.assertEqual(self.rep["verify"]["actual_recalc"]["mismatch"], [])
        self.assertTrue(self.rep["verify"]["ok"])

    def test_identity_rows_are_formulas_of_subq(self):
        ws = self.wb["subQ"]
        L = self.lay.letter("2027Q1")
        r_cogs, r_rev, r_gp = self._subq_row("매출원가"), self._subq_row("매출액"), self._subq_row("매출총이익")
        self.assertEqual(ws["%s%d" % (L, r_cogs)].value, '=IFERROR(%s%d-%s%d,"")' % (L, r_rev, L, r_gp))
        r_a, r_l, r_e = self._subq_row("자산총계"), self._subq_row("부채총계"), self._subq_row("자본총계")
        self.assertEqual(ws["%s%d" % (L, r_a)].value, '=IFERROR(%s%d+%s%d,"")' % (L, r_l, L, r_e))
        r_eps, r_ni, r_sh = self._subq_row("EPS"), self._subq_row("지배주주순이익"), self._subq_row("주식수")
        self.assertEqual(ws["%s%d" % (L, r_eps)].value, '=IFERROR(%s%d*100/%s%d,"")' % (L, r_ni, L, r_sh))

    def test_tax_is_max_pretax_times_scalar_rate(self):
        """법인세 = MAX(세전,0) × 세율 — 드라이버 셀은 3절 '세율' 스칼라 참조(연녹), 스칼라를 바꾸면 전 분기 법인세·순이익·EPS 가 따라간다."""
        ws, vr = self.wb["subQ"], self.wb["변수"]
        r_tax, r_pt = self._subq_row("법인세비용"), self._subq_row("세전이익")
        drv, sc = _row_of(vr, "drv:법인세비용"), _row_of(vr, "sc:세율")
        L = self.lay.letter("2026Q3")
        self.assertEqual(ws["%s%d" % (L, r_tax)].value, '=IFERROR(MAX(%s%d,0)*변수!%s%d,"")' % (L, r_pt, L, drv))
        self.assertEqual(vr["%s%d" % (L, drv)].value, "=$D$%d" % sc)
        self.assertEqual(vr["%s%d" % (L, drv)].fill.fgColor.rgb[-6:], "E2EFDA")
        self.assertAlmostEqual(vr["D%d" % sc].value, 0.22, places=4)                 # 표본 세율 22%
        self.assertTrue(self.rep["links"]["세율"]["linked"])
        probes = {p["name"]: p for p in self.rep["verify"]["sensitivity"]}
        p = probes["세율 +1%p → 법인세비용"]
        self.assertTrue(p["ok"], p)
        self.assertTrue(all(c["moved"] for c in p["chain"]))                          # 당기순이익·EPS 연쇄
        # 스칼라 연결: 판관비율·지배주주비중 드라이버 셀도 =$D$n
        for key, name in (("drv:판관비", "sc:판관비율"), ("drv:지배주주순이익", "sc:지배주주비중")):
            self.assertEqual(vr["%s%d" % (L, _row_of(vr, key))].value, "=$D$%d" % _row_of(vr, name))

    def test_opm_sensitivity_chain(self):
        probes = {p["name"]: p for p in self.rep["verify"]["sensitivity"]}
        p = probes["OPM(영업이익) +1%p → 영업이익"]
        self.assertTrue(p["ok"], p)
        self.assertAlmostEqual(p["delta"], p["expected_delta"], delta=0.05)
        self.assertTrue(all(c["moved"] for c in p["chain"]))

    def test_annual_columns(self):
        ws = self.wb["subQ"]
        r = self._subq_row("매출액")
        # 2022(실적·계정 매핑) → BS연결 연간 열 VLOOKUP, 2027(추정) → 4분기 SUM, 자산총계 연간 → 4Q 값
        self.assertIn("VLOOKUP($D%d,BS연결" % r, ws["%s%d" % (self.lay.letter("2022"), r)].value)
        q1, q4 = self.lay.letter("2027Q1"), self.lay.letter("2027Q4")
        self.assertEqual(ws["%s%d" % (self.lay.letter("2027"), r)].value, "=SUM(%s%d:%s%d)" % (q1, r, q4, r))
        ra = self._subq_row("자산총계")
        self.assertEqual(ws["%s%d" % (self.lay.letter("2027"), ra)].value, "=%s%d" % (q4, ra))
        em = X.Emulator(self.wb)
        got = em.value("subQ", "%s%d" % (self.lay.letter("2022"), r))
        self.assertAlmostEqual(got, 59446.67, places=1)           # golden 2022.12A 매출액(수익) 5,944,667.45 백만원

    def test_styles_boundary_and_estimate_fill(self):
        ws = self.wb["subQ"]
        r = self._subq_row("영업이익")
        c_last = ws.cell(r, self.lay.col_of("2026Q2"))
        self.assertEqual(c_last.border.right.style, "medium")
        c_est = ws.cell(r, self.lay.col_of("2026Q3"))
        self.assertEqual(c_est.fill.fgColor.rgb[-6:], "FFF9DB")
        c_act = ws.cell(r, self.lay.col_of("2026Q1"))
        self.assertNotEqual(c_act.fill.fgColor.rgb[-6:], "FFF9DB")
        self.assertEqual(c_est.number_format, "#,##0")
        # 갭필 추정(실적 구간)도 노랑 + 메모
        c_gap = ws.cell(r, self.lay.col_of("2024Q2"))
        self.assertEqual(c_gap.fill.fgColor.rgb[-6:], "FFF9DB")
        self.assertIn("갭필", c_gap.comment.text)

    def test_views_lookup_subq(self):
        q = self.wb["분기"]
        self.assertEqual(q.cell(1, 5).value, "4Q21")
        self.assertIsNone(next((c for c in range(5, q.max_column + 1) if isinstance(q.cell(1, c).value, int)), None))  # 분기만
        r = next(rr for rr in range(5, q.max_row + 1) if q.cell(rr, 1).value == "영업이익")
        self.assertEqual(q.cell(r, 5).value, '=IFERROR(VLOOKUP($A%d,subQ,MATCH(E$1,SUBQH,0),0),"")' % r)
        em = X.Emulator(self.wb)
        col = next(c for c in range(5, q.max_column + 1) if q.cell(1, c).value == "2Q26")
        self.assertAlmostEqual(em.value("분기", "%s%d" % (X.get_column_letter(col), r)), 3250.24, places=2)
        a = self.wb["연간예상"]
        self.assertEqual(a.cell(1, 5).value, 2021)
        ra = next(rr for rr in range(5, a.max_row + 1) if a.cell(rr, 1).value == "매출액")
        col = next(c for c in range(5, a.max_column + 1) if a.cell(1, c).value == 2025)
        self.assertAlmostEqual(em.value("연간예상", "%s%d" % (X.get_column_letter(col), ra)), 106500.11, places=1)

    def test_tp_and_fx_sheets(self):
        tp = self.wb["TP"]
        self.assertEqual(tp["B1"].value, "종가")
        self.assertTrue(str(tp["C1"].value).startswith("=변수!D"))
        formulas = [c.value for row in tp.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(any("PER_" in f or "ROUND(" in f for f in formulas))
        fx = self.wb["외화"]
        self.assertIn("VLOOKUP($D5,환율,MATCH(E$1,변수H,0),0)", fx["E5"].value)

    def test_size_and_counts(self):
        v = self.rep["verify"]
        self.assertTrue(v["size_ok"])
        self.assertLess(v["size_bytes"], 2 * 1024 * 1024)
        self.assertGreater(v["formulas_total"], 1500)
        self.assertGreater(v["sheets"]["subQ"]["comments"], 100)
        self.assertNotIn("error", self.rep)


class TestSlsSheet(unittest.TestCase):
    def test_sls_from_inline_dict(self):
        model = _load_sample()
        model["sls"] = {"stock": "010140", "origin": "2026Q2", "unit": "USD_million",
                        "contracts": [{"rcp": "r1", "type": "LNGC", "ships": 2, "amt_usd_m": 500, "cohort": "⑤초호황", "schedule": {"2026Q3": 50},
                                       "fx_at_sign": 1350, "counted": True, "signed_by_origin": True, "signed": "2024-01-01"}],
                        "hedge": {"hedge_ratio": 0.7, "hedge_rate": None, "kind": "estimate"},
                        "by_quarter": {"2026Q3": {"usd_m": 50, "by_type": {"LNGC": 50}, "by_cohort": {"⑤초호황": 50}, "hedged_krw_m": 69000,
                                                  "applied_rate": 1380, "hedge_ratio": 0.7, "hedge_rate": 1350, "spot_assumed": 1450}},
                        "by_year": {"2026": {"usd_m": 50}}, "target_opm": {"2026Q3": {"opm": 0.12, "basis": "test"}},
                        "warnings": ["테스트"]}
        tmp = tempfile.mkdtemp(prefix="kship_sls_")
        try:
            b = X.Builder(model)
            wb = b.build()
            self.assertIn("SLS", wb.sheetnames)
            self.assertEqual(wb.sheetnames.index("SLS"), wb.sheetnames.index("subQ") + 1)
            ws = wb["SLS"]
            self.assertEqual(ws.cell(1, 5).value, "3Q26")
            self.assertEqual(ws.cell(1, 6).value, 2026)
            keys = [ws.cell(r, 1).value for r in range(5, ws.max_row + 1)]
            for k in ("usd_m", "type:LNGC", "cohort:⑤초호황", "hedge_ratio", "applied_rate", "target_opm", "spot_krw", "krw_link", "usd_incl"):
                self.assertIn(k, keys)
            # 계약 원화 수식: 스케줄 × 포함 × 캡 × (h × 헤지환율 + (1−h) × spot) — spot 은 변수 환율 참조(없으면 sls spot_assumed)
            g = b._sls_geometry()
            kf = ws.cell(g["krw_first"], 5).value
            self.assertTrue(kf.startswith("=IFERROR(E%d*$" % g["usd_first"]), kf)
            self.assertIn("$D$%d" % g["h_row"], kf)
            self.assertIn('VLOOKUP(" 원/달러(평균)",환율', ws.cell(g["spot_row"], 5).value)
            self.assertEqual(ws.cell(g["h_row"], 4).value, "=변수!$D$%d" % b.scalar_row["헤지비율"])
            em = X.Emulator(wb)
            # 표본에 2026Q3 원/달러(평균) 가정 1380 → 50 × (0.7×1350 + 0.3×1380) = 67,950 백만원
            self.assertAlmostEqual(em.value("SLS", "E%d" % g["krw_link_row"]), 50 * (0.7 * 1350 + 0.3 * 1380), places=3)
            # 표본 모델은 sls 형 세그먼트 드라이버가 없어 subQ 매출조선은 연결되지 않는다(사유 기록)
            self.assertFalse((b.links.get("선표") or {}).get("linked"))
            out = os.path.join(tmp, "x.xlsx")
            X.save_atomic(wb, out)
            self.assertTrue(os.path.getsize(out) > 0)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestReadmeAndDeterminism(unittest.TestCase):
    """V5 검증에서 잡은 결함의 회귀 테스트 — README 재현 명령·가정 요약·추천 아님, UPDATE 스탬프·문서 속성·zip 시각 = built_at, 두 번 생성 바이트 동일."""

    @classmethod
    def setUpClass(cls):
        cls.model = _load_sample()
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_det_")
        none = os.path.join(cls.tmp, "none.json")
        cls.out1 = os.path.join(cls.tmp, "a.xlsx")
        cls.out2 = os.path.join(cls.tmp, "b.xlsx")
        X.build_one(SAMPLE, out_path=cls.out1, fin_path=none, sls_path=none, fx_path=none, prices_path=none)
        X.build_one(SAMPLE, out_path=cls.out2, fin_path=none, sls_path=none, fx_path=none, prices_path=none)
        cls.wb = load_workbook(cls.out1)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_readme_has_source_disclaimer_assumptions_and_repro(self):
        txt = "\n".join(str(c.value) for row in self.wb["README"].iter_rows() for c in row if c.value is not None)
        self.assertIn("출처: DART", txt)
        self.assertIn("추천 아님", txt)
        self.assertIn("투자조언 아님", txt)
        self.assertIn("재현: cd argus/kship/tools", txt)
        self.assertIn("kship_model_xlsx.py --model assets/models/010140.json --verify", txt)
        self.assertIn("kship_model.py --build --all --xlsx", txt)
        self.assertIn("가정 요약", txt)
        self.assertIn("· 세율 22.00%", txt)
        self.assertIn("연결된 가정: 세율, 판관비율, 지배주주비중", txt)
        self.assertIn("· 환율 — 변수 1절 8행", txt)
        self.assertIn("검증(--verify)", txt)

    def test_bs_update_stamp_is_built_at_not_today(self):
        self.assertTrue(str(self.model["built_at"]).startswith("2026-09-30"))
        for sn in ("BS연결", "BS별도"):
            self.assertEqual(self.wb[sn]["C3"].value, "UPDATE: 26-09-30", sn)

    def test_core_modified_equals_created_and_zip_times_fixed(self):
        import hashlib
        import re
        import zipfile
        with zipfile.ZipFile(self.out1) as z:
            core = z.read("docProps/core.xml").decode("utf-8")
            infos = z.infolist()
        created = re.search(r"<dcterms:created[^>]*>([^<]*)<", core).group(1)
        modified = re.search(r"<dcterms:modified[^>]*>([^<]*)<", core).group(1)
        self.assertEqual(created, "2026-09-30T00:00:00Z")
        self.assertEqual(modified, created)
        self.assertEqual({i.date_time for i in infos}, {(2026, 9, 30, 0, 0, 0)})
        with open(self.out1, "rb") as f1, open(self.out2, "rb") as f2:
            h1, h2 = hashlib.sha256(f1.read()).hexdigest(), hashlib.sha256(f2.read()).hexdigest()
        self.assertEqual(h1, h2, "같은 입력인데 바이트가 다르다")


class TestBsMissingAndIdentity(unittest.TestCase):
    """BS 시트에 값이 없는 확정 셀은 VLOOKUP(→0) 대신 모델값/빈 칸 · BS 값 ≠ 모델값(부호 반전)이면 모델값 직접 · 역산 기준은 시트 값 · 세전 항등식에 환관련손익 포함."""

    def _fin(self, model, q1_offset_million):
        rev = {x["key"]: x for x in model["rows"]}["매출액"]
        return {"cons": {"bs": {}, "cf": {}, "is": {
            "2026Q2": {"매출액(수익)": round(rev["q"]["2026Q2"]["v"] * 100, 2)},
            "2026Q1": {"매출액(수익)": round(rev["q"]["2026Q1"]["v"] * 100 + q1_offset_million, 2)}}}}

    def test_actual_cells_without_bs_value(self):
        model = _load_sample()
        rows = {x["key"]: x for x in model["rows"]}
        rows["금융손익"]["q"]["2026Q1"] = {"v": None, "kind": "actual", "src": "테스트 — 값 없음"}
        # 2026Q1 매출 BS 값 = 모델 v + 0.9 백만원(반올림 전 값: 모델 v 는 0.01억 = 1 백만원 반올림) → VLOOKUP 셀, 역산 기준은 그 값
        b = X.Builder(model, fin=self._fin(model, 0.9))
        wb = b.build()
        ws = wb["subQ"]
        lay = X.Layout(model)
        L2, L1 = lay.letter("2026Q2"), lay.letter("2026Q1")
        self.assertIn("VLOOKUP($D", ws["%s%d" % (L2, _row_of(ws, "매출액"))].value)              # BS 에 있음 → VLOOKUP
        self.assertIn("VLOOKUP($D", ws["%s%d" % (L1, _row_of(ws, "매출액"))].value)
        c = ws["%s%d" % (L2, _row_of(ws, "지배주주순이익"))]
        self.assertEqual(c.value, rows["지배주주순이익"]["q"]["2026Q2"]["v"])                 # BS 에 없음 → 모델값 직접
        self.assertIn("BS연결 시트에 (지배주주지분)당기순이익 2026.06 값 없음", c.comment.text)
        c1 = ws["%s%d" % (L1, _row_of(ws, "금융손익"))]
        self.assertIsNone(c1.value)                                                        # 모델도 BS 도 없음 → 빈 칸(0 표시 금지)
        # 2022 연간: BS 에 연간 값 없음 + 4분기 있음 → SUM 수식(VLOOKUP 0 아님)
        q1, q4 = lay.letter("2022Q1"), lay.letter("2022Q4")
        r = _row_of(ws, "매출액")
        self.assertEqual(ws["%s%d" % (lay.letter("2022"), r)].value, "=SUM(%s%d:%s%d)" % (q1, r, q4, r))
        # 역산 기준 = 시트 값(BS/100, 반올림 전) — 모델 반올림값이 아니다
        v_prev = rows["매출액"]["q"]["2026Q1"]["v"]
        self.assertAlmostEqual(b._sheet_base(rows["매출액"], "2026Q1"), v_prev + 0.009, places=9)
        self.assertNotEqual(b._sheet_base(rows["매출액"], "2026Q1"), v_prev)
        em = X.Emulator(wb)
        self.assertAlmostEqual(em.value("subQ", "%s%d" % (lay.letter("2027Q1"), r)), rows["매출액"]["q"]["2027Q1"]["v"], delta=0.05)

    def test_bs_value_far_from_model_uses_model_value(self):
        """BS 값이 모델값과 0.4억 다르면(비용 부호 반전·재분류) VLOOKUP 을 쓰지 않고 모델값 직접 + 메모 — 역산 기준도 모델값."""
        model = _load_sample()
        rows = {x["key"]: x for x in model["rows"]}
        b = X.Builder(model, fin=self._fin(model, 40.0))
        wb = b.build()
        ws = wb["subQ"]
        lay = X.Layout(model)
        c = ws["%s%d" % (lay.letter("2026Q1"), _row_of(ws, "매출액"))]
        self.assertEqual(c.value, rows["매출액"]["q"]["2026Q1"]["v"])
        self.assertIn("≠ 모델값", c.comment.text)
        self.assertEqual(b._sheet_base(rows["매출액"], "2026Q1"), rows["매출액"]["q"]["2026Q1"]["v"])
        rep = X.verify(self._save(wb), model)
        self.assertEqual(rep["estimate_recalc"]["mismatch"], [])

    def _save(self, wb):
        self.tmp = tempfile.mkdtemp(prefix="kship_xlsx_bs_")
        self.addCleanup(shutil.rmtree, self.tmp, True)
        out = os.path.join(self.tmp, "x.xlsx")
        X.save_atomic(wb, out)
        return out

    def test_pretax_identity_includes_fx_pnl_row(self):
        model = _load_sample()
        rows = {x["key"]: x for x in model["rows"]}
        i = [x["key"] for x in model["rows"]].index("세전이익")
        fx_row = {"key": "환관련손익", "label": "환관련손익", "group": "손익", "unit": "억원",
                  "q": {q: {"v": 132.71, "kind": "estimate", "basis": "테스트 — 외화 순노출 × Δ환율"} for q in model["periods"]["quarters"] if q > "2026Q2"}}
        model["rows"].insert(i, fx_row)
        for q in fx_row["q"]:
            rows["세전이익"]["q"][q]["v"] = round(rows["세전이익"]["q"][q]["v"] + 132.71, 2)
        plan = X.plan_drivers(model["rows"])
        self.assertIn(("+", "환관련손익"), plan["세전이익"]["expr"])
        b = X.Builder(model)
        wb = b.build()
        ws = wb["subQ"]
        lay = X.Layout(model)
        L = lay.letter("2026Q3")
        r_pt, r_fx, r_op = b.subq_row["세전이익"], b.subq_row["환관련손익"], b.subq_row["영업이익"]
        f = ws["%s%d" % (L, r_pt)].value
        self.assertTrue(f.startswith("=IFERROR(%s%d+" % (L, r_op)), f)
        self.assertIn("+%s%d" % (L, r_fx), f)
        em = X.Emulator(wb)
        self.assertAlmostEqual(em.value("subQ", "%s%d" % (L, r_pt)), rows["세전이익"]["q"]["2026Q3"]["v"], delta=0.05)
        # 노출(fx_pnl.net_usd_exposure_m)이 없으면 환관련손익은 abs 드라이버(연결 안 됨)로 남고 사유가 기록된다
        self.assertEqual(b.plan["환관련손익"]["type"], "abs")
        self.assertFalse(b.links["환관련손익"]["linked"])


class TestFinDetailAndDerivedRows(unittest.TestCase):
    """T4 영업외 세부 4행·EBITDA·OPM·PER/PBR·신규수주 행이 xlsx 에 수식으로 들어가는지 — 표본에 행을 합성해 확인(5차 검증)."""

    @classmethod
    def setUpClass(cls):
        m = _load_sample()
        rows = {x["key"]: x for x in m["rows"]}
        qs = m["periods"]["quarters"]
        la = m["periods"]["last_actual"]
        fq = [q for q in qs if q > la]
        aq = [q for q in qs if q <= la]
        order = [x["key"] for x in m["rows"]]

        def put(after, key, label, unit, cells, group="손익"):
            m["rows"].insert([x["key"] for x in m["rows"]].index(after) + 1, {"key": key, "label": label, "group": group, "unit": unit, "q": cells})
            rows[key] = m["rows"][[x["key"] for x in m["rows"]].index(key)]
        # 이자발생자산·총차입금(표본에 총차입금은 있음) → 이자손익 = (IA(q−1)×0.012 − DEBT(q−1)×0.05)/4
        ia = {}
        for i, q in enumerate(qs):
            ia[q] = {"v": round(30000.0 + 100.0 * i, 2), "kind": "actual" if q <= la else "estimate", "src": "t"} if q <= la else {"v": round(30000.0 + 100.0 * i, 2), "kind": "estimate", "basis": "t"}
        put("순차입금", "이자발생자산", "이자발생자산", "억원", ia, group="재무상태")
        ra, rd = 0.012, 0.05
        intc, fxc, derc, othc = {}, {}, {}, {}
        for q in qs:
            if q <= la:
                intc[q] = {"v": -40.0, "kind": "actual", "src": "fin.cons.is.이자수익 − fin.cons.is.이자비용"}
                fxc[q] = {"v": -213.88, "kind": "actual", "src": "fin.cons.is.외환차익 − 외환차손 + 외화환산이익 − 외화환산손실"}
                derc[q] = {"v": -16.0, "kind": "actual", "src": "fin.cons.is.파생상품이익 − 파생상품손실"}
                othc[q] = {"v": round(rows["금융손익"]["q"][q]["v"] + 40.0 + 213.88 + 16.0, 2), "kind": "actual", "src": "잔차"} if rows["금융손익"]["q"].get(q, {}).get("v") is not None else {}
            else:
                pq = qs[qs.index(q) - 1]
                debt = rows["총차입금"]["q"][pq]["v"]
                iv = round((ia[pq]["v"] * ra - debt * rd) / 4, 2)
                intc[q] = {"v": iv, "kind": "estimate", "basis": "이자발생자산 × r_a − 총차입금 × r_d (연율) ÷ 4"}
                fxc[q] = {"v": 0.0, "kind": "estimate", "basis": "0 가정"}
                derc[q] = {"v": 0.0, "kind": "estimate", "basis": "0 가정"}
                othc[q] = {"v": -0.01, "kind": "estimate", "basis": "잔차 중위"}
                rows["금융손익"]["q"][q]["v"] = round(iv + 0.0 + 0.0 - 0.01, 2)
                # 세전·법인세·순이익·지배NI 연쇄 갱신(항등식 유지)
                pt = round(rows["영업이익"]["q"][q]["v"] + rows["금융손익"]["q"][q]["v"] + rows["기타영업외손익"]["q"][q]["v"], 2)
                rows["세전이익"]["q"][q]["v"] = pt
                tax = round(max(pt, 0) * 0.22, 2)
                rows["법인세비용"]["q"][q]["v"] = tax
                rows["당기순이익"]["q"][q]["v"] = round(pt - tax, 2)
                rows["지배주주순이익"]["q"][q]["v"] = round((pt - tax) * 1.005, 2)
                rows["EPS"]["q"][q]["v"] = round(rows["지배주주순이익"]["q"][q]["v"] * 100 / rows["주식수"]["q"][q]["v"], 1)
        put("금융손익", "이자손익", "이자손익", "억원", intc)
        put("이자손익", "외환손익", "외환손익", "억원", fxc)
        put("외환손익", "파생상품손익", "파생상품손익", "억원", derc)
        put("파생상품손익", "기타금융손익", "기타금융손익", "억원", {q: c for q, c in othc.items() if c})
        # OPM(파생)·EBITDA(= 영업이익 + 감가상각비)·매출조선신규/OP조선신규(정의상 0 → 추정 패널 값)
        put("영업이익", "OPM", "OPM", "%", {q: {"v": round(rows["영업이익"]["q"][q]["v"] / rows["매출액"]["q"][q]["v"], 4), "kind": ("actual" if q <= la else "estimate"), "basis": "영업이익 ÷ 매출액"}
                                        for q in qs if rows["영업이익"]["q"].get(q, {}).get("v") is not None and rows["매출액"]["q"].get(q, {}).get("v")})
        put("감가상각비", "EBITDA", "EBITDA", "억원", {q: {"v": round(rows["영업이익"]["q"][q]["v"] + rows["감가상각비"]["q"][q]["v"], 2), "kind": "estimate", "basis": "영업이익 + 감가상각비"}
                                               for q in fq if rows["감가상각비"]["q"].get(q, {}).get("v") is not None})
        new = {q: {"v": (0.0 if q <= la else round(100.0 * max(0, qs.index(q) - qs.index(la) - 1), 2)), "kind": "estimate", "basis": "forecast_panel base(테스트)"} for q in qs if q >= "2026Q1"}
        put("OP조선", "매출조선신규", "매출 조선 신규수주", "억원", new, group="사업부")
        opm_seg = {q: rows["OP조선"]["q"][q]["v"] / rows["매출조선"]["q"][q]["v"] for q in fq}
        put("매출조선신규", "OP조선신규", "OP 조선 신규수주", "억원", {q: {"v": (0.0 if q <= la else round(new[q]["v"] * opm_seg[q], 2)), "kind": "estimate", "basis": "매출조선신규 × 타겟 OPM"} for q in new}, group="사업부")
        # fin: 외환손익 복합 계정·이자손익·금융손익 BS 값(2026Q2)
        q = la
        cls.fin = {"cons": {"bs": {}, "cf": {}, "is": {q: {"매출액(수익)": round(rows["매출액"]["q"][q]["v"] * 100, 2), "금융손익": round(rows["금융손익"]["q"][q]["v"] * 100, 2),
                                                           "이자손익": -4000.0, "외환차익": 3973.4, "외환차손": 9687.51, "외화환산이익": 4099.38, "외화환산손실": 19773.27,
                                                           "파생상품이익": 104767.46, "파생상품손실": 106367.46}}}}        # 백만원: 외환 합·차 = −21,388 → −213.88억, 파생 = −16.0억
        cls.model, cls.rows, cls.qs, cls.la, cls.fq = m, rows, qs, la, fq
        cls.b = X.Builder(copy.deepcopy(m), fin=cls.fin)
        cls.wb = cls.b.build()
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_detail_")
        cls.out = os.path.join(cls.tmp, "d.xlsx")
        X.save_atomic(cls.wb, cls.out)
        cls.rep = X.verify(cls.out, m)
        cls.lay = X.Layout(m)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_plan_types(self):
        p = self.b.plan
        self.assertEqual(p["금융손익"]["type"], "identity")
        self.assertEqual([k for _, k in p["금융손익"]["expr"]], ["이자손익", "외환손익", "파생상품손익", "기타금융손익"])
        self.assertEqual(p["이자손익"]["type"], "interest")
        self.assertEqual(p["외환손익"]["type"], "abs")
        self.assertEqual(p["기타금융손익"]["type"], "abs")
        self.assertEqual(p["EBITDA"]["type"], "identity")
        self.assertEqual(p["OPM"]["type"], "divide")
        self.assertEqual(p["PER"]["type"] if "PER" in p else "valuation", "valuation")
        self.assertEqual(p["매출조선신규"]["type"], "abs")
        self.assertEqual(p["OP조선신규"], {"type": "ratio_parent", "base": "매출조선신규", "parent": "OP조선", "note": p["OP조선신규"]["note"]})
        self.assertEqual(p["법인세비용"]["type"], "tax")
        self.assertTrue(self.b.links["이자율"]["linked"], self.b.links["이자율"])
        self.assertAlmostEqual(self.b.fit["이자율_자산"], 0.012, places=4)
        self.assertAlmostEqual(self.b.fit["이자율_부채"], 0.05, places=4)

    def test_formulas(self):
        ws, vr, lay = self.wb["subQ"], self.wb["변수"], self.lay
        L, Lp = lay.letter("2026Q3"), lay.letter("2026Q2")
        r = {k: _row_of(ws, k) for k in ("금융손익", "이자손익", "외환손익", "파생상품손익", "기타금융손익", "이자발생자산", "총차입금", "EBITDA", "영업이익", "감가상각비", "OPM", "매출액", "매출조선신규", "OP조선신규")}
        self.assertEqual(ws["%s%d" % (L, r["금융손익"])].value, '=IFERROR(%s%d+%s%d+%s%d+%s%d,"")' % (L, r["이자손익"], L, r["외환손익"], L, r["파생상품손익"], L, r["기타금융손익"]))
        sc_a, sc_d = _row_of(vr, "sc:이자율_자산"), _row_of(vr, "sc:이자율_부채")
        self.assertEqual(ws["%s%d" % (L, r["이자손익"])].value, '=IFERROR((%s%d*변수!$D$%d-%s%d*변수!$D$%d)/4,"")' % (Lp, r["이자발생자산"], sc_a, Lp, r["총차입금"], sc_d))
        self.assertEqual(ws["%s%d" % (L, r["EBITDA"])].value, '=IFERROR(%s%d+%s%d,"")' % (L, r["영업이익"], L, r["감가상각비"]))
        self.assertEqual(ws["%s%d" % (L, r["OPM"])].value, '=IFERROR(%s%d/%s%d,"")' % (L, r["영업이익"], L, r["매출액"]))
        self.assertEqual(ws["%s%d" % (Lp, r["OPM"])].value, '=IFERROR(%s%d/%s%d,"")' % (Lp, r["영업이익"], Lp, r["매출액"]))     # 실적 구간도 파생 수식
        self.assertTrue(ws["%s%d" % (L, r["외환손익"])].value.startswith("=변수!%s" % L))                                      # abs 드라이버(죽은 값 아님)
        self.assertEqual(ws["%s%d" % (L, r["OP조선신규"])].value, '=IFERROR(%s%d*변수!%s%d,"")' % (L, r["매출조선신규"], L, _row_of(vr, "drv:OP조선")))
        # 실적 구간: 외환손익 = BS 계정 4개 합·차 VLOOKUP, 기타금융손익 = 잔차 항등식, 이자손익 = VLOOKUP(BS 값이 모델값과 맞을 때)
        f = ws["%s%d" % (Lp, r["외환손익"])].value
        self.assertTrue(f.startswith('=IFERROR((VLOOKUP("외환차익",BS연결,MATCH(%s$1,BS연결H,0),0)-VLOOKUP("외환차손"' % Lp), f)
        self.assertEqual(ws["%s%d" % (Lp, r["기타금융손익"])].value, '=IFERROR(%s%d-%s%d-%s%d-%s%d,"")' % (Lp, r["금융손익"], Lp, r["이자손익"], Lp, r["외환손익"], Lp, r["파생상품손익"]))
        self.assertIn("VLOOKUP($D%d,BS연결" % r["이자손익"], ws["%s%d" % (Lp, r["이자손익"])].value)
        # 파생상품손익 BS 값(-16.0) == 모델값 → 복합 VLOOKUP
        self.assertIn('VLOOKUP("파생상품이익"', ws["%s%d" % (Lp, _row_of(ws, "파생상품손익"))].value)

    def test_recalc_and_sensitivity(self):
        e, a = self.rep["estimate_recalc"], self.rep["actual_recalc"]
        self.assertEqual(e["mismatch"], [], e["mismatch"][:5])
        self.assertEqual(a["mismatch"], [], a["mismatch"][:5])
        self.assertEqual(self.rep["links"]["value_cells"], 0, self.rep["links"]["value_keys"])
        names = {p["name"]: p for p in self.rep["sensitivity"]}
        self.assertIn("자산이자율 +1%p → 이자손익", names)
        p = names["자산이자율 +1%p → 이자손익"]
        self.assertTrue(p["ok"], p)
        ia_la = self.rows["이자발생자산"]["q"][self.la]["v"]
        self.assertAlmostEqual(p["delta"], ia_la * 0.01 / 4, delta=0.05)
        self.assertTrue(p["chain"][0]["moved"])                                            # 금융손익 연쇄
        # 합성 모델은 선표 연결이 없어(매출조선 = 매출액 비중) 신규수주 프로브는 돌지 않는다 — 실데이터 테스트(TestRealModelLinks)가 담당
        self.assertFalse([p for p in self.rep["sensitivity"] if p["name"].startswith("매출조선신규")])
        self.assertTrue(self.rep["ok"], {k: v for k, v in self.rep.items() if k in ("checks", "estimate_recalc", "actual_recalc", "sensitivity")})

    def test_duplicate_vlookup_key_not_created(self):
        """모델에 OPM 행이 있으면 '비율(참고)' 블록은 OPM 을 다시 만들지 않는다(VLOOKUP 키 중복 → 분기·연간예상이 엉뚱한 행을 읽음)."""
        ws = self.wb["subQ"]
        keys = [ws.cell(r, 1).value for r in range(X.DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value]
        self.assertEqual(keys.count("OPM"), 1)
        self.assertEqual(keys.count("NPM"), 1)


class TestYoyFallbackAndZeroBase(unittest.TestCase):
    def test_yoy_without_prior_year_falls_back_to_qoq(self):
        """전년동기가 없는 회사(상장 직후)의 매출 추정은 전분기 대비 수식으로 — 죽은 값 셀을 남기지 않는다."""
        m = _load_sample()
        qs = m["periods"]["quarters"]
        keep = [q for q in qs if q >= "2025Q4"]
        m["periods"]["quarters"] = keep
        m["periods"]["annual"] = sorted({q[:4] for q in keep})
        for row in m["rows"]:
            row["q"] = {q: c for q, c in row.get("q", {}).items() if q in keep}
            row["a"] = {y: a for y, a in (row.get("a") or {}).items() if y in m["periods"]["annual"]}
        b = X.Builder(m)
        wb = b.build()
        ws, vr, lay = wb["subQ"], wb["변수"], X.Layout(m)
        r = _row_of(ws, "매출액")
        L, Lp = lay.letter("2026Q3"), lay.letter("2026Q2")
        self.assertEqual(ws["%s%d" % (L, r)].value, '=IFERROR(%s%d*(1+변수!%s%d),"")' % (Lp, r, L, _row_of(vr, "drv:매출액")))
        self.assertIn("전년동기 없음 → 전분기 대비", vr["%s%d" % (L, _row_of(vr, "drv:매출액"))].comment.text)
        em = X.Emulator(wb)
        mv = {x["key"]: x for x in m["rows"]}["매출액"]["q"]["2026Q3"]["v"]
        self.assertAlmostEqual(em.value("subQ", "%s%d" % (L, r)), mv, delta=0.05)

    def test_zero_base_zero_value_gives_zero_param(self):
        row = {"key": "총차입금", "q": {"2026Q2": {"v": 0.0}, "2026Q3": {"v": 0.0}}}
        lay = X.Layout({"periods": {"quarters": ["2026Q2", "2026Q3"], "annual": [], "last_actual": "2026Q2"}})
        self.assertEqual(X.implied_param(row, "2026Q3", {"type": "qoq"}, {}, lay), 0.0)
        self.assertIsNone(X.implied_param({"key": "x", "q": {"2026Q2": {"v": 0.0}, "2026Q3": {"v": 5.0}}}, "2026Q3", {"type": "qoq"}, {}, lay))


def _scenario_model(meta_extra=None):
    """표본 모델에 매출조선신규 행과 4개 시나리오(기존만·보수·기준·낙관)를 붙인다. meta_extra 는 scenarios.meta 에 덮어쓸 필드."""
    m = _load_sample()
    rows = {x["key"]: x for x in m["rows"]}
    qs = m["periods"]["quarters"]
    la = m["periods"]["last_actual"]
    fq = [q for q in qs if q > la]
    new_base = {q: round(50.0 * (i + 1), 2) for i, q in enumerate(fq)}
    m["rows"].insert([x["key"] for x in m["rows"]].index("OP조선") + 1,
                     {"key": "매출조선신규", "label": "매출 조선 신규", "group": "사업부", "unit": "억원",
                      "q": {q: {"v": new_base[q], "kind": "estimate", "basis": "panel"} for q in fq}})
    opm = {q: rows["영업이익"]["q"][q]["v"] / rows["매출액"]["q"][q]["v"] for q in fq}
    sc = {"meta": {"source": "test", "calibrated": False, "in_rows": "base", "fiscal_years": ["2026", "2027", "2028"], "note": "t"}}
    sc["meta"].update(meta_extra or {})
    for case, mult in (("existing_only", 0.0), ("conservative", 0.5), ("base", 1.0), ("optimistic", 1.5)):
        qd = {}
        for q in fq:
            ns = new_base[q] * mult
            qd[q] = {"rev": round(rows["매출액"]["q"][q]["v"] - new_base[q] + ns, 2), "op": round(rows["영업이익"]["q"][q]["v"] + (ns - new_base[q]) * opm[q], 2),
                     "new_order_revenue": round(ns, 2), "kind": "estimate"}
        ann = {}
        for y in ("2026", "2027", "2028"):
            ks = ["%sQ%d" % (y, k) for k in range(1, 5)]
            rv = sum(qd[k]["rev"] if k in qd else rows["매출액"]["q"][k]["v"] for k in ks)
            op = sum(qd[k]["op"] if k in qd else rows["영업이익"]["q"][k]["v"] for k in ks)
            ann[y] = {"rev": round(rv, 2), "op": round(op, 2)}
        sc[case] = {"in_rows": case == "base", "quarterly": qd, "annual": ann}
    m["scenarios"] = sc
    return m, fq


class TestScenarioSheet(unittest.TestCase):
    """시나리오 시트 — 모델 scenarios(보수/기준/낙관/기존만) 가 수식으로 들어가고 base 는 subQ 와 같다."""

    def test_scenarios_sheet(self):
        m, fq = _scenario_model()
        b = X.Builder(m)
        wb = b.build()
        self.assertIn("시나리오", wb.sheetnames)
        self.assertEqual(wb.sheetnames.index("시나리오"), wb.sheetnames.index("subQ") + 1)
        ws = wb["시나리오"]
        self.assertEqual(ws.cell(1, 5).value, "3Q26")
        self.assertEqual(ws.cell(1, 7).value, 2026)                   # 4Q26 뒤 연간 열
        rows_s = {ws.cell(r, 1).value: r for r in range(X.DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value}
        self.assertEqual(ws.cell(rows_s["scn:base:신규수주매출"], 5).value, "=subQ!%s%d" % (X.Layout(m).letter("2026Q3"), b.subq_row["매출조선신규"]))
        self.assertTrue(ws.cell(rows_s["scn:optimistic:매출액"], 5).value.startswith("=IFERROR(subQ!"))
        self.assertEqual(ws.cell(rows_s["scn:optimistic:신규수주매출"], 5).fill.fgColor.rgb[-6:], "FFF2AB")
        tmp = tempfile.mkdtemp(prefix="kship_xlsx_scn_")
        try:
            out = os.path.join(tmp, "s.xlsx")
            X.save_atomic(wb, out)
            rep = X.verify(out, m)
            self.assertEqual(rep["scenarios"]["mismatch"], [], rep["scenarios"]["mismatch"][:5])
            self.assertEqual(rep["scenarios"]["cells"], 4 * (len(fq) + 3) * 2)
            self.assertEqual(rep["estimate_recalc"]["mismatch"], [])
            # 표본은 단일 세그먼트(매출조선 = 매출액 비중)라 신규 행이 매출액에 더해지는 구조가 아니다 → 신규수주 프로브는 돌지 않아야 한다(실데이터 테스트가 담당)
            self.assertFalse([p for p in rep["sensitivity"] if p["name"].startswith("매출조선신규")])
            # base 시나리오 매출액 == subQ 매출액(행에 포함된 시나리오)
            em = X.Emulator(load_workbook(out))
            lay = X.Layout(m)
            r_base = [r for r in range(X.DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value == "scn:base:매출액"][0]
            self.assertAlmostEqual(em.value("시나리오", "E%d" % r_base), em.value("subQ", "%s%d" % (lay.letter("2026Q3"), b.subq_row["매출액"])), places=6)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_scenarios_sheet_marks_fallback_low_confidence(self):
        """covered_scope 폴백(한화오션·HJ)이면 시나리오 시트 2행에 폴백 문구, 신규수주 셀 메모에 source_field·폴백·저신뢰."""
        note = "폴백(저신뢰): 패널이 전범위 new_order_revenue 를 비웠다 → covered_scope_new_revenue 사용"
        for extra, fb in (({"source_field": "covered_scope_new_revenue", "fallback": True, "confidence": "low", "fallback_note": note}, True),
                          ({"source_field": "new_order_revenue", "fallback": False, "confidence": "panel_default"}, False)):
            m, fq = _scenario_model(extra)
            b = X.Builder(m)
            wb = b.build()
            ws = wb["시나리오"]
            self.assertEqual("※ " + note in str(ws.cell(2, 1).value), fb)
            self.assertEqual("covered_scope_new_revenue 폴백" in str(ws.cell(2, 1).value), False)       # fallback_note 가 있으면 일반 문구를 대신한다
            rows_s = {ws.cell(r, 1).value: r for r in range(X.DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value}
            comments = [ws.cell(rows_s["scn:%s:신규수주매출" % c], 5).comment for c in ("conservative", "optimistic")]
            texts = [c.text for c in comments if c is not None]
            self.assertTrue(texts)
            for t in texts:
                self.assertIn(extra["source_field"], t)
                self.assertEqual("폴백·저신뢰" in t, fb)
        m, _ = _scenario_model({"source_field": "covered_scope_new_revenue", "fallback": True})     # fallback_note 없으면 기본 문구
        self.assertIn("※ covered_scope_new_revenue 폴백(저신뢰)", str(X.Builder(m).build()["시나리오"].cell(2, 1).value))



@unittest.skipUnless(os.path.exists(REAL_MODEL) and os.path.exists(REAL_SLS) and os.path.exists(os.path.join(X.ASSETS, "fx.json")),
                     "실데이터(assets/models/010140.json·sls·fx) 없음")
class TestRealModelLinks(unittest.TestCase):
    """실데이터(삼성重 010140): 선표 연결(환율 forward → SLS 건조시점 환율 → 매출조선 → 매출액), 시나리오·세율·이자율 연결, 전 분기 재계산.
    수치는 고정하지 않는다 — 구조와 '재계산 == 모델' 만 단언(데이터 표류에 강건)."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_real_")
        cls.out = os.path.join(cls.tmp, "010140_model.xlsx")
        cls.rep = X.build_one(REAL_MODEL, out_path=cls.out, do_verify=True)
        cls.model = X.load_json(REAL_MODEL)
        cls.wb = load_workbook(cls.out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_links_and_sheets(self):
        lk = self.rep["links"]
        self.assertTrue(lk["선표"]["linked"], lk["선표"])
        self.assertLessEqual(lk["선표"]["max_diff_m"], X.TOL_SLS_M)
        for k in ("세율", "판관비율", "지배주주비중", "이자율"):
            self.assertTrue(lk[k]["linked"], (k, lk[k]))
        self.assertEqual(self.wb.sheetnames, ["README", "변수", "BS연결", "BS별도", "subQ", "SLS", "시나리오", "분기", "연간예상", "TP", "외화"])

    def test_sales_formula_references_sls_and_fx(self):
        ws, lay = self.wb["subQ"], X.Layout(self.model)
        r = _row_of(ws, "매출조선")
        q = lay.est_quarters[0]
        f = ws["%s%d" % (lay.letter(q), r)].value
        self.assertTrue(f.startswith("=IFERROR(SLS!"), f)
        self.assertIn("+변수!%s" % lay.letter(q), f)
        self.assertIn("+%s%d" % (lay.letter(q), _row_of(ws, "매출조선신규")), f)
        # 매출액 = 매출조선 + 매출기타(세그먼트 합 항등식), 영업이익 = OP조선 + OP기타
        self.assertEqual(ws["%s%d" % (lay.letter(q), _row_of(ws, "매출액"))].value, '=IFERROR(%s%d+%s%d,"")' % (lay.letter(q), r, lay.letter(q), _row_of(ws, "매출기타")))
        sl = self.wb["SLS"]
        srow = {sl.cell(rr, 1).value: rr for rr in range(X.DATA_ROW, sl.max_row + 1) if sl.cell(rr, 1).value}
        self.assertIn('VLOOKUP(" 원/달러(평균)",환율', sl.cell(srow["spot_krw"], 5).value)

    def test_full_recalc_and_sensitivity(self):
        v = self.rep["verify"]
        self.assertTrue(v["ok"], {k: v[k] for k in ("estimate_recalc", "actual_recalc", "scenarios") if k in v})
        self.assertEqual(v["links"]["value_cells"], 0, v["links"]["value_keys"])
        self.assertEqual(v["estimate_recalc"]["mismatch"], [])
        self.assertEqual(v["actual_recalc"]["mismatch"], [])
        self.assertEqual(v["scenarios"]["mismatch"], [])
        names = {p["name"].split(" ")[0]: p for p in v["sensitivity"]}
        for k in ("세율", "OPM(OP조선)", "원/달러(평균)", "자산이자율", "매출조선신규"):
            self.assertIn(k, names, list(names))
            self.assertTrue(names[k]["ok"], names[k])
        fxp = names["원/달러(평균)"]
        self.assertGreater(fxp["delta"], 0)                        # 환율 +100원 → 매출조선 증가(헤지 안 된 30%)
        self.assertTrue(all(c["moved"] for c in fxp["chain"]))      # 매출액·영업이익 연쇄


# ── 라운드 7 (p) — 조정EPS 행: 추정 = EPS 참조(identity, 변수 드라이버 행 없음), 의심 분기 = 세후 차감 산식 셀 ──────────────

def _adjusted_eps_model(q="2026Q2", excess_after_tax=1000.0, ctrl_share=1.0):
    """표본 모델에 kship_model.py 라운드 6 꼴의 조정EPS 행·assumptions.one_offs_detected 를 얹는다. 의심 분기 eps_adj = (지배NI − 세후 초과 × 지배 비중)×100/주식수(1자리)."""
    m = _load_sample()
    rows = {x["key"]: x for x in m["rows"]}
    eps, la = rows["EPS"], m["periods"]["last_actual"]
    ni, sh = rows["지배주주순이익"]["q"][q]["v"], rows["주식수"]["q"][q]["v"]
    eps_adj = round((ni - excess_after_tax * ctrl_share) * 100 / sh, 1)
    row = {"key": "조정EPS", "label": "조정 EPS(일회성 의심 분기의 초과 비영업손익 세후 차감 — 모델 추정)", "group": "주당", "unit": "원", "q": {}, "a": {}}
    for qq, c in eps["q"].items():
        if qq == q:
            row["q"][qq] = {"v": eps_adj, "kind": "estimate", "basis": "지배NI %s억 − 세후 초과 %s억 × 지배 비중 %.2f ÷ 유통주식수 — 모의" % (ni, excess_after_tax, ctrl_share)}
        elif qq <= la:
            row["q"][qq] = {"v": c["v"], "kind": "actual", "src": "= EPS(일회성 의심 없음)"}
        else:
            row["q"][qq] = {"v": c["v"], "kind": "estimate", "basis": "= EPS(추정 구간은 일회성 미가정)"}
    for y, c in (eps.get("a") or {}).items():
        row["a"][y] = dict(c, v=(round(c["v"] - eps["q"][q]["v"] + eps_adj, 1) if (y == q[:4] and isinstance(c.get("v"), (int, float))) else c.get("v")))
    m["rows"].insert(m["rows"].index(eps) + 1, row)
    for v in (m.get("views") or {}).values():                       # 표본의 보고서 뷰는 키 목록 — 실물(kship_model.py)처럼 EPS 뒤에 조정EPS
        if isinstance(v, list) and v and isinstance(v[0], list) and "EPS" in v[0]:
            v[0].insert(v[0].index("EPS") + 1, "조정EPS")
    m["assumptions"]["one_offs_detected"] = [{
        "q": q, "nonop": 1500.0, "op": 300.0, "excess_nonop": 1282.05, "baseline_nonop_12q": 217.95, "tax_rate_applied": 0.22, "tax_rate_source": "분기 유효세율(모의)",
        "ctrl_share_applied": ctrl_share, "ctrl_share_source": "그 분기 지배NI/NI", "excess_after_tax": excess_after_tax, "ni_ctrl_adj": round(ni - excess_after_tax * ctrl_share, 2),
        "eps_reported": eps["q"][q]["v"], "eps_adj": eps_adj, "note": "모의 일회성"}]
    return m, eps_adj


class TestAdjustedEpsRow(unittest.TestCase):
    """합성(표본 + 조정EPS 행): plan identity · 변수 drv 행 없음 · 추정/비의심 셀 '=EPS 셀' · 의심 셀 산식+FILL_EST+메모 · verify 0 mismatch · README 산식 줄 · VLOOKUP 키 1개."""

    @classmethod
    def setUpClass(cls):
        cls.model, cls.eps_adj = _adjusted_eps_model()
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_adj_")
        cls.out = os.path.join(cls.tmp, "adj.xlsx")
        cls.b = X.Builder(cls.model, fin=None)
        cls.wb_built = cls.b.build()
        X.save_atomic(cls.wb_built, cls.out)
        cls.rep = X.verify(cls.out, cls.model)
        cls.wb = load_workbook(cls.out)
        cls.lay = X.Layout(cls.model)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_plan_identity_and_no_driver_row(self):
        d = self.b.plan["조정EPS"]
        self.assertEqual((d["type"], d["expr"]), ("identity", [("+", "EPS")]))
        self.assertIn("일회성", d["note"])
        vs = self.wb["변수"]
        keys = [vs.cell(r, 1).value for r in range(X.DATA_ROW, vs.max_row + 1)]
        self.assertNotIn("drv:조정EPS", keys)
        self.assertIn("drv:매출액", keys)

    def test_cells(self):
        ws = self.wb["subQ"]
        r, r_eps, r_ni, r_sh = _row_of(ws, "조정EPS"), _row_of(ws, "EPS"), _row_of(ws, "지배주주순이익"), _row_of(ws, "주식수")
        lay = self.lay
        for q in lay.est_quarters[:3]:                                                   # 추정 구간 = EPS 셀 참조
            L = lay.letter(q)
            self.assertEqual(ws["%s%d" % (L, r)].value, '=IFERROR(%s%d,"")' % (L, r_eps))
        Lp = lay.letter("2026Q1")                                                        # 비의심 실적 분기도 = EPS 셀(값 직접 아님)
        self.assertEqual(ws["%s%d" % (Lp, r)].value, '=IFERROR(%s%d,"")' % (Lp, r_eps))
        self.assertNotEqual(ws["%s%d" % (Lp, r)].fill.fgColor.rgb[-6:], "FFF9DB")
        Lq = lay.letter("2026Q2")                                                        # 의심 분기 = 산식 셀
        c = ws["%s%d" % (Lq, r)]
        self.assertTrue(c.value.startswith('=IFERROR((%s%d-(' % (Lq, r_ni)), c.value)
        self.assertTrue(c.value.endswith('*100/%s%d,"")' % (Lq, r_sh)), c.value)
        self.assertEqual(c.value, '=IFERROR((%s%d-(1000.0)*(1.0))*100/%s%d,"")' % (Lq, r_ni, Lq, r_sh))
        self.assertEqual(c.fill.fgColor.rgb[-6:], "FFF9DB")
        for needle in ("12분기 기준선", "보고 EPS", "재계산되지 않음", "조정EPS 산식(모델 추정", "1000.0·1.0"):
            self.assertIn(needle, c.comment.text)
        em = X.Emulator(self.wb)
        self.assertAlmostEqual(em.value("subQ", "%s%d" % (Lq, r)), self.eps_adj, delta=X.tol_for("조정EPS", "원", self.eps_adj))
        self.assertAlmostEqual(em.value("subQ", "%s%d" % (Lq, r)), (2243.33 - 1000.0) * 100 / 854.15, places=6)
        self.assertEqual(ws["%s%d" % (Lq, r_eps)].value, '=IFERROR(%s%d*100/%s%d,"")' % (Lq, r_ni, Lq, r_sh))   # 보고 EPS 행은 그대로

    def test_verify_readme_and_views(self):
        self.assertEqual(self.rep["estimate_recalc"]["mismatch"], [])
        self.assertEqual(self.rep["actual_recalc"]["mismatch"], [])
        self.assertNotIn("조정EPS", self.rep["links"]["value_keys"])
        txt = "\n".join(str(c.value) for row in self.wb["README"].iter_rows() for c in row if c.value is not None)
        self.assertIn("· 조정EPS(일회성 의심 1분기: 2026Q2 보고 263→조정 %s원)" % format(round(self.eps_adj), ",d"), txt)
        self.assertIn("조정EPS 는 EPS 참조(일회성 의심 분기만 산식 셀)", txt)
        for name in ("분기", "연간예상"):
            ws = self.wb[name]
            self.assertEqual(sum(1 for r in range(X.DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value == "조정EPS"), 1, name)

    def test_without_one_off_constants_falls_back_to_value(self):
        m, eps_adj = _adjusted_eps_model()
        del m["assumptions"]["one_offs_detected"][0]["excess_after_tax"]
        wb = X.Builder(m, fin=None).build()
        ws = wb["subQ"]
        c = ws["%s%d" % (X.Layout(m).letter("2026Q2"), _row_of(ws, "조정EPS"))]
        self.assertEqual(c.value, eps_adj)                                              # 상수 없음 → 기존 값 직접 경로(메모 '갭필')
        self.assertIn("추정(실적 구간 갭필)", c.comment.text)


REAL_ADJ = os.path.join(X.MODELS_DIR, "002380.json")


@unittest.skipUnless(os.path.exists(REAL_ADJ) and os.path.exists(os.path.join(X.FIN_DIR, "002380.json")) and os.path.exists(os.path.join(X.ASSETS, "prices.json")),
                     "실데이터(assets/models/002380.json·fin·prices) 없음")
class TestRealAdjustedEps(unittest.TestCase):
    """실데이터(케이씨씨 002380 — 2026Q2 일회성 의심): build_one → tmp, verify ok, 의심 분기 수식 셀 에뮬레이트 ≈ one_offs_detected.eps_adj, 다음 분기 '=EPS 셀'. 운영 xlsx 는 쓰지 않는다."""

    @classmethod
    def setUpClass(cls):
        cls.model = X.load_json(REAL_ADJ)
        det = (cls.model.get("assumptions") or {}).get("one_offs_detected") or []
        if not det or "조정EPS" not in {r["key"] for r in cls.model["rows"]}:
            raise unittest.SkipTest("002380 모델에 조정EPS 행/one_offs_detected 없음")
        cls.det = det
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_real_adj_")
        cls.out = os.path.join(cls.tmp, "002380_model.xlsx")
        cls.rep = X.build_one(REAL_ADJ, out_path=cls.out, do_verify=True)
        cls.wb = load_workbook(cls.out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_verify_and_cells(self):
        v = self.rep["verify"]
        self.assertTrue(v["ok"], {k: v[k] for k in ("estimate_recalc", "actual_recalc") if k in v})
        self.assertEqual(v["actual_recalc"]["mismatch"], [])
        self.assertEqual(v["estimate_recalc"]["mismatch"], [])
        ws, lay = self.wb["subQ"], X.Layout(self.model)
        r, r_eps = _row_of(ws, "조정EPS"), _row_of(ws, "EPS")
        em = X.Emulator(self.wb)
        for o in self.det:
            L = lay.letter(o["q"])
            c = ws["%s%d" % (L, r)]
            self.assertTrue(str(c.value).startswith('=IFERROR((%s%d-(' % (L, _row_of(ws, "지배주주순이익"))), c.value)
            self.assertAlmostEqual(em.value("subQ", "%s%d" % (L, r)), o["eps_adj"], delta=X.tol_for("조정EPS", "원", o["eps_adj"]))
            self.assertIn("재계산되지 않음", c.comment.text)
        q1 = lay.est_quarters[0]
        self.assertEqual(ws["%s%d" % (lay.letter(q1), r)].value, '=IFERROR(%s%d,"")' % (lay.letter(q1), r_eps))
        vs = self.wb["변수"]
        self.assertNotIn("drv:조정EPS", [vs.cell(rr, 1).value for rr in range(X.DATA_ROW, vs.max_row + 1)])


# ── 라운드 7 수정 레인 xlsx_links — 이월결손 램프(assumptions.tax_schedule) 세율 링크·드라이버 셀·민감도 ① · 신규수주 문구(scenarios.meta) ──────────────

TAX_LOW, TAX_TERMINAL = 0.0068, 0.22          # 합성 이월결손 램프 — 한화오션 042660 과 같은 시작·종착 세율


def _tax_schedule_model(nonop_override=None):
    """표본 모델에 kship_model.py 라운드 7 L2 꼴의 이월결손 램프(carryforward_ramp)를 얹는다 — origin 연도 분기는 저세율 유지, 이후 8분기 선형 램프 → 22%(r4).
    추정 분기 법인세 = round(max(세전,0) × r4 세율, 2) 로 다시 계산하고 당기순이익·지배주주순이익(×1.005)·EPS·자본총계/지배주주지분 롤·자산총계·BPS·순차입금(−NI×50%)·연간을 따라 고친다.
    nonop_override = {q: 기타영업외손익} — 세전 ≤ 0 분기를 만들 때(세전 항등식도 다시 계산). 반환 (모델, r4 스케줄)."""
    m = _load_sample()
    rows = {x["key"]: x for x in m["rows"]}
    la = m["periods"]["last_actual"]
    fq = [q for q in m["periods"]["quarters"] if q > la]
    hold = [q for q in fq if q[:4] == la[:4]]
    ramp = [q for q in fq if q[:4] > la[:4]]
    sched = {q: TAX_LOW for q in hold}
    sched.update({q: round(TAX_LOW + (TAX_TERMINAL - TAX_LOW) * (i + 1) / len(ramp), 4) for i, q in enumerate(ramp)})
    stage = {q: ("유지" if q in hold else "램프 %d/%d" % (ramp.index(q) + 1, len(ramp))) for q in fq}
    m["assumptions"].update({
        "tax_rate": TAX_LOW, "tax_rate_terminal": TAX_TERMINAL, "tax_path": "carryforward_ramp", "tax_schedule": sched,
        "tax_basis": "최근 8분기 중 세전 > 0 인 7분기(2024Q4~2026Q2) 법인세/세전 중위 0.7%·최근 4분기 중위 -1.2% 가 모두 5% 미만 → 이월결손 공제 중으로 보고 0.7% 를 {} 까지 유지, {}~{} 선형 램프로 법정세율 근사 22% 수렴 — 모의".format(hold[-1], ramp[0], ramp[-1]),
        "tax_carryforward": {"median_pos8": TAX_LOW, "median_recent4": -0.0121, "n_pos8": 7, "quarters": ["2024Q4", "2025Q1", "2025Q2", "2025Q3", "2025Q4", "2026Q1", "2026Q2"],
                             "hold_until": hold[-1], "ramp_from": ramp[0], "ramp_to": ramp[-1]}})

    def v(k, q):
        return rows[k]["q"][q]["v"]

    def put(k, q, x, basis=None):
        rows[k]["q"][q]["v"] = x
        if basis:
            rows[k]["q"][q]["basis"] = basis
    prev = la
    for q in fq:
        if nonop_override and q in nonop_override:
            put("기타영업외손익", q, nonop_override[q], "모의: 일회성 비영업손실(세전 ≤ 0 분기)")
            put("세전이익", q, round(v("영업이익", q) + v("금융손익", q) + v("기타영업외손익", q), 2))
        pt = v("세전이익", q)
        put("법인세비용", q, round(max(pt, 0.0) * sched[q], 2), "max(세전, 0) × 세율 %.1f%%(이월결손 경로 %s — 모의)" % (sched[q] * 100, stage[q]))
        put("당기순이익", q, round(pt - v("법인세비용", q), 2))
        put("지배주주순이익", q, round(v("당기순이익", q) * 1.005, 2))
        put("EPS", q, round(v("지배주주순이익", q) * 100 / v("주식수", q), 2))
        put("자본총계", q, round(v("자본총계", prev) + v("당기순이익", q), 2))
        put("지배주주지분", q, round(v("지배주주지분", prev) + v("지배주주순이익", q), 2))
        put("자산총계", q, round(v("부채총계", q) + v("자본총계", q), 2))
        put("BPS", q, round(v("지배주주지분", q) * 100 / v("주식수", q), 2))
        put("순차입금", q, round(v("순차입금", prev) - v("당기순이익", q) * 0.5, 2))
        prev = q
    for y in sorted({q[:4] for q in fq}):                                             # 연간: 흐름 = 분기 합, 잔액 = Q4
        qs = [q for q in m["periods"]["quarters"] if q[:4] == y]
        for k in ("기타영업외손익", "세전이익", "법인세비용", "당기순이익", "지배주주순이익", "EPS"):
            rows[k]["a"][y]["v"] = round(sum(v(k, q) for q in qs), 2)
        for k in ("자본총계", "지배주주지분", "자산총계", "BPS", "순차입금"):
            rows[k]["a"][y]["v"] = v(k, qs[-1])
    return m, sched


def _readme_text(wb):
    return "\n".join(str(c.value) for row in wb["README"].iter_rows() for c in row if c.value is not None)


class TestWordingHelpers(unittest.TestCase):
    """순수 함수 — scenarios.meta → 신규수주 문구 키, tax_schedule 요약(회사 페이지 fmt_pct_r4 와 같은 자릿수)."""

    def test_new_orders_wording(self):
        self.assertEqual(X.new_orders_wording(None), ("panel", "panel"))
        self.assertEqual(X.new_orders_wording({"source_field": "new_order_revenue"}), ("panel", "panel"))
        self.assertEqual(X.new_orders_wording({"source_field": "covered_scope_new_revenue", "post_origin_mode": "net_panel"}), ("net_panel", "net_panel"))
        self.assertEqual(X.new_orders_wording({"source_field": "ledger_signing_rate", "post_origin_mode": "net_panel"}), ("ledger_signing_rate", "net_panel"))
        self.assertEqual(X.new_orders_wording({"source_field": "ledger_signing_rate"}), ("ledger_signing_rate", "panel"))     # 원장 폴백이 상계 모드보다 우선
        self.assertEqual(set(X.NEW_ORDERS_LABEL_KO), {"ledger_signing_rate", "net_panel", "panel"})
        self.assertEqual(set(X.NEW_ORDERS_POST_KO), {"net_panel", "panel"})

    def test_pct_r4_and_tax_schedule_text(self):
        self.assertEqual([X.pct_r4(x) for x in (0.0045, 0.0068, 0.22, 0.1134, 0.0)], ["0.45%", "0.68%", "22%", "11.34%", "0%"])
        asm = {"tax_schedule": {"2026Q3": 0.0068, "2026Q4": 0.0068, "2027Q1": 0.0334, "2028Q4": 0.22}, "tax_rate_terminal": 0.22,
               "tax_carryforward": {"hold_until": "2026Q4", "ramp_from": "2027Q1", "ramp_to": "2028Q4"}}
        self.assertEqual(X.tax_schedule_text(asm), ("0.68%", "22%", "2026Q4 까지 유지 → 2027Q1~2028Q4 선형 램프"))
        self.assertEqual(X.tax_schedule_text({"tax_schedule": {"2026Q3": 0.0045, "2028Q4": 0.2}}), ("0.45%", "20%", "2026Q3~2028Q4 선형 램프"))   # terminal·구간 없으면 스케줄 끝값·첫~끝


class TestTaxScheduleLinks(unittest.TestCase):
    """합성(표본 + 이월결손 램프): '세율' 스칼라 연결 해제 + tax_schedule 사유 · 드라이버 셀 = 분기별 세율(역산, r4 스케줄과 0.01억 반올림 안) · 법인세 = MAX(세전,0)×드라이버 ·
    민감도 ① 은 첫 추정 분기 드라이버 셀 프로브 · README/3절 '세율' 에 terminal·램프 구간 병기 · verify ok."""

    @classmethod
    def setUpClass(cls):
        cls.model, cls.sched = _tax_schedule_model()
        cls.rows = {r["key"]: r for r in cls.model["rows"]}
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_tax_")
        cls.out = os.path.join(cls.tmp, "tax.xlsx")
        cls.b = X.Builder(cls.model, fin=None)
        X.save_atomic(cls.b.build(), cls.out)
        cls.rep = X.verify(cls.out, cls.model)
        cls.wb = load_workbook(cls.out)
        cls.lay = X.Layout(cls.model)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_links_reason_names_schedule_not_guess(self):
        lk = self.b.links["세율"]
        self.assertFalse(lk["linked"])
        self.assertTrue(lk["schedule"])
        self.assertEqual(lk["reason"], "이월결손 램프(assumptions.tax_schedule, 0.68%→22%) — 분기별 세율 드라이버 행 사용")
        self.assertNotIn("지주 자회사 합산", lk["reason"])
        self.assertEqual(lk["n"], len(self.lay.est_quarters))
        self.assertGreater(lk["max_err"], X.TOL_EOK)                                     # 단일 세율로는 재현 불가 — 그래도 사유는 스케줄
        self.assertNotIn("세율", self.b.fit)
        d = self.b.plan["법인세비용"]
        self.assertEqual((d["type"], d.get("scalar"), d["schedule"]), ("tax", None, self.sched))
        for k in ("판관비율", "지배주주비중"):                                           # 다른 스칼라 연결은 그대로
            self.assertTrue(self.b.links[k]["linked"], k)

    def test_driver_cells_follow_schedule(self):
        vr, ws = self.wb["변수"], self.wb["subQ"]
        drv, sc = _row_of(vr, "drv:법인세비용"), _row_of(vr, "sc:세율")
        r_tax, r_pt = _row_of(ws, "법인세비용"), _row_of(ws, "세전이익")
        self.assertEqual(vr.cell(drv, 2).value, "tax")
        self.assertIn("분기별 세율(이월결손 램프, assumptions.tax_schedule)", vr.cell(drv, 3).value)
        for q in self.lay.est_quarters:
            L = self.lay.letter(q)
            c = vr["%s%d" % (L, drv)]
            self.assertIsInstance(c.value, float, (q, c.value))                           # =$D$n 스칼라 참조가 아니라 분기별 숫자
            pt = self.rows["세전이익"]["q"][q]["v"]
            self.assertLessEqual(abs(c.value - self.sched[q]), 0.005 / pt + 1e-9, (q, c.value, self.sched[q]))   # 법인세 0.01억 반올림만큼만 어긋난다
            self.assertEqual(c.fill.fgColor.rgb[-6:], "FFF2AB")
            self.assertIn("모델 가정 tax_schedule %s(r4)" % self.sched[q], c.comment.text)
            self.assertEqual(ws["%s%d" % (L, r_tax)].value, '=IFERROR(MAX(%s%d,0)*변수!%s%d,"")' % (L, r_pt, L, drv))
        self.assertAlmostEqual(vr["D%d" % sc].value, TAX_LOW, places=6)                 # 3절 '세율' 은 참고(시작 세율), 수식이 참조하지 않는다
        note = vr.cell(sc, 5).value
        self.assertTrue(note.startswith("참고(연결 안 됨: 이월결손 램프(assumptions.tax_schedule, 0.68%→22%) — 분기별 세율 드라이버 행 사용) — 모델 assumptions 값(시작 0.68% → terminal 22%, "
                                        "2026Q4 까지 유지 → 2027Q1~2028Q4 선형 램프 — 분기별 값은 2절 drv:법인세비용 행)"), note)
        self.assertNotIn("=$D$%d" % sc, [str(vr["%s%d" % (self.lay.letter(q), drv)].value) for q in self.lay.est_quarters])

    def test_sensitivity_probes_driver_cell(self):
        names = {p["name"]: p for p in self.rep["sensitivity"]}
        self.assertNotIn("세율 +1%p → 법인세비용", names)                                  # 스칼라 프로브 대신
        p = names["세율 2026Q3 드라이버 셀 +1%p → 법인세비용"]
        self.assertTrue(p["ok"], p)
        L = self.lay.letter("2026Q3")
        self.assertEqual(p["changed"], ["변수!%s%d" % (L, _row_of(self.wb["변수"], "drv:법인세비용"))])
        self.assertAlmostEqual(p["expected_delta"], self.rows["세전이익"]["q"]["2026Q3"]["v"] * 0.01, places=6)
        self.assertAlmostEqual(p["delta"], p["expected_delta"], delta=X.TOL_EOK)
        self.assertEqual([c["cell"].split("!")[0] for c in p["chain"]], ["subQ", "subQ"])
        self.assertTrue(all(c["moved"] for c in p["chain"]))                               # 당기순이익·EPS 연쇄
        self.assertEqual([x for x in self.rep["sensitivity"] if x.get("skipped")], [])
        self.assertEqual(self.rep["estimate_recalc"]["mismatch"], [])
        self.assertEqual(self.rep["actual_recalc"]["mismatch"], [])
        self.assertEqual(self.rep["links"]["value_cells"], 0, self.rep["links"]["value_keys"])
        self.assertTrue(self.rep["ok"], {k: self.rep[k] for k in ("checks", "estimate_recalc", "sensitivity")})
        compact = X._compact({"out": self.out, "size_bytes": self.rep["size_bytes"], "links": self.b.links, "verify": self.rep})
        self.assertEqual(compact["links"]["세율"], "n")
        self.assertEqual(compact["sensitivity"], "%d/%d" % (len(self.rep["sensitivity"]), len(self.rep["sensitivity"])))

    def test_readme_wording(self):
        txt = _readme_text(self.wb)
        self.assertIn("· 세율 0.68% → 22%(terminal, 2026Q4 까지 유지 → 2027Q1~2028Q4 선형 램프) — 이월결손 램프(assumptions.tax_schedule): 변수 drv:법인세비용 추정 분기 셀이 분기별 세율", txt)
        self.assertIn("· 근거: " + self.model["assumptions"]["tax_basis"], txt)
        self.assertIn("연결된 가정: 판관비율, 지배주주비중", txt)
        self.assertIn("세율(이월결손 램프(assumptions.tax_schedule, 0.68%→22%) — 분기별 세율 드라이버 행 사용)", txt)
        for bad in ("· 세율 22.00%", "· 세율 0.68% —", "지주 자회사 합산", "22.00%"):
            self.assertNotIn(bad, txt)

    def test_nonpositive_pretax_quarter_uses_schedule_value_and_probe_moves_on(self):
        """세전 ≤ 0 분기는 법인세 0 → 역산 불가 → 드라이버 셀 = tax_schedule 값(메모), 민감도 ① 은 다음 양(+)세전 분기 드라이버 셀로."""
        m, sched = _tax_schedule_model(nonop_override={"2026Q3": -3000.0})
        rows = {r["key"]: r for r in m["rows"]}
        self.assertLess(rows["세전이익"]["q"]["2026Q3"]["v"], 0)
        self.assertEqual(rows["법인세비용"]["q"]["2026Q3"]["v"], 0.0)
        tmp = tempfile.mkdtemp(prefix="kship_xlsx_tax_neg_")
        try:
            out = os.path.join(tmp, "neg.xlsx")
            b = X.Builder(m, fin=None)
            X.save_atomic(b.build(), out)
            rep = X.verify(out, m)
            wb, lay = load_workbook(out), X.Layout(m)
            vr = wb["변수"]
            drv = _row_of(vr, "drv:법인세비용")
            c = vr["%s%d" % (lay.letter("2026Q3"), drv)]
            self.assertEqual(c.value, sched["2026Q3"])
            self.assertIn("모델 assumptions.tax_schedule(세전 ≤ 0 분기 — 법인세 0 이라 역산 불가)", c.comment.text)
            self.assertAlmostEqual(vr["%s%d" % (lay.letter("2026Q4"), drv)].value, sched["2026Q4"], places=5)
            names = {p["name"]: p for p in rep["sensitivity"]}
            p = names["세율 2026Q4 드라이버 셀 +1%p → 법인세비용"]
            self.assertTrue(p["ok"], p)
            self.assertAlmostEqual(p["expected_delta"], rows["세전이익"]["q"]["2026Q4"]["v"] * 0.01, places=6)
            self.assertEqual([x for x in rep["sensitivity"] if x.get("skipped")], [])
            self.assertEqual(rep["estimate_recalc"]["mismatch"], [])
            self.assertTrue(rep["ok"])
            self.assertTrue(b.links["세율"]["schedule"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_all_nonpositive_pretax_marks_probe_skipped(self):
        """추정 전 분기 세전 ≤ 0 이면 세율 프로브를 조용히 빼지 않고 skipped 항목으로 남기고(ok 집계 제외) --all 요약에 '+1 skipped'."""
        fq = [q for q in self.model["periods"]["quarters"] if q > self.model["periods"]["last_actual"]]
        m, sched = _tax_schedule_model(nonop_override={q: -5000.0 for q in fq})
        tmp = tempfile.mkdtemp(prefix="kship_xlsx_tax_skip_")
        try:
            out = os.path.join(tmp, "skip.xlsx")
            b = X.Builder(m, fin=None)
            X.save_atomic(b.build(), out)
            rep = X.verify(out, m)
            sk = [p for p in rep["sensitivity"] if p.get("skipped")]
            self.assertEqual([p["name"] for p in sk], ["세율 +1%p → 법인세비용"])
            self.assertIn("추정 전 분기 세전 ≤ 0", sk[0]["why"])
            self.assertNotIn("ok", sk[0])
            self.assertEqual(rep["estimate_recalc"]["mismatch"], [])
            self.assertTrue(rep["ok"], [p for p in rep["sensitivity"] if not p.get("ok")])
            self.assertEqual(b.links["세율"], {"linked": False, "schedule": True, "n": len(fq), "max_err": None,
                                                "reason": "이월결손 램프(assumptions.tax_schedule, 0.68%→22%) — 분기별 세율 드라이버 행 사용"})
            vr = load_workbook(out)["변수"]
            drv = _row_of(vr, "drv:법인세비용")
            self.assertEqual([vr["%s%d" % (X.Layout(m).letter(q), drv)].value for q in fq], [sched[q] for q in fq])   # 전부 스케줄 값
            compact = X._compact({"out": out, "size_bytes": rep["size_bytes"], "links": b.links, "verify": rep})
            n_ok = sum(1 for p in rep["sensitivity"] if p.get("ok"))
            self.assertEqual(compact["sensitivity"], "%d/%d +1 skipped" % (n_ok, len(rep["sensitivity"]) - 1))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestNewOrdersWording(unittest.TestCase):
    """합성 시나리오 모델: 변수 매출조선신규 라벨·README 신규수주 줄·시나리오 셀 메모가 scenarios.meta(source_field·post_origin_mode)를 따른다."""

    def test_label_readme_and_comment_follow_meta(self):
        cases = ((None, "forecast_panel", "forecast_panel base 신규수주 매출", "통계 흐름, 공시 계약 개별 반영 아님"),
                 ({"source_field": "covered_scope_new_revenue", "post_origin_mode": "net_panel", "fallback": True}, "forecast_panel",
                  "forecast_panel base · 공시 상계 신규수주 매출", "공시 계약은 체결 분기별로 패널 신규수주와 상계(SLS 일정 유지)"),
                 ({"source_field": "ledger_signing_rate", "post_origin_mode": "net_panel", "fallback": True}, "공시 계약 원장 체결 속도 폴백(저신뢰)",
                  "공시 계약 원장 체결 속도 폴백 신규수주 매출", "공시 계약은 체결 분기별로 패널 신규수주와 상계(SLS 일정 유지)"))
        for extra, src_txt, label, post in cases:
            m, fq = _scenario_model(extra)
            m["new_orders"] = {"scenario_in_rows": "base", "calibrated": False, "panel_status": "partial"}
            b = X.Builder(m)
            wb = b.build()
            vr = wb["변수"]
            self.assertEqual(vr.cell(_row_of(vr, "drv:매출조선신규"), 3).value, "매출조선신규 (%s, 억원 — 가정)" % label, extra)
            txt = _readme_text(wb)
            self.assertIn("· 신규수주 — %s base 시나리오를 매출조선신규 행에 포함(calibrated=False, status partial — %s). 보수/낙관/기존만 은 시나리오 시트에만(합산 안 함)" % (src_txt, post), txt)
            self.assertEqual("공시 계약 개별 반영 아님" in txt, extra is None, extra)
            ws = wb["시나리오"]
            rows_s = {ws.cell(r, 1).value: r for r in range(X.DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value}
            cm = ws.cell(rows_s["scn:conservative:신규수주매출"], 5).comment.text
            self.assertTrue(cm.startswith(("공시 계약 원장 체결 속도 conservative " if extra and extra["source_field"] == "ledger_signing_rate" else "forecast_panel conservative ")), cm)
            self.assertEqual(" · 폴백·저신뢰" in cm, bool(extra and extra.get("fallback")))


REAL_CF = os.path.join(X.MODELS_DIR, "042660.json")
REAL_LEDGER = os.path.join(X.MODELS_DIR, "097230.json")


@unittest.skipUnless(os.path.exists(REAL_CF) and os.path.exists(REAL_LEDGER) and os.path.exists(os.path.join(X.ASSETS, "prices.json")),
                     "실데이터(assets/models/042660.json·097230.json·prices) 없음")
class TestRealTaxScheduleAndLedgerWording(unittest.TestCase):
    """실데이터(한화오션 042660 — carryforward_ramp · HJ중공업 097230 — 원장 체결 속도 폴백): build_one → tmp(do_verify), verify ok, 세율 링크 사유 = tax_schedule,
    드라이버 셀 ≈ tax_schedule(r4 5e-5 + 0.01억 반올림 안), 세율 드라이버 프로브 ok, 신규수주 문구 = scenarios.meta. 운영 xlsx 는 쓰지 않고 수치는 고정하지 않는다."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="kship_xlsx_real_tax_")
        cls.reps, cls.models, cls.wbs = {}, {}, {}
        for p in (REAL_CF, REAL_LEDGER):
            m = X.load_json(p)
            if not isinstance((m.get("assumptions") or {}).get("tax_schedule"), dict):
                raise unittest.SkipTest("%s 모델에 tax_schedule 없음" % m["stock"])
            out = os.path.join(cls.tmp, "%s_model.xlsx" % m["stock"])
            cls.reps[m["stock"]] = X.build_one(p, out_path=out, do_verify=True)
            cls.models[m["stock"]] = m
            cls.wbs[m["stock"]] = load_workbook(out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_tax_schedule_links_and_probe(self):
        for stock, rep in self.reps.items():
            m, wb, lay = self.models[stock], self.wbs[stock], X.Layout(self.models[stock])
            asm, rows = m["assumptions"], {r["key"]: r for r in m["rows"]}
            v = rep["verify"]
            self.assertTrue(v["ok"], (stock, {k: v[k] for k in ("estimate_recalc", "actual_recalc") if k in v}, [p for p in v["sensitivity"] if not p.get("ok")]))
            s0, st, _ = X.tax_schedule_text(asm)
            lk = rep["links"]["세율"]
            self.assertEqual((lk["linked"], lk["schedule"]), (False, True), stock)
            self.assertEqual(lk["reason"], "이월결손 램프(assumptions.tax_schedule, %s→%s) — 분기별 세율 드라이버 행 사용" % (s0, st))
            self.assertEqual(st, X.pct_r4(asm["tax_rate_terminal"]))
            vr = wb["변수"]
            drv = _row_of(vr, "drv:법인세비용")
            for q in lay.est_quarters:
                pt = rows["세전이익"]["q"][q]["v"]
                c = vr["%s%d" % (lay.letter(q), drv)]
                self.assertIsInstance(c.value, float, (stock, q, c.value))
                tol = 5e-5 + (0.005 / pt if pt > 0 else 0.0)                               # r4 스케줄 반올림 + 법인세 0.01억 반올림
                self.assertLessEqual(abs(c.value - asm["tax_schedule"][q]), tol, (stock, q, c.value, asm["tax_schedule"][q]))
            probes = [p for p in v["sensitivity"] if p["name"].startswith("세율 ")]
            self.assertEqual(len(probes), 1, [p["name"] for p in v["sensitivity"]])
            self.assertIn("드라이버 셀 +1%p", probes[0]["name"])
            self.assertTrue(probes[0]["ok"], probes[0])
            self.assertTrue(all(c["moved"] for c in probes[0]["chain"]))
            txt = _readme_text(wb)
            self.assertIn("→ %s(terminal, " % st, txt)
            self.assertIn("이월결손 램프(assumptions.tax_schedule)", txt)
            self.assertNotIn("지주 자회사 합산", txt)                                        # 옛 추측 사유 — 백테스트 문구의 '지주 자회사 모델' 은 무관

    def test_new_orders_wording_follows_meta(self):
        for stock, wb in self.wbs.items():
            meta = (self.models[stock].get("scenarios") or {}).get("meta") or {}
            src_key, post_key = X.new_orders_wording(meta)
            vr = wb["변수"]
            self.assertEqual(vr.cell(_row_of(vr, "drv:매출조선신규"), 3).value, "매출조선신규 (%s, 억원 — 가정)" % X.NEW_ORDERS_LABEL_KO[src_key], stock)
            txt = _readme_text(wb)
            self.assertIn(" — %s). 보수/낙관/기존만 은 시나리오 시트에만" % X.NEW_ORDERS_POST_KO[post_key], txt)
            self.assertEqual("공시 계약 원장 체결 속도 폴백(저신뢰) base 시나리오" in txt, src_key == "ledger_signing_rate", stock)
        self.assertEqual(X.new_orders_wording(self.models["097230"]["scenarios"]["meta"])[0], "ledger_signing_rate")     # HJ = 원장 폴백(2026-10-09 모델)



if __name__ == "__main__":
    unittest.main()
