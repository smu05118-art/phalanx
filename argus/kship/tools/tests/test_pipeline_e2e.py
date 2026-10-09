#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""파이프라인 끝-대-끝 재현성 테스트(V10 reproducibility, 2026-10-05).

실행: cd argus/kship/tools && python3 -m unittest tests.test_pipeline_e2e -v
      KSHIP_E2E=0          재생성 테스트(≈30~60초)를 건너뛴다 — 검사기 단위 테스트만
      KSHIP_E2E_GIT=1      커밋 산출 재현 검사(summary.json 을 마지막으로 바꾼 커밋 트리를 꺼내 다시 만들어 바이트 비교, git 필요)

세 겹으로 본다.
  ① 검사기(selfcheck_models) 단위 — src 문법 파서·숫자 역파싱·스코프 규칙·보고 종류(consistency/freshness)·매출조선신규 재계산(원장 창·전환 조건·horizon 가드)·오류 주입(검사기가 공허하지 않음).
  ② 오프라인 재생성 — argus/kship(fin_cache 제외)+argus/kce/tools 를 임시 디렉터리에 두 벌 복사해 sls→model+xlsx→page→parts→hub 를
     죽은 프록시(네트워크 시도 즉시 실패) 아래 돌린다. 두 벌이 바이트까지 같아야 하고(sls 의 built_at 만 날짜), 재생성 트리는 교차 정합 실패 0
     (다른 레인의 알려진 결함 KNOWN_DEFECTS 패턴만 허용 — 패턴 밖은 실패) · 신선도 실패 0 이어야 한다. 그 뒤 사본에 오류를 주입해 검사기가
     각 그룹에서 잡는지 확인한다. 레포 산출물과의 차이는 **보고만** 한다(일일 갱신이 시세·환율·계약만 받고 모델은 주 1회라 어긋남이 정상 — MODEL.md §8).
  ③ 커밋 산출 재현(옵트인) — 커밋 트리의 입력으로 다시 만들면 커밋된 sls(built_at 제외)·모델 json·xlsx·models.html 과 바이트가 같아야 한다.
     회사 페이지는 조선사 5장이 '오늘' 날짜를 품어(kship_page/kship_orders_views — --today 없음) 날짜가 다르면 달라진다 → 조선사 외는 같아야 한다.

레포 파일은 읽기만 하고 쓰지 않는다. 네트워크·DART 접근 없음.
"""
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.dirname(HERE)
KSHIP = os.path.dirname(TOOLS)
ARGUS = os.path.dirname(KSHIP)
KCE_TOOLS = os.path.join(ARGUS, "kce", "tools")
sys.path.insert(0, TOOLS)

import selfcheck_models as SC                                   # noqa: E402

PY = sys.executable
E2E_ON = os.environ.get("KSHIP_E2E", "1") != "0"
OFFLINE_ENV = {"http_proxy": "http://127.0.0.1:9", "https_proxy": "http://127.0.0.1:9", "HTTP_PROXY": "http://127.0.0.1:9",
               "HTTPS_PROXY": "http://127.0.0.1:9", "no_proxy": "", "PYTHONDONTWRITEBYTECODE": "1"}
STEPS = [("sls", ["kship_sls.py", "--all"]),
         ("model", ["kship_model.py", "--build", "--all", "--xlsx", "--today", "{today}"]),
         ("page", ["kship_page.py", "--all"]),
         ("parts", ["kship_parts.py", "--all"]),
         ("hub", ["kship_model_section.py", "--hub"])]
YARDS = {"010140", "042660", "329180", "439260", "097230"}

# 다른 레인의 알려진 결함 — 검사기 실패 중 이 패턴에 맞는 것만 허용한다(패턴 밖 실패는 회귀). 고쳐지면 자연히 0 이 되고 이 표를 지우면 된다.
KNOWN_DEFECTS = {
    "fin_model_mislabel": [(r"\d{6} 지배주주순이익 \d{4}Q\d: fin face 에 \(지배주주지분\)당기순이익",
                            "kship_model.py:_actual_series — 지배주주순이익 face 누락 분기를 당기순이익 − 비지배로 파생하면서 src 는 fin face 계정으로 표기(경고만)")],
    "xlsx": [(r"\d{6} subQ 이자손익 \d{4}Q\d: VLOOKUP",
              "kship_fin.py:synth_is 이자손익 = 이자수익 − 이자비용 — 비용을 음수로 저장한 회사(002380·073010·272210)는 부호가 틀리고, "
              "kship_model_xlsx.py:_subq_quarter 가 그 계정을 VLOOKUP 해 모델(부호 반전 T4)과 어긋남")],
}


def _sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def _json_wo(path, drop=("built_at",)):
    with open(path, encoding="utf-8") as f:
        d = json.load(f)
    for k in drop:
        d.pop(k, None)
    return json.dumps(d, sort_keys=True, ensure_ascii=False)


def _stage(dst, src_kship=KSHIP, src_kce=KCE_TOOLS):
    """argus/kship(fin_cache 제외) + argus/kce/tools(assets 제외) 를 dst/argus/ 아래 복사한다."""
    shutil.copytree(src_kship, os.path.join(dst, "argus", "kship"), ignore=shutil.ignore_patterns("fin_cache", "__pycache__", "*.pyc", ".collect.lock"))
    shutil.copytree(src_kce, os.path.join(dst, "argus", "kce", "tools"), ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "assets", "tests"))
    return os.path.join(dst, "argus", "kship")


def _run_pipeline(kship_root, today):
    """사본 안에서 5단계를 순서대로 돈다. (단계명, rc, 마지막 출력, 삼킨 오류 줄) 목록.

    kship_model.py 는 회사별 xlsx 생성 실패를 `xlsx <stock>: <오류>` 한 줄로만 적고 종료 코드 0 을 돌려준다(build_all 의 try/except) —
    rc 만 보면 xlsx 가 하나도 안 만들어져도 통과하므로 그 줄을 따로 모아 단계 실패로 친다."""
    tools = os.path.join(kship_root, "tools")
    env = dict(os.environ)
    env.update(OFFLINE_ENV)
    out = []
    for name, args in STEPS:
        cmd = [PY] + [a.format(today=today) for a in args]
        p = subprocess.run(cmd, cwd=tools, env=env, capture_output=True, text=True, timeout=900)
        text = p.stdout + p.stderr
        swallowed = [ln.strip() for ln in text.splitlines()
                     if re.match(r"\s*xlsx \d{6}: ", ln) and not re.search(r"\d+ bytes$", ln.strip())]
        swallowed += [ln.strip() for ln in text.splitlines() if re.match(r"\s*\[warn\]", ln)]        # 페이지 훅의 섹션 렌더 실패 경고
        out.append((name, p.returncode, text[-1500:], swallowed))
        if p.returncode != 0:
            break
    return out


def _tree_files(root):
    """비교 대상 산출물: 상대경로 → 절대경로."""
    files = {}
    for sub in ("tools/assets/sls", "tools/assets/models", "models"):
        d = os.path.join(root, sub)
        if os.path.isdir(d):
            for f in os.listdir(d):
                if f.endswith((".json", ".xlsx")):
                    files["%s/%s" % (sub, f)] = os.path.join(d, f)
    for d in os.listdir(root):
        if re.fullmatch(r"\d{6}", d) and os.path.isfile(os.path.join(root, d, "index.html")):
            files["%s/index.html" % d] = os.path.join(root, d, "index.html")
    for f in ("models.html", "index.html", "parts.html", "coverage.html"):
        if os.path.isfile(os.path.join(root, f)):
            files[f] = os.path.join(root, f)
    return files


def _compare(a_root, b_root):
    """두 트리의 산출물 차이: {'same': n, 'differ': [...], 'only_a': [...], 'only_b': [...]} — sls json 은 built_at 을 뺀 뒤 비교."""
    fa, fb = _tree_files(a_root), _tree_files(b_root)
    res = {"same": 0, "differ": [], "only_a": sorted(set(fa) - set(fb)), "only_b": sorted(set(fb) - set(fa))}
    for rel in sorted(set(fa) & set(fb)):
        if _sha(fa[rel]) == _sha(fb[rel]):
            res["same"] += 1
            continue
        if rel.startswith("tools/assets/sls/") and rel.endswith(".json") and _json_wo(fa[rel]) == _json_wo(fb[rel]):
            res["same"] += 1
            continue
        res["differ"].append(rel)
    return res


def _unknown_failures(rep):
    """검사기 보고에서 KNOWN_DEFECTS 패턴 밖의 consistency 실패 — {그룹: [실패...]}. 알려진 결함은 {그룹: 건수} 로 따로."""
    unknown, known = {}, {}
    for c in rep.checks.values():
        if c.kind != "consistency":
            continue
        pats = [p for p, _ in KNOWN_DEFECTS.get(c.key, [])]
        for f in c.fails:
            if any(re.match(p, f) for p in pats):
                known[c.key] = known.get(c.key, 0) + 1
            else:
                unknown.setdefault(c.key, []).append(f)
    return unknown, known


# ── ① 검사기 단위 ────────────────────────────────────────────

class TestParsers(unittest.TestCase):
    def test_parse_num(self):
        self.assertEqual(SC.parse_num("127,235"), 127235.0)
        self.assertEqual(SC.parse_num("-76"), -76.0)
        self.assertEqual(SC.parse_num("1.7%"), 1.7)
        self.assertEqual(SC.parse_num("14.3배"), 14.3)
        self.assertEqual(SC.parse_num("1,031원"), 1031.0)
        self.assertIsNone(SC.parse_num("—"))
        self.assertIsNone(SC.parse_num(""))
        self.assertIsNone(SC.parse_num(None))

    def test_tol_a_follows_fmt_a_rounding(self):
        self.assertAlmostEqual(SC.tol_a(9.96), 0.05, places=6)        # 소수 1자리 구간
        self.assertAlmostEqual(SC.tol_a(10.0), 0.5, places=6)         # 정수 구간
        self.assertAlmostEqual(SC.tol_a(-123456.7), 0.5, places=6)

    def test_qlabel(self):
        self.assertEqual(SC.qlabel_to_q("2Q26"), "2026Q2")
        self.assertEqual(SC.qlabel_to_q("4Q21"), "2021Q4")
        self.assertEqual(SC.qlabel_to_q(2025), "2025")
        self.assertIsNone(SC.qlabel_to_q("YQ"))
        self.assertIsNone(SC.qlabel_to_q(None))

    def test_q_add(self):
        self.assertEqual(SC.q_add("2026Q2", -1), "2026Q1")
        self.assertEqual(SC.q_add("2026Q1", -1), "2025Q4")
        self.assertEqual(SC.q_add("2025Q4", 1), "2026Q1")

    def test_fin_scope_rule(self):
        fin = {"cons": {"is": {"2026Q1": {"매출액(수익)": 10.0}, "2026Q2": {}}}, "sep": {"is": {"2026Q2": {"매출액(수익)": 7.0}, "2026Q3": {}}}}
        self.assertEqual(SC.fin_scope_of(fin, "2026Q1"), "cons")
        self.assertEqual(SC.fin_scope_of(fin, "2026Q2"), "sep")          # 연결 매출 없음 → 별도 보충
        self.assertEqual(SC.fin_scope_of(fin, "2026Q3"), "cons")         # 둘 다 없음 → cons 유지
        self.assertEqual(SC.fin_scope_of({"cons": {"is": {}}, "sep": {"is": {"2026Q1": {"매출액(수익)": 1}}}}, "2026Q1"), "sep")


class _FakeTree:
    """SrcEval 이 필요한 최소 트리(prices·sls·종속사 모델)."""

    def __init__(self, prices=None, sls=None, models=None):
        self.prices = prices or {}
        self._sls, self._models = sls or {}, models or {}

    def sls(self, stock):
        return self._sls.get(stock)

    def model(self, stock):
        return self._models.get(stock)

    def fin(self, stock):
        return None


def _mini_model(rows):
    return {"stock": "000001", "rows": [{"key": k, "q": {q: c for q, c in v.items()}} for k, v in rows.items()]}


class TestSrcEval(unittest.TestCase):
    Q = "2026Q2"

    def setUp(self):
        self.fin = {"stock": "000001", "quarters": [self.Q],
                    "cons": {"is": {self.Q: {"매출액(수익)": 123456.789, "영업이익": 12345.678, "이자수익": 100.0, "이자비용": -40.0,
                                            "외환차익": 10.0, "외환차손": 3.0, "외화환산이익": 2.0, "외화환산손실": 1.0, "금융손익": 500.0,
                                            "당기순이익": 1000.0, "(비지배주주지분)당기순이익": 100.0}},
                             "bs": {self.Q: {"지배주주지분": 500000.0}}, "cf": {self.Q: {"감가상각비": 200.0}}},
                    "sep": {"is": {}}, "shares": {self.Q: {"common_outstanding": 2000000}}}
        rows = {"매출액": {self.Q: {"v": 1234.57, "kind": "actual", "src": "fin.cons.is.매출액(수익)"}},
                "영업이익": {self.Q: {"v": 123.46, "kind": "actual", "src": "fin.cons.is.영업이익"}},
                "감가상각비": {self.Q: {"v": 2.0, "kind": "actual", "src": "fin.cons.cf.감가상각비"}},
                "지배주주지분": {self.Q: {"v": 5000.0, "kind": "actual", "src": "fin.cons.bs.지배주주지분"}},
                "지배주주순이익": {self.Q: {"v": 9.0, "kind": "actual", "src": "fin.cons.is.(지배주주지분)당기순이익"}},
                "주식수": {self.Q: {"v": 2.0, "kind": "actual", "src": "fin.shares.common_outstanding"}},
                "EPS": {self.Q: {"v": 450.0, "kind": "actual", "src": "지배주주순이익 ÷ 유통주식수(파생)"}},
                "이자손익": {self.Q: {"v": 0.6, "kind": "actual", "src": "fin.cons.is.이자수익 − fin.cons.is.이자비용(음수 저장 → 부호 반전)"}},
                "외환손익": {self.Q: {"v": 0.08, "kind": "actual", "src": "fin.cons.is.외환차익 − fin.cons.is.외환차손 + fin.cons.is.외화환산이익 − fin.cons.is.외화환산손실"}},
                "파생상품손익": {self.Q: {"v": 0.0, "kind": "actual", "src": "x"}},
                "기타금융손익": {self.Q: {"v": 4.32, "kind": "actual", "src": "fin.cons.is.금융손익 − 이자손익 − 외환손익 − 파생상품손익(잔차)"}},
                "조정EPS": {self.Q: {"v": 450.0, "kind": "actual", "src": "= EPS(일회성 의심 없음)"}}}
        self.model = _mini_model(rows)
        self.ev = SC.SrcEval(_FakeTree(), "000001", self.model, self.fin)

    def test_single_fin_token_and_unrounded_base(self):
        v, cls = self.ev.eval("fin.cons.is.매출액(수익)", self.Q)
        self.assertAlmostEqual(v, 1234.56789)
        self.assertEqual(cls, "r2")
        # OPM 은 반올림 전 fin 값으로: 12345.678/123456.789
        v, cls = self.ev.eval("영업이익 ÷ 매출액(파생)", self.Q, row_key="OPM")
        self.assertAlmostEqual(v, 12345.678 / 123456.789)
        self.assertEqual(cls, "r4")

    def test_compound_with_sign_flip_and_residual(self):
        v, _ = self.ev.eval("fin.cons.is.이자수익 − fin.cons.is.이자비용(음수 저장 → 부호 반전)", self.Q)
        self.assertAlmostEqual(v, (100.0 - 40.0) / 100)                 # 이자비용 −40 → 부호 반전 → 40 크기
        v, _ = self.ev.eval("fin.cons.is.외환차익 − fin.cons.is.외환차손 + fin.cons.is.외화환산이익 − fin.cons.is.외화환산손실", self.Q)
        self.assertAlmostEqual(v, (10 - 3 + 2 - 1) / 100)
        v, _ = self.ev.eval("fin.cons.is.금융손익 − 이자손익 − 외환손익 − 파생상품손익(잔차)", self.Q)
        self.assertAlmostEqual(v, 5.0 - 0.6 - 0.08 - 0.0)
        v, _ = self.ev.eval("fin.cons.is.매출액(수익) (연결 손익 없는 분기 → 별도 보충)".replace("cons", "cons"), self.Q)
        self.assertAlmostEqual(v, 1234.56789)                              # 꼬리 주석 제거

    def test_copy_row_and_per_share(self):
        v, cls = self.ev.eval("= EPS(일회성 의심 없음)", self.Q)
        self.assertEqual((v, cls), (450.0, "exact"))
        v, cls = self.ev.eval("지배주주지분 ÷ 유통주식수(파생)", self.Q)
        self.assertAlmostEqual(v, 500000.0 / 100 * 1e8 / 2000000)
        self.assertEqual(cls, "r1")

    def test_mislabel_and_unverifiable(self):
        # face 에 (지배주주지분)당기순이익 이 없고 당기순이익·비지배가 있으면 Mislabel(파생값 동봉)
        with self.assertRaises(SC.Mislabel) as cm:
            self.ev.eval("fin.cons.is.(지배주주지분)당기순이익", self.Q, row_key="지배주주순이익")
        self.assertAlmostEqual(cm.exception.expected, (1000.0 - 100.0) / 100)
        with self.assertRaises(SC.Unverifiable):
            self.ev.eval("fin.cons.is.없는계정", self.Q)
        with self.assertRaises(SC.Unverifiable):
            self.ev.eval("듣도 보도 못한 문구", self.Q)
        with self.assertRaises(SC.Unverifiable):
            self.ev.eval("", self.Q)

    def test_toss_history_and_sls_reconcile(self):
        prices = {"rows": {"000001": {"history_quarterly": {self.Q: {"avg": 9000.0}}}}}
        sls = {"000001": {"reconcile": {self.Q: {"reported_segment_rev_m": 55555.5}}}}
        rows = {"EPS": {SC.q_add(self.Q, -i): {"v": 100.0, "kind": "actual"} for i in range(4)},
                "BPS": {self.Q: {"v": 4500.0, "kind": "actual"}}}
        ev = SC.SrcEval(_FakeTree(prices=prices, sls=sls), "000001", _mini_model(rows), self.fin)
        self.assertAlmostEqual(ev.eval("토스 일봉 분기 평균 ÷ TTM EPS", self.Q)[0], 9000.0 / 400.0)
        self.assertAlmostEqual(ev.eval("토스 일봉 분기 평균 ÷ BPS", self.Q)[0], 2.0)
        self.assertAlmostEqual(ev.eval("sls.reconcile.reported_segment_rev_m ← yards_cache 부문표(3개월분)", self.Q)[0], 555.555)


def _panel_c(N_by_scn, K, quarters, pfld="covered_scope_new_revenue", excluded=("EP및특수선",), drop=()):
    """합성 forecast_panel 회사 — N(KRW_million, 시나리오별 상수 또는 분기 목록) · 누적 인식 K(h)(매출 = N × K, 시나리오 공통 커널) · horizon = 1.. · coverage.excluded_segments."""
    sc = {}
    for scn, n in N_by_scn.items():
        rows = []
        for i, q in enumerate(quarters):
            if q in drop:
                continue
            nv = n[i] if isinstance(n, (list, tuple)) else n
            rows.append({"quarter": q, "horizon": i + 1, "new_orders": nv, pfld: nv * K[i]})
        sc[scn] = {"quarterly": rows}
    return {"stock": "000001", "origin": "2026Q2", "scenarios": sc, "coverage": {"excluded_segments": list(excluded)}}


def _sls_c(origin="2026Q2", ledger=(), post=(), T_lens=()):
    """합성 sls — ledger: (체결일, 억원, type) origin 까지 체결 · post: (체결일, 억원, type) origin 이후 체결. 일정 길이는 T_lens 순서대로 돌려 붙인다."""
    cs = []
    for i, (d, eok, t) in enumerate(list(ledger) + list(post)):
        n = T_lens[i % len(T_lens)] if T_lens else 0
        cs.append({"rcp": "r%d" % i, "type": t, "counted": True, "signed": d, "amt_krw_m": eok * 100.0, "schedule": {"q%d" % j: 1.0 for j in range(n)}})
    cs.append({"rcp": "other", "type": "OTHER", "counted": True, "signed": "2026-08-20", "amt_krw_m": 99999.0, "schedule": {"a": 1.0}})   # OTHER 는 어디에도 안 든다
    return {"origin": origin, "contracts": cs}


class TestPanelNewOrdersExpected(unittest.TestCase):
    """selfcheck panel_new_orders_expected — 모델 자기신고(ledger_rate.levels·source_field·kernel_mode)를 읽지 않고 패널·sls 에서 다시 센다(2026-10-09)."""
    Q = ["2026Q3", "2026Q4", "2027Q1", "2027Q2"]
    K = [0.0, 0.2, 0.5, 0.9]                                 # 누적 인식 → 역산 커널 k = {0:0, 1:0.2, 2:0.3, 3:0.4}
    POST = [("2026-08-01", 400.0, "CONT"), ("2026-09-01", 300.0, "NAVAL")]     # 상계 D — NAVAL 은 폴백 회사(EP및특수선 제외)면 범위 밖

    def _panel_mode(self, excluded=("EP및특수선",)):
        panel = _panel_c({"conservative": 80000.0, "base": 100000.0, "optimistic": 120000.0}, self.K, self.Q, excluded=excluded)
        sls = _sls_c(ledger=[("2026-05-01", 1500.0, "CONT")], post=self.POST, T_lens=(10,))      # 창 [2026Q2 1,500] → 비율 1,000/1,500 ≥ 0.5 · 체결 > 0 1분기 → 전환 없음
        no = {"panel_field": "covered_scope_new_revenue", "source_field": "covered_scope_new_revenue", "kernel_mode": "backsolved"}
        return panel, sls, no

    def test_panel_mode_netting_excludes_out_of_scope_type(self):
        panel, sls, no = self._panel_mode()
        exp, tol, note, fails = SC.panel_new_orders_expected(panel, sls, no, "net_panel", self.Q)
        self.assertEqual(fails, [])
        self.assertEqual([round(exp[q], 6) for q in self.Q], [0.0, 120.0, 380.0, 740.0])          # net(2026Q3) = 1,000 − 400(NAVAL 300 제외)
        self.assertIn("NAVAL", note)
        panel2, _, _ = self._panel_mode(excluded=())                                                # 제외 부문 없으면 NAVAL 도 상계 → net 300
        exp2 = SC.panel_new_orders_expected(panel2, sls, no, "net_panel", self.Q)[0]
        self.assertEqual([round(exp2[q], 6) for q in self.Q], [0.0, 60.0, 290.0, 620.0])
        exp3 = SC.panel_new_orders_expected(panel, sls, no, "exclude_sls", self.Q)[0]              # exclude_sls 는 D = 0
        self.assertEqual([round(exp3[q], 6) for q in self.Q], [0.0, 200.0, 500.0, 900.0])
        self.assertLess(tol, 0.02)

    def test_panel_mode_flags_self_reported_switch_and_field(self):
        panel, sls, no = self._panel_mode()
        no2 = dict(no, source_field="ledger_signing_rate", panel_field="new_order_revenue", kernel_mode="aggregate_ratio",
                   ledger_rate={"levels": {"conservative": 1500.0, "base": 1500.0, "optimistic": 1500.0}})
        exp, _, _, fails = SC.panel_new_orders_expected(panel, sls, no2, "net_panel", self.Q)
        self.assertIsNotNone(exp)                                                                   # None 반환 금지 — 재계산값으로 간다
        self.assertTrue(any(f.startswith("source_field 'ledger_signing_rate' ≠ 전환 조건 재계산") for f in fails), fails)
        self.assertTrue(any(f.startswith("panel_field 'new_order_revenue' ≠ 패널 재판정 'covered_scope_new_revenue'") for f in fails), fails)
        self.assertTrue(any(f.startswith("kernel_mode 'aggregate_ratio' ≠ 재계산 'backsolved'") for f in fails), fails)

    def _ledger_mode(self, opt_N=500000.0, cons_N=(7900.0,) * 4):
        # 패널 base N 110억(상수) · 원장 창 2025Q4~2026Q2 [3,000 · 2,000 · 4,000](NAVAL 500 제외) → 비율 110/3,000 = 0.0367 → 전환. 낙관 패널 N 5,000억 > max 4,000 → 하한
        panel = _panel_c({"conservative": list(cons_N), "base": 11000.0, "optimistic": opt_N}, self.K, self.Q)
        sls = _sls_c(ledger=[("2025-11-15", 3000.0, "CONT"), ("2025-12-20", 500.0, "NAVAL"), ("2026-02-10", 2000.0, "LNGC"), ("2026-05-05", 4000.0, "CONT")],
                     post=self.POST, T_lens=(10, 8, 12, 14, 6, 9))                                 # 일정 길이 10·8·12·14·6·9 → 중위 9.5 → T 10
        no = {"panel_field": "covered_scope_new_revenue", "source_field": "ledger_signing_rate", "kernel_mode": "backsolved",
              "ledger_rate": {"levels": {"conservative": 2000.0, "base": 3000.0, "optimistic": 5000.0}, "window_quarters": ["2025Q4", "2026Q1", "2026Q2"],
                              "n_quarters_used": 3, "panel_to_ledger_ratio": 0.0367, "curve": "panel_kernel", "alt_linear_T": 10}}
        return panel, sls, no

    def test_ledger_mode_recomputes_levels_from_contracts(self):
        panel, sls, no = self._ledger_mode()
        exp, _, note, fails = SC.panel_new_orders_expected(panel, sls, no, "net_panel", self.Q)
        self.assertEqual(fails, [])
        self.assertEqual([round(exp[q], 6) for q in self.Q], [0.0, 520.0, 1380.0, 2540.0])       # N 3,000 · net(2026Q3) 2,600 · 커널 0/0.2/0.3/0.4
        self.assertIn("N 3,000억", note)
        # 자기신고 변조 — 수준·창·비율·T 모두 FAIL 로 드러나고 N 은 재계산값(exp 불변)
        bad = json.loads(json.dumps(no))
        bad["ledger_rate"]["levels"]["base"] = 3100.0
        bad["ledger_rate"]["window_quarters"] = ["2026Q1", "2026Q2"]
        bad["ledger_rate"]["panel_to_ledger_ratio"] = 0.05
        bad["ledger_rate"]["alt_linear_T"] = 12
        exp2, _, _, fails2 = SC.panel_new_orders_expected(panel, sls, bad, "net_panel", self.Q)
        self.assertEqual(exp2, exp)
        heads = [f.split(" ")[0] for f in fails2]
        self.assertEqual(heads, ["ledger_rate.levels.base", "ledger_rate", "ledger_rate.panel_to_ledger_ratio", "ledger_rate"], fails2)
        self.assertIn("3100.0 ≠ 원장 재계산 3000.00(median", fails2[0])
        # 패널 모드라고 자기신고하면 전환 조건 재계산이 잡는다
        pan = dict(no, source_field="covered_scope_new_revenue", ledger_rate=None)
        exp3, _, _, fails3 = SC.panel_new_orders_expected(panel, sls, pan, "net_panel", self.Q)
        self.assertEqual([round(exp3[q], 6) for q in self.Q], [0.0, 0.0, 22.0, 55.0])              # 패널 N 110 − 400 → 0 · 110 × 커널
        self.assertEqual(len(fails3), 1)
        self.assertIn("source_field 'covered_scope_new_revenue' ≠ 전환 조건 재계산(패널/원장 비율 0.0367 · 체결 > 0 3분기 · 창 2025Q4~2026Q2(3분기) · curve panel_kernel → ledger_signing_rate)", fails3[0])

    def test_ledger_mode_linear_curve_when_a_scenario_kernel_is_not_backsolvable(self):
        panel, sls, no = self._ledger_mode(cons_N=(7900.0, 7900.0, 8000.0, 7900.0))               # 보수 N 비상수 → 세 커널 중 하나 결손 → 선형 T=10
        no["kernel_mode"], no["ledger_rate"]["curve"] = "linear_ledger_median", "linear_ledger_median"
        exp, _, _, fails = SC.panel_new_orders_expected(panel, sls, no, "net_panel", self.Q)
        self.assertEqual(fails, [])
        self.assertEqual([round(exp[q], 6) for q in self.Q], [0.0, 260.0, 560.0, 860.0])         # (2,600 · 3,000 · 3,000)/10 누적
        no["ledger_rate"]["alt_linear_T"] = 12
        fails2 = SC.panel_new_orders_expected(panel, sls, no, "net_panel", self.Q)[3]
        self.assertEqual(len(fails2), 1)
        self.assertIn("curve/alt_linear_T linear_ledger_median/12 ≠ 재계산 linear_ledger_median/10(계약 일정 길이 6건 중위)", fails2[0])

    def test_horizon_gap_skips_without_keyerror(self):
        panel, sls, no = self._panel_mode()
        gap = _panel_c({"conservative": 80000.0, "base": 100000.0, "optimistic": 120000.0}, self.K, self.Q, drop=("2027Q1",))
        exp, tol, note, fails = SC.panel_new_orders_expected(gap, sls, no, "net_panel", self.Q)
        self.assertEqual((exp, tol), (None, None))
        self.assertIn("패널 horizon 결손 2027Q1", note)
        self.assertIn("재계산 생략", note)
        self.assertEqual(fails, ["kernel_mode 'backsolved' ≠ 재계산 'aggregate_ratio'"])        # 모델이 집계 비율 근사라고 적었어야 한다
        for r in gap["scenarios"]["base"]["quarterly"]:                                             # horizon 오프셋(2..)도 역산 불가
            r["horizon"] += 1
        exp2, _, note2, _ = SC.panel_new_orders_expected(gap, sls, dict(no, kernel_mode="aggregate_ratio"), "net_panel", self.Q)
        self.assertIsNone(exp2)
        self.assertIn("재계산 생략", note2)


class TestReport(unittest.TestCase):
    def test_kinds_and_exit_semantics(self):
        rep = SC.Report()
        a = rep.c("a", "A", "consistency")
        b = rep.c("b", "B", "freshness")
        a.ok(3)
        b.fail("stale")
        self.assertEqual(rep.failed(("consistency",)), 0)
        self.assertEqual(rep.failed(("consistency", "freshness")), 1)
        d = rep.as_dict()
        self.assertEqual((d["failed_consistency"], d["failed_freshness"]), (0, 1))
        self.assertIn("stale", rep.table())
        self.assertIn("consistency 실패 0 · freshness 실패 1", rep.table())


# ── ② 오프라인 재생성 → 결정론 → 교차 정합 → 오류 주입 ─────────────────────

@unittest.skipUnless(E2E_ON, "KSHIP_E2E=0")
class TestRegenerationE2E(unittest.TestCase):
    """두 벌 재생성(같은 입력·같은 --today) → 바이트 동일 · 재생성 트리 교차 정합 · 오류 주입."""

    @classmethod
    def setUpClass(cls):
        if not os.path.isdir(KCE_TOOLS):
            raise unittest.SkipTest("argus/kce/tools 없음 — kship_lib 가 import 하는 공용 층")
        if not os.path.isdir(os.path.join(TOOLS, "assets", "fin")):
            raise unittest.SkipTest("assets/fin 없음 — 재생성 입력 부족")
        s = SC._load(os.path.join(TOOLS, "assets", "models", "summary.json")) or {}
        cls.today = str(s.get("built_at") or "")[:10] or "2026-10-02"
        cls.tmp = tempfile.mkdtemp(prefix="kship_e2e_")
        cls.roots = []
        cls.runs = []
        for i in (1, 2):
            root = _stage(os.path.join(cls.tmp, "run%d" % i))
            cls.roots.append(root)
            cls.runs.append(_run_pipeline(root, cls.today))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_01_steps_succeed_offline(self):
        for i, run in enumerate(self.runs, 1):
            names = [n for n, _, _, _ in run]
            self.assertEqual(names, [n for n, _ in STEPS], "run%d 단계 누락: %s" % (i, run[-1]))
            for name, rc, tail, swallowed in run:
                self.assertEqual(rc, 0, "run%d %s rc=%d\n%s" % (i, name, rc, tail))
                self.assertNotIn("Traceback", tail, "run%d %s 트레이스백\n%s" % (i, name, tail))
                self.assertNotIn("Connection refused", tail, "run%d %s 네트워크 시도(오프라인 위반)\n%s" % (i, name, tail))
                self.assertEqual(swallowed, [], "run%d %s 가 종료 코드 0 으로 삼킨 오류 %d건(처음 3): %s" % (i, name, len(swallowed), swallowed[:3]))
        # xlsx 가 실제로 이번 재생성에서 다시 쓰였는지 — README 스탬프(생성 날짜) = --today. rc 0 이어도 xlsx 가 안 만들어지는 경로를 막는다.
        from openpyxl import load_workbook
        root = self.roots[0]
        xl = sorted(f for f in os.listdir(os.path.join(root, "models")) if f.endswith("_model.xlsx"))
        self.assertGreaterEqual(len(xl), 50, xl)
        for f in xl[:3] + xl[-3:]:
            ws = load_workbook(os.path.join(root, "models", f), read_only=True)["README"]
            a2 = str(ws.cell(2, 1).value or "")
            self.assertIn("생성 %s" % self.today, a2, "%s README A2 %r — 이번 재생성(--today %s) 산출이 아님" % (f, a2, self.today))

    def test_01b_model_step_reports_each_xlsx(self):
        """model 단계 출력의 `xlsx <stock>: … bytes` 줄 수 = 모델 수(xlsx 생성이 통째로 빠지면 여기서 드러난다)."""
        name, rc, tail, swallowed = self.runs[0][1]
        self.assertEqual(name, "model")
        n_models = len([f for f in os.listdir(os.path.join(self.roots[0], "tools", "assets", "models")) if re.fullmatch(r"\d{6}\.json", f)])
        n_xlsx = len([f for f in os.listdir(os.path.join(self.roots[0], "models")) if f.endswith("_model.xlsx")])
        self.assertEqual(n_xlsx, n_models, "xlsx %d ≠ 모델 json %d" % (n_xlsx, n_models))

    def test_02_two_runs_are_byte_identical(self):
        res = _compare(self.roots[0], self.roots[1])
        self.assertEqual(res["only_a"] + res["only_b"], [], "파일 집합이 다름")
        self.assertEqual(res["differ"], [], "두 번 재생성이 다른 파일: %s" % res["differ"][:10])
        self.assertGreaterEqual(res["same"], 58 + 58 + 56 + 6)           # 모델 json·xlsx·회사 페이지·sls 최소치

    def test_03_regenerated_tree_has_same_file_set_as_repo(self):
        res = _compare(KSHIP, self.roots[0])
        self.assertEqual(res["only_a"], [], "레포에만 있는 산출물 %s" % res["only_a"][:10])
        self.assertEqual(res["only_b"], [], "재생성에만 있는 산출물 %s" % res["only_b"][:10])
        # 값 차이는 보고만 — 레포 산출이 입력(시세·환율·계약)보다 오래됐으면 정상적으로 다르다(MODEL.md §8)
        by = {}
        for rel in res["differ"]:
            by[rel.split("/")[0] if "/" in rel else rel] = by.get(rel.split("/")[0] if "/" in rel else rel, 0) + 1
        sys.stderr.write("\n[e2e] 레포 vs 재생성(--today %s) 다른 파일 %d / 같은 파일 %d — %s\n" % (self.today, len(res["differ"]), res["same"], by))

    def test_04_regenerated_tree_is_cross_consistent(self):
        rep = SC.run(self.roots[0])
        sys.stderr.write("\n[e2e] 재생성 트리 교차 정합\n" + rep.table(max_detail=3) + "\n")
        unknown, known = _unknown_failures(rep)
        for key, n in known.items():
            sys.stderr.write("[e2e] 알려진 결함(허용) %s ×%d — %s\n" % (key, n, KNOWN_DEFECTS[key][0][1]))
        self.assertEqual(unknown, {}, "알려진 결함 패턴 밖의 consistency 실패:\n" +
                         "\n".join("%s: %s" % (k, v[:3]) for k, v in unknown.items()))
        # 검사가 실제로 돌았는지 — 전수 규모의 하한(공허 통과 방지)
        chk = {c.key: c.n for c in rep.checks.values()}
        self.assertGreater(chk.get("fin_model", 0), 20000)
        self.assertGreater(chk.get("fin_coverage", 0), 15000)
        self.assertGreater(chk.get("xlsx", 0), 100000)
        self.assertGreater(chk.get("page", 0), 20000)
        self.assertGreater(chk.get("hub", 0), 500)
        self.assertGreater(chk.get("sls_model", 0), 300)
        self.assertGreater(chk.get("summary", 0), 2000)

    def test_05_regenerated_tree_is_fresh(self):
        rep = SC.run(self.roots[0], xlsx=False)
        fresh = {c.key: c.fails for c in rep.checks.values() if c.kind == "freshness" and c.fails}
        self.assertEqual(fresh, {}, "같은 입력으로 다시 만든 트리에 신선도 실패:\n%s" % "\n".join("%s: %s" % (k, v[:3]) for k, v in fresh.items()))

    def test_06_checker_catches_injected_faults(self):
        """사본 2 에 결함을 심고 검사기가 그룹별로 잡는지 — 검사기가 자기신고를 읽지 않음을 확인(게이트 공허 통과 방지)."""
        root = self.roots[1]
        st = "010140"
        mp = os.path.join(root, "tools", "assets", "models", "%s.json" % st)
        m = SC._load(mp)
        la = m["periods"]["last_actual"]
        for r in m["rows"]:
            if r["key"] == "매출액":
                r["q"][la]["v"] += 1.0                                      # (a) 모델 actual 셀 ≠ fin
        with open(mp, "w", encoding="utf-8") as f:
            json.dump(m, f, ensure_ascii=False, indent=1)
        sp = os.path.join(root, "tools", "assets", "models", "summary.json")
        s = SC._load(sp)
        for r in s["rows"]:
            if r["stock"] == st:
                r["fy"]["2026E"]["rev"] += 1.0                               # (b) summary ≠ 모델 → 허브도 어긋남
        with open(sp, "w", encoding="utf-8") as f:
            json.dump(s, f, ensure_ascii=False, indent=1)
        lp = os.path.join(root, "tools", "assets", "sls", "%s.json" % st)
        sls = SC._load(lp)
        fq = [q for q in m["periods"]["quarters"] if q > la][0]
        sls["by_quarter"][fq]["marine_hedged_krw_m"] += 1000.0   # (d) sls ≠ 모델 매출조선(net_panel 모드는 모델이 marine 전체를 읽음) · 내부 항등식(marine ≠ signed_by_origin + post_origin)
        with open(lp, "w", encoding="utf-8") as f:
            json.dump(sls, f, ensure_ascii=False, indent=1)
        pp = os.path.join(root, "tools", "assets", "prices.json")
        pr = SC._load(pp)
        pr["rows"][st]["close"] += 10                                        # (e) prices ≠ valuation.price
        with open(pp, "w", encoding="utf-8") as f:
            json.dump(pr, f, ensure_ascii=False)
        pg = os.path.join(root, st, "index.html")
        with open(pg, encoding="utf-8") as f:
            h = f.read()
        h2 = re.sub(r"(<b>)([^<]*)(<small>억</small></b>\s*<span>FY2027E 매출)", r"\g<1>1\g<3>", h, count=1)      # (c) 페이지 KPI
        self.assertNotEqual(h, h2, "페이지 KPI 타일을 못 찾음 — 주입 실패")
        with open(pg, "w", encoding="utf-8") as f:
            f.write(h2)
        # (f) 조정EPS 병기(라운드 7) — 사본 2 에서 KPI '(조정 …원)' 괄호가 있는 회사를 고른다(KCC 002380 우선). 페이지 괄호 삭제 → page, 허브 data-adj 변조 → hub
        adj = {}
        for d in sorted(os.listdir(root)):
            if re.fullmatch(r"\d{6}", d) and os.path.isfile(os.path.join(root, d, "index.html")):
                ys = [y for _, y, _, _, _, _, a in SC.KPI_TILE_RE.findall(SC._section_html(SC._read(os.path.join(root, d, "index.html")), d) or "") if a]
                if ys:
                    adj[d] = ys[0]
        self.assertTrue(adj, "사본 2 에 조정EPS 병기 회사가 없음 — 주입 불가")
        st2 = "002380" if "002380" in adj else sorted(adj)[0]
        pg2 = os.path.join(root, st2, "index.html")
        h = SC._read(pg2)
        h2 = re.sub(r"(· EPS [^<(]+?원) \(조정 [^<)]+?원\)(</i>)", r"\1\2", h, count=1)
        self.assertNotEqual(h, h2, "%s 페이지 KPI 조정EPS 괄호를 못 찾음 — 주입 실패" % st2)
        with open(pg2, "w", encoding="utf-8") as f:
            f.write(h2)
        hp = os.path.join(root, "models.html")
        hh = SC._read(hp)

        def _tamper_adj(mt):
            tr = mt.group(0)
            if st2 in tr and 'data-adj="' in tr:
                return re.sub(r'data-adj="(-?[\d.]+)"', lambda a: 'data-adj="%s"' % (float(a.group(1)) + 1.0), tr, count=1)
            return tr
        hh2 = re.sub(r"<tr>.*?</tr>", _tamper_adj, hh, flags=re.S)
        self.assertNotEqual(hh, hh2, "%s 허브 행 data-adj 를 못 찾음 — 주입 실패" % st2)
        with open(hp, "w", encoding="utf-8") as f:
            f.write(hh2)
        # (g) HJ重 모델의 원장 수준 자기신고(new_orders.ledger_rate.levels.base) 변조 → panel_model(검사기가 sls.contracts 로 수준을 다시 센다 — 2026-10-09)
        hj = "097230"
        mp2 = os.path.join(root, "tools", "assets", "models", "%s.json" % hj)
        m2 = SC._load(mp2)
        self.assertEqual((m2.get("new_orders") or {}).get("source_field"), "ledger_signing_rate", "HJ 가 ledger_signing_rate 모드가 아님 — 주입 전제 깨짐")
        m2["new_orders"]["ledger_rate"]["levels"]["base"] += 100.0
        with open(mp2, "w", encoding="utf-8") as f:
            json.dump(m2, f, ensure_ascii=False, indent=1)
        rep = SC.run(root, xlsx=True, emulate=False)
        fails = {c.key: c.fails for c in rep.checks.values()}

        def has(key, *needles):
            return any(all(n in f for n in needles) for f in fails.get(key, []))
        self.assertTrue(has("fin_model", st, "매출액", la), fails.get("fin_model", [])[:3])
        self.assertTrue(has("fin_coverage", st, "매출액", la), fails.get("fin_coverage", [])[:3])
        self.assertTrue(has("xlsx", st, "subQ 매출액", la), fails.get("xlsx", [])[:3])
        self.assertTrue(has("page", st, "손익표 매출액", la), fails.get("page", [])[:3])
        self.assertTrue(has("page", st, "KPI FY2027E 매출"), fails.get("page", [])[:3])
        self.assertTrue(has("summary", st, "fy.2026E.rev"), fails.get("summary", [])[:3])
        self.assertTrue(has("hub", st, "FY2026E 매출(억)"), fails.get("hub", [])[:3])
        self.assertTrue(has("sls_model", st, fq, "매출조선"), fails.get("sls_model", [])[:3])
        self.assertTrue(has("sls_internal", st, fq), fails.get("sls_internal", [])[:3])
        self.assertTrue(has("prices_model", st), fails.get("prices_model", [])[:3])
        self.assertTrue(has("page", st2, "KPI FY%sE EPS조정" % adj[st2]), fails.get("page", [])[:3])
        self.assertTrue(has("hub", st2, "EPS data-adj"), fails.get("hub", [])[:3])
        self.assertTrue(has("panel_model", hj, "ledger_rate.levels.base"), fails.get("panel_model", [])[:3])
        self.assertFalse(has("panel_model", hj, "매출조선신규"), fails.get("panel_model", [])[:3])         # N 은 재계산값 — 행 자체는 여전히 맞다


# ── ③ 커밋 산출 재현(옵트인) ────────────────────────────────

@unittest.skipUnless(E2E_ON and os.environ.get("KSHIP_E2E_GIT") == "1", "KSHIP_E2E_GIT=1 일 때만(git 필요)")
class TestCommittedReproducibility(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        repo = subprocess.run(["git", "-C", KSHIP, "rev-parse", "--show-toplevel"], capture_output=True, text=True)
        if repo.returncode != 0:
            raise unittest.SkipTest("git 저장소 아님")
        cls.repo = repo.stdout.strip()
        rel = os.path.relpath(os.path.join(TOOLS, "assets", "models", "summary.json"), cls.repo)
        commit = subprocess.run(["git", "-C", cls.repo, "log", "-1", "--format=%H", "--", rel], capture_output=True, text=True).stdout.strip()
        if not commit:
            raise unittest.SkipTest("summary.json 커밋 없음")
        cls.commit = commit
        cls.tmp = tempfile.mkdtemp(prefix="kship_e2e_git_")
        rel_kship, rel_kce = os.path.relpath(KSHIP, cls.repo), os.path.relpath(KCE_TOOLS, cls.repo)
        for name in ("work", "ref"):
            d = os.path.join(cls.tmp, name)
            os.makedirs(d)
            ar = subprocess.run(["git", "-C", cls.repo, "archive", commit, rel_kship, rel_kce], capture_output=True)
            subprocess.run(["tar", "-x", "-C", d], input=ar.stdout, check=True)
        cls.work = os.path.join(cls.tmp, "work", rel_kship)
        cls.ref = os.path.join(cls.tmp, "ref", rel_kship)
        s = SC._load(os.path.join(cls.work, "tools", "assets", "models", "summary.json")) or {}
        cls.today = str(s.get("built_at") or "")[:10]
        cls.pipe = _run_pipeline(cls.work, cls.today)        # 이름 주의: TestCase.run 을 가리면 unittest 가 깨진다

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_committed_outputs_reproduce(self):
        for name, rc, tail, swallowed in self.pipe:
            self.assertEqual(rc, 0, "%s rc=%d\n%s" % (name, rc, tail))
            self.assertEqual(swallowed, [], "%s 가 삼킨 오류: %s" % (name, swallowed[:3]))
        res = _compare(self.work, self.ref)
        non_page = [r for r in res["differ"] if not re.fullmatch(r"\d{6}/index\.html", r)]
        self.assertEqual(non_page, [], "커밋 %s 의 입력으로 다시 만든 산출이 커밋과 다름: %s" % (self.commit[:8], non_page[:10]))
        pages = {r[:6] for r in res["differ"] if re.fullmatch(r"\d{6}/index\.html", r)}
        sys.stderr.write("\n[e2e-git] 커밋 %s 재현: 같음 %d · 다른 페이지 %s(조선사 페이지는 '오늘' 날짜를 품음 — kship_page --today 없음)\n"
                         % (self.commit[:8], res["same"], sorted(pages)))
        self.assertTrue(pages <= YARDS, "조선사 외 페이지가 다름: %s" % sorted(pages - YARDS))


if __name__ == "__main__":
    unittest.main()
